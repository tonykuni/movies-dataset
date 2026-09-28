# -*- coding: utf-8 -*-
"""
VIA | via_fwdper_free4_fetch.py  v001  (2026-09-28)
================================================================
免費四源 Forward PER + Trailing PBR 日頻擷取器（附來源 / 口徑 / 可比較性標記）

四條免費 forward PER 線（其他一律不抓，避免混口徑）：
  1. FACTSET   Earnings Insight       S&P 500          週五   NTM analyst consensus     (直接值 + 隱含 fwd EPS 倒算)
  2. SPDJI     sp-500-eps-est.xlsx    S&P 500          週     bottom-up operating EPS   (close ÷ 未來4季 EPS 合計 → 日頻倒算)
  3. MSCI      indexes/index/<code>   MSCI USA/TW/KR/JP/CN/EM/World/ACWI  月底   NTM consensus (直接值)
  4. NIKKEI    日経平均プロフィル      Nikkei 225       日     公司自估 (予想PER)         (直接值；非 consensus)

Trailing PBR（免費層沒有 forward PBR，一律 trailing）：
  SPDJI BVPS(季) → close ÷ BVPS 日頻 ; MSCI P/BV(月) ; NIKKEI PBR(日)

可比較性（comparable_group）：
  MSCI_NTM     : MSCI 各指數之間可直接橫向比較（同方法、同 as_of）  ← 唯一跨市場可比組
  US_CONSENSUS : FactSet vs SPDJI 可對帳（同指數、皆分析師/官方估值，但 NTM vs 4季合計、口徑不同 → 只准對帳不准平均）
  JP_COMPANY   : Nikkei 予想PER 公司自估，只能跟自己歷史比，不得與 consensus 類橫比

執行：
  python via_fwdper_free4_fetch.py --selftest            # 離線 mock 閘門
  python via_fwdper_free4_fetch.py --activate            # 真跑（本機帶網路）
  python via_fwdper_free4_fetch.py --activate --days 250 --out D:/VIA/out
證據：forward 一律 Estimated_ThirdParty / Inferred；trailing PBR 官方值 = Confirmed_Official；抓不到 = Pending（不杜撰）
"""
from __future__ import annotations
import argparse, datetime as dt, io, json, math, os, re, sys, time, traceback
from dataclasses import dataclass, field, asdict
from typing import Optional, Callable

VERSION = "v001"
TODAY = dt.date.today()
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) VIA-fwdper-free4/" + VERSION}

# ------------------------------------------------------------------ config
MSCI_CODES = {   # code → (顯示名, 標題必須包含, 對應本地指數)
    "984000": ("MSCI USA",              "MSCI USA Index",              "S&P 500 proxy"),
    "915800": ("MSCI Taiwan",           "MSCI Taiwan Index",           "^TWII proxy"),   # 已驗證 2026-09
    "935600": ("MSCI Korea",            "MSCI Korea Index",            "KOSPI proxy"),
    "939200": ("MSCI Japan",            "MSCI Japan Index",            "TOPIX proxy"),
    "929700": ("MSCI China",            "MSCI China Index",            "HSI/A proxy"),
    "891800": ("MSCI Emerging Markets", "MSCI Emerging Markets Index", "EM"),
    "990100": ("MSCI World",            "MSCI World Index",            "DM"),
    "892400": ("MSCI ACWI",             "MSCI ACWI Index",             "Global"),
}
URL_FACTSET_TOPIC = "https://insight.factset.com/topic/earnings"
URL_FACTSET_PDF   = "https://advantage.factset.com/hubfs/Website/Resources%20Section/Research%20Desk/Earnings%20Insight/EarningsInsight_{mmddyy}{sfx}.pdf"
URL_SPDJI_XLSX    = "https://www.spglobal.com/spdji/en/documents/additional-material/sp-500-eps-est.xlsx"
URL_MSCI_INDEX    = "https://www.msci.com/indexes/index/{code}"
URL_NIKKEI_CANDS  = ["https://indexes.nikkei.co.jp/nkave/archives/data?list=per",
                     "https://indexes.nikkei.co.jp/nkave/statistics/per",
                     "https://indexes.nikkei.co.jp/en/nkave/archives/data?list=per"]
URL_STOOQ         = "https://stooq.com/q/d/l/?s={sym}&i=d"
PRICE_SYMS        = {"S&P 500": ("^spx", "^GSPC"), "Nikkei 225": ("^nkx", "^N225")}

# ------------------------------------------------------------------ record
@dataclass
class Row:
    date: str                 # 資料日期 (daily grid)
    as_of_source: str         # 來源公布/基準日
    index: str
    source: str               # FACTSET / SPDJI / MSCI / NIKKEI
    metric: str               # FWD_PER / FWD_EPS_IMPLIED / TRAIL_PBR / BVPS / FWD_EPS_12M
    value: Optional[float]
    unit: str
    method: str               # DIRECT / IMPLIED_EPS_FFILL / CLOSE_DIV_FWDEPS / CLOSE_DIV_BVPS / FFILL_MONTHLY
    basis: str                # NTM_CONSENSUS / BOTTOMUP_OPERATING_4Q / COMPANY_FORECAST / TRAILING_BOOK
    comparable_group: str     # MSCI_NTM / US_CONSENSUS / JP_COMPANY
    evidence: str             # Confirmed_Official / Estimated_ThirdParty / Inferred / Pending
    staleness_days: int
    quality_flag: str
    source_url: str

def R(**k) -> Row:
    return Row(**k)

# ------------------------------------------------------------------ transport (injectable → 可離線測)
class Transport:
    def __init__(self, timeout=25, tries=3):
        self.timeout, self.tries = timeout, tries
        self.log = []
    def get(self, url: str, binary=False):
        import urllib.request, urllib.error
        last = None
        for i in range(self.tries):
            try:
                req = urllib.request.Request(url, headers=UA)
                with urllib.request.urlopen(req, timeout=self.timeout) as r:
                    data = r.read()
                self.log.append((url, "OK", len(data)))
                return data if binary else data.decode("utf-8", "replace")
            except Exception as e:  # noqa
                last = e; self.log.append((url, "ERR", str(e)[:80])); time.sleep(0.8 * (i + 1))
        raise RuntimeError(f"GET failed {url}: {last}")

class MockTransport(Transport):
    def __init__(self, table: dict):
        super().__init__(); self.table = table
    def get(self, url, binary=False):
        for k, v in self.table.items():
            if k in url:
                self.log.append((url, "MOCK", 0))
                return v if (binary or not isinstance(v, bytes)) else v.decode()
        raise RuntimeError("mock miss " + url)

# ------------------------------------------------------------------ helpers
def strip_html(s: str) -> str:
    s = re.sub(r"<script.*?</script>|<style.*?</style>", " ", s, flags=re.S | re.I)
    s = re.sub(r"<[^>]+>", "\n", s)
    s = re.sub(r"&nbsp;|&#160;", " ", s); s = re.sub(r"&amp;", "&", s)
    return re.sub(r"\n\s*\n+", "\n", s)

def fnum(x) -> Optional[float]:
    try:
        v = float(str(x).replace(",", "").strip())
        return v if math.isfinite(v) else None
    except Exception:
        return None

_MON = {m: i for i, m in enumerate(["jan","feb","mar","apr","may","jun","jul","aug","sep","oct","nov","dec"], 1)}
def parse_en_date(s: str) -> Optional[dt.date]:
    m = re.search(r"([A-Za-z]{3})[a-z.]*\s+(\d{1,2}),?\s+(\d{4})", s)
    if not m: return None
    mo = _MON.get(m.group(1).lower()[:3])
    return dt.date(int(m.group(3)), mo, int(m.group(2))) if mo else None

def daily_grid(days: int) -> list[dt.date]:
    return [TODAY - dt.timedelta(days=i) for i in range(days, -1, -1)]

# ------------------------------------------------------------------ 1. FACTSET
def parse_factset_article(html: str) -> dict:
    txt = strip_html(html)
    out = {}
    m = re.search(r"forward 12-month P/E ratio for the S&P 500 (?:is|was|stands at|of) (\d{1,2}\.\d)", txt, re.I)
    if m: out["fwd_pe"] = float(m.group(1))
    m = re.search(r"forward 12-month EPS estimate(?: of)? \(?\$?(\d{2,3}\.\d{2})", txt, re.I) \
        or re.search(r"forward 12-month EPS estimate of \$(\d{2,3}\.\d{2})", txt, re.I)
    if m: out["fwd_eps"] = float(m.group(1))
    m = re.search(r"closing price(?: of)? \(?(\d{1,2},?\d{3}\.\d{2})", txt, re.I)
    if m: out["close_ref"] = fnum(m.group(1))
    d = None
    for line in txt.splitlines()[:400]:
        d = parse_en_date(line)
        if d and abs((TODAY - d).days) < 400: break
        d = None
    if d: out["as_of"] = d.isoformat()
    return out

def fetch_factset(tp: Transport) -> tuple[dict, str]:
    """回傳 (parsed, url)。先 HTML 文章，失敗再試最近 4 個週五 PDF（需 pypdf）。"""
    try:
        idx = tp.get(URL_FACTSET_TOPIC)
        links = re.findall(r'href="(https://insight\.factset\.com/[^"]*(?:earnings-season-update|earnings-insight|sp-500)[^"]*)"', idx, re.I)
        for u in links[:6]:
            try:
                p = parse_factset_article(tp.get(u))
                if "fwd_pe" in p: return p, u
            except Exception: pass
    except Exception: pass
    # PDF fallback
    try:
        import pypdf  # type: ignore
    except Exception:
        return {}, URL_FACTSET_TOPIC
    d = TODAY
    for _ in range(35):
        if d.weekday() == 4:
            for sfx in ("", "A"):
                u = URL_FACTSET_PDF.format(mmddyy=d.strftime("%m%d%y"), sfx=sfx)
                try:
                    raw = tp.get(u, binary=True)
                    txt = "\n".join((pg.extract_text() or "") for pg in pypdf.PdfReader(io.BytesIO(raw)).pages[:4])
                    p = parse_factset_article(txt); p.setdefault("as_of", d.isoformat())
                    if "fwd_pe" in p: return p, u
                except Exception: pass
        d -= dt.timedelta(days=1)
    return {}, URL_FACTSET_TOPIC

# ------------------------------------------------------------------ 2. S&P DJI xlsx
def parse_spdji_xlsx(raw: bytes) -> dict:
    """容錯掃描：找 ESTIMATES 區塊的季度 operating EPS 估值，與 QUARTERLY DATA 的 BOOK VALUE PER SHARE。"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    out = {"est_quarters": [], "bvps": []}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        name = (ws.title or "").upper()
        # --- 估值區塊
        if "EST" in name or any("ESTIMATES" in str(c).upper() for r in rows[:60] for c in r if c):
            in_est, seen_actual = False, False
            for r in rows:
                cells = [c for c in r]
                head = " ".join(str(c).upper() for c in cells if c is not None)
                if "ESTIMATES" in head and "ACTUAL" not in head: in_est = True; continue
                if "ACTUAL" in head and in_est: seen_actual = True; in_est = False; continue
                if in_est and cells and isinstance(cells[0], (dt.date, dt.datetime)):
                    nums = [fnum(c) for c in cells[1:] if fnum(c) is not None]
                    if nums:
                        q = cells[0].date() if isinstance(cells[0], dt.datetime) else cells[0]
                        out["est_quarters"].append((q.isoformat(), nums[0]))   # 第一個數值欄 = operating EPS 估值
        # --- BVPS
        hdr_idx = None
        for i, r in enumerate(rows[:80]):
            for j, c in enumerate(r):
                if c and "BOOK VALUE" in str(c).upper(): hdr_idx = (i, j)
        if hdr_idx:
            i0, j0 = hdr_idx
            for r in rows[i0 + 1:]:
                if r and isinstance(r[0], (dt.date, dt.datetime)) and j0 < len(r) and fnum(r[j0]):
                    q = r[0].date() if isinstance(r[0], dt.datetime) else r[0]
                    out["bvps"].append((q.isoformat(), fnum(r[j0])))
    out["est_quarters"] = sorted(set(out["est_quarters"]))
    out["bvps"] = sorted(set(out["bvps"]))
    return out

def spdji_fwd12m_eps(est_quarters: list[tuple[str, float]], on: dt.date) -> Optional[float]:
    fut = [v for d, v in est_quarters if dt.date.fromisoformat(d) > on]
    return round(sum(fut[:4]), 2) if len(fut) >= 4 else None

# ------------------------------------------------------------------ 3. MSCI
def parse_msci_page(html: str, must_contain: str) -> dict:
    txt = strip_html(html)
    out = {"title_ok": must_contain.lower() in txt.lower()}
    for key, lab in (("pe_fwd", r"P/E Fwd"), ("pe", r"P/E(?! Fwd)"), ("pbv", r"P/BV"), ("dy", r"Div Yld \(%\)")):
        m = re.search(lab + r"\s*\n\s*([\d.]+|na)", txt)
        if m: out[key] = None if m.group(1) == "na" else float(m.group(1))
    m = re.search(r"Data as of\s+([A-Za-z.]+\s+\d{1,2},\s+\d{4})", txt)
    if m:
        d = parse_en_date(m.group(1))
        if d: out["as_of"] = d.isoformat()
    return out

# ------------------------------------------------------------------ 4. NIKKEI
def parse_nikkei_per(html: str) -> list[tuple[str, float, Optional[float]]]:
    """回傳 [(date, PER, PBR)]；容忍 HTML 表格或 CSV。日期 YYYY/MM/DD 或 YYYY-MM-DD。"""
    txt = strip_html(html) if "<" in html[:2000] else html
    rows = []
    for m in re.finditer(r"(20\d{2})[/\-.](\d{1,2})[/\-.](\d{1,2})\D{1,40}?(\d{1,3}\.\d{1,2})(?:\D{1,40}?(\d{1,2}\.\d{1,2}))?", txt):
        try:
            d = dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            continue
        per, pbr = float(m.group(4)), (float(m.group(5)) if m.group(5) else None)
        if 5 <= per <= 80:
            rows.append((d.isoformat(), per, pbr if (pbr and 0.3 <= pbr <= 6) else None))
    return sorted(set(rows))

# ------------------------------------------------------------------ prices
def parse_stooq_csv(csv: str) -> list[tuple[str, float]]:
    out = []
    for line in csv.splitlines()[1:]:
        p = line.split(",")
        if len(p) >= 5 and fnum(p[4]) and re.match(r"\d{4}-\d{2}-\d{2}", p[0]):
            out.append((p[0], fnum(p[4])))
    return out

def fetch_prices(tp: Transport, name: str) -> tuple[list[tuple[str, float]], str]:
    stq, yf = PRICE_SYMS[name]
    u = URL_STOOQ.format(sym=stq)
    try:
        px = parse_stooq_csv(tp.get(u))
        if px: return px, u
    except Exception: pass
    try:
        import yfinance  # type: ignore
        h = yfinance.Ticker(yf).history(period="2y")
        return [(d.strftime("%Y-%m-%d"), float(v)) for d, v in h["Close"].items()], "yfinance:" + yf
    except Exception:
        return [], u

def price_on(px: dict[str, float], d: dt.date, back=7) -> tuple[Optional[str], Optional[float]]:
    for i in range(back):
        k = (d - dt.timedelta(days=i)).isoformat()
        if k in px: return k, px[k]
    return None, None

# ------------------------------------------------------------------ assemble daily long table
def assemble(days: int, fs: dict, fs_url: str, sp: dict, sp_url: str, msci: dict, nk: list, nk_url: str,
             px_spx: dict, px_nk: dict, px_urls: dict) -> list[Row]:
    rows: list[Row] = []
    grid = daily_grid(days)
    # ---- 1 FACTSET : 直接值 + 隱含 EPS 前推
    if fs.get("fwd_pe"):
        asof = dt.date.fromisoformat(fs.get("as_of", TODAY.isoformat()))
        k0, c0 = price_on(px_spx, asof)
        implied_eps = fs.get("fwd_eps") or (round(c0 / fs["fwd_pe"], 2) if c0 else None)
        rows.append(R(date=asof.isoformat(), as_of_source=asof.isoformat(), index="S&P 500", source="FACTSET",
                      metric="FWD_PER", value=fs["fwd_pe"], unit="x", method="DIRECT", basis="NTM_CONSENSUS",
                      comparable_group="US_CONSENSUS", evidence="Estimated_ThirdParty", staleness_days=0,
                      quality_flag="OK", source_url=fs_url))
        if implied_eps:
            rows.append(R(date=asof.isoformat(), as_of_source=asof.isoformat(), index="S&P 500", source="FACTSET",
                          metric="FWD_EPS_IMPLIED", value=implied_eps, unit="USD", method="CLOSE_DIV_PE" if not fs.get("fwd_eps") else "DIRECT",
                          basis="NTM_CONSENSUS", comparable_group="US_CONSENSUS",
                          evidence="Inferred" if not fs.get("fwd_eps") else "Estimated_ThirdParty",
                          staleness_days=0, quality_flag="OK", source_url=fs_url))
            for d in grid:
                if d <= asof: continue
                k, c = price_on(px_spx, d, back=1)
                if not c: continue
                rows.append(R(date=d.isoformat(), as_of_source=asof.isoformat(), index="S&P 500", source="FACTSET",
                              metric="FWD_PER", value=round(c / implied_eps, 2), unit="x", method="IMPLIED_EPS_FFILL",
                              basis="NTM_CONSENSUS", comparable_group="US_CONSENSUS", evidence="Inferred",
                              staleness_days=(d - asof).days, quality_flag="FFILL_EPS_FROM_LAST_REPORT", source_url=fs_url))
    else:
        rows.append(R(date=TODAY.isoformat(), as_of_source="", index="S&P 500", source="FACTSET", metric="FWD_PER",
                      value=None, unit="x", method="DIRECT", basis="NTM_CONSENSUS", comparable_group="US_CONSENSUS",
                      evidence="Pending", staleness_days=-1, quality_flag="FAIL_NO_SOURCE", source_url=fs_url))
    # ---- 2 SPDJI : close ÷ 未來4季 EPS（現行 vintage）；close ÷ BVPS
    eq, bv = sp.get("est_quarters", []), sp.get("bvps", [])
    if eq:
        for d in grid:
            k, c = price_on(px_spx, d, back=1)
            if not c: continue
            e12 = spdji_fwd12m_eps(eq, d)
            if e12:
                rows.append(R(date=d.isoformat(), as_of_source=TODAY.isoformat(), index="S&P 500", source="SPDJI",
                              metric="FWD_PER", value=round(c / e12, 2), unit="x", method="CLOSE_DIV_FWDEPS",
                              basis="BOTTOMUP_OPERATING_4Q", comparable_group="US_CONSENSUS", evidence="Inferred",
                              staleness_days=0, quality_flag="OK" if d >= TODAY - dt.timedelta(days=7) else "CURRENT_VINTAGE_APPLIED_TO_HISTORY",
                              source_url=sp_url))
        rows.append(R(date=TODAY.isoformat(), as_of_source=TODAY.isoformat(), index="S&P 500", source="SPDJI",
                      metric="FWD_EPS_12M", value=spdji_fwd12m_eps(eq, TODAY), unit="USD", method="SUM_NEXT_4Q",
                      basis="BOTTOMUP_OPERATING_4Q", comparable_group="US_CONSENSUS", evidence="Estimated_ThirdParty",
                      staleness_days=0, quality_flag="OK", source_url=sp_url))
    else:
        rows.append(R(date=TODAY.isoformat(), as_of_source="", index="S&P 500", source="SPDJI", metric="FWD_PER",
                      value=None, unit="x", method="CLOSE_DIV_FWDEPS", basis="BOTTOMUP_OPERATING_4Q",
                      comparable_group="US_CONSENSUS", evidence="Pending", staleness_days=-1,
                      quality_flag="FAIL_NO_SOURCE_OR_PARSE", source_url=sp_url))
    if bv:
        past = [(d, v) for d, v in bv if dt.date.fromisoformat(d) <= TODAY]
        if past:
            bd, bvps = past[-1]
            for d in grid:
                k, c = price_on(px_spx, d, back=1)
                if not c: continue
                rows.append(R(date=d.isoformat(), as_of_source=bd, index="S&P 500", source="SPDJI", metric="TRAIL_PBR",
                              value=round(c / bvps, 3), unit="x", method="CLOSE_DIV_BVPS", basis="TRAILING_BOOK",
                              comparable_group="US_CONSENSUS", evidence="Inferred",
                              staleness_days=(d - dt.date.fromisoformat(bd)).days, quality_flag="TRAILING_ONLY", source_url=sp_url))
            rows.append(R(date=bd, as_of_source=bd, index="S&P 500", source="SPDJI", metric="BVPS", value=bvps, unit="USD",
                          method="DIRECT", basis="TRAILING_BOOK", comparable_group="US_CONSENSUS", evidence="Confirmed_Official",
                          staleness_days=0, quality_flag="OK", source_url=sp_url))
    # ---- 3 MSCI : 月底直接值 + 日頻 ffill 標記
    for code, (name, _, proxy) in MSCI_CODES.items():
        p = msci.get(code) or {}
        url = URL_MSCI_INDEX.format(code=code)
        if not p or not p.get("title_ok") or p.get("pe_fwd") is None:
            rows.append(R(date=TODAY.isoformat(), as_of_source=p.get("as_of", ""), index=name, source="MSCI", metric="FWD_PER",
                          value=None, unit="x", method="DIRECT", basis="NTM_CONSENSUS", comparable_group="MSCI_NTM",
                          evidence="Pending", staleness_days=-1,
                          quality_flag=("FAIL_CODE_MISMATCH" if p and not p.get("title_ok") else "FAIL_NO_SOURCE_OR_NA"), source_url=url))
            continue
        asof = dt.date.fromisoformat(p.get("as_of", TODAY.isoformat()))
        for metric, key, basis, ev, flag in (("FWD_PER", "pe_fwd", "NTM_CONSENSUS", "Estimated_ThirdParty", "OK"),
                                             ("TRAIL_PBR", "pbv", "TRAILING_BOOK", "Confirmed_Official", "TRAILING_ONLY"),
                                             ("TRAIL_PER", "pe", "TRAILING_EARNINGS", "Confirmed_Official", "TRAILING_ONLY")):
            if p.get(key) is None: continue
            rows.append(R(date=asof.isoformat(), as_of_source=asof.isoformat(), index=name, source="MSCI", metric=metric,
                          value=p[key], unit="x", method="DIRECT", basis=basis, comparable_group="MSCI_NTM", evidence=ev,
                          staleness_days=0, quality_flag=flag, source_url=url))
            for d in grid:
                if d <= asof or d.weekday() >= 5: continue
                rows.append(R(date=d.isoformat(), as_of_source=asof.isoformat(), index=name, source="MSCI", metric=metric,
                              value=p[key], unit="x", method="FFILL_MONTHLY", basis=basis, comparable_group="MSCI_NTM",
                              evidence="Inferred", staleness_days=(d - asof).days,
                              quality_flag="FFILL_MONTHLY_NO_PRICE_ADJ", source_url=url))
    # ---- 4 NIKKEI : 日頻直接值
    if nk:
        for d, per, pbr in nk:
            if dt.date.fromisoformat(d) < grid[0]: continue
            rows.append(R(date=d, as_of_source=d, index="Nikkei 225", source="NIKKEI", metric="FWD_PER", value=per, unit="x",
                          method="DIRECT", basis="COMPANY_FORECAST", comparable_group="JP_COMPANY",
                          evidence="Estimated_ThirdParty", staleness_days=0, quality_flag="COMPANY_FORECAST_NOT_CONSENSUS", source_url=nk_url))
            if pbr:
                rows.append(R(date=d, as_of_source=d, index="Nikkei 225", source="NIKKEI", metric="TRAIL_PBR", value=pbr, unit="x",
                              method="DIRECT", basis="TRAILING_BOOK", comparable_group="JP_COMPANY",
                              evidence="Confirmed_Official", staleness_days=0, quality_flag="TRAILING_ONLY", source_url=nk_url))
    else:
        rows.append(R(date=TODAY.isoformat(), as_of_source="", index="Nikkei 225", source="NIKKEI", metric="FWD_PER",
                      value=None, unit="x", method="DIRECT", basis="COMPANY_FORECAST", comparable_group="JP_COMPANY",
                      evidence="Pending", staleness_days=-1, quality_flag="FAIL_NO_SOURCE_VERIFY_URL", source_url=nk_url))
    # ---- price rows (audit)
    for name, px, in (("S&P 500", px_spx), ("Nikkei 225", px_nk)):
        for d in grid:
            k = d.isoformat()
            if k in px:
                rows.append(R(date=k, as_of_source=k, index=name, source="STOOQ/YF", metric="CLOSE", value=px[k], unit="pt",
                              method="DIRECT", basis="PRICE", comparable_group="-", evidence="Confirmed_Official",
                              staleness_days=0, quality_flag="OK", source_url=px_urls.get(name, "")))
    return rows

# ------------------------------------------------------------------ gates
def gate(rows: list[Row]) -> list[tuple[str, bool, str]]:
    g = []
    fwd = [r for r in rows if r.metric.startswith("FWD")]
    g.append(("G1 forward 不得 Confirmed", all(r.evidence != "Confirmed_Official" for r in fwd), f"{len(fwd)} fwd rows"))
    g.append(("G2 每列有 source_url", all(r.source_url for r in rows), ""))
    g.append(("G3 無 forward PBR 混入", not any(r.metric == "FWD_PBR" for r in rows), ""))
    mix = {(r.index, r.source, r.metric, r.basis) for r in rows if r.metric == "FWD_PER"}
    g.append(("G4 同源同指標單一 basis", len(mix) == len({(a, b, c) for a, b, c, _ in mix}), str(len(mix))))
    g.append(("G5 Nikkei 不在 consensus 組", all(r.comparable_group == "JP_COMPANY" for r in rows if r.source == "NIKKEI"), ""))
    g.append(("G6 MSCI 皆 MSCI_NTM 組", all(r.comparable_group == "MSCI_NTM" for r in rows if r.source == "MSCI"), ""))
    g.append(("G7 Pending 值為 null", all(r.value is None for r in rows if r.evidence == "Pending"), ""))
    g.append(("G8 ffill 列 staleness>0", all(r.staleness_days > 0 for r in rows if r.method in ("FFILL_MONTHLY", "IMPLIED_EPS_FFILL")), ""))
    return g

# ------------------------------------------------------------------ output
def write_outputs(rows: list[Row], gates, out_dir: str, tlog) -> dict:
    os.makedirs(out_dir, exist_ok=True)
    stamp = TODAY.strftime("%Y%m%d")
    base = os.path.join(out_dir, f"VIA_fwdper_free4_{stamp}")
    recs = [asdict(r) for r in rows]
    # CSV (utf-8 no BOM)
    import csv
    with open(base + ".csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(recs[0].keys())); w.writeheader(); w.writerows(recs)
    paths = {"csv": base + ".csv"}
    try:
        import pandas as pd
        pd.DataFrame(recs).to_parquet(base + ".parquet", index=False); paths["parquet"] = base + ".parquet"
    except Exception as e:
        paths["parquet"] = f"(skipped: {str(e)[:60]})"
    with open(base + "_audit.json", "w", encoding="utf-8") as f:
        json.dump({"version": VERSION, "run": TODAY.isoformat(), "transport": tlog, "gates": gates}, f, ensure_ascii=False, indent=1)
    paths["audit"] = base + "_audit.json"
    paths["html"] = write_html(rows, gates, base + ".html")
    return paths

def latest_matrix(rows: list[Row]):
    best = {}
    for r in rows:
        if r.metric not in ("FWD_PER", "TRAIL_PBR"): continue
        k = (r.index, r.source, r.metric)
        if k not in best or (r.value is not None and r.date > best[k].date): best[k] = r
    return sorted(best.values(), key=lambda r: (r.comparable_group, r.index, r.source, r.metric))

def write_html(rows, gates, path):
    mx = latest_matrix(rows)
    def td(v): return "—" if v is None else (f"{v:.2f}" if isinstance(v, float) else str(v))
    trs = "".join(
        f"<tr class='{'pend' if r.evidence=='Pending' else ''}'><td>{r.comparable_group}</td><td>{r.index}</td><td>{r.source}</td>"
        f"<td>{r.metric}</td><td class='num'>{td(r.value)}</td><td>{r.basis}</td><td>{r.method}</td><td>{r.as_of_source or '—'}</td>"
        f"<td>{r.staleness_days if r.staleness_days>=0 else '—'}</td><td>{r.evidence}</td><td>{r.quality_flag}</td>"
        f"<td><a href='{r.source_url}' target='_blank'>src</a></td></tr>" for r in mx)
    gts = "".join(f"<li class='{'ok' if ok else 'ng'}'>{'PASS' if ok else 'FAIL'} · {n} <span>{m}</span></li>" for n, ok, m in gates)
    html = f"""<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>VIA · 免費四源 Forward PER / Trailing PBR · {TODAY}</title>
<link href="https://fonts.googleapis.com/css2?family=Syne:wght@600;800&family=DM+Sans:wght@400;500&family=DM+Mono&display=swap" rel="stylesheet">
<style>
:root{{--bg:#f5f4f0;--paper:#fff;--ink:#1e1d1a;--line:#dbd9d3;--blue:#4c78a8;--teal:#439a9a;--up:#c96b5a;--dn:#5a9e6f}}
body{{margin:0;background:var(--bg);color:var(--ink);font-family:'DM Sans',system-ui,sans-serif;font-size:14px}}
.wrap{{max-width:1380px;margin:0 auto;padding:28px 22px}}
h1{{font-family:Syne,sans-serif;font-weight:800;font-size:24px;margin:0 0 4px}} .sub{{color:#6b6963;font-family:'DM Mono',monospace;font-size:12px}}
.strip{{height:4px;margin:14px 0 18px;background:linear-gradient(90deg,#4c78a8,#439a9a,#5a9e6f,#d9a441,#c96b5a,#8b5fa8,#1e1d1a)}}
.seal{{display:inline-block;border:2px solid var(--up);color:var(--up);font-family:Syne;font-weight:800;padding:2px 7px;border-radius:2px;margin-left:8px}}
.card{{background:var(--paper);border:1px solid var(--line);border-radius:3px;padding:16px 18px;margin-bottom:16px}}
h2{{font-family:Syne;font-size:15px;margin:0 0 10px}}
table{{border-collapse:collapse;width:100%;font-size:12.5px}} th,td{{border-bottom:1px solid var(--line);padding:6px 8px;text-align:left;white-space:nowrap}}
th{{font-family:'DM Mono',monospace;font-weight:500;color:#6b6963;font-size:11px}} td.num{{font-family:'DM Mono',monospace;text-align:right}}
tr.pend td{{color:#9a9890;font-style:italic}} .tbl{{overflow-x:auto}}
ul.g{{list-style:none;padding:0;margin:0;font-family:'DM Mono',monospace;font-size:12px}} ul.g li{{padding:4px 0;border-bottom:1px dashed var(--line)}}
li.ok{{color:var(--dn)}} li.ng{{color:var(--up)}} li span{{color:#9a9890;margin-left:8px}}
.note{{font-size:12.5px;color:#4a4944;line-height:1.55}} code{{font-family:'DM Mono',monospace;background:var(--bg);padding:1px 4px;border-radius:2px}}
</style></head><body><div class="wrap">
<h1>免費四源 Forward PER / Trailing PBR 日頻矩陣 <span class="seal">庫</span></h1>
<div class="sub">VIA · via_fwdper_free4_fetch {VERSION} · run {TODAY} · rows {len(rows)}</div><div class="strip"></div>
<div class="card"><h2>最新值矩陣（每 指數×來源×指標 取最新日）</h2><div class="tbl"><table>
<tr><th>可比組</th><th>指數</th><th>來源</th><th>指標</th><th>值</th><th>口徑 basis</th><th>取法 method</th><th>as_of</th><th>stale(d)</th><th>證據</th><th>quality</th><th>URL</th></tr>{trs}</table></div></div>
<div class="card"><h2>口徑閘門</h2><ul class="g">{gts}</ul></div>
<div class="card"><h2>可比較性規則</h2><div class="note">
<b>MSCI_NTM</b>：唯一可跨市場橫向比較的組（USA / Taiwan / Korea / Japan / China / EM / World / ACWI 同方法、同月底 as_of）。月頻，日頻列為 <code>FFILL_MONTHLY</code>（未按價格調整），staleness 已標。<br>
<b>US_CONSENSUS</b>：FactSet（NTM consensus，週）與 S&amp;P DJI（未來 4 季 operating EPS 合計，官方估值表）同指數可<b>對帳</b>不可平均；S&amp;P DJI 日頻為 <code>CLOSE_DIV_FWDEPS</code> 以現行 vintage 回套歷史，已標 <code>CURRENT_VINTAGE_APPLIED_TO_HISTORY</code>。<br>
<b>JP_COMPANY</b>：日經 予想PER 為公司自估（非分析師 consensus），只能與自身歷史比，不得與 MSCI Japan / FactSet 橫比。<br>
<b>PBR</b>：免費層全為 trailing（<code>TRAILING_ONLY</code>），無 forward PBR。<br>
證據：forward 皆 <code>Estimated_ThirdParty</code> 或 <code>Inferred</code>；抓不到為 <code>Pending</code>、值為 null，不杜撰。非投資建議。</div></div>
</div></body></html>"""
    with open(path, "w", encoding="utf-8") as f: f.write(html)
    return path

# ------------------------------------------------------------------ run
def run(tp: Transport, days: int, out_dir: str, open_html=True) -> tuple[list[Row], list, dict]:
    print(f"[VIA] fwdper-free4 {VERSION}  days={days}  out={out_dir}")
    px_spx_l, u1 = fetch_prices(tp, "S&P 500"); px_nk_l, u2 = fetch_prices(tp, "Nikkei 225")
    px_spx, px_nk = dict(px_spx_l), dict(px_nk_l)
    print(f"  prices  SPX={len(px_spx)}  NKX={len(px_nk)}")
    fs, fs_url = fetch_factset(tp); print(f"  FACTSET {fs or 'PENDING'}")
    try:
        sp = parse_spdji_xlsx(tp.get(URL_SPDJI_XLSX, binary=True))
    except Exception as e:
        sp = {}; print(f"  SPDJI   ERR {str(e)[:70]}")
    print(f"  SPDJI   est_q={len(sp.get('est_quarters', []))}  bvps={len(sp.get('bvps', []))}")
    msci = {}
    for code, (name, must, _) in MSCI_CODES.items():
        try:
            msci[code] = parse_msci_page(tp.get(URL_MSCI_INDEX.format(code=code)), must)
        except Exception as e:
            msci[code] = {}
        p = msci[code]; print(f"  MSCI    {name:<22} fwd={p.get('pe_fwd')} pbv={p.get('pbv')} asof={p.get('as_of')} title_ok={p.get('title_ok')}")
    nk, nk_url = [], URL_NIKKEI_CANDS[0]
    for u in URL_NIKKEI_CANDS:
        try:
            nk = parse_nikkei_per(tp.get(u))
            if nk: nk_url = u; break
        except Exception: pass
    print(f"  NIKKEI  rows={len(nk)}")
    rows = assemble(days, fs, fs_url, sp, URL_SPDJI_XLSX, msci, nk, nk_url, px_spx, px_nk, {"S&P 500": u1, "Nikkei 225": u2})
    gates = gate(rows)
    paths = write_outputs(rows, gates, out_dir, tp.log)
    for n, ok, m in gates: print(f"  {'PASS' if ok else 'FAIL'} {n} {m}")
    print("  outputs:", json.dumps(paths, ensure_ascii=False, indent=1))
    if open_html and sys.platform.startswith("win"):
        try: os.startfile(paths["html"])  # type: ignore[attr-defined]
        except Exception: pass
    return rows, gates, paths

# ------------------------------------------------------------------ selftest (offline)
def _mock_xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "ESTIMATES&PEs"
    ws.append(["ESTIMATES", None, None]); 
    q = [("2027-06-30", 78.1), ("2027-03-31", 76.0), ("2026-12-31", 74.5), ("2026-09-30", 71.2), ("2026-06-30", 69.9)]
    for d, v in q: ws.append([dt.datetime.fromisoformat(d), v, v * 0.95])
    ws.append(["ACTUALS", None, None]); ws.append([dt.datetime(2026, 3, 31), 66.0, 63.0])
    ws2 = wb.create_sheet("QUARTERLY DATA"); ws2.append(["QUARTER END", "BOOK VALUE PER SHARE", "X"])
    ws2.append([dt.datetime(2026, 6, 30), 1312.5, 1]); ws2.append([dt.datetime(2026, 3, 31), 1288.0, 1])
    b = io.BytesIO(); wb.save(b); return b.getvalue()

def selftest() -> int:
    global TODAY
    TODAY = dt.date(2026, 9, 28)
    px = "Date,Open,High,Low,Close,Volume\n" + "\n".join(f"{(TODAY-dt.timedelta(days=i)).isoformat()},1,1,1,{6900-i},1" for i in range(60, -1, -1) if (TODAY-dt.timedelta(days=i)).weekday() < 5)
    nkpx = "Date,Open,High,Low,Close,Volume\n" + "\n".join(f"{(TODAY-dt.timedelta(days=i)).isoformat()},1,1,1,{48000-i},1" for i in range(60, -1, -1) if (TODAY-dt.timedelta(days=i)).weekday() < 5)
    fs_idx = '<a href="https://insight.factset.com/sp-500-earnings-season-update-september-25-2026">x</a>'
    fs_art = "<html><body><p>September 25, 2026</p><p>The forward 12-month P/E ratio for the S&P 500 is 20.3. This P/E ratio is above the five-year average.</p></body></html>"
    msci_tw = "<html><title>MSCI Taiwan Index</title><body>MSCI Taiwan Index\nIndex code\n915800\nDiv Yld (%)\n1.44\nP/E\n31.29\nP/E Fwd\n19.83\nP/BV\n5.81\nData as of July 31, 2026</body></html>"
    msci_fm = "<html><body>MSCI Frontier Markets Index\nP/E\n13.38\nP/E Fwd\nna\nP/BV\n1.98\nData as of Apr. 30, 2026</body></html>"
    nk = "<table><tr><td>2026/09/25</td><td>17.21</td><td>1.92</td></tr><tr><td>2026/09/24</td><td>17.05</td><td>1.90</td></tr></table>"
    table = {"stooq.com/q/d/l/?s=^spx": px, "stooq.com/q/d/l/?s=^nkx": nkpx, "insight.factset.com/topic": fs_idx,
             "sp-500-earnings-season-update": fs_art, "sp-500-eps-est.xlsx": _mock_xlsx(),
             "indexes/index/915800": msci_tw, "indexes/index/984000": msci_fm, "archives/data?list=per": nk}
    tp = MockTransport(table)
    checks = []
    def C(name, cond, info=""): checks.append((name, bool(cond), info)); print(f"  {'PASS' if cond else 'FAIL'} {name} {info}")
    # unit parses
    p = parse_factset_article(fs_art); C("factset parse pe", p.get("fwd_pe") == 20.3, str(p))
    C("factset parse date", p.get("as_of") == "2026-09-25")
    m = parse_msci_page(msci_tw, "MSCI Taiwan Index"); C("msci parse", m.get("pe_fwd") == 19.83 and m.get("pbv") == 5.81 and m.get("as_of") == "2026-07-31" and m["title_ok"], str(m))
    m2 = parse_msci_page(msci_fm, "MSCI USA Index"); C("msci na + mismatch", m2.get("pe_fwd") is None and not m2["title_ok"], str(m2))
    sp = parse_spdji_xlsx(_mock_xlsx()); C("spdji est quarters", len(sp["est_quarters"]) == 5, str(sp["est_quarters"][:2]))
    C("spdji bvps", sp["bvps"] and sp["bvps"][-1][1] == 1312.5, str(sp["bvps"]))
    e12 = spdji_fwd12m_eps(sp["est_quarters"], TODAY); C("spdji fwd12m sum next4", e12 == round(71.2 + 74.5 + 76.0 + 78.1, 2), str(e12))
    nkr = parse_nikkei_per(nk); C("nikkei parse", len(nkr) == 2 and nkr[-1] == ("2026-09-25", 17.21, 1.92), str(nkr))
    C("stooq parse", len(parse_stooq_csv(px)) > 30)
    # end-to-end (mock; 其他 MSCI code 走 mock miss → Pending)
    out = os.path.join(os.getcwd(), "_selftest_out")
    rows, gates, paths = run(tp, 60, out, open_html=False)
    for n, ok, msg in gates: C(n, ok, msg)
    src = {r.source for r in rows}; C("all 4 sources present", {"FACTSET", "SPDJI", "MSCI", "NIKKEI"} <= src, str(src))
    fs_rows = [r for r in rows if r.source == "FACTSET" and r.metric == "FWD_PER"]
    C("factset implied ffill has staleness", any(r.method == "IMPLIED_EPS_FFILL" and r.staleness_days > 0 for r in fs_rows))
    C("pending rows for unmocked msci", any(r.source == "MSCI" and r.evidence == "Pending" for r in rows))
    C("csv written", os.path.exists(paths["csv"])); C("html written", os.path.exists(paths["html"]))
    C("no FWD_PBR anywhere", not any(r.metric == "FWD_PBR" for r in rows))
    n_ok = sum(1 for _, ok, _ in checks if ok)
    print(f"\n[SELFTEST] {n_ok}/{len(checks)} PASS")
    return 0 if n_ok == len(checks) else 1

# ------------------------------------------------------------------ main
if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--selftest", action="store_true"); ap.add_argument("--activate", action="store_true")
    ap.add_argument("--days", type=int, default=400); ap.add_argument("--out", default=os.path.join(os.getcwd(), "VIA_out_fwdper"))
    ap.add_argument("--no-open", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        sys.exit(selftest())
    try:
        run(Transport(), a.days, a.out, open_html=not a.no_open)
    except Exception:
        traceback.print_exc(); sys.exit(2)
