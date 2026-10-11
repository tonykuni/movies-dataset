#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""VRN_SystemManager v0153 — 薄尾:prep 動詞 = VRN 第一步(操作員 2026-10-10)。只動 VRN 自己的範圍(VRN 樹 · VIA_Reports/vrn · TEMP);VDF 個股清單唯讀;VCGC 不參與。
  prep [--dir <夾>] [--no-vdf] [--max-pages 60] [--workers 4]
   ① 掃描:所有輸入檔的類型 · 大小 · 規格(PDF 頁數 / 頁面尺寸 / 前 3 頁文字量 / 加密;影像 px / 原 dpi / 色彩)
   ② 轉換:Word(.docx/.doc)· TXT · 影像 → PDF,放 VIA_Reports/vrn/prep/converted/
        影像:不重採樣、原像素 · 標 325 dpi(PIL resolution=325)
        Word:Word COM(win32com → PowerShell COM)→ LibreOffice → docx2pdf → python-docx 文字版(最後一階,標黃)
        TXT :PyMuPDF → reportlab(繁中 CID 字型)· 自動判編碼(utf-8 / utf-16 / cp950 / gb18030)
        Word/TXT 輸出向量 PDF(文字可選取 = 比 325 dpi 點陣更好讀);--raster-all 才一律轉 325 dpi 影像
   ③ 讀取失敗補救:pdfplumber 開不了 → pikepdf 修復到 TEMP 再讀;仍開不了或前 3 頁文字 < 40 字(掃描件)→ 每頁 325 dpi 點陣到 TEMP(PyMuPDF → pypdfium2 → pdftoppm)
   ④ 檔名判讀:報告日 · 代號 · yfinance 代號 · Bloomberg 代號 · 名稱 · 券商
        日期 regex:YYYYMMDD · 民國 1YYMMDD · YYMMDD · CTBC+MMDD(年取檔案時間)→ 都沒有才用檔案時間(標黃)
        代號 regex:分隔符包住的四碼 · 「(3706 TT)」/「3014TT」Bloomberg 形 · 「2330.TW(O)」yfinance 形 · 年份/日期不算
        名冊:VRN 自家名冊(VRN_TWRoster_Offline · VRN_TWTicker_RIE)+ VDF 個股清單(唯讀;--no-vdf 略過)→ 名稱 · 市場別 → yfinance(.TW 上市 / .TWO 上櫃)· Bloomberg(<代號> TT)
        無代號時:用名冊中英名稱反查(先剝掉券商字樣,避免「統一投顧」誤中「統一」)
        券商:VRN_Broker_Dict(brokers + brokers_extended 的 aliases / abbr)· 冊沒收的用內建種子並標黃(= 該補進冊)
   ⑤ 產出:VIA_Reports/vrn/prep/PREP_MATRIX_latest.html(跳出)· PREP_MANIFEST_latest.json(每檔要讀的 PDF 路徑,下一步 extract 用)· PREP_latest.csv · VRN_Prep_Ledger.jsonl
        燈:紅 = 讀不了也轉不了 · 黃 = 待 OCR / 欄位缺或來源弱 · 綠 = 可讀且欄位齊 · 灰 = 不支援的檔型
其餘動詞照前版鏈。
"""
from __future__ import annotations

# ===== [VIA:ACCEL-BRIDGE:v0100] SuperAccel 加速器橋(批102 全樹導入令;graceful 零行為變更) =====
try:
    import sys as _sa_sys
    from pathlib import Path as _sa_Path
    _sa_p = _sa_Path(__file__).resolve()
    while _sa_p.parent != _sa_p:
        if (_sa_p / "supportive modules" / "VIA_SuperAccel_Module.py").exists():
            _sa_sys.path.insert(0, str(_sa_p / "supportive modules"))
            break
        _sa_p = _sa_p.parent
    import VIA_SuperAccel_Module as VIA_ACCEL  # noqa: F401
except ImportError:
    VIA_ACCEL = None
# ===== [VIA:ACCEL-BRIDGE:END] =====

import csv
import datetime
import hashlib
import html
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

HERE = Path(__file__).resolve().parent
_STEM = "VRN_SystemManager"
TAG = "v0153"


def _vnum_v0153(path) -> int:
    m = re.search(r"_v(\d{4})$", Path(path).stem)
    return int(m.group(1)) if m else -1


def _load_v0153(path: Path, name: str):
    if name not in sys.modules:
        spec = importlib.util.spec_from_file_location(name, path)
        mod = importlib.util.module_from_spec(spec)
        sys.modules[name] = mod
        spec.loader.exec_module(mod)
    return sys.modules[name]


PRIOR_PATH = max((p for p in HERE.glob(_STEM + "_v*.py") if 0 <= _vnum_v0153(p) < _vnum_v0153(__file__)), key=_vnum_v0153)
PRIOR = _load_v0153(PRIOR_PATH, _STEM + "_prior_for_" + Path(__file__).stem)


def __getattr__(name):
    return getattr(PRIOR, name)


def _resolve(name):
    import types
    mod, seen = PRIOR, set()
    while isinstance(mod, types.ModuleType) and id(mod) not in seen:
        seen.add(id(mod))
        if name in vars(mod):
            return vars(mod)[name]
        mod = vars(mod).get("PRIOR")
    return None


def _home() -> Path:
    return Path(os.environ.get("VIA_VRN_SSOT_HOME") or HERE)


def _rep() -> Path:
    return Path(os.environ.get("VIA_VRN_HEALTH_OUT") or (_home().parents[1] / "VIA_Reports" / "vrn"))


def _vdf_home() -> Path:
    return Path(os.environ.get("VIA_VDF_HOME") or (_home().parent / "VDF"))


_PDF = {".pdf"}
_WORD = {".docx", ".doc", ".rtf", ".odt"}
_TXT = {".txt", ".md", ".csv"}
_IMG = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp", ".gif"}
DPI = 325


def _has(mod: str) -> bool:
    try:
        return importlib.util.find_spec(mod) is not None
    except (ImportError, ValueError):
        return False


def _soffice() -> str | None:
    for c in ("soffice", "libreoffice"):
        p = shutil.which(c)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if Path(p).exists():
            return p
    return None


def tools() -> dict:
    t = {k: _has(k) for k in ("pdfplumber", "fitz", "pypdfium2", "PIL", "docx", "reportlab", "pikepdf", "win32com", "docx2pdf")}
    t["soffice"] = bool(_soffice())
    t["pdftoppm"] = bool(shutil.which("pdftoppm"))
    t["ps_word_com"] = os.name == "nt" and bool(shutil.which("powershell") or shutil.which("pwsh"))
    return t


# ───────── 名冊(VRN 自家 + VDF 唯讀)與券商冊 ─────────
_CODE_KEYS = ("code", "ticker", "stock_id", "stockid", "symbol", "id", "代號", "證券代號", "股票代號", "公司代號", "sid")
_NAME_KEYS = ("name", "name_zh", "zh", "short_name", "shortname_zh", "公司簡稱", "證券名稱", "名稱", "cname", "股票名稱", "公司名稱")
_EN_KEYS = ("name_en", "en", "english", "longname", "shortname", "英文簡稱", "english_name", "en_name")
_MKT_KEYS = ("market", "exchange", "市場別", "board", "type", "suffix", "yf", "yfinance", "yf_symbol", "market_type")


def _norm_mkt(v) -> str:
    s = str(v or "").strip().upper()
    if not s:
        return ""
    if s.endswith(".TWO") or any(k in s for k in ("上櫃", "TPEX", "OTC", "GRETAI", "TWO")):
        return "TPEx"
    if "興櫃" in s or "ESB" in s or "EMERGING" in s:
        return "ESB"
    if s.endswith(".TW") or any(k in s for k in ("上市", "TWSE", "TSE", "LISTED")):
        return "TWSE"
    return ""


def _add_rec(roster: dict, code, name="", en="", mkt="", src=""):
    code = str(code or "").strip().upper()
    m = re.fullmatch(r"(\d{4,6}[A-Z]?)(?:\.TWO?)?", code)
    if not m:
        return
    c = m.group(1)
    mk = _norm_mkt(mkt) or _norm_mkt(code if "." in code else "")
    e = roster.setdefault(c, {"name": "", "en": "", "market": "", "src": src})
    if name and not e["name"]:
        e["name"] = str(name).strip()
    if en and not e["en"]:
        e["en"] = str(en).strip()
    if mk and not e["market"]:
        e["market"] = mk
    if src and src not in e["src"]:
        e["src"] = (e["src"] + "+" + src).strip("+")


def _pick(d: dict, keys) -> str:
    low = {str(k).lower(): v for k, v in d.items()}
    for k in keys:
        v = low.get(k.lower())
        if isinstance(v, (str, int)) and str(v).strip():
            return str(v)
    return ""


def _walk_json(obj, roster, src, depth=0):
    if depth > 6:
        return
    if isinstance(obj, dict):
        code = _pick(obj, _CODE_KEYS)
        if code:
            _add_rec(roster, code, _pick(obj, _NAME_KEYS), _pick(obj, _EN_KEYS), _pick(obj, _MKT_KEYS), src)
        for k, v in obj.items():
            if re.fullmatch(r"\d{4,6}[A-Z]?(\.TWO?)?", str(k)):
                if isinstance(v, str):
                    _add_rec(roster, k, v, "", "", src)
                elif isinstance(v, dict):
                    _add_rec(roster, k, _pick(v, _NAME_KEYS), _pick(v, _EN_KEYS), _pick(v, _MKT_KEYS) or k, src)
            if isinstance(v, (dict, list)):
                _walk_json(v, roster, src, depth + 1)
    elif isinstance(obj, list):
        for v in obj[:200000]:
            if isinstance(v, (dict, list)):
                _walk_json(v, roster, src, depth + 1)


def _load_table(p: Path, roster: dict, src: str) -> int:
    before = len(roster)
    try:
        if p.suffix.lower() == ".json":
            _walk_json(json.loads(p.read_text(encoding="utf-8-sig", errors="replace")), roster, src)
        elif p.suffix.lower() == ".csv":
            with p.open(encoding="utf-8-sig", errors="replace") as fh:
                for i, row in enumerate(csv.DictReader(fh)):
                    if i > 200000:
                        break
                    code = _pick(row, _CODE_KEYS)
                    if code:
                        _add_rec(roster, code, _pick(row, _NAME_KEYS), _pick(row, _EN_KEYS), _pick(row, _MKT_KEYS), src)
    except (OSError, ValueError, csv.Error):
        pass
    return len(roster) - before


def load_roster(use_vdf: bool = True) -> tuple:
    roster, used = {}, []
    vrn = _home()
    cands = sorted(set(list(vrn.glob("VRN_TWRoster_Offline_v*.json")) + list(vrn.glob("VRN_TWTicker_RIE_v*.json")) + list((vrn / "knowledge").glob("*Roster*.json")) + list((vrn / "SSOT").glob("*Roster*.json")) + list((vrn / "knowledge").glob("*Ticker_Master*.json"))))
    for p in cands:
        n = _load_table(p, roster, "VRN")
        used.append({"file": p.name, "sys": "VRN", "added": n})
    if use_vdf and _vdf_home().is_dir():
        rx = re.compile(r"(?i)(stock_?list|twstock|tw_?universe|universe|listing|roster|stock_?master|company_basic|instrument_?identity|ticker_?master)")
        vc = []
        for p in _vdf_home().rglob("*"):
            if p.is_file() and p.suffix.lower() in (".json", ".csv") and rx.search(p.name) and "_superseded" not in p.parts and p.stat().st_size < 50 * 1024 * 1024:
                vc.append(p)
            if len(vc) >= 40:
                break
        for p in sorted(vc, key=lambda q: q.stat().st_mtime, reverse=True)[:25]:
            n = _load_table(p, roster, "VDF")
            if n:
                used.append({"file": p.name, "sys": "VDF(唯讀)", "added": n})
    return roster, used


_SEED_BROKERS = {"Goldman Sachs": ["GS", "Goldman", "高盛"], "Morgan Stanley": ["MS", "Morgan Stanley", "摩根士丹利", "大摩"], "UBS": ["UBS", "瑞銀"], "Citi": ["Citi", "Citigroup", "花旗"], "Daiwa": ["Daiwa", "大和"],
                 "J.P. Morgan": ["JP", "JPM", "J.P. Morgan", "摩根大通", "小摩"], "CLSA": ["CLSA", "CLST", "里昂"], "Macquarie": ["MQ", "Macquarie", "麥格理"], "GF Securities": ["GF", "GFHK", "廣發"],
                 "兆豐": ["兆豐", "Mega"], "華南": ["華南投顧", "華南永昌", "華南"], "群益": ["群益"], "統一": ["統一投顧", "統一證券"], "國泰": ["國泰證期", "國泰"], "中信": ["CTBC", "中信"], "凱基": ["凱基", "KGI"], "台新": ["台新"]}


def load_brokers() -> tuple:
    """book[canon] = {"aliases": {alias: "冊" | "內建種子"}};冊沒收的縮寫(JP · CLST · MQ …)由種子補,判到時標種子 → 該補進 Broker_Dict。"""
    book, used = {}, []
    for p in sorted(set(list((_home() / "knowledge").glob("VRN_Broker_Dict_v*.json")) + list(_home().glob("VRN_Broker_Dict_v*.json")) + list((_home() / "registry").glob("VRN_BROKER_LIST_v*.json")))):
        try:
            d = json.loads(p.read_text(encoding="utf-8-sig"))
        except (OSError, ValueError):
            continue
        n0 = len(book)
        secs = []
        if isinstance(d, dict):
            for k, v in d.items():
                if isinstance(v, dict) and v and all(isinstance(x, dict) for x in v.values()):
                    secs.append(v)
                elif isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
                    secs.append({str(x.get("name") or x.get("canonical") or x.get("broker") or x.get("zh") or i): x for i, x in enumerate(v)})
        for sec in secs:
            for canon, ent in sec.items():
                al = [str(a) for a in (ent.get("aliases") or ent.get("alias") or []) if str(a).strip()]
                for k in ("abbr", "en", "zh", "short"):
                    if isinstance(ent.get(k), str) and ent[k].strip():
                        al.append(ent[k])
                al.append(str(canon))
                e = book.setdefault(str(canon), {"aliases": {}})
                for a in al:
                    if len(a) >= 2:
                        e["aliases"].setdefault(a, "冊")
        used.append({"file": p.name, "brokers": len(book) - n0})
    owner = {a.lower(): c for c, e in book.items() for a in e["aliases"]}
    for canon, al in _SEED_BROKERS.items():
        tgt = next((owner[a.lower()] for a in al if a.lower() in owner), None)
        e = book.setdefault(tgt or canon, {"aliases": {}})
        for a in al:
            if a.lower() not in owner:
                e["aliases"][a] = "內建種子"
                owner[a.lower()] = tgt or canon
    return book, used


def _alias_rx(a: str) -> str:
    if re.fullmatch(r"[A-Za-z][A-Za-z.& ]*", a):
        return r"(?<![A-Za-z])" + re.escape(a) + r"(?![A-Za-z])"
    return re.escape(a)


def match_broker(name: str, book: dict) -> tuple:
    best = None
    for canon, e in book.items():
        for a, src in e["aliases"].items():
            if re.search(_alias_rx(a), name, flags=re.I if a.isascii() else 0):
                if best is None or len(a) > len(best[1]):
                    best = (canon, a, src)
    return best or ("", "", "")


# ───────── 檔名判讀 ─────────
_TK_RX = re.compile(r"(?:^|[\s\-_(（,【】])([1-9]\d{3}|00\d{2,4})(?=TT|[\s\-_)）,.【】]|$)")
_BBG_RX = re.compile(r"(?<!\d)(\d{4,6})\s?TT(?![A-Za-z])")
_YF_RX = re.compile(r"(?<!\d)(\d{4,6})\.(TWO|TW)(?![A-Za-z])", re.I)
_FOREIGN_RX = re.compile(r"(?<!\d)(\d{6}\.(?:SZ|SH)|\d{4,5}\.(?:HK|T))(?![A-Za-z])", re.I)
_NAME_BEFORE = re.compile(r"([\u4e00-\u9fff]{2,8}(?:-KY)?)\s*[\(（]\s*\d{4}")
_NAME_AFTER = re.compile(r"(?<!\d)\d{4}(?:\s?TT)?[ _\-]?([\u4e00-\u9fff]{2,8}(?:-KY)?)")
_NAME_AFTER_EN = re.compile(r"(?<!\d)\d{4}_([A-Z][A-Za-z]{1,14}(?:-KY)?)")
_TK_HEAD = re.compile(r"^([1-9]\d{3})(?=[\u4e00-\u9fff])")


def _cut_broker(nm: str, book: dict) -> str:
    """名稱裡夾到券商字樣(健策凱基投顧)→ 在非開頭位置切掉;開頭就是券商字(中信金 · 國泰金)不切。"""
    cut = len(nm)
    for e in book.values():
        for a in e["aliases"]:
            if len(a) >= 2 and not a.isascii():
                i = nm.find(a)
                if 0 < i < cut:
                    cut = i
    return nm[:cut]


def _valid_date(y, m, d):
    try:
        dt = datetime.date(int(y), int(m), int(d))
    except ValueError:
        return None
    return dt if datetime.date(2010, 1, 1) <= dt <= datetime.date.today() + datetime.timedelta(days=370) else None


def parse_date(name: str, mtime: float) -> tuple:
    for m in re.finditer(r"(?<!\d)(20\d{2})[-_.]?(\d{2})[-_.]?(\d{2})(?!\d)", name):
        dt = _valid_date(*m.groups())
        if dt:
            return dt.isoformat(), "檔名 YYYYMMDD"
    for m in re.finditer(r"(?<!\d)(1[01]\d)(\d{2})(\d{2})(?!\d)", name):
        dt = _valid_date(int(m.group(1)) + 1911, m.group(2), m.group(3))
        if dt:
            return dt.isoformat(), "檔名 民國 1YYMMDD"
    for m in re.finditer(r"(?<!\d)(\d{2})(\d{2})(\d{2})(?!\d)", name):
        if 15 <= int(m.group(1)) <= 35:
            dt = _valid_date(2000 + int(m.group(1)), m.group(2), m.group(3))
            if dt:
                return dt.isoformat(), "檔名 YYMMDD"
    fy = datetime.date.fromtimestamp(mtime).year
    for m in re.finditer(r"(?:CTBC|中信)(\d{2})(\d{2})(?!\d)", name):
        dt = _valid_date(fy, m.group(1), m.group(2))
        if dt:
            return dt.isoformat(), "檔名 MMDD + 檔案年"
    return datetime.date.fromtimestamp(mtime).isoformat(), "檔案時間(非報告日)"


def parse_name_fields(name: str, roster: dict, book: dict) -> dict:
    stem = Path(name).stem
    canon, alias, bsrc = match_broker(stem, book)
    stripped = re.sub(_alias_rx(alias), " ", stem, flags=re.I if alias.isascii() else 0) if alias else stem
    out = {"broker": canon, "broker_alias": alias, "broker_src": bsrc or "", "code": "", "code_src": "", "name_file": "", "name": "", "name_check": "", "market": "", "yf": "", "bbg": "", "foreign": ""}
    m = _YF_RX.search(stem)
    if m:
        out.update(code=m.group(1), code_src="檔名 yfinance 形", market="TPEx" if m.group(2).upper() == "TWO" else "TWSE")
    if not out["code"]:
        m = _BBG_RX.search(stem)
        if m:
            out.update(code=m.group(1), code_src="檔名 Bloomberg 形")
    if not out["code"]:
        m = _TK_RX.search(stem) or _TK_HEAD.search(stem)
        if m:
            out.update(code=m.group(1), code_src="檔名四碼")
    if out["code"] and roster and out["code"] not in roster and 1990 <= int(out["code"][:4]) <= 2035 and len(out["code"]) == 4:
        out.update(code="", code_src="")          # 像年份又不在名冊 → 不算代號
    m = _FOREIGN_RX.search(stem)
    if m:
        out["foreign"] = m.group(1).upper()
    for rx in (_NAME_BEFORE, _NAME_AFTER, _NAME_AFTER_EN):
        m = rx.search(stem)
        if m:
            nm = _cut_broker(m.group(1), book)
            if len(nm.replace("-KY", "")) >= 2:
                out["name_file"] = nm
                break
    nonstock = _resolve("_nonstock_type") or (lambda n: "")
    if not out["code"] and not nonstock(name):
        hit = None
        toks = [t for t in re.split(r"[\s_\-()（）,.【】｜|]+", stripped) if t]
        for nm, c in sorted(((r["name"], c) for c, r in roster.items() if r.get("name") and len(r["name"]) >= 2), key=lambda x: -len(x[0])):
            if nm in stripped:
                hit = (c, "名冊反查(中文名 %s)" % nm)
                break
        if not hit:
            en_first = {}
            for c, r in roster.items():
                w = re.split(r"[\s,.\-]+", r.get("en", "").strip())
                if w and len(w[0]) >= 4:
                    en_first.setdefault(w[0].lower(), c)
            for t in toks:
                if t.lower() in en_first and len(t) >= 4:
                    hit = (en_first[t.lower()], "名冊反查(英文名 %s)" % t)
                    break
        if hit:
            out.update(code=hit[0], code_src=hit[1])
    if out["code"]:
        r = roster.get(out["code"]) or {}
        out["name"] = r.get("name") or out["name_file"]
        if not out["market"]:
            out["market"] = r.get("market", "")
        if out["name_file"] and r.get("name"):
            nf, nr = out["name_file"].replace("-KY", ""), r["name"].replace("-KY", "")
            out["name_check"] = "✓" if (nf in nr or nr in nf) else "≠ 名冊 %s" % r["name"]
        elif r.get("name"):
            out["name_check"] = "名冊"
        elif out["name_file"]:
            out["name_check"] = "只有檔名"
        out["yf"] = out["code"] + {"TWSE": ".TW", "TPEx": ".TWO"}.get(out["market"], "")
        out["bbg"] = out["code"] + " TT"
    return out


# ───────── 掃描 · 轉換 · 讀取 ─────────
def _sha8(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as fh:
        for ch in iter(lambda: fh.read(1 << 20), b""):
            h.update(ch)
    return h.hexdigest()[:8]


def _pdf_spec(p: Path) -> dict:
    sp = {"pages": 0, "page_mm": "", "text3": 0, "encrypted": False, "err": ""}
    try:
        import pdfplumber
        with pdfplumber.open(str(p)) as pdf:
            sp["pages"] = len(pdf.pages)
            if pdf.pages:
                w, h = pdf.pages[0].width, pdf.pages[0].height
                sp["page_mm"] = "%dx%d" % (round(w / 72 * 25.4), round(h / 72 * 25.4))
            for pg in pdf.pages[:3]:
                sp["text3"] += len((pg.extract_text() or "").strip())
    except Exception as exc:  # noqa: BLE001
        sp["err"] = "%s:%s" % (type(exc).__name__, str(exc)[:80])
        sp["encrypted"] = "encrypt" in str(exc).lower() or "password" in str(exc).lower()
    return sp


def _img_spec(p: Path) -> dict:
    try:
        from PIL import Image
        with Image.open(p) as im:
            return {"px": "%dx%d" % im.size, "mode": im.mode, "dpi_in": "%s" % (round(im.info["dpi"][0]) if im.info.get("dpi") else "—"), "frames": getattr(im, "n_frames", 1)}
    except Exception as exc:  # noqa: BLE001
        return {"err": "%s:%s" % (type(exc).__name__, str(exc)[:60])}


def img_to_pdf(src: Path, dst: Path) -> tuple:
    from PIL import Image
    with Image.open(src) as im:
        frames = []
        for i in range(getattr(im, "n_frames", 1)):
            im.seek(i)
            f = im.copy()
            if f.mode in ("RGBA", "LA", "P"):
                f = f.convert("RGBA")
                bg = Image.new("RGB", f.size, (255, 255, 255))
                bg.paste(f, mask=f.split()[-1])
                f = bg
            elif f.mode != "RGB" and f.mode != "L":
                f = f.convert("RGB")
            frames.append(f)
    frames[0].save(str(dst), "PDF", resolution=float(DPI), save_all=len(frames) > 1, append_images=frames[1:])
    return "PIL(原像素 · %d dpi)" % DPI, ""


def _read_text_any(p: Path) -> str:
    raw = p.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp950", "big5hkscs", "gb18030", "latin-1"):
        try:
            t = raw.decode(enc)
            if enc == "utf-16" and not raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
                continue
            return t
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def text_to_pdf(text: str, dst: Path) -> tuple:
    if _has("fitz"):
        import fitz  # noqa: WPS433
        doc = fitz.open()
        rect = fitz.paper_rect("a4")
        lines = text.splitlines() or [""]
        chunk = 60
        for i in range(0, len(lines), chunk):
            pg = doc.new_page(width=rect.width, height=rect.height)
            pg.insert_textbox(fitz.Rect(50, 50, rect.width - 50, rect.height - 50), "\n".join(lines[i:i + chunk]), fontsize=9.5, fontname="china-t")
        doc.save(str(dst))
        return "PyMuPDF(向量 · 繁中)", ""
    if _has("reportlab"):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfgen import canvas
        pdfmetrics.registerFont(UnicodeCIDFont("MSung-Light"))
        c = canvas.Canvas(str(dst), pagesize=A4)
        w, h = A4
        y, size, width_chars = h - 50, 9.5, 52
        c.setFont("MSung-Light", size)
        for para in text.splitlines() or [""]:
            segs = [para[i:i + width_chars] for i in range(0, max(len(para), 1), width_chars)] or [""]
            for s in segs:
                if y < 50:
                    c.showPage()
                    c.setFont("MSung-Light", size)
                    y = h - 50
                c.drawString(50, y, s)
                y -= size * 1.45
        c.save()
        return "reportlab(向量 · 繁中 CID)", ""
    return "", "缺 PyMuPDF / reportlab(pip install pymupdf)"


def word_to_pdf(src: Path, dst: Path, tmp: Path) -> tuple:
    errs = []
    if os.name == "nt" and _has("win32com"):
        try:
            import win32com.client  # noqa: WPS433
            w = win32com.client.DispatchEx("Word.Application")
            w.Visible, w.DisplayAlerts = False, 0
            try:
                d = w.Documents.Open(str(src), False, True)
                d.ExportAsFixedFormat(str(dst), 17)
                d.Close(0)
            finally:
                w.Quit()
            if dst.exists():
                return "Word COM(win32com)", ""
        except Exception as exc:  # noqa: BLE001
            errs.append("win32com:%s" % str(exc)[:50])
    if os.name == "nt":
        ps = shutil.which("powershell") or shutil.which("pwsh")
        if ps:
            cmd = "$ErrorActionPreference='Stop'; $w=New-Object -ComObject Word.Application; $w.Visible=$false; $w.DisplayAlerts=0; try { $d=$w.Documents.Open('%s',$false,$true); $d.ExportAsFixedFormat('%s',17); $d.Close(0) } finally { $w.Quit() }" % (str(src).replace("'", "''"), str(dst).replace("'", "''"))
            try:
                subprocess.run([ps, "-NoProfile", "-NonInteractive", "-Command", cmd], capture_output=True, timeout=180)
                if dst.exists():
                    return "Word COM(PowerShell)", ""
                errs.append("ps-com:無輸出(可能沒裝 Word)")
            except (subprocess.TimeoutExpired, OSError) as exc:
                errs.append("ps-com:%s" % type(exc).__name__)
    so = _soffice()
    if so:
        try:
            prof = tmp / "lo_profile"
            subprocess.run([so, "-env:UserInstallation=" + prof.as_uri(), "--headless", "--convert-to", "pdf", "--outdir", str(dst.parent), str(src)], capture_output=True, timeout=240)
            made = dst.parent / (src.stem + ".pdf")
            if made.exists():
                if made != dst:
                    made.replace(dst)
                return "LibreOffice", ""
            errs.append("soffice:無輸出")
        except (subprocess.TimeoutExpired, OSError) as exc:
            errs.append("soffice:%s" % type(exc).__name__)
    if _has("docx2pdf") and os.name == "nt":
        try:
            from docx2pdf import convert  # noqa: WPS433
            convert(str(src), str(dst))
            if dst.exists():
                return "docx2pdf", ""
        except Exception as exc:  # noqa: BLE001
            errs.append("docx2pdf:%s" % str(exc)[:40])
    if src.suffix.lower() == ".docx" and _has("docx"):
        try:
            import docx  # noqa: WPS433
            d = docx.Document(str(src))
            parts = [pp.text for pp in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            how, err = text_to_pdf("\n".join(parts), dst)
            if how:
                return "python-docx 文字版(版面遺失)+" + how, ""
            errs.append(err)
        except Exception as exc:  # noqa: BLE001
            errs.append("python-docx:%s" % str(exc)[:40])
    return "", ";".join(errs) or "沒有可用的 Word 轉換工具"


def raster325(src: Path, dst: Path, max_pages: int) -> tuple:
    if _has("fitz"):
        try:
            import fitz  # noqa: WPS433
            s = fitz.open(str(src))
            out = fitz.open()
            for i, pg in enumerate(s):
                if i >= max_pages:
                    break
                pix = pg.get_pixmap(dpi=DPI)
                img = fitz.open("png", pix.tobytes("png"))
                out.insert_pdf(fitz.open("pdf", img.convert_to_pdf()))
            out.save(str(dst))
            return "PyMuPDF %d dpi" % DPI, ""
        except Exception as exc:  # noqa: BLE001
            err1 = "fitz:%s" % str(exc)[:40]
    else:
        err1 = "無 fitz"
    if _has("pypdfium2") and _has("PIL"):
        try:
            import pypdfium2 as pdfium  # noqa: WPS433
            doc = pdfium.PdfDocument(str(src))
            imgs = []
            for i in range(min(len(doc), max_pages)):
                imgs.append(doc[i].render(scale=DPI / 72).to_pil().convert("RGB"))
            if imgs:
                imgs[0].save(str(dst), "PDF", resolution=float(DPI), save_all=len(imgs) > 1, append_images=imgs[1:])
                return "pypdfium2 %d dpi" % DPI, ""
        except Exception as exc:  # noqa: BLE001
            err1 += ";pdfium:%s" % str(exc)[:40]
    if shutil.which("pdftoppm") and _has("PIL"):
        try:
            from PIL import Image
            pref = dst.with_suffix("")
            subprocess.run(["pdftoppm", "-r", str(DPI), "-l", str(max_pages), "-png", str(src), str(pref)], capture_output=True, timeout=600)
            pngs = sorted(dst.parent.glob(pref.name + "-*.png"))
            if pngs:
                ims = [Image.open(q).convert("RGB") for q in pngs]
                ims[0].save(str(dst), "PDF", resolution=float(DPI), save_all=len(ims) > 1, append_images=ims[1:])
                for q in pngs:
                    q.unlink()
                return "pdftoppm %d dpi" % DPI, ""
        except Exception as exc:  # noqa: BLE001
            err1 += ";pdftoppm:%s" % str(exc)[:40]
    return "", err1


def _repair_pdf(src: Path, dst: Path) -> bool:
    if not _has("pikepdf"):
        return False
    try:
        import pikepdf  # noqa: WPS433
        with pikepdf.open(str(src), allow_overwriting_input=False) as pdf:
            pdf.save(str(dst))
        return dst.exists()
    except Exception:  # noqa: BLE001
        return False


def prep(d: Path, use_vdf: bool = True, max_pages: int = 60, workers: int = 4, raster_all: bool = False) -> dict:
    now = datetime.datetime.now()
    stamp = now.strftime("%Y%m%dT%H%M%S")
    rep = _rep() / "prep"
    conv = rep / "converted"
    tmp = Path(os.environ.get("VIA_SPILL_DIR") or (Path(tempfile.gettempdir()) / "VIA_progress")) / ("vrn_prep_" + stamp)
    for q in (rep, conv, tmp):
        q.mkdir(parents=True, exist_ok=True)
    tl = tools()
    roster, rsrc = load_roster(use_vdf)
    book, bsrc = load_brokers()
    files = sorted(p for p in d.rglob("*") if p.is_file() and not p.name.startswith("~$")) if d.is_dir() else []
    rows = []
    nonstock = _resolve("_nonstock_type") or (lambda n: "")
    for p in files:
        ext = p.suffix.lower()
        st = p.stat()
        r = {"file": p.name, "path": str(p), "ext": ext, "kb": round(st.st_size / 1024, 1), "sha8": "", "type": "", "spec": "", "convert": "", "convert_err": "", "read_pdf": "", "read": "", "lamp": "GRAY", "notes": []}
        if ext in _PDF:
            r["type"] = "PDF"
        elif ext in _WORD:
            r["type"] = "Word"
        elif ext in _TXT:
            r["type"] = "TXT"
        elif ext in _IMG:
            r["type"] = "影像"
        else:
            r["type"] = "不支援"
            r["read"] = "略過(檔型不在範圍)"
            rows.append(r)
            continue
        r["sha8"] = _sha8(p)
        dst = conv / (p.stem + ".pdf")
        if r["type"] == "影像":
            sp = _img_spec(p)
            r["spec"] = "%s · %s · 原 dpi %s%s" % (sp.get("px", "?"), sp.get("mode", "?"), sp.get("dpi_in", "?"), (" · %d 幀" % sp["frames"]) if sp.get("frames", 1) > 1 else "") if "err" not in sp else sp["err"]
            if _has("PIL"):
                try:
                    r["convert"], r["convert_err"] = img_to_pdf(p, dst)
                except Exception as exc:  # noqa: BLE001
                    r["convert_err"] = "%s:%s" % (type(exc).__name__, str(exc)[:60])
            else:
                r["convert_err"] = "缺 Pillow"
        elif r["type"] == "TXT":
            t = _read_text_any(p)
            r["spec"] = "%d 字 · %d 行" % (len(t), t.count("\n") + 1)
            r["convert"], r["convert_err"] = text_to_pdf(t, dst)
        elif r["type"] == "Word":
            r["spec"] = "%s" % ext
            r["convert"], r["convert_err"] = word_to_pdf(p, dst, tmp)
        if r["type"] != "PDF":
            if dst.exists() and not r["convert_err"]:
                if raster_all and r["type"] != "影像":
                    rd = conv / (p.stem + "_r325.pdf")
                    how, err = raster325(dst, rd, max_pages)
                    if how:
                        r["convert"] += " → " + how
                        dst = rd
                r["read_pdf"] = str(dst)
            else:
                r["read"] = "轉換失敗"
                r["lamp"] = "RED"
        else:
            r["read_pdf"] = str(p)
        r["_mtime"] = st.st_mtime
        rows.append(r)

    def check(r):
        if not r.get("read_pdf") or r["type"] == "不支援":
            return r
        pdfp = Path(r["read_pdf"])
        sp = _pdf_spec(pdfp)
        if sp["err"]:
            fx = tmp / ("repair_" + pdfp.name)
            if _repair_pdf(pdfp, fx):
                sp2 = _pdf_spec(fx)
                if not sp2["err"]:
                    r["notes"].append("pikepdf 修復")
                    pdfp, sp = fx, sp2
                    r["read_pdf"] = str(fx)
        spec_pdf = "%s 頁 · %s mm · 前3頁 %d 字%s" % (sp["pages"], sp["page_mm"] or "?", sp["text3"], " · 加密" if sp["encrypted"] else "")
        r["spec"] = (r["spec"] + " → " if r["type"] != "PDF" else "") + spec_pdf
        thin = sp["text3"] < 40 if r["type"] == "PDF" else sp["text3"] == 0      # 轉出來的向量 PDF 有字就算可讀
        if sp["err"] or thin:
            if r["type"] == "影像" and not sp["err"]:
                r["read"] = "影像 PDF %d dpi(待 OCR)" % DPI
            else:
                rd = tmp / (pdfp.stem + "_r325.pdf")
                how, err = raster325(pdfp, rd, max_pages)
                if how:
                    r["read"] = ("讀取失敗 → " if sp["err"] else "掃描件 → ") + how + " 到 TEMP(待 OCR)"
                    r["read_pdf"] = str(rd)
                    r["notes"].append(sp["err"] or "前 3 頁文字 %d" % sp["text3"])
                else:
                    r["read"] = "讀不了且點陣失敗:" + (sp["err"] or "") + ";" + err
                    r["lamp"] = "RED"
        else:
            r["read"] = "可讀(文字層)"
        return r

    with ThreadPoolExecutor(max_workers=max(1, workers)) as ex:
        rows = list(ex.map(check, rows))
    for r in rows:
        if r["type"] == "不支援":
            continue
        dt, dsrc = parse_date(r["file"], r.pop("_mtime", 0) or 0)
        f = parse_name_fields(r["file"], roster, book)
        kind = nonstock(r["file"]) if not f["code"] else ""
        r.update(date=dt, date_src=dsrc, kind=kind, **f)
        yel = []
        if r["lamp"] == "RED":
            continue
        if "待 OCR" in r["read"]:
            yel.append("待 OCR")
        if dsrc.startswith("檔案時間"):
            yel.append("報告日用檔案時間")
        if r["broker_src"] != "冊":            # 種子 = 冊沒收這個寫法,該補進 Broker_Dict
            yel.append("券商冊未收" if r["broker"] else "券商未判出")
        if not kind:
            if not r["code"]:
                yel.append("個股代號未判出")
            else:
                if not r["market"]:
                    yel.append("市場別未知(.TW/.TWO 待定)")
                if not r["name"]:
                    yel.append("名稱未知")
                if r["name_check"].startswith("≠"):
                    yel.append("檔名名稱與名冊不符")
        if "python-docx 文字版" in r["convert"]:
            yel.append("Word 文字版(版面遺失)")
        r["lamp"] = "YELLOW" if yel else "GREEN"
        r["notes"] = r["notes"] + yel
    ok_types = [r for r in rows if r["type"] != "不支援"]
    cov = lambda k: sum(1 for r in ok_types if r.get(k)) if ok_types else 0  # noqa: E731
    stock = [r for r in ok_types if not r.get("kind")]
    summary = {"ts": now.isoformat(timespec="seconds"), "dir": str(d), "files": len(rows), "types": dict(Counter(r["type"] for r in rows)), "lamps": dict(Counter(r["lamp"] for r in rows)),
               "converted": sum(1 for r in rows if r["convert"] and not r["convert_err"]), "convert_fail": sum(1 for r in rows if r["convert_err"]), "readable": sum(1 for r in rows if r["read"] == "可讀(文字層)"),
               "raster": sum(1 for r in rows if "dpi" in r.get("read", "") and "TEMP" in r.get("read", "")), "ocr_pending": sum(1 for r in rows if "待 OCR" in r.get("read", "")),
               "date_from_name": sum(1 for r in ok_types if str(r.get("date_src", "")).startswith("檔名")), "broker_in_book": sum(1 for r in ok_types if r.get("broker_src") == "冊"), "broker_seed": sum(1 for r in ok_types if r.get("broker_src") == "內建種子"),
               "stock_reports": len(stock), "with_code": sum(1 for r in stock if r.get("code")), "with_yf": sum(1 for r in stock if str(r.get("yf", "")).endswith((".TW", ".TWO"))), "with_name": sum(1 for r in stock if r.get("name")),
               "nonstock": len(ok_types) - len(stock), "roster": len(roster), "roster_src": rsrc, "broker_book": len(book), "broker_src": bsrc, "tools": tl, "temp": str(tmp), "converted_dir": str(conv), "dpi": DPI}
    written = [rep, conv, tmp]
    allowed = [_home().resolve(), _rep().resolve(), Path(tempfile.gettempdir()).resolve()] + ([Path(os.environ["VIA_SPILL_DIR"]).resolve()] if os.environ.get("VIA_SPILL_DIR") else [])
    summary["writes_ok"] = all(any(str(w.resolve()).startswith(str(a)) for a in allowed) for w in written)
    man = {"schema": "VIA.VRN.Prep.v1", "summary": summary, "rows": rows}
    (rep / "PREP_MANIFEST_latest.json").write_text(json.dumps(man, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    cols = ["lamp", "file", "type", "kb", "spec", "convert", "read", "date", "date_src", "code", "code_src", "yf", "bbg", "name", "name_check", "broker", "broker_src", "kind", "read_pdf", "notes"]
    with (rep / "PREP_latest.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(dict(r, notes=";".join(r.get("notes", []))))
    with (rep / "VRN_Prep_Ledger.jsonl").open("a", encoding="utf-8") as fh:
        fh.write(json.dumps({k: summary[k] for k in ("ts", "dir", "files", "types", "lamps", "converted", "convert_fail", "readable", "raster", "stock_reports", "with_code", "with_yf", "date_from_name")}, ensure_ascii=False) + "\n")
    page = rep / "PREP_MATRIX_latest.html"
    page.write_text(_prep_html(man), encoding="utf-8")
    summary["html"] = str(page)
    lamp = "RED" if summary["lamps"].get("RED") else ("YELLOW" if summary["lamps"].get("YELLOW") else "GREEN")
    return {"verb": "prep", "summary": summary, "rows": rows, "lamp": lamp if files else "GRAY", "html": str(page)}


def _prep_html(man: dict) -> str:
    s, rows = man["summary"], man["rows"]
    L = {"GREEN": "var(--lamp-green,#16a34a)", "YELLOW": "var(--lamp-yellow,#eab308)", "RED": "var(--lamp-red,#dc2626)", "GRAY": "var(--lamp-gray,#9ca3af)"}
    order = {"RED": 0, "YELLOW": 1, "GREEN": 2, "GRAY": 3}
    e = html.escape

    def lp(l):
        return "<i class='lp %s' style='background:%s'></i>" % (l, L[l])

    def chip(k, v):
        return "<span class='chip %s'>%s</span>" % ("on" if v else "off", e(k))
    trs = []
    for r in sorted(rows, key=lambda x: (order.get(x["lamp"], 9), x["file"])):
        code = r.get("code", "")
        trs.append("<tr data-l='%s'><td>%s</td><td class='f' title='%s'>%s</td><td>%s</td><td class='n'>%s</td><td class='d'>%s</td><td class='d'>%s</td><td class='d'>%s</td><td>%s<div class='d'>%s</div></td><td>%s<div class='d'>%s</div></td><td>%s</td><td>%s</td><td>%s<div class='d'>%s</div></td><td>%s<div class='d'>%s</div></td><td class='d'>%s</td></tr>" % (
            r["lamp"], lp(r["lamp"]), e(r.get("path", "")), e(r["file"]), e(r["type"]), r["kb"], e(r.get("spec", "")), e(r.get("convert", "") or r.get("convert_err", "") or "—"), e(r.get("read", "")),
            e(r.get("date", "") or "—"), e(r.get("date_src", "")), e(code or ("n/a · " + r["kind"] if r.get("kind") else "—")), e(r.get("code_src", "")), e(r.get("yf", "") or "—"), e(r.get("bbg", "") or "—"),
            e(r.get("name", "") or "—"), e(r.get("name_check", "")), e(r.get("broker", "") or "—"), e(r.get("broker_src", "")), e(";".join(r.get("notes", [])))))
    pct = lambda a, b: ("%.0f%%" % (100.0 * a / b)) if b else "—"  # noqa: E731
    cards = [("檔", s["files"], " · ".join("%s %d" % kv for kv in s["types"].items())), ("轉 PDF", s["converted"], "失敗 %d" % s["convert_fail"]), ("可讀(文字層)", s["readable"], "點陣 325 %d · 待 OCR %d" % (s["raster"], s["ocr_pending"])),
             ("報告日來自檔名", pct(s["date_from_name"], s["files"] - s["types"].get("不支援", 0)), "%d 檔" % s["date_from_name"]), ("個股報告", s["stock_reports"], "非個股 %d" % s["nonstock"]),
             ("代號", pct(s["with_code"], s["stock_reports"]), "yfinance %s · 名稱 %s" % (pct(s["with_yf"], s["stock_reports"]), pct(s["with_name"], s["stock_reports"]))), ("券商(冊)", s["broker_in_book"], "種子 %d" % s["broker_seed"]), ("名冊", s["roster"], " · ".join("%s %s +%d" % (x["sys"], x["file"], x["added"]) for x in s["roster_src"])[:120])]
    card_html = "".join("<div class='card'><div class='k'>%s</div><div class='v'>%s</div><div class='d'>%s</div></div>" % (e(k), e(str(v)), e(str(d))) for k, v, d in cards)
    tool_html = "".join(chip(k, v) for k, v in s["tools"].items())
    lamps = " · ".join("%s%s %d" % (lp(k), k, v) for k, v in sorted(s["lamps"].items(), key=lambda kv: order.get(kv[0], 9)))
    css = ("body{font-family:'Microsoft JhengHei UI','Segoe UI',Arial;font-size:12px;color:#1f2937;background:var(--bg,#fafafa);margin:0;padding:12px}h1{font-size:16px;margin:0 0 4px}.meta{color:#6b7280;font-size:11px;margin-bottom:8px}"
           ".cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:6px;margin:6px 0 10px}.card{background:#fff;border:1px solid #e5e7eb;border-radius:8px;padding:6px 8px}.card .k{color:#6b7280;font-size:11px}.card .v{font-size:17px;font-weight:700}"
           ".chip{display:inline-block;border-radius:10px;padding:1px 7px;margin:0 3px 3px 0;font-size:11px;border:1px solid #e5e7eb}.chip.on{background:#ecfdf5;color:#065f46}.chip.off{background:#f3f4f6;color:#9ca3af;text-decoration:line-through}"
           ".bar button{font-size:11px;margin-right:4px;padding:2px 8px;border:1px solid #d1d5db;background:#fff;border-radius:6px;cursor:pointer}.wrap{overflow-x:auto}table{border-collapse:collapse;width:100%;background:#fff}th,td{border:1px solid #eceff3;padding:2px 5px;text-align:left;vertical-align:top}"
           "th{background:#111827;color:#fff;position:sticky;top:0;font-weight:600}td.n{text-align:right}td.f{max-width:260px;word-break:break-all}.d{color:#6b7280;font-size:10.5px}.lp{display:inline-block;width:11px;height:11px;border-radius:50%;vertical-align:middle;margin-right:3px}"
           ".lp.RED{animation:bl 2.4s ease-in-out infinite}@keyframes bl{0%,100%{opacity:1}50%{opacity:.25}}")
    js = "function f(l){document.querySelectorAll('tbody tr').forEach(function(t){t.style.display=(!l||t.dataset.l===l)?'':'none'})}"
    head = "<!doctype html><html lang='zh-Hant'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>VRN prep · 第一步</title><style>" + css + "</style><script>" + js + "</script></head><body>"
    return head + ("<h1>VRN 第一步 · 輸入掃描 / 轉 PDF(%d dpi)/ 檔名判讀</h1><div class='meta'>%s · 夾 %s · %s · 轉換放 %s · 點陣放 TEMP %s · 寫入範圍 %s</div>"
            "<div class='cards'>%s</div><div class='meta'>工具:%s</div>"
            "<div class='bar'><button onclick=\"f('')\">全部</button><button onclick=\"f('RED')\">只看紅</button><button onclick=\"f('YELLOW')\">只看黃</button><button onclick=\"f('GREEN')\">只看綠</button></div>"
            "<div class='wrap'><table><thead><tr><th>燈</th><th>檔</th><th>類型</th><th>KB</th><th>規格</th><th>轉換</th><th>讀取</th><th>報告日</th><th>代號</th><th>yfinance</th><th>Bloomberg</th><th>名稱</th><th>券商</th><th>備註</th></tr></thead><tbody>%s</tbody></table></div>"
            "<div class='meta' style='margin-top:8px'>燈:紅 = 讀不了也轉不了 · 黃 = 待 OCR / 欄位缺或來源弱(檔案時間 · 券商冊未收 · 市場別未知)· 綠 = 可讀且欄位齊 · 灰 = 不支援檔型。非個股文件(晨會 · 早報 · 產業 · ETF …)代號欄顯示 n/a,不算缺。下一步:extract 讀 PREP_MANIFEST_latest.json 的 read_pdf。</div></body></html>") % (
        DPI, e(s["ts"]), e(s["dir"]), lamps, e(s["converted_dir"]), e(s["temp"]), "VRN 範圍 ✓" if s["writes_ok"] else "⚠ 超出 VRN 範圍", card_html, tool_html, "".join(trs))


def _ui_prep_card(page: Path) -> dict:
    if not page.exists():
        return {"patched": False}
    h = page.read_text(encoding="utf-8", errors="replace")
    cfg = (_resolve("config_get") or (lambda: {}))()
    indir = cfg.get("input_dir") or r"C:\測試樣本報告"
    ln = (_resolve("_launcher_name") or (lambda: "Invoke-VIA-Launch-v0119.ps1"))()
    mf = _rep() / "prep" / "PREP_MANIFEST_latest.json"
    last, mat = "尚未跑", _rep() / "prep" / "PREP_MATRIX_latest.html"
    lamp = "GRAY"
    if mf.exists():
        try:
            sm = json.loads(mf.read_text(encoding="utf-8")).get("summary", {})
            lamps = sm.get("lamps", {})
            lamp = "RED" if lamps.get("RED") else ("YELLOW" if lamps.get("YELLOW") else "GREEN")
            last = "%s · 檔 %s · 轉 PDF %s · 可讀 %s · 待 OCR %s · 個股 %s(代號 %s · yfinance %s)· %s" % (sm.get("ts", ""), sm.get("files"), sm.get("converted"), sm.get("readable"), sm.get("ocr_pending"), sm.get("stock_reports"), sm.get("with_code"), sm.get("with_yf"), " ".join("%s=%s" % kv for kv in sorted(lamps.items())))
        except ValueError:
            pass
    col = {"GREEN": "var(--lamp-green,#16a34a)", "YELLOW": "var(--lamp-yellow,#eab308)", "RED": "var(--lamp-red,#dc2626)", "GRAY": "var(--lamp-gray,#9ca3af)"}[lamp]
    cmd = 'pwsh -ExecutionPolicy Bypass -File "$env:USERPROFILE\\Downloads\\%s" -Sub VRN -Verb prep' % ln
    card = ("<div id='vrn-prep-job' style='border:2px solid var(--border,#d1d5db);border-radius:8px;padding:10px 12px;margin:0 0 10px;background:var(--card,#fff);font-size:12px'>"
            "<div style='font-weight:700;font-size:13px;margin-bottom:4px'>① 第一步 · 輸入掃描 → Word/TXT/影像轉 PDF(%d dpi)→ 讀取失敗點陣 → 檔名判讀(日期 · 代號 · yfinance · Bloomberg · 名稱 · 券商)</div>"
            "<div>輸入夾:<code>%s</code></div><div style='margin:4px 0'>指令:<code style='user-select:all;word-break:break-all'>%s</code></div>"
            "<div><a href='via://VRN/prep' style='display:inline-block;padding:4px 10px;border:1px solid var(--border,#e0e0e0);border-radius:6px;text-decoration:none;margin-right:6px'>執行(滑鼠)</a>"
            "<a href='via://VRN/prep?dir=1' style='display:inline-block;padding:4px 10px;border:1px solid var(--border,#e0e0e0);border-radius:6px;text-decoration:none;margin-right:6px'>選夾執行</a>"
            "<a href='%s' target='_blank' style='display:inline-block;padding:4px 10px;border:1px solid var(--border,#e0e0e0);border-radius:6px;text-decoration:none'>開矩陣結果</a></div>"
            "<div style='margin-top:6px'><i style='display:inline-block;width:10px;height:10px;border-radius:50%%;background:%s;vertical-align:middle'></i> 上次:%s</div></div>") % (
        DPI, html.escape(indir), html.escape(cmd), mat.as_uri() if mat.exists() else "#", col, html.escape(last))
    h = re.sub(r"<div id='vrn-prep-job'.*?</div></div>", "", h, flags=re.S)
    if "<div id='vrn-main-job'" in h:
        h = h.replace("<div id='vrn-main-job'", card + "<div id='vrn-main-job'", 1)
    else:
        m = re.search(r"(<body[^>]*>)", h)
        h = h[:m.end()] + card + h[m.end():] if m else card + h
    page.write_text(h, encoding="utf-8")
    return {"patched": True, "lamp": lamp}


def _print_prep(o: dict) -> None:
    s = o.get("summary") or {}
    if not s:
        print("[計] VRN prep · 夾不在或空 · GRAY")
        return
    print("[計] VRN prep · 夾 %s · 檔 %d(%s)· 燈 %s · %s" % (s["dir"], s["files"], " ".join("%s=%d" % kv for kv in s["types"].items()), " ".join("%s=%d" % kv for kv in sorted(s["lamps"].items())), o["lamp"]))
    print("[計] 轉 PDF %d(失敗 %d)· 可讀 %d · 點陣 %d dpi %d · 待 OCR %d · 報告日來自檔名 %d · 個股 %d(代號 %d · yfinance %d · 名稱 %d)· 非個股 %d · 券商冊 %d(種子 %d)" % (s["converted"], s["convert_fail"], s["readable"], s["dpi"], s["raster"], s["ocr_pending"], s["date_from_name"], s["stock_reports"], s["with_code"], s["with_yf"], s["with_name"], s["nonstock"], s["broker_in_book"], s["broker_seed"]))
    print("[計] 名冊 %d 檔代號(%s)· 券商冊 %d 家 · 工具 %s · 寫入範圍 %s" % (s["roster"], " · ".join("%s %s +%d" % (x["sys"], x["file"], x["added"]) for x in s["roster_src"]) or "無", s["broker_book"], " ".join(k for k, v in s["tools"].items() if v), "VRN ✓" if s["writes_ok"] else "⚠"))
    for r in [x for x in o["rows"] if x["lamp"] == "RED"][:15]:
        print("  [RED] %s · %s · %s" % (r["file"][:60], r["read"] or r["convert_err"], r.get("convert_err", "")[:60]))
    yel = Counter(n for x in o["rows"] if x["lamp"] == "YELLOW" for n in x.get("notes", []) if not n.startswith(("Pdf", "pikepdf")))
    for k, v in yel.most_common(12):
        print("  [YEL] %s · %d 檔" % (k, v))
    print("  [U/I] %s" % o["html"])
    print("NEXT: %s" % ("第一步全綠 → extract 讀 PREP_MANIFEST" if o["lamp"] == "GREEN" else "看矩陣黃紅:券商冊未收 → 補 Broker_Dict;市場別未知 → 補名冊;待 OCR → OCR 閘;報告日用檔案時間 → 人看"))


def main(argv=None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    os.environ.setdefault("VIA_FROM_VCGC", "YES")
    if "--selftest" in args[:2]:
        return selftest()

    def opt(flag, default=None):
        return args[args.index(flag) + 1] if flag in args and args.index(flag) + 1 < len(args) else default
    if args[:1] in (["prep"], ["extract"]) and "--dir" not in args and len(args) >= 2 and Path(args[-1]).is_dir():
        args = args[:-1] + ["--dir", args[-1]]          # via:// 選夾執行把路徑裸附在最後 → 當 --dir
    if args[:1] == ["prep"] and len(args) >= 2 and "--dir" not in args and Path(args[1]).is_dir():
        args = ["prep", "--dir", args[1]] + args[2:]
    if args[:1] == ["ui"]:
        rc = PRIOR.main(args)
        r = _ui_prep_card(_rep() / "VRN_UI_latest.html")
        print("[計] VRN ui v0153 · 左面板 ① prep 卡 %s · 上次 %s" % ("已放" if r.get("patched") else "頁不在", r.get("lamp", "—")))
        return rc
    if args[:1] == ["prep"]:
        cfg = (_resolve("config_get") or (lambda: {}))()
        d = Path(opt("--dir") or cfg.get("input_dir") or r"C:\測試樣本報告")
        o = prep(d, use_vdf=("--no-vdf" not in args), max_pages=int(opt("--max-pages", "60") or 60), workers=int(opt("--workers", "4") or 4), raster_all=("--raster-all" in args))
        _print_prep(o)
        return 1 if o["lamp"] == "RED" and not o["summary"].get("readable") else 0
    return PRIOR.main(args)


def selftest() -> int:
    p = f = 0

    def chk(name, cond):
        nonlocal p, f
        if cond:
            p += 1
            print("  [OK] %s" % name)
        else:
            f += 1
            print("  [FAIL] %s" % name)

    td = Path(tempfile.mkdtemp(prefix="vrnprep-"))
    home = td / "functional modules" / "VRN"
    (home / "knowledge").mkdir(parents=True)
    (home / "SSOT").mkdir()
    vdf = td / "functional modules" / "VDF" / "output_hub" / "export"
    vdf.mkdir(parents=True)
    rep = td / "VIA_Reports" / "vrn"
    inp = td / "inbox"
    inp.mkdir()
    saved = {k: os.environ.get(k) for k in ("VIA_VRN_SSOT_HOME", "VIA_VRN_HEALTH_OUT", "VIA_VDF_HOME", "VIA_SPILL_DIR")}
    os.environ.update({"VIA_VRN_SSOT_HOME": str(home), "VIA_VRN_HEALTH_OUT": str(rep), "VIA_VDF_HOME": str(td / "functional modules" / "VDF"), "VIA_SPILL_DIR": str(td / "spill")})
    (home / "VRN_TWRoster_Offline_v0100.json").write_text(json.dumps({"rows": [{"code": "2330", "name": "台積電", "market": "上市", "name_en": "TSMC"}, {"code": "4171", "name": "瑞基", "market": "上櫃"}, {"code": "3653", "name": "健策", "market": "上市"}, {"code": "3038", "name": "全台", "market": "上市"}, {"code": "1216", "name": "統一", "market": "上市"}]}, ensure_ascii=False), encoding="utf-8")
    (vdf / "vdf_tw_stock_list.csv").write_text("stock_id,name,market,name_en\n5274,信驊,TPEx,ASPEED Technology\n", encoding="utf-8")
    (home / "knowledge" / "VRN_Broker_Dict_v0100.json").write_text(json.dumps({"brokers": {"凱基": {"abbr": "KGI", "aliases": ["凱基", "凱基投顧"]}, "中信": {"abbr": "CTBC", "aliases": ["CTBC", "中信投顧"]}, "華南": {"abbr": "HN", "aliases": ["華南投顧", "華南"]}, "統一": {"abbr": "PSC", "aliases": ["統一投顧"]}},
                                                                           "brokers_extended": {"Goldman Sachs": {"abbr": "GS", "aliases": ["GS", "Goldman Sachs"]}}}, ensure_ascii=False), encoding="utf-8")
    roster, rs = load_roster(True)
    book, bs = load_brokers()
    chk("① 名冊:VRN 自家 + VDF 唯讀(5274 信驊 TPEx)· 券商冊 + 種子(Daiwa 冊沒收 → 種子)", roster["5274"]["market"] == "TPEx" and roster["4171"]["market"] == "TPEx" and any(x["sys"].startswith("VDF") for x in rs) and set(book["Daiwa"]["aliases"].values()) == {"內建種子"} and book["凱基"]["aliases"]["凱基"] == "冊" and book["Goldman Sachs"]["aliases"].get("高盛") == "內建種子")
    cases = {"瑞基(4171,NR_未評等)-CTBC251208.pdf": ("2025-12-08", "4171", "4171.TWO", "4171 TT", "瑞基", "中信"), "GS-2330 20251203.pdf": ("2025-12-03", "2330", "2330.TW", "2330 TT", "台積電", "Goldman Sachs"),
             "3653健策凱基投顧 向子慧 (3).txt": (None, "3653", "3653.TW", "3653 TT", "健策", "凱基"), "260914_daiwa_aspeed.pdf": ("2026-09-14", "5274", "5274.TWO", "5274 TT", "信驊", "Daiwa"),
             "華南投顧-3038-全台-Memo-20251209.docx": ("2025-12-09", "3038", "3038.TW", "3038 TT", "全台", "華南"), "第一場 2026年投資大趨勢 - 華南投顧 -1141201.pdf": ("2025-12-01", "", "", "", "", "華南"),
             "統一投顧-20251209投資早報.pdf": ("2025-12-09", "", "", "", "", "統一"), "【國泰證期研究部】神達(3706 TT)-初次評等買進(+30.4_)-大顯神威.pdf": (None, "3706", "3706", "3706 TT", "神達", "國泰"),
             "凱基投顧_2891 中信金_施志鴻_20260519.pdf": ("2026-05-19", "2891", "2891", "2891 TT", "中信金", "凱基"), "JP-2330 20250718.pdf": ("2025-07-18", "2330", "2330.TW", "2330 TT", "台積電", "J.P. Morgan"),
             "第二場 2026海外投資展望 - 華南永昌海外商品部.pdf": (None, "", "", "", "", "華南")}
    bad = []
    for fn, (dt, code, yf, bbg, nm, br) in cases.items():
        got = parse_name_fields(fn, roster, book)
        d0, _ = parse_date(fn, datetime.datetime(2026, 9, 15).timestamp())
        if (dt and d0 != dt) or got["code"] != code or got["yf"] != yf or got["bbg"] != bbg or (nm and got["name"] != nm) or got["broker"] != br:
            bad.append("%s→%s/%s/%s/%s/%s/%s" % (fn[:20], d0, got["code"], got["yf"], got["bbg"], got["name"], got["broker"]))
    chk("② 檔名判讀 11 例(中信金不被剝 · JP 縮寫靠種子 · 「2026海外」不算代號 · 民國日期 · CTBC 日期 · Bloomberg 形 · 名冊反查英文 aspeed→5274 · 「統一投顧」不誤中統一 1216)%s" % ("" if not bad else " · " + " | ".join(bad)), not bad)
    from reportlab.pdfgen import canvas as _cv
    c = _cv.Canvas(str(inp / "GS-2330 20251203.pdf"))
    c.drawString(72, 720, "Income statement Revenue 1000 Gross profit 400 Operating income 250 Net income 200 " * 2)
    c.save()
    (inp / "3653健策凱基投顧 向子慧 (3).txt").write_text("健策 3653 法說摘要\n營收成長 25%\n毛利率 45%\n" * 5, encoding="cp950")
    from PIL import Image
    Image.new("RGB", (650, 975), (255, 255, 255)).save(inp / "926708.jpg", dpi=(96, 96))
    Image.new("RGB", (400, 500), (250, 250, 250)).save(inp / "scan_only.pdf", "PDF", resolution=100.0)
    (inp / "notes.ini").write_text("x", encoding="utf-8")
    o = prep(inp, max_pages=3, workers=2)
    R = {r["file"]: r for r in o["rows"]}
    import pdfplumber
    with pdfplumber.open(R["926708.jpg"]["read_pdf"]) as pdf:
        w_in = pdf.pages[0].width / 72
    chk("③ 影像 → PDF 標 325 dpi 原像素(650 px = 2.0 in)· TXT(cp950)→ 向量 PDF 可讀 · 不支援檔型 = 灰", abs(w_in - 650 / DPI) < 0.02 and "影像 PDF" in R["926708.jpg"]["read"] and R["3653健策凱基投顧 向子慧 (3).txt"]["read"] == "可讀(文字層)" and R["notes.ini"]["lamp"] == "GRAY")
    chk("④ 掃描件 PDF → 325 dpi 點陣到 TEMP(待 OCR)· 文字 PDF 可讀 · GS-2330 綠", "325" in R["scan_only.pdf"]["read"] and "TEMP" in R["scan_only.pdf"]["read"] and Path(R["scan_only.pdf"]["read_pdf"]).exists() and str(td / "spill") in R["scan_only.pdf"]["read_pdf"] and R["GS-2330 20251203.pdf"]["lamp"] == "GREEN")
    chk("⑤ 產出:矩陣 HTML · MANIFEST · CSV · 帳 · 寫入只在 VRN 範圍", all((rep / "prep" / n).exists() for n in ("PREP_MATRIX_latest.html", "PREP_MANIFEST_latest.json", "PREP_latest.csv", "VRN_Prep_Ledger.jsonl")) and o["summary"]["writes_ok"])
    (rep / "VRN_UI_latest.html").write_text("<html><body><div id='vrn-main-job' x></div></body></html>", encoding="utf-8")
    _ui_prep_card(rep / "VRN_UI_latest.html")
    _ui_prep_card(rep / "VRN_UI_latest.html")
    hh = (rep / "VRN_UI_latest.html").read_text(encoding="utf-8")
    chk("⑧ VRN 頁左面板 ① prep 卡在主作業卡之前 · 重跑不重複 · via://VRN/prep 與選夾", hh.count("vrn-prep-job") == 1 and hh.index("vrn-prep-job") < hh.index("vrn-main-job") and "via://VRN/prep?dir=1" in hh)
    print("  ── 前版鏈自測(原樣印出)──")
    prc = 0 if os.environ.get("VIA_SKIP_PRIOR_SELFTEST") == "1" else PRIOR.selftest()
    chk("⑥ 前版鏈 %s 自測 rc 0" % PRIOR_PATH.stem, prc == 0)
    chk("⑦ 帶加速器橋", "[VIA:ACCEL-BRIDGE:v0100]" in Path(__file__).read_text(encoding="utf-8"))
    for k, v in saved.items():
        if v is None:
            os.environ.pop(k, None)
        else:
            os.environ[k] = v
    shutil.rmtree(td, ignore_errors=True)
    print("[計] VRN_SystemManager_v0153 自測 %d/%d · %s" % (p, p + f, "PASS" if f == 0 else "FAIL"))
    return 0 if f == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
