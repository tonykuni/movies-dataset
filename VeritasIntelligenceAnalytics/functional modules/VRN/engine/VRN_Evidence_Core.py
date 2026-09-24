#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
VRN_Evidence_Core v0102 -- one evidence rule set for both VRN engines.

v0101 -> v0102 (mother 批729; operator 2026-09-24: "報告後小字體不相關附錄可抓到局部識別券商" and
  "報告後方小字級附錄可抓到所有評等法可增加到同義字"):
  appendix_lines() reads the small print of the LAST pages (page 1 excluded; small = glyph size below 0.9 x the
  page-1 body size).  appendix_evidence() grades the issuer there with the SAME broker_evidence() and the same
  deny gate; only the appendix call widens the disclosure test with the Chinese forms (版權所有 / 本報告由 /
  免責聲明 ...), so page-1 grading is byte-identical to v0101 (the 106-report truth set was scored on it).
  rating_scale() harvests the rating definitions printed there (買進(Buy):預期… / Outperform (OP): …) as
  CANDIDATES: each word is KNOWN or NEW against the merged rating words; a NEW word gets a suggested key only when
  its paired word is KNOWN (買進 ~ Buy).  Nothing is ever written to a book here (synonyms are only-add and go
  through the hub).  CLI: python VRN_Evidence_Core.py appendix <pdf|folder>... [--out DIR] [--last N].
  adj_basis() (operator 2026-09-24: "所有目標價及各前一日的價格都要換成 ADJ CLOSE,上漲空間都要用最新的 ADJ CLOSE"):
  the target price and the day-before price in ADJ CLOSE terms and the upside over the latest ADJ CLOSE, delegated
  to the ENG073 tail's adj_quote() (the mother's one L99 formula, 批659); the market database is opened read-only
  and every miss has its own state (ADJ_NO_ENGINE / ADJ_NO_DB / ADJ_NO_DUCKDB / ADJ_DB_BUSY / the engine's own).

v0100 -> v0101 (mother 批728: the sister's 2026-09-23 work brought home to its proper place):
  the two deny-listed CJK names are no longer spelled out in GENERIC_BROKER_ALIASES -- the mother's
  purge gate (CGC_MDL177 verify, 批679/批702) counts any spelled-out denied name as live data.
  generic_aliases() now adds the deny list's own CN_ALIAS at run time instead, so the behaviour is the
  same wherever CGC_MDL177 resolves (and the deny list still runs first in broker_evidence).
  One deny gate for everything (L05): deny_phrases() = the mother overlay's deny_keys (operator rulings 批413 /
  批679 -- the source of truth of the mother's gates; it also denies 摩根 / 摩通, which CGC_MDL177's China list
  does not carry) UNION CGC_MDL177 CN_ALIAS / CN_CANON.  broker_evidence now runs ONE longest-first competition
  between denied names and legitimate aliases per line (the mother's 批681 rule, SUP_MDL015): a bare 摩根 is
  denied while 摩根士丹利 still resolves to MS, and 中信 inside 中信證券 is never CTBC.  deny_shadowed() lets both
  engines apply the same rule to their own alias matchers; compare_broker() denies overlay keys too.

Spec v0.2.0 section 7 names this module.  VIA_VRN_FirstPageEngine and
VRN_Integrated_ReportDatabase_Engine both call it, so a rule lives in one place (LL404: no
second head).  Where a rule exists in the mother books it is read from there:

  supportive modules/registry/VRN_FieldRules_SSOT_v0100.json            rating / target_price / broker /
                                                                        date / contact (批634, ENG086 canon)
  supportive modules/registry/VIA_VRN_ReportFieldRules_SSOT_v0100.json  filename rating slot, analyst titles,
                                                                        broker e-mail domains, phones (批713)
  supportive modules/registry/VIA_VRN_FieldSpec_SSOT_v0100.json         15 columns + six states (批712)
  supportive modules/ssot/VIA_Financial_Institution_SSOT_v0100.json     28 canonical brokers, 6 rating keys

and through the mother's own VCGC tools, mirrored byte-for-byte by scripts/VIA_VCGC_Sync.py:
  CGC_MDL177 ChinaBrokerPurge   deny list (批679: 廣發 / 中信證券·中信建投 / 中金 / 海通 / 國泰君安) runs first
  CGC_MDL182 ReportFieldRulers  per-e-mail contact windows (批713 anchor backtracking), domain verdicts
  CGC_MDL181 VRNFieldSpecGate   the six states of the 15 columns
  SUP_MDL749 VRNFieldRuleHub    the single rule-book reader; its ENG086 answers are kept as a cross-check

The tables in this file are the offline fallback plus additions measured on the 106 real reports
of 2026-09-23 (English sell-side layouts, CTBC filler spaces, Daiwa-Cathay e-mail domain ...).
Every addition names the report family that needed it.

Stdlib only.  LIVE off, no network, never sets VIA_NET_CONSENT / VIA_SCRAPE_CONSENT.
"""

# ===== [VIA:ACCEL-BRIDGE:v0100] SuperAccel 加速器橋(批102 全樹導入令;graceful 零行為變更) =====
try:
    import sys as _sa_sys
    from pathlib import Path as _sa_Path
    _sa_p = _sa_Path(__file__).resolve()
    while _sa_p.parent != _sa_p:
        if (_sa_p / "supportive modules" / "VIA_SuperAccel_Module.py").exists():
            _sa_sys.path.insert(0, str(_sa_p / "supportive modules"))
            break
        _sa_p = _sa_p.parent
    import VIA_SuperAccel_Module as VIA_ACCEL  # noqa: N816
except Exception:
    VIA_ACCEL = None  # graceful:加速器缺席零影響
# ===== [VIA:ACCEL-BRIDGE:END] =====

import copy
import datetime
import functools
import glob
import importlib.util
import json
import os
import re
import statistics
import sys
import unicodedata
from collections import OrderedDict, defaultdict

CORE_VERSION = "v0102"

# =====================================================================
# PARAMETERS (all tunables live here; measured values say where they came from)
# =====================================================================
# Glyph gap (in em) above which two neighbouring glyphs are separated by a space.  Measured on
# page 1 of the 97 real PDFs: letter-letter word gaps sit at 0.15-0.30 em and intra-word gaps below
# 0.05; CJK tracking reaches 0.20; CTBC digit tracking ("3 10 . 0 0") sits at 0.12-0.14.
GAP_EM = {"CC": 0.45, "DD": 0.20, "LL": 0.12, "LD": 0.12, "DL": 0.12, "LP": 0.12, "PL": 0.12}
GAP_EM_DEFAULT = 0.15
ROW_TOL_EM = 0.40            # two glyphs share a row when their baselines differ by less than this
OVERLAY_EM = 0.30            # a glyph overlapping its left neighbour by more than this is an overlay
HEADER_LINES = 14            # first N lines = header zone
FOOTER_LINES = 8             # last N lines = footer zone
DATE_SANE_YEARS = (2015, 2035)
PRICE_SANE = (0.5, 200000.0)

STOCK_TYPES = ("STOCK",)
NON_STOCK_TYPES = ("INDUSTRY", "MARKET", "MACRO", "ETF", "EVENT", "OTHER")

MOTHER_BOOKS = {
    "field_rules": "VRN_FieldRules_SSOT_v0100.json",
    "report_field_rules": "VIA_VRN_ReportFieldRules_SSOT_v0100.json",
    "field_spec": "VIA_VRN_FieldSpec_SSOT_v0100.json",
    "broker_list": "VRN_BROKER_LIST_v01.json",
}
# VCGC tools mirrored from the mother (scripts/VIA_VCGC_Sync.py); newest version wins (tail version, L54)
VCGC_TOOLS = {
    "hub": ("supportive modules/70_VRN_Rules", "SUP_MDL749_VRNFieldRuleHub_v*.py"),
    "rulers": ("supportive modules/registry", "CGC_MDL182_ReportFieldRulers_v*.py"),
    "gate": ("supportive modules/registry", "CGC_MDL181_VRNFieldSpecGate_v*.py"),
    "purge": ("supportive modules/registry", "CGC_MDL177_ChinaBrokerPurge_v*.py"),
}
INSTITUTION_SSOT_GLOB = ("supportive modules/ssot", "VIA_Financial_Institution_SSOT_v*.json")

# =====================================================================
# 1 . RULE BOOKS
# =====================================================================
_RULES_CACHE = {}


def repo_root(start=None):
    """The tree root that holds 'functional modules' (sister repo root = mother VeritasIntelligenceAnalytics/)."""
    return _repo_root_from(os.path.dirname(os.path.abspath(start or __file__)))


@functools.lru_cache(maxsize=64)
def _repo_root_from(probe):
    """批730 speed: the walk up is file-system calls (slow on a synced Windows folder) and ran ~30,000 times per
    batch from the deny check; the tree does not move during a run, so each starting folder is walked once."""
    start_dir = probe
    while True:
        if os.path.isdir(os.path.join(probe, "functional modules")):
            return probe
        parent = os.path.dirname(probe)
        if parent == probe:
            return start_dir
        probe = parent


def find_rule_books(start=None):
    """Locate the mother books: the VCGC mirror first, then VRN/registry, then the sealed intake baseline."""
    here = os.path.dirname(os.path.abspath(start or __file__))
    vrn_root = os.path.dirname(here) if os.path.basename(here) == "engine" else here
    candidates = [
        os.path.join(repo_root(start), "supportive modules", "registry"),       # VCGC mirror (scripts/VIA_VCGC_Sync.py)
        os.path.join(vrn_root, "registry"),
        os.path.join(vrn_root, "references", "intake", "VIA_SSOT_Additive_Audit_v0100", "baseline"),
    ]
    found = {}
    for key, name in MOTHER_BOOKS.items():
        for folder in candidates:
            path = os.path.join(folder, name)
            if os.path.isfile(path):
                found[key] = path
                break
    return found


def _read_json(path):
    try:
        with open(path, "r", encoding="utf-8-sig") as handle:
            return json.load(handle)
    except Exception:
        return None


def _version_of(path):
    m = re.search(r"_v(\d+)\.", os.path.basename(path))
    return int(m.group(1)) if m else -1


_VCGC_CACHE = {}


def vcgc(name, start=None):
    """Import one mirrored VCGC tool: hub (SUP_MDL749) / rulers (CGC_MDL182) / gate (CGC_MDL181) /
    purge (CGC_MDL177).  None when the mirror is absent or the tool fails to load (graceful)."""
    root = repo_root(start)
    key = (name, root)
    if key in _VCGC_CACHE:
        return _VCGC_CACHE[key]
    mod = None
    folder, pattern = VCGC_TOOLS.get(name, (None, None))
    if folder:
        try:
            hits = sorted(glob.glob(os.path.join(root, folder, pattern)), key=_version_of)
            if hits:
                modname = "_vrn_vcgc_" + name
                spec = importlib.util.spec_from_file_location(modname, hits[-1])
                mod = importlib.util.module_from_spec(spec)
                sys.modules[modname] = mod
                spec.loader.exec_module(mod)
                mod.__vcgc_path__ = hits[-1]
        except Exception:
            mod = None
    _VCGC_CACHE[key] = mod
    return mod


OVERLAY_GLOB = ("supportive modules/ssot", "VIA_FinancialInstitution_Overlay_v*.json")
_DENY_CACHE = {}


@functools.lru_cache(maxsize=256)
def _deny_norm(value):
    # 批730 speed: the same page text is checked once per broker alias (hundreds per file) -- normalise it once
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value or ""))).strip().casefold()


_DENY_NORM_CACHE = {}


def _deny_table(start=None):
    """(normalised deny phrases, longest first · the same as a set): normalised once per tree.
    批730 speed: the check used to re-normalise every deny phrase on every call (1.17 million times for 19 one-page
    files); the answers are the same, only computed once."""
    root = repo_root(start)
    hit = _DENY_NORM_CACHE.get(root)
    if hit is None:
        norm = tuple(_deny_norm(k) for k in deny_phrases(start))
        hit = (norm, frozenset(norm))
        _DENY_NORM_CACHE[root] = hit
    return hit


def deny_phrases(start=None):
    """The mother's broker deny list (v0101): overlay deny_keys (newest VIA_FinancialInstitution_Overlay_v*.json;
    operator rulings 批413/批679) UNION CGC_MDL177 CN_ALIAS / CN_CANON.  Original spellings, longest first.
    Empty when neither source resolves (the sister repo without its mirror): nothing is invented here."""
    root = repo_root(start)
    if root in _DENY_CACHE:
        return _DENY_CACHE[root]
    out = set()
    try:
        hits = sorted(glob.glob(os.path.join(root, OVERLAY_GLOB[0], OVERLAY_GLOB[1])), key=_version_of)
        if hits:
            for k in (_read_json(hits[-1]) or {}).get("deny_keys", []) or []:
                if str(k).strip():
                    out.add(str(k).strip())
    except Exception:
        pass
    purge = vcgc("purge", start)
    for attr in ("CN_ALIAS", "CN_CANON"):
        for k in (getattr(purge, attr, ()) or ()) if purge is not None else ():
            if str(k).strip():
                out.add(str(k).strip())
    res = tuple(sorted(out, key=lambda x: (-len(x), x)))
    _DENY_CACHE[root] = res
    return res


def is_denied_token(value, start=None):
    """True when the whole value IS a denied name (normalised exact match) -- the filename / canonical-key layer."""
    n = _deny_norm(value)
    return bool(n) and n in _deny_table(start)[1]


def deny_shadowed(alias, text, start=None):
    """True when EVERY occurrence of a legitimate alias in `text` sits inside a denied name that is at least as
    long (longest wins, the mother's 批681 rule): 中信 inside 中信證券 -> shadowed; 摩根士丹利 contains the denied
    bare 摩根 but is longer -> not shadowed; '中信證券 ... 中信投顧' keeps the standalone hit.  An alias that is
    itself a denied name is always shadowed."""
    a = _deny_norm(alias)
    if not a:
        return False
    if is_denied_token(a, start):
        return True
    low = _deny_norm(text)
    a_spans = list(_alias_spans(a, low))
    if not a_spans:
        return False
    d_spans = []
    for dn in _deny_table(start)[0]:
        if len(dn) >= len(a) and a in dn:
            d_spans.extend(_alias_spans(dn, low))
    return bool(d_spans) and all(any(ds <= s0 and t0 <= dt for ds, dt in d_spans) for s0, t0 in a_spans)


def institution_ssot(start=None):
    """Mother institution SSOT (28 canonical brokers, 6 rating keys); None when absent.  The mother path is
    tried first; in this repo the book lives once, in the sealed audit package (never copied twice)."""
    root = repo_root(start)
    folders = [os.path.join(root, INSTITUTION_SSOT_GLOB[0]),
               os.path.join(root, "functional modules", "VRN", "references", "intake", "VIA_SSOT_Additive_Audit_v0100", "baseline")]
    for folder in folders:
        hits = sorted(glob.glob(os.path.join(folder, INSTITUTION_SSOT_GLOB[1])), key=_version_of)
        hits = [h for h in hits if "_sha" not in os.path.basename(h)]
        if hits:
            return _read_json(hits[-1])
    return None


def vcgc_status(start=None):
    """Which mirrored tools and books this run actually used (shown in reports; LL402)."""
    out = {name: (os.path.relpath(getattr(vcgc(name, start), "__vcgc_path__", ""), repo_root(start))
                  if vcgc(name, start) else None) for name in VCGC_TOOLS}
    out["institution_ssot"] = bool(institution_ssot(start))
    out["books"] = sorted(find_rule_books(start))
    return out


def load_rules(start=None):
    """Read the books once; every consumer gets the same dict (paths + raw books)."""
    key = os.path.abspath(start or __file__)
    if key in _RULES_CACHE:
        return _RULES_CACHE[key]
    paths = find_rule_books(start)
    books = {k: _read_json(p) for k, p in paths.items()}
    rules = {"paths": paths, "books": {k: v for k, v in books.items() if v is not None}}
    _RULES_CACHE[key] = rules
    return rules


def _rule(rules, *path, default=None):
    node = rules.get("books", {}) if isinstance(rules, dict) else {}
    for part in path:
        if isinstance(node, dict) and part in node:
            node = node[part]
        else:
            return default
    return node


# =====================================================================
# 2 . TEXT LINES FROM GLYPHS (PDF) / PARAGRAPHS (DOCX) / FILE (TXT)
# =====================================================================
def char_class(ch):
    """C = CJK ideograph, D = digit or decimal mark, L = Latin letter, P = anything else."""
    if not ch:
        return "P"
    o = ord(ch[0])
    if 0x4E00 <= o <= 0x9FFF or 0x3400 <= o <= 0x4DBF or 0xF900 <= o <= 0xFAFF:
        return "C"
    if ch[0].isdigit() or ch[0] in ".,":
        return "D"
    if ch[0].isalpha():
        return "L"
    return "P"


def _is_wide(ch):
    return bool(ch) and unicodedata.east_asian_width(ch[0]) in ("W", "F")


def _glyph_right(c):
    if c.get("x1") is not None:
        return float(c["x1"])
    size = float(c.get("size") or 10.0)
    return float(c["x0"]) + size * (1.0 if _is_wide(c.get("text", "")) else 0.5)


def _baseline(c):
    if c.get("bottom") is not None:
        return float(c["bottom"])
    return float(c.get("top", 0.0)) + float(c.get("size") or 10.0)


def join_row(chars):
    """Glyphs of one row -> text.  Space glyphs are ignored; a space is written only where the
    geometric gap says so (CTBC filler spaces vanish, English words stop gluing together)."""
    out = []
    prev = None
    solid = [float(g.get("size") or 0.0) for g in chars if g.get("text", "").strip()]
    median = statistics.median(solid) if solid else 0.0
    for c in sorted(chars, key=lambda g: float(g.get("x0", 0.0))):
        t = c.get("text", "")
        if not t or not t.strip():
            continue
        if prev is not None:
            if t == prev.get("text") and abs(float(c["x0"]) - float(prev["x0"])) < 1.0:
                continue                         # faux-bold overprint: same glyph within 1pt
            size = max(float(prev.get("size") or 0.0), float(c.get("size") or 0.0), 1.0)
            gap = (float(c["x0"]) - _glyph_right(prev)) / size
            if gap < -OVERLAY_EM:
                # a glyph drawn on top of its neighbour belongs to another layer (chart label over text:
                # KGI '九4月' / a name split by a stray '%'); keep the one whose size matches the row
                if abs(float(c.get("size") or 0.0) - median) >= abs(float(prev.get("size") or 0.0) - median):
                    continue
                out.pop()
                if out and out[-1] == " ":
                    out.pop()
                prev = None
                out.append(t)
                prev = c
                continue
            pair = char_class(prev["text"][-1]) + char_class(t[0])
            if gap > GAP_EM.get(pair, GAP_EM_DEFAULT):
                out.append(" ")
        out.append(t)
        prev = c
    return "".join(out)


def group_rows(chars):
    """Cluster glyphs into rows by baseline (tolerance relative to the glyph size)."""
    items = [c for c in (chars or []) if c.get("text")]
    items.sort(key=lambda c: (_baseline(c), float(c.get("x0", 0.0))))
    rows = []
    for c in items:
        base = _baseline(c)
        size = float(c.get("size") or 10.0) or 10.0
        if rows and abs(base - rows[-1]["base"]) < ROW_TOL_EM * min(size, rows[-1]["size"]):
            rows[-1]["chars"].append(c)
        else:
            rows.append({"base": base, "size": size, "chars": [c]})
    return [r["chars"] for r in rows]


def lines_from_chars(chars):
    """Rows of glyphs -> line dicts (text, size, bold, top, x0, x1, bottom, chars), top to bottom."""
    lines = []
    for cs in group_rows(chars):
        text = join_row(cs)
        if not text.strip():
            continue
        solid = [c for c in cs if c.get("text", "").strip()] or cs
        sizes = [float(c.get("size") or 0.0) for c in solid]
        bold = sum(1 for c in solid if "bold" in str(c.get("fontname") or "").lower()) > len(solid) / 2
        lines.append({
            "text": text,
            "size": round(statistics.median(sizes), 2) if sizes else 0.0,
            "bold": bold,
            "top": min(float(c.get("top", 0.0)) for c in solid),
            "bottom": max(_baseline(c) for c in solid),
            "x0": min(float(c.get("x0", 0.0)) for c in solid),
            "x1": max(_glyph_right(c) for c in solid),
            "chars": sorted(solid, key=lambda g: float(g.get("x0", 0.0))),
        })
    return lines


# 批730 speed (operator 2026-09-24「剛剛跑好慢」): one report's page 1 and its last pages were parsed again by every
# reader (first page, appendix body size, appendix, column tables) and again in G07 / G08 -- 24 pdfplumber parses for a
# 15-page file.  The same file (path · size · modified time) and page give the same glyphs, so they are parsed once;
# callers get copies, and a file edited in place gets a new key.
_PAGE_CACHE = OrderedDict()
_PAGE_CACHE_MAX = 24
_APPENDIX_CACHE = OrderedDict()
_APPENDIX_CACHE_MAX = 512
_PAGE_COUNT_CACHE = {}


def _file_key(path):
    """(absolute path, size, modified time in ns); None when the file cannot be stat'ed (then nothing is cached)."""
    try:
        st = os.stat(path)
    except OSError:
        return None
    return os.path.abspath(str(path)), st.st_size, st.st_mtime_ns


def pdf_page_chars(path, page_index=0, order=("pdfplumber", "fitz")):
    """Glyphs of one PDF page as dicts (text, x0, x1, top, bottom, size, fontname) plus (width, height).
    Returns None when no PDF library is installed or the file cannot be opened."""
    key = _file_key(path)
    k = key + (int(page_index), tuple(order)) if key is not None else None
    if k is not None and k in _PAGE_CACHE:
        _PAGE_CACHE.move_to_end(k)
        chars, size = _PAGE_CACHE[k]
        return [dict(c) for c in chars], size
    got = _pdf_page_chars_parse(path, page_index, order)
    if k is not None and got is not None:
        _PAGE_CACHE[k] = ([dict(c) for c in got[0]], got[1])
        while len(_PAGE_CACHE) > _PAGE_CACHE_MAX:
            _PAGE_CACHE.popitem(last=False)
    return got


def _pdf_page_chars_parse(path, page_index=0, order=("pdfplumber", "fitz")):
    """The uncached reader behind pdf_page_chars()."""
    for engine in order:
        if engine == "pdfplumber":
            try:
                import pdfplumber
            except ImportError:
                continue
            try:
                with pdfplumber.open(path) as pdf:
                    if len(pdf.pages) <= page_index:
                        return [], (595.0, 842.0)
                    page = pdf.pages[page_index]
                    # rotated glyphs (vertical chart labels) are left out: they interleave with the rows
                    # they cross ('九月 15, 2026' became '九4月 15, 2026' on KGI 2637)
                    chars = [{"text": c.get("text", ""), "x0": float(c.get("x0", 0)), "x1": float(c.get("x1", 0)),
                              "top": float(c.get("top", 0)), "bottom": float(c.get("bottom", 0)),
                              "size": float(c.get("size", 0)), "fontname": str(c.get("fontname", ""))}
                             for c in page.chars if c.get("upright", True)]
                    return chars, (float(page.width), float(page.height))
            except Exception:
                continue
        elif engine == "fitz":
            try:
                import fitz
            except ImportError:
                continue
            try:
                doc = fitz.open(path)
                try:
                    if len(doc) <= page_index:
                        return [], (595.0, 842.0)
                    page = doc[page_index]
                    chars = []
                    for block in page.get_text("rawdict").get("blocks", []):
                        for line in block.get("lines", []):
                            direction = line.get("dir", (1.0, 0.0))
                            if abs(direction[0] - 1.0) > 0.01 or abs(direction[1]) > 0.01:
                                continue            # rotated text line
                            for span in line.get("spans", []):
                                for ch in span.get("chars", []):
                                    x0, y0, x1, y1 = ch["bbox"]
                                    chars.append({"text": ch.get("c", ""), "x0": float(x0), "x1": float(x1),
                                                  "top": float(y0), "bottom": float(y1),
                                                  "size": float(span.get("size", 0)), "fontname": str(span.get("font", ""))})
                    return chars, (float(page.rect.width), float(page.rect.height))
                finally:
                    doc.close()
            except Exception:
                continue
    return None


def pdf_page_lines(path, page_index=0):
    """Text lines of one PDF page (list of str); [] when the page has no text layer."""
    got = pdf_page_chars(path, page_index)
    if not got:
        return []
    chars, _size = got
    return [l["text"] for l in lines_from_chars(chars)]


def _docx_textbox_paragraphs(element):
    """Text boxes (w:txbxContent) are outside document.paragraphs; the HuaNan memo date lives there."""
    out = []
    ns_w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for box in element.iter(ns_w + "txbxContent"):
        for para in box.iter(ns_w + "p"):
            text = "".join(t.text or "" for t in para.iter(ns_w + "t")).strip()
            if text:
                out.append(text)
    return out


def docx_lines(path):
    """Header, text boxes, body paragraphs, tables and footer of a DOCX, in reading order.
    Returns (lines, sizes) where sizes[i] is the largest run size of line i (None when unknown)."""
    try:
        import docx
    except ImportError:
        return None
    try:
        document = docx.Document(path)
    except Exception:
        return None
    lines, sizes = [], []

    def add(text, size=None):
        text = (text or "").strip()
        if text:
            lines.append(text)
            sizes.append(size)

    seen = set()
    for section in document.sections:
        for part in (section.header, section.first_page_header):
            try:
                for para in part.paragraphs:
                    if para.text.strip() and para.text.strip() not in seen:
                        seen.add(para.text.strip())
                        add(para.text)
                for text in _docx_textbox_paragraphs(part._element):
                    if text not in seen:
                        seen.add(text)
                        add(text)
            except Exception:
                continue
    for text in _docx_textbox_paragraphs(document.element.body):
        if text not in seen:
            seen.add(text)
            add(text)
    for para in document.paragraphs:
        run_sizes = [run.font.size.pt for run in para.runs if run.font is not None and run.font.size is not None]
        add(para.text, max(run_sizes) if run_sizes else None)
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                t = cell.text.strip()
                if t and (not cells or cells[-1] != t):
                    cells.append(t)
            add(" ".join(cells))
    for section in document.sections:
        for part in (section.footer, section.first_page_footer):
            try:
                for para in part.paragraphs:
                    if para.text.strip() and para.text.strip() not in seen:
                        seen.add(para.text.strip())
                        add(para.text)
            except Exception:
                continue
    return lines, sizes


def txt_lines(path):
    """Plain-text report: try UTF-8 (with BOM), UTF-16, CP950 / Big5."""
    raw = None
    try:
        with open(path, "rb") as handle:
            raw = handle.read()
    except Exception:
        return None
    for enc in ("utf-8-sig", "utf-16", "cp950", "big5", "latin-1"):
        try:
            text = raw.decode(enc)
        except Exception:
            continue
        if enc == "utf-16" and not raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
            continue
        return [l.strip() for l in text.splitlines() if l.strip()]
    return None


def document_first_page_lines(path):
    """First-page lines of any supported report: PDF page 1, DOCX whole document, TXT whole file.
    Returns {"lines": [...], "source": "pdf|docx|txt", "text_layer": bool}."""
    suffix = os.path.splitext(path)[1].lower()
    if suffix == ".pdf":
        lines = pdf_page_lines(path, 0)
        return {"lines": lines, "source": "pdf", "text_layer": bool(lines)}
    if suffix == ".docx":
        got = docx_lines(path)
        lines = got[0] if got else []
        return {"lines": lines, "source": "docx", "text_layer": bool(lines)}
    if suffix == ".txt":
        lines = txt_lines(path) or []
        return {"lines": lines, "source": "txt", "text_layer": bool(lines)}
    return {"lines": [], "source": suffix.lstrip("."), "text_layer": False}


# =====================================================================
# 3 . VOCABULARY (fallback tables; the books extend them, never replace)
# =====================================================================
# Rating words -> canonical.  Mother rating.canon_map + rating.local_scale are merged in at runtime.
RATING_WORDS = {
    "strong buy": "BUY", "conviction buy": "BUY", "high-conviction outperform": "BUY", "buy": "BUY",
    "outperform": "BUY", "overweight": "BUY", "add": "ADD", "accumulate": "ADD",
    "neutral": "HOLD", "hold": "HOLD", "market perform": "HOLD", "sector perform": "HOLD",
    "equal-weight": "HOLD", "equalweight": "HOLD", "equal weight": "HOLD", "market weight": "HOLD",
    "sell": "SELL", "underperform": "SELL", "underweight": "SELL", "reduce": "SELL",
    "not rated": "NOT_RATED", "not-rated": "NOT_RATED", "no rating": "NOT_RATED", "unrated": "NOT_RATED",
    "強力買進": "BUY", "積極買進": "BUY", "買進": "BUY", "增加持股": "BUY", "增持": "BUY", "加碼": "BUY",
    "優於大盤": "BUY", "強於大盤": "BUY", "表現優於大盤": "BUY",
    "逢低買進": "ADD", "逢低": "ADD",
    "中立": "HOLD", "持有": "HOLD", "區間持有": "HOLD", "區間操作": "HOLD", "區間": "HOLD", "持平": "HOLD",
    "標準配置": "HOLD", "同步大盤": "HOLD", "與大盤同步": "HOLD",
    "減碼": "SELL", "賣出": "SELL", "減少持股": "SELL", "減持": "SELL", "劣於大盤": "SELL", "弱於大盤": "SELL",
    "表現劣於大盤": "SELL",
    "未評等": "NOT_RATED", "不予評等": "NOT_RATED", "暫不評等": "NOT_RATED", "未評級": "NOT_RATED", "尚未評級": "NOT_RATED",
}
RATING_FINE = {"strong buy": "STRONG_BUY", "conviction buy": "STRONG_BUY", "high-conviction outperform": "STRONG_BUY",
               "強力買進": "STRONG_BUY", "積極買進": "STRONG_BUY", "逢低買進": "BUY_ON_DIPS"}
RATING_SHORT = {"OW": "BUY", "EW": "HOLD", "UW": "SELL", "B": "BUY", "N": "HOLD", "H": "HOLD", "S": "SELL",
                "NR": "NOT_RATED", "OP": "BUY", "UP": "SELL", "MP": "HOLD", "SB": "BUY"}
RATING_ACTION_WORDS = {
    "UPGRADE": ("調升", "升評", "上調", "upgrade", "upgraded"),
    "DOWNGRADE": ("調降", "降評", "下調", "downgrade", "downgraded"),
    "MAINTAIN": ("維持", "重申", "maintain", "maintained", "reiterate", "reiterated", "reaffirm", "keep"),
    "INITIATE": ("初次評等", "首次評等", "初評", "初次", "重啟", "重新覆蓋", "initiate", "initiating", "initiation"),
}

# E-mail domain -> broker.  Mother ReportFieldRules broker_domains_addendum is merged in at runtime
# (its labels are bridged by DOMAIN_LABEL_BRIDGE).  Added 2026-09-23 from the real corpus:
# gfgroup.com.hk (GFHK), cl-sec.com (CLSA), uni-psg.com (統一), daiwacm-cathay.com.tw is Daiwa research.
DOMAIN_BROKER = {
    "gs.com": "GS", "morganstanley.com": "MS", "ubs.com": "UBS", "citi.com": "CITI", "jpmorgan.com": "JPM",
    "jpmchase.com": "JPM", "daiwacm-cathay.com.tw": "DAIWA", "daiwacm.com": "DAIWA", "cl-sec.com": "CLSA",
    "clsa.com": "CLSA", "macquarie.com": "MACQUARIE", "kgi.com": "KGI", "kgi.com.tw": "KGI", "kgieworld.com.tw": "KGI",
    "ctbcsec.com": "CTBC", "ctbcsecurities.com": "CTBC", "ctbcbank.com": "CTBC", "ctbcsis.com": "CTBC",
    "megabank.com.tw": "MEGA", "emega.com.tw": "MEGA", "megasec.com.tw": "MEGA", "entrust.com.tw": "HUANAN",
    "cathaysec.com.tw": "CATHAY", "cathayfut.com.tw": "CATHAY", "capital.com.tw": "CAPITAL",
    "cim.capital.com.tw": "CAPITAL", "pscnet.com.tw": "PRESIDENT", "uni-psg.com": "PRESIDENT",
    "gfgroup.com.hk": "GF", "yuanta.com": "YUANTA", "yuanta.com.tw": "YUANTA", "fubon.com": "FUBON",
    "sinopac.com": "SINOPAC", "nomura.com": "NOMURA", "bofa.com": "BOFA", "hsbc.com": "HSBC", "hsbc.com.tw": "HSBC",
    "masterlink.com.tw": "MASTERLINK", "esunsec.com.tw": "ESUN", "taishinbank.com.tw": "TAISHIN",
}
DOMAIN_LABEL_BRIDGE = {"MORGAN STANLEY": "MS", "GOLDMAN SACHS": "GS", "J.P. MORGAN": "JPM", "DAIWA-CATHAY": "DAIWA",
                       "JPMORGAN": "JPM", "MORGANSTANLEY": "MS", "FIRSTSEC": "FIRST", "CREDITSUISSE": "CS"}

# Aliases that are ordinary words or company names: never broker evidence outside an issuer line.
# Mother broker.generic_latin is merged in at runtime (add, capital, concord, first, mega, oriental, president).
GENERIC_BROKER_ALIASES = {
    "first", "capital", "mega", "president", "add", "concord", "oriental", "jp", "ms", "gs", "gf", "mq", "mcq",
    "ml", "boa", "db", "cs", "nmr", "citic", "megabank", "guotai junan", "大華", "第一", "統一",
    "中信金", "國泰金控", "富邦金控", "元大金控", "永豐金", "第一金", "兆豐金", "華南金", "台新金", "玉山金", "凱基金",
}
# Joint-venture / legal names that must win over their parts (大和國泰 is Daiwa research, not Cathay).
EXTRA_BROKER_ALIASES = {
    "DAIWA": ["daiwa-cathay", "daiwa cathay", "大和國泰", "daiwa capital markets"],
    "CTBC": ["ctbc securities investment service", "中國信託證券投顧", "中國信託證券投資顧問"],
    "KGI": ["凱基證券投資顧問"],
    "GF": ["gf securities (hong kong)"],
    "JPM": ["j.p. morgan securities"],
    "MS": ["morgan stanley taiwan", "morgan stanley asia"],
    "GS": ["goldman sachs (asia)"],
    "CATHAY": ["國泰期貨", "國泰證券投資顧問"],
    "MASTERLINK": ["元富投顧", "元富證券投資顧問"],
}
ISSUER_AFTER_RX = re.compile(
    r"^[\s,.\-]*(?:\((?:asia|taiwan|hong kong|singapore|asia pacific)\)\s*|taiwan\s+|asia\s+|hong kong\s+)?"
    r"(?:securities|證券|投顧|投資顧問|證期|期貨|research|capital\s+markets|global\s+markets|l\.?\s?l\.?\s?c|limited|ltd|"
    r"inc\b|investment\s+consulting|investment\s+service|brokerage|研究部|研究中心|綜合證券)", re.I)
ISSUER_WORD_RX = re.compile(r"(證券|投顧|投資顧問|證期|期貨|securities|capital markets|l\.l\.c|limited)", re.I)
SOURCE_LINE_RX = re.compile(r"(資料來源|來源\s*[:：]|source\s*[:：]|sources\s*[:：])", re.I)
DISCLOSURE_RX = re.compile(r"does and seeks to do business with companies covered", re.I)
COPYRIGHT_RX = re.compile(r"(©|copyright|著作權)", re.I)
RECIPIENT_RX = re.compile(r"(exclusive use|prepared for|for the use of|distributed to|供.{0,6}使用)", re.I)
EMAIL_RX = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9.'-]+\.[A-Za-z]{2,}")
URL_RX = re.compile(r"(?:https?://|www\.)\S+", re.I)

EN_MONTHS = {"jan": 1, "january": 1, "feb": 2, "february": 2, "mar": 3, "march": 3, "apr": 4, "april": 4,
             "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7, "aug": 8, "august": 8, "sep": 9, "sept": 9,
             "september": 9, "oct": 10, "october": 10, "nov": 11, "november": 11, "dec": 12, "december": 12}
ZH_MONTHS = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10, "十一": 11, "十二": 12}
_MON = r"(Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)"

# Report-type cues (filename first; page second).  Calibrated on the 106 real names.
TYPE_FILENAME_CUES = [
    ("ETF", re.compile(r"ETF", re.I)),
    ("MACRO", re.compile(r"總經|聯準會|央行|FOMC|macro", re.I)),
    ("EVENT", re.compile(r"論壇|研討會|第[一二三四五六七八九十]場|summit|forum|conference", re.I)),
    ("MARKET", re.compile(r"晨會|早報|盤勢|盤後|市場觀察|日股|港股|美股|期貨|解盤|晨間|操作策略|daily|morning", re.I)),
    ("INDUSTRY", re.compile(r"產業|專題|industry|sector|databook|thermal|memory|automation|pcb|ccl|abf|power|hardware|"
                            r"components|optical|iphone|tpu|cpo|insights|功耗", re.I)),
]
STOCK_FILENAME_CUE = re.compile(r"個股|訪談|memo|takeaway|介紹|初次評等|note", re.I)


def merged_rating_words(rules=None):
    words = dict(RATING_WORDS)
    canon = _rule(rules or {}, "field_rules", "rules", "rating", "canon_map", default={}) or {}
    for level, terms in canon.items():
        for term in terms:
            words.setdefault(str(term).lower(), level)
    local = _rule(rules or {}, "field_rules", "rules", "rating", "local_scale", default={}) or {}
    for term, level in local.items():
        words.setdefault(str(term).lower(), level)
    return words


def merged_domain_map(rules=None):
    domains = dict(DOMAIN_BROKER)
    add = _rule(rules or {}, "report_field_rules", "broker_domains_addendum", "map", default={}) or {}
    for dom, label in add.items():
        key = DOMAIN_LABEL_BRIDGE.get(str(label).upper(), str(label).upper())
        domains.setdefault(str(dom).lower(), key)
    return domains


def generic_aliases(rules=None):
    out = set(GENERIC_BROKER_ALIASES)
    for word in _rule(rules or {}, "field_rules", "rules", "broker", "generic_latin", default=[]) or []:
        out.add(str(word).lower())
    # v0101: the deny list's own aliases are never broker evidence either -- read from CGC_MDL177, not spelled here
    purge = vcgc("purge")
    for word in (getattr(purge, "CN_ALIAS", ()) or ()) if purge is not None else ():
        out.add(str(word).lower())
    return out


def _alternation(words):
    return "|".join(re.escape(w) for w in sorted(set(words), key=len, reverse=True))


def canonical_rating_key(raw, coarse, fine=None, inst=None):
    """Mother rule: rating must land on the hub's canonical keys (institution SSOT broker_ratings:
    NOT_RATED / STRONG_BUY / BUY / HOLD / SELL / STRONG_SELL).  Alias lookup first; ADD -> BUY there."""
    reg = ((inst or {}).get("registries") or {}).get("broker_ratings") or {}
    low = re.sub(r"\s*\((?:\d|unchanged|on\s*cl)\)$", "", (raw or "").strip().lower())
    for key, spec in reg.items():
        names = [spec.get("chinese_name", ""), spec.get("english_name", "")] + list(spec.get("aliases") or [])
        if any(low == str(a).strip().lower() for a in names if a):
            return key, "INSTITUTION_ALIAS"
    if fine in ("STRONG_BUY", "STRONG_SELL"):
        return fine, "FINE"
    key = {"BUY": "BUY", "ADD": "BUY", "HOLD": "HOLD", "SELL": "SELL", "NOT_RATED": "NOT_RATED"}.get(coarse or "")
    if key and (not reg or key in reg):
        return key, "COARSE"
    return None, "NOT_IN_KEYS"


def normalize_rating_word(word, words=None):
    """'Buy (1)' / '買進(Buy)' / 'OW' / 'Equal-weight' -> (canonical, fine) or (None, None)."""
    if not word:
        return None, None
    words = words or RATING_WORDS
    w = re.sub(r"\s+", " ", str(word)).strip().strip("「」『』“”\"'.,;:()（）").lower()
    w = re.sub(r"\s*\((?:\d|unchanged|on cl)\)$", "", w)
    if w.upper() in RATING_SHORT:
        return RATING_SHORT[w.upper()], None
    if w in words:
        return words[w], RATING_FINE.get(w)
    for term in sorted(words, key=len, reverse=True):
        if re.search(r"[一-鿿]", term) and term in w:
            return words[term], RATING_FINE.get(term)
    return None, None


# =====================================================================
# 4 . E-MAILS AND BROKER EVIDENCE (tiers: EMAIL > DISCLOSURE > ISSUER > HEADER > BODY)
# =====================================================================
def find_emails(lines):
    """Analyst e-mails with their line index; recipient stamps ('exclusive use of ...') are excluded."""
    out = []
    for i, line in enumerate(lines):
        if RECIPIENT_RX.search(line):
            continue
        for m in EMAIL_RX.finditer(line):
            email = m.group(0).strip("'.")
            local, _, domain = email.partition("@")
            domain = domain.strip("'").lower()
            if not local or not domain:
                continue
            out.append({"email": local + "@" + domain, "local": local, "domain": domain, "line": i, "start": m.start()})
    return out


def domain_broker(domain, domains=None):
    """Suffix match (hk.daiwacm.com -> daiwacm.com); a glued tail (…com.twequity) still matches its prefix."""
    domains = domains or DOMAIN_BROKER
    dom = (domain or "").lower().strip(".'")
    best = None
    for known, broker in domains.items():
        if dom == known or dom.endswith("." + known) or (dom.startswith(known) and dom[len(known):len(known) + 1].isalpha()):
            if best is None or len(known) > len(best[0]):
                best = (known, broker)
    return best[1] if best else None


def _alias_spans(alias, low):
    if re.search(r"[一-鿿]", alias):
        start = 0
        while True:
            idx = low.find(alias, start)
            if idx < 0:
                return
            yield idx, idx + len(alias)
            start = idx + 1
    else:
        for m in re.finditer(r"(?<![a-z0-9])" + re.escape(alias) + r"(?![a-z0-9])", low):
            yield m.start(), m.end()


_CODE_AFTER_RX = re.compile(r"^\s*(?:-?\s*ky)?\s*[（(]?\s*[1-9]\d{3}(?![0-9])")
_CODE_BEFORE_RX = re.compile(r"[1-9]\d{3}(?:\.twO?|\s*tt)?\s*[)）]?\s*$", re.I)


def _adjacent_to_code(low, start, end):
    """LL67: a name glued to a stock code is the covered company, not the issuer."""
    return bool(_CODE_AFTER_RX.match(low[end:end + 12]) or _CODE_BEFORE_RX.search(low[max(0, start - 10):start]))


def broker_evidence(lines, alias_table, filename_broker=None, rules=None, disclosure_rx=None, issuer_guard=None):
    """Who issued this report, graded by evidence tier.

    alias_table: {canonical: [aliases]} from the calling engine (its dictionary + knowledge).
    Returns {"broker", "tier", "strong", "candidates": {broker: {tier: count}}, "evidence": [...]}.
    strong = the page itself proves the issuer (e-mail domain, disclosure/copyright, issuer line);
    only a strong page broker may contradict the filename broker."""
    domains = merged_domain_map(rules)
    generic = generic_aliases(rules)
    purge = vcgc("purge")

    deny_list = deny_phrases()

    def denied_alias(alias):
        try:
            return bool(purge and purge.is_cn_alias(alias)) or is_denied_token(alias)
        except Exception:
            return False

    def denied_key(key):
        try:
            return bool(purge and purge.is_cn_canon(key)) or is_denied_token(key)
        except Exception:
            return False

    cn_domains = tuple(getattr(purge, "CN_DOMAIN", ()) or ())
    denied = defaultdict(int)
    table = defaultdict(set)
    for canon, aliases in (alias_table or {}).items():
        for a in aliases:
            if a:
                table[str(canon).upper()].add(str(a).lower())
    for canon, aliases in EXTRA_BROKER_ALIASES.items():
        table[canon].update(a.lower() for a in aliases)
    flat = sorted(((a, c) for c, al in table.items() for a in al), key=lambda x: -len(x[0]))

    tiers = defaultdict(lambda: defaultdict(int))
    evidence = []
    for e in find_emails(lines):
        b = domain_broker(e["domain"], domains)
        if any(e["domain"] == d or e["domain"].endswith("." + d) for d in cn_domains) or (b and denied_key(b)):
            denied[b or e["domain"]] += 1
            evidence.append({"broker": b, "tier": "DENIED", "line": e["line"], "text": "@" + e["domain"]})
            continue
        if b:
            tiers[b]["EMAIL"] += 1
            evidence.append({"broker": b, "tier": "EMAIL", "line": e["line"], "text": "@" + e["domain"]})
    n = len(lines)
    for i, raw in enumerate(lines):
        line = URL_RX.sub(" ", EMAIL_RX.sub(" ", raw))
        low = line.lower()
        if not low.strip() or RECIPIENT_RX.search(line):
            continue
        taken = []
        hits = []
        # v0101: ONE longest-first competition between denied names and legitimate aliases (the mother's 批681
        # rule).  A denied span closes to every shorter alias ('China International Capital' never surfaces as
        # the generic 'capital' of 群益; 中信 inside 中信證券 is never CTBC), while a longer legitimate alias keeps
        # its span (摩根士丹利 is MS although the bare 摩根 is denied).  Equal length: the deny list wins.
        cands = []
        for phrase in deny_list:
            for s0, t0 in _alias_spans(str(phrase).lower(), low):
                cands.append((-(t0 - s0), 0, s0, t0, str(phrase), None))
        for alias, canon in flat:
            for s, t in _alias_spans(alias, low):
                cands.append((-(t - s), 1, s, t, alias, canon))
        for _neg, kind, s, t, word, canon in sorted(cands):
            if any(not (t <= a or s >= b) for a, b in taken):
                continue
            taken.append((s, t))
            if kind == 0:
                denied[word] += 1
                evidence.append({"broker": None, "tier": "DENIED", "line": i, "text": line[max(0, s - 12):t + 24].strip()})
            else:
                hits.append((s, t, word, canon))
        hits.sort()
        if not hits:
            continue
        # v0102: disclosure_rx is passed only by appendix_evidence(); page 1 keeps the v0101 test exactly
        disclosure = bool(DISCLOSURE_RX.search(line)) or bool(disclosure_rx is not None and disclosure_rx.search(line))
        copyright_ = bool(COPYRIGHT_RX.search(line))
        source = bool(SOURCE_LINE_RX.search(line))
        zone = i < HEADER_LINES or i >= n - FOOTER_LINES
        for s, t, alias, canon in hits:
            if _adjacent_to_code(low, s, t):
                continue
            if denied_alias(alias) or denied_key(canon):
                denied[canon] += 1          # 批679 deny list runs first: evidence kept, never a broker
                evidence.append({"broker": canon, "tier": "DENIED", "line": i, "text": line[max(0, s - 12):t + 24].strip()})
                continue
            issuer = bool(ISSUER_WORD_RX.search(alias)) or bool(ISSUER_AFTER_RX.match(low[t:t + 40]))
            is_generic = alias in generic
            if disclosure or copyright_:
                tier = "DISCLOSURE"
            elif issuer and not source and (issuer_guard is None or issuer_guard(line)):
                tier = "ISSUER"
            elif is_generic:
                continue
            elif source:
                tier = "HEADER"
            elif zone:
                tier = "HEADER"
            else:
                tier = "BODY"
            tiers[canon][tier] += 1
            evidence.append({"broker": canon, "tier": tier, "line": i, "text": line[max(0, s - 12):t + 24].strip()})
    order = ("EMAIL", "DISCLOSURE", "ISSUER", "HEADER", "BODY")

    def rank(b):
        t = tiers[b]
        best = next((k for k, name in enumerate(order) if t.get(name)), len(order))
        return (best, -sum(t.values()), 0 if b == filename_broker else 1, b)

    denied_out = dict(denied)
    if not tiers:
        return {"broker": None, "tier": "NONE", "strong": False, "candidates": {}, "evidence": evidence[:40],
                "denied": denied_out}
    ranked = sorted(tiers, key=rank)
    top = ranked[0]
    top_tier = order[rank(top)[0]]
    # the filename broker keeps the page when it has page evidence of the same best tier
    if filename_broker and filename_broker in tiers and filename_broker != top:
        if order[rank(filename_broker)[0]] == top_tier:
            top = filename_broker
    return {"broker": top, "tier": top_tier, "strong": top_tier in ("EMAIL", "DISCLOSURE", "ISSUER"),
            "candidates": {b: dict(tiers[b]) for b in ranked}, "evidence": evidence[:40], "denied": denied_out}


def compare_broker(filename_broker, page):
    """PASS / MISMATCH / INSUFFICIENT_EVIDENCE / N/A.  Only strong page evidence can contradict the name."""
    pb = (page or {}).get("broker")
    cands = (page or {}).get("candidates") or {}
    purge = vcgc("purge")
    try:
        if filename_broker and ((purge and purge.is_cn_canon(filename_broker)) or is_denied_token(filename_broker)):
            return "DENIED"
    except Exception:
        pass
    if not filename_broker and not pb:
        return "N/A"
    if not filename_broker or not pb:
        return "INSUFFICIENT_EVIDENCE"
    if filename_broker == pb:
        return "PASS"
    if filename_broker in cands:
        return "PASS" if not (page or {}).get("strong") else "AMBIGUOUS"
    return "MISMATCH" if (page or {}).get("strong") else "INSUFFICIENT_EVIDENCE"


# =====================================================================
# APPENDIX SMALL PRINT (v0102, mother 批729)
# =====================================================================
# The last pages of a sell-side report carry the issuer's disclosures and the rating definitions in small print
# (operator 2026-09-24).  Page 1 is never re-read here; the appendix only answers what page 1 could not.
APPENDIX_PAGES = 3          # the last N pages (page 1 excluded) are read as the appendix
SMALL_RATIO = 0.9           # small print = glyph size below SMALL_RATIO x the page-1 body size
# Chinese / English disclosure forms, used ONLY for appendix lines (page-1 grading keeps DISCLOSURE_RX as it was)
APPENDIX_DISCLOSURE_RX = re.compile(
    r"(版權所有|著作權所有|本(?:研究)?報告(?:係|是)?由|免責聲明|重要聲明|揭露事項|利益衝突|分析師聲明|"
    r"analyst\s+certification|important\s+disclosures?|disclaimer|all\s+rights\s+reserved|"
    r"this\s+(?:research\s+)?report\s+(?:is|was|has\s+been)\s+(?:prepared|issued|published|produced)\s+by)", re.I)
# a rating definition says what the word expects: return, relative performance, benchmark, a percentage
RATING_DEF_SIGNAL_RX = re.compile(
    r"(預期|報酬|表現|漲幅|跌幅|超越|落後|優於|劣於|大盤|指數|基準|%|％|expect|return|outperform|underperform|"
    r"benchmark|relative|upside|downside|appreciat|depreciat)", re.I)
# 批730 (measured on 21 realistic appendix fixtures): a CJK word may carry an inner hyphen or middle dot
# (元大 持有-超越同業); an English alt may read "O or Over" (Morgan Stanley); an ASCII hyphen separates only after a
# space or a closing bracket, so Equal-weight is never cut into "Equal" + "weight (E or Equal) - ..."
# 批730 review: each part of a CJK word stays 2-6 characters (a 10-character run let headings through:
# 本報告投資評等之定義:...); an em dash always separates (Buy—expected ...), an en dash when a space touches it
_CJK_WORD = r"[\u4e00-\u9fff]{2,6}(?:[\-‐–・·][\u4e00-\u9fff]{2,6})?"
_RS_CJK = re.compile(r"^\s*(?P<word>" + _CJK_WORD + r")\s*(?:[(（]\s*(?P<alt>[A-Za-z][A-Za-z\- /]{0,24}|[\u4e00-\u9fff]{2,6})\s*[)）])?"
                     r"\s*[:：]\s*(?P<def>.+)$")
_RS_LAT = re.compile(r"^\s*(?P<word>[A-Za-z][A-Za-z\-]{0,20}(?:\s[A-Za-z\-]{1,12}){0,2})\s*"
                     r"(?:[(（]\s*(?P<alt>[A-Za-z0-9+\-]{1,10}(?:\s+or\s+[A-Za-z0-9+\-]{1,10})?|[\u4e00-\u9fff]{2,6})\s*[)）])?"
                     r"(?:\s*[:：]|\s*—|\s*–(?=\s)|\s+–|(?:\s+|(?<=[)）])\s*)-)\s*(?P<def>.+)$")
# a table row without a colon ("買進  預期未來12個月報酬率大於15%"): taken when the word is a known rating word,
# or when the row sits next to an accepted definition row (the same table)
_RS_BARE = re.compile(r"^\s*(?P<word>" + _CJK_WORD + r")\s+(?P<def>\S.+)$")
_RS_ROW = re.compile(r"^\s*(?P<word>[\u4e00-\u9fff]{2,6})\s+(?P<alt>[A-Za-z][A-Za-z\-]{1,20}(?:\s[A-Za-z\-]{1,12})?)\s+(?P<def>.+)$")
# headers and field labels that look like a definition row but are not a rating word
_RS_NOT_A_RATING = re.compile(
    r"^(?:投資)?評[等級](?:定義|說明|標準|制度|方式)?$|^(?:stock\s+|investment\s+)?ratings?(?:\s+(?:definitions?|system|scale|key|guide))?$|"
    r"^(?:註|說明|備註|資料來源|來源|note|notes|source|sources|disclaimer|analyst|分析師|定義|definition)$|"
    r"target|price|eps|目標|股價|價格|營收|盈餘|本益比|殖利率|yield|upside|downside|報酬率|"
    # 批730: section labels that also carry return words (Valuation: ... implying 15% upside)
    r"valuation|估值|評價|^risks?$|風險|methodology|方法|catalysts?|催化|investment\s+thesis|投資論點|"
    # 批730 review: table headers and footnotes next to a rating table (投資建議 未來12個月預期報酬率 · 評等類別 ... ·
    # 過去績效 不代表未來表現 · 上述評等 ... · 本公司 ... · 投資評等之定義如下)
    r"定義|如下|投資建議|類別|績效|上述|本公司|本報告|投資人", re.I)
# 批730 review: a colon-less row taken only for sitting next to an accepted row must read like a definition (a threshold
# or a comparison); a header or a footnote next to the table does not
_RS_THRESHOLD_RX = re.compile(
    r"(\d+(?:\.\d+)?\s*[%％]|大於|小於|超過|低於|高於|介於|優於|劣於|落後|相當|以上|以下|領先|"
    r"more\s+than|less\s+than|above|below|exceed|between|within|in\s+line|outperform|underperform)", re.I)
# 批730: an ISSUER line in the appendix names the company; a peer-table row names a broker next to a rating word or a
# target price (同業評等彙整:富邦證券 中立 目標價1000元) and is not the issuer unless it carries the legal name
_APPENDIX_TABLE_ROW_RX = re.compile(
    r"(目標價|target\s*price|\bTP\b|NT\$|\d[\d,]*(?:\.\d+)?\s*元(?![\u4e00-\u9fff])|買進|賣出|中立|持有|增加持股|減少持股|降低持股|區間操作|"
    r"(?<![A-Za-z])(?:buy|sell|hold|neutral|outperform|underperform|overweight|underweight)(?![A-Za-z]))", re.I)
_LEGAL_ENTITY_RX = re.compile(r"(股份有限公司|有限公司|co\.?,?\s*ltd|limited|l\.?\s?l\.?\s?c\b|\binc\b|pte\.?\s*ltd|\bplc\b)", re.I)


def _appendix_issuer_line(line):
    """批730: the appendix-only guard on the ISSUER tier (page-1 grading does not pass it)."""
    return bool(_LEGAL_ENTITY_RX.search(line)) or not _APPENDIX_TABLE_ROW_RX.search(line)


def pdf_page_count(path):
    """Number of pages of a PDF; 0 when no PDF library is installed or the file cannot be opened."""
    key = _file_key(path)
    if key is not None and key in _PAGE_COUNT_CACHE:
        return _PAGE_COUNT_CACHE[key]
    n = _pdf_page_count_open(path)
    if key is not None and n:
        _PAGE_COUNT_CACHE[key] = n
    return n


def _pdf_page_count_open(path):
    """The uncached counter behind pdf_page_count()."""
    try:
        if importlib.util.find_spec("pdfplumber") is not None:
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                return len(pdf.pages)
        if importlib.util.find_spec("fitz") is not None:
            import fitz
            doc = fitz.open(path)
            try:
                return len(doc)
            finally:
                doc.close()
    except Exception:
        return 0
    return 0


def _body_size(lines):
    """Page-1 body size: the glyph size carrying the most characters (rounded to 0.5 pt); None without sizes."""
    weight = defaultdict(int)
    for l in lines or []:
        size = float(l.get("size") or 0.0)
        if size > 0:
            weight[round(size * 2) / 2] += len(str(l.get("text", "")).strip())
    return max(weight, key=lambda k: (weight[k], k)) if weight else None


def appendix_lines(path, last_pages=APPENDIX_PAGES, small_only=True):
    """Small-print lines of the report's last pages (page 1 excluded).
    Returns (lines, info): lines = [{"page", "text", "size"}]; info says which pages were read, the page-1 body size
    and the small-print threshold.  Not a PDF / one page only / no text layer = ([], info with the honest reason).
    批730 speed: read once per file version (the batch, the first-page gate and the idempotency pass all ask)."""
    key = _file_key(path)
    k = key + (int(last_pages), bool(small_only)) if key is not None else None
    if k is not None and k in _APPENDIX_CACHE:
        _APPENDIX_CACHE.move_to_end(k)
        return copy.deepcopy(_APPENDIX_CACHE[k])
    got = _appendix_lines_read(path, last_pages, small_only)
    if k is not None:
        _APPENDIX_CACHE[k] = copy.deepcopy(got)
        while len(_APPENDIX_CACHE) > _APPENDIX_CACHE_MAX:
            _APPENDIX_CACHE.popitem(last=False)
    return got


def _appendix_lines_read(path, last_pages=APPENDIX_PAGES, small_only=True):
    """The uncached reader behind appendix_lines()."""
    info = {"state": "NODATA", "pages": [], "n_pages": 0, "body_size": None, "threshold": None, "why": ""}
    if not str(path).lower().endswith(".pdf"):
        info["why"] = "appendix is read from PDF only"
        return [], info
    n = pdf_page_count(path)
    info["n_pages"] = n
    if n < 2:
        info["why"] = "one page: no appendix" if n == 1 else "PDF unreadable or no PDF library"
        return [], info
    first = pdf_page_chars(path, 0)
    body = _body_size(lines_from_chars(first[0])) if first else None
    pages = list(range(max(1, n - int(last_pages)), n))
    info["pages"] = [i + 1 for i in pages]
    rows = []
    for i in pages:
        got = pdf_page_chars(path, i)
        if not got:
            continue
        rows += [(i + 1, l) for l in lines_from_chars(got[0])]
    if body is None:                         # page 1 without a text layer: fall back to the appendix pages themselves
        body = _body_size([l for _, l in rows])
    thr = round(body * SMALL_RATIO, 2) if body else None
    info.update({"body_size": body, "threshold": thr})
    out = []
    for page, l in rows:
        text = str(l.get("text", "")).strip()
        if not text:
            continue
        size = float(l.get("size") or 0.0)
        if small_only and (thr is None or size <= 0 or size >= thr):
            continue
        out.append({"page": page, "text": text, "size": size})
    info["state"] = "OK" if out else "EMPTY"
    if not out:
        info["why"] = "no small print on the last pages" if rows else "no text layer on the last pages"
    return out, info


def _rating_word_is_label(word):
    return bool(_RS_NOT_A_RATING.search(str(word or "").strip()))


def rating_scale(lines, rules=None, words=None):
    """Rating definitions printed in the appendix -> candidates [{word, alt, word_known, alt_known, canonical,
    suggested, definition, page, line}].  KNOWN / NEW is judged against the merged rating words (the same words the
    engines use); a NEW word is suggested a key ONLY through its KNOWN pair (買進(Buy) · Outperform (增加持股)).
    Candidates only: this function never writes a book."""
    words = words or merged_rating_words(rules if rules is not None else load_rules())
    out, seen = [], set()
    texts = [(item.get("text", "") if isinstance(item, dict) else str(item)) for item in (lines or [])]
    matches = {}
    for idx, text in enumerate(texts):
        for rx in (_RS_CJK, _RS_LAT, _RS_ROW):
            m = rx.match(text)
            if m and RATING_DEF_SIGNAL_RX.search(m.group("def")) and not _rating_word_is_label(m.group("word")):
                matches[idx] = m
                break
    # 批730: colon-less table rows, anchored on a known rating word, then on a neighbour row of the same table
    bare = {}
    for idx, text in enumerate(texts):
        if idx in matches:
            continue
        m = _RS_BARE.match(text)
        if m and RATING_DEF_SIGNAL_RX.search(m.group("def")) and not _rating_word_is_label(m.group("word")):
            bare[idx] = m
    for idx, m in list(bare.items()):
        if normalize_rating_word(m.group("word"), words)[0] is not None:
            matches[idx] = bare.pop(idx)
    grew = True
    while grew:                              # 批730 review: a neighbour must also read like a definition row
        grew = False
        for idx, m in list(bare.items()):
            if ((idx - 1) in matches or (idx + 1) in matches) and _RS_THRESHOLD_RX.search(m.group("def")):
                matches[idx] = bare.pop(idx)
                grew = True
    for idx in sorted(matches):
        item = (lines or [])[idx]
        page = item.get("page") if isinstance(item, dict) else None
        m = matches[idx]
        word = re.sub(r"\s+", " ", m.group("word")).strip()
        alt = re.sub(r"\s+", " ", ((m.groupdict().get("alt") or ""))).strip() or None
        definition = m.group("def").strip()
        if not RATING_DEF_SIGNAL_RX.search(definition) or _rating_word_is_label(word) or (alt and _rating_word_is_label(alt)):
            continue
        key = (word.lower(), (alt or "").lower())
        if key in seen:
            continue
        seen.add(key)
        wc, _ = normalize_rating_word(word, words)
        ac, _ = normalize_rating_word(alt, words) if alt else (None, None)
        suggested = None
        if wc is None and ac is not None:
            suggested = ac
        elif ac is None and wc is not None and alt:
            suggested = wc
        out.append({"word": word, "alt": alt, "word_known": wc is not None, "alt_known": ac is not None,
                    "canonical": wc or ac, "suggested": suggested, "definition": definition[:160],
                    "page": page, "line": idx})
    return out


def appendix_evidence(path, alias_table=None, filename_broker=None, rules=None, last_pages=APPENDIX_PAGES):
    """Issuer and rating scale from the appendix small print.
    broker = broker_evidence() on the small-print lines with the SAME deny gate; the appendix-only disclosure forms
    make '本報告由凱基證券投資顧問…發佈' / '© … 版權所有' strong.  vs_filename = compare_broker() on that result."""
    lines, info = appendix_lines(path, last_pages)
    texts = [l["text"] for l in lines]
    if texts:
        ev = broker_evidence(texts, alias_table or {}, filename_broker=filename_broker, rules=rules,
                             disclosure_rx=APPENDIX_DISCLOSURE_RX, issuer_guard=_appendix_issuer_line)
        for e in ev.get("evidence", []):
            li = e.get("line")
            if isinstance(li, int) and 0 <= li < len(lines):
                e["page"] = lines[li]["page"]
    else:
        ev = {"broker": None, "tier": "NONE", "strong": False, "candidates": {}, "evidence": [], "denied": {}}
    return {"state": info["state"], "info": info, "broker": ev,
            "vs_filename": compare_broker(filename_broker, ev) if filename_broker else "N/A",
            "rating_scale": rating_scale(lines, rules=rules), "n_lines": len(lines)}


def ssot_alias_table(start=None):
    """{canonical: [names]} from the institution SSOT (for callers without an engine dictionary, e.g. the CLI)."""
    inst = institution_ssot(start) or {}
    regs = inst.get("registries") or {}
    out = defaultdict(set)
    for grp in ("domestic_brokers", "foreign_brokers"):
        for key, spec in (regs.get(grp) or {}).items():
            spec = spec if isinstance(spec, dict) else {}
            for name in [key, spec.get("english_name"), spec.get("chinese_name")] + list(spec.get("aliases") or []):
                if name and len(str(name).strip()) >= 2:
                    out[str(key).upper()].add(str(name).strip())
    return {k: sorted(v) for k, v in out.items()}


# =====================================================================
# 4c . ADJ CLOSE BASIS (mother 批729; operator 2026-09-24: "所有目標價及各前一日的價格都要換成 ADJ CLOSE,
#      上漲空間都要用最新的 ADJ CLOSE")
# =====================================================================
# The formula is the mother's L99 (批659 operator ruling) and it lives in ONE place, the ENG073 tail's
# adj_quote(): factor = adj_close / close on the trading day before the report date, target x factor,
# upside over the latest adj_close.  This section only finds that tail and the market database and passes the
# answer on (LL404: no second head).  Read-only: the database is opened read_only and nothing is written.
ADJ_ENGINE_GLOB = ("functional modules/VRN", "VRN_ENG073_ReportStructuredDB_v*.py")
ADJ_MISS_STATES = ("ADJ_NO_ENGINE", "ADJ_NO_DB", "ADJ_NO_DUCKDB", "ADJ_DB_BUSY", "ADJ_ERROR")
_ADJ_CACHE = {}


def adj_engine(start=None):
    """(module | None, why): the newest ENG073 that carries adj_quote() (v0137+); loaded once per tree."""
    root = repo_root(start)
    key = ("engine", root)
    if key in _ADJ_CACHE:
        return _ADJ_CACHE[key]
    folder, pattern = ADJ_ENGINE_GLOB
    hits = sorted(glob.glob(os.path.join(root, folder, pattern)), key=_version_of)
    mod, why = None, "no %s under %s" % (pattern, folder)
    if hits:
        name = os.path.basename(hits[-1])
        try:
            spec = importlib.util.spec_from_file_location("_vrn_adj_eng073", hits[-1])
            mod = importlib.util.module_from_spec(spec)
            sys.modules[spec.name] = mod
            spec.loader.exec_module(mod)
            why = name
            if not hasattr(mod, "adj_quote"):
                mod, why = None, "%s has no adj_quote() (v0137+)" % name
        except Exception as exc:
            sys.modules.pop("_vrn_adj_eng073", None)
            mod, why = None, "%s failed to load: %s: %s" % (name, type(exc).__name__, str(exc)[:80])
    _ADJ_CACHE[key] = (mod, why)
    return mod, why


def adj_db(db=None, start=None):
    """(path | None, why).  An explicit db wins; otherwise the ENG073 resolver (VIA_DB_VDF_TW_MARKET >
    VIA_DATA_HOME > the mega path > a tree search), resolved once per tree and environment; its console lines are
    kept in `why` instead of printed."""
    if db:
        return (str(db), "explicit") if os.path.isfile(str(db)) else (None, "explicit db missing: %s" % db)
    mod, why = adj_engine(start)
    if mod is None or not hasattr(mod, "_resolve_db"):
        return None, why if mod is None else "the engine has no database resolver"
    key = ("db", repo_root(start), os.environ.get("VIA_DB_VDF_TW_MARKET"), os.environ.get("VIA_DATA_HOME"))
    if key in _ADJ_CACHE:
        return _ADJ_CACHE[key]
    import contextlib
    import io
    buf = io.StringIO()
    try:
        with contextlib.redirect_stdout(buf):
            path = mod._resolve_db(None)
    except Exception as exc:
        path = None
        buf.write("%s: %s" % (type(exc).__name__, exc))
    said = " ".join(buf.getvalue().split())[:240]
    _ADJ_CACHE[key] = (str(path), said or "resolved") if path else (None, said or "not found")
    return _ADJ_CACHE[key]


def _adj_ticker(ticker):
    """'2330' / '2330.TW' / '6147.TWO' / '2330 TT' / '00878' -> the bare code; None without one."""
    m = re.search(r"(?<![0-9A-Za-z])(\d{4,6}[A-Z]?)(?=$|[^0-9A-Za-z]|\.TWO?\b|TT\b)", str(ticker or "").strip())
    return m.group(1) if m else None


def _adj_date(value):
    """'2026-05-19' / '2026/5/19' / '20260519' / '2026年5月19日' -> '2026-05-19'; None when it is not a date."""
    s = str(value or "").strip()
    m = re.fullmatch(r"((?:19|20)\d{2})(\d{2})(\d{2})", s) or \
        re.match(r"((?:19|20)\d{2})\s*[-/.年]\s*(\d{1,2})\s*[-/.月]\s*(\d{1,2})", s)
    if not m:
        return None
    try:
        return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
    except ValueError:
        return None


def adj_basis(ticker, report_date, target_price, db=None, start=None, page_price=None):
    """Target price and the day-before price in ADJ CLOSE terms; upside over the latest ADJ CLOSE.
    Delegates to ENG073 adj_quote() and returns its fields (adj_factor, target_price_adj, price_prev_close /
    price_prev_adj / price_prev_date, price_latest_adj / price_latest_date, upside_adj, upside_adj_state, ...).
    state = the engine's upside_adj_state (ADJ_OK / ADJ_NO_KEY / ADJ_NO_TARGET / ADJ_NO_FACTOR / ADJ_NO_LATEST /
    ADJ_ABSURD(...) / ADJ_EVENT_UNADJUSTED(...)), or one of ADJ_MISS_STATES when the answer could not be asked for;
    never a made-up 1.0.  page_price = the price printed on the report page: ENG073 v0138+ divides the adj close by the
    exchange close of the day before the report, and by this page price when the exchange has no close for those days
    (adj_factor_basis says which)."""
    tk, rd = _adj_ticker(ticker), _adj_date(report_date)
    try:
        tp = float(target_price) if target_price not in (None, "") else None
    except (TypeError, ValueError):
        tp = None
    out = {"state": None, "why": "", "engine": None, "db": None, "ticker": tk, "report_date": rd, "target_price": tp}
    mod, why = adj_engine(start)
    if mod is None:
        out.update(state="ADJ_NO_ENGINE", why=why)
        return out
    out["engine"] = why
    path, why = adj_db(db, start)
    out["db"] = path
    if not path:
        out.update(state="ADJ_NO_DB", why=why)
        return out
    try:
        import duckdb
    except ImportError:
        out.update(state="ADJ_NO_DUCKDB", why="duckdb is not installed")
        return out
    try:
        con = duckdb.connect(path, read_only=True)
    except Exception as exc:          # a writer holds the file, or this process opened it read-write
        out.update(state="ADJ_DB_BUSY", why="%s: %s" % (type(exc).__name__, str(exc)[:160]))
        return out
    try:
        try:
            q = mod.adj_quote(con, tk, rd, tp, page_price=page_price)
        except TypeError:             # an ENG073 before v0138 has no page_price
            q = mod.adj_quote(con, tk, rd, tp)
    except Exception as exc:
        out.update(state="ADJ_ERROR", why="%s: %s" % (type(exc).__name__, str(exc)[:160]))
        return out
    finally:
        con.close()
    out.update(q)
    out["state"] = q.get("upside_adj_state")
    out["why"] = q.get("adj_block_reason") or ""
    return out


# =====================================================================
# 5 . STOCK CODES ON THE PAGE, REPORT TYPE, PRIMARY TICKER, TRI-CODE
# =====================================================================
_YF_RX = re.compile(r"(?<![0-9A-Za-z])([1-9]\d{3}|00\d{2,4}[A-Z]?)\.(TWO|TW)(?![A-Za-z])")
_BB_RX = re.compile(r"(?<![0-9A-Za-z])([1-9]\d{3}|00\d{2,4}[A-Z]?)\s?TT(?![A-Za-z])")
_PAREN_RX = re.compile(r"[（(]\s*([1-9]\d{3})(?:\s*(?:TT|\.TWO?|-KY))?\s*(?:[,，][^)）]{0,24})?[)）]")
_LABEL_CODE_RX = re.compile(r"(?:股票代號|證券代號|代號|代碼)\s*[:：]?\s*([1-9]\d{3})(?![0-9])")
_NAME_BEFORE_RX = re.compile(r"([一-鿿A-Za-z][一-鿿A-Za-z0-9&.'’\- ]{0,40}?)\s*[（(]?\s*$")


def _year_like(code):
    return code.isdigit() and 2019 <= int(code) <= 2035


def page_codes(lines):
    """Every Taiwan code form on the page: yfinance 2330.TW, Bloomberg 2330 TT, 南亞(1303), 代號 1303."""
    out = []
    for i, line in enumerate(lines):
        spans = []
        for kind, rx in (("yfinance", _YF_RX), ("bloomberg", _BB_RX), ("paren", _PAREN_RX), ("label", _LABEL_CODE_RX)):
            for m in rx.finditer(line):
                code = m.group(1)
                if any(not (m.end() <= a or m.start() >= b) for a, b in spans) and kind in ("paren", "label"):
                    continue
                if kind in ("paren", "label") and _year_like(code):
                    continue
                before = line[:m.start()]
                name = ""
                nm = _NAME_BEFORE_RX.search(before)
                if nm:
                    name = nm.group(1).strip()
                spans.append((m.start(), m.end()))
                out.append({"core": code, "kind": kind, "raw": m.group(0), "line": i, "name_before": name,
                            "market": ("TPEX" if kind == "yfinance" and m.group(2).upper() == "TWO" else
                                       ("TWSE" if kind == "yfinance" else ""))})
    return out


def detect_report_type(filename, lines, filename_ticker=None, codes=None, signals=None):
    """STOCK / INDUSTRY / MARKET / MACRO / ETF / EVENT / OTHER.
    A ticker in the filename means STOCK; otherwise filename cues, then single-company page signals."""
    suffix = os.path.splitext(filename or "")[1].lower()
    if suffix in (".jpg", ".jpeg", ".png", ".tif", ".tiff"):
        return {"type": "OTHER", "why": "IMAGE_FILE"}
    if filename_ticker:
        return {"type": "STOCK", "why": "FILENAME_TICKER"}
    stem = os.path.splitext(os.path.basename(filename or ""))[0]
    signals = signals or {}
    page_stock = bool((signals.get("labeled_tp") and signals.get("labeled_cp")) or
                      (signals.get("labeled_rating") and (signals.get("labeled_tp") or signals.get("labeled_cp"))))
    for kind, rx in TYPE_FILENAME_CUES:
        if rx.search(stem):
            if kind == "INDUSTRY" and page_stock and signals.get("single_company"):
                return {"type": "STOCK", "why": "PAGE_SINGLE_COMPANY_OVER_" + kind}
            return {"type": kind, "why": "FILENAME_CUE"}
    if page_stock and signals.get("single_company"):
        return {"type": "STOCK", "why": "PAGE_SINGLE_COMPANY"}
    if STOCK_FILENAME_CUE.search(stem) and page_stock:
        return {"type": "STOCK", "why": "FILENAME_CUE+PAGE"}
    return {"type": "INDUSTRY", "why": "NO_SINGLE_COMPANY_EVIDENCE"}


def primary_ticker(codes, filename_ticker=None, report_type="STOCK"):
    """The covered company's code.  Filename ticker wins when the page shows it; otherwise the first
    header code of a single-company report.  Non-stock reports have none (peers stay in `others`)."""
    cores = []
    for c in codes or []:
        if c["core"] not in cores:
            cores.append(c["core"])
    if report_type not in STOCK_TYPES:
        return {"ticker": None, "source": "NOT_APPLICABLE", "others": cores}
    if filename_ticker:
        src = "FILENAME+PAGE" if filename_ticker in cores else "FILENAME"
        return {"ticker": filename_ticker, "source": src, "others": [c for c in cores if c != filename_ticker]}
    header = [c for c in codes or [] if c["line"] < HEADER_LINES]
    if header:
        return {"ticker": header[0]["core"], "source": "PAGE_HEADER", "others": [c for c in cores if c != header[0]["core"]]}
    if codes:
        return {"ticker": codes[0]["core"], "source": "PAGE_FIRST_CODE", "others": [c for c in cores if c != codes[0]["core"]]}
    return {"ticker": None, "source": "NONE", "others": cores}


def tri_code(filename_ticker, codes, primary=None, official_market=None):
    """Spec section 2 verdict.  PASS when the filename core is among the page platform codes (peers may
    also be printed); MISMATCH only when the page's own primary code differs and the filename code is absent."""
    cores = sorted({c["core"] for c in codes or []})
    platform = [c for c in codes or [] if c["kind"] in ("yfinance", "bloomberg", "paren", "label")]
    if not filename_ticker and not platform:
        return {"verdict": "N/A", "cores": cores}
    if not filename_ticker or not platform:
        return {"verdict": "INSUFFICIENT_EVIDENCE", "cores": cores, "filename": filename_ticker or "",
                "page_primary": primary or ""}
    if filename_ticker in cores:
        verdict = "PASS"
        yf = next((c for c in platform if c["core"] == filename_ticker and c["kind"] == "yfinance"), None)
        if official_market and yf and yf["market"] and yf["market"] != str(official_market).upper():
            verdict = "MARKET_MISMATCH"
        return {"verdict": verdict, "cores": cores, "filename": filename_ticker}
    header = [c for c in platform if c["line"] < HEADER_LINES]
    if header:
        return {"verdict": "MISMATCH", "cores": cores, "filename": filename_ticker, "page_primary": header[0]["core"]}
    return {"verdict": "INSUFFICIENT_EVIDENCE", "cores": cores, "filename": filename_ticker}


# =====================================================================
# 6 . RATING (label > box/title > quoted > prose > filename slot; previous values never count)
# =====================================================================
_PREV_LINE_RX = re.compile(r"前次投資建議\s*[:：]?\s*[\d./\-]*\s*$|前次評等\s*[:：]?\s*$")
_EXCLUDE_RATING_LINE_RX = re.compile(r"評等分級|評等定義|評等說明|rating definition|rating system|ratings? distribution|"
                                     r"industry view|前次評等|前一次評等|previous rating|prior rating", re.I)


def _rating_patterns(words):
    zh = [w for w in words if re.search(r"[一-鿿]", w)]
    en = [w for w in words if not re.search(r"[一-鿿]", w)]
    ZH = "(" + _alternation(zh) + ")"
    EN = "(" + _alternation(en) + r")(?:\s*\((?:\d|unchanged|on\s*CL)\))?"
    ANY = "(" + _alternation(zh + en) + ")"
    act = r"(?:維持|調升|調降|上調|下調|重啟評等|初次評等|首次評等|重新覆蓋|新增)"
    return [
        # priority 0: labelled fields
        (0, "LABEL_EN", re.compile(r"(?<![A-Za-z])(?:stock\s*rating|12[- ]?month\s*rating|investment\s*rating|rating|"
                                   r"recommendation|citi[’']?s\s*take)\s*[:：]?\s*" + EN + r"(?![A-Za-z])", re.I)),
        (0, "LABEL_ZH", re.compile(r"(?<!前次)(?<!前一次)(?<![原舊])(?:投資評等|投資評級|投資建議|評等|評級)\s*[:：]?\s*"
                                   r"(?:維持|為|建議|調升至|調降至|調升為|調降為|調整為|重啟|初次)?\s*[「『“\"]?\s*" + ZH)),
        # priority 1: rating boxes and title lines
        (1, "BOX_ZH", re.compile(r"^\s*" + ZH + r"\s*(?:[‧・·•.\-–—]\s*" + act + r")?\s*$")),
        (1, "BOX_EN", re.compile(r"^\s*" + EN + r"\s*\.?\s*$", re.I)),
        (1, "TITLE_CODE", re.compile(r"[（(]\s*[1-9]\d{3}(?:\s*TT|\.TWO?)?\s*[)）]\s*[-–—:：]?\s*[「]?" + ANY, re.I)),
        (1, "PAREN_SLOT", re.compile(r"[（(]\s*[1-9]\d{3}\s*TT\s*[,，]\s*(?:NT\$|\$)\s*[0-9][0-9,]*(?:\.\d+)?\s*[,，]\s*" + ANY + r"\s*[)）]", re.I)),
        (1, "PRICE_DASH", re.compile(r"(?:NT\$|TWD|US\$|HK\$)\s*[0-9][0-9,]*(?:\.\d+)?\s*[-–]\s*(?:high-conviction\s+)?" + EN, re.I)),
        (1, "UPDOWNSIDE", re.compile(r"up/downside\s*:?\s*[+\-−–]?\s*[0-9.]+\s*%\s*" + EN, re.I)),
        (1, "MAINTAINED_EN", re.compile(r"^\s*(?:maintained|upgraded|downgraded|initiated|reiterated)\s+(?:to\s+)?" + EN + r"\s*$", re.I)),
        (1, "COMPANY_EN", re.compile(r"^[A-Z][\w&.,'’\- ]{1,60}?\s+(overweight|neutral|underweight|equal-weight|outperform|underperform|buy|sell|hold)\s*$", re.I)),
        (2, "TITLE_SEMI", re.compile(r"[;；]\s*(?:(?:reiterate|maintain|upgrade\s+to|downgrade\s+to|initiate\s+(?:at|with)|initiating\s+(?:at|with))\s+)?"
                                     + EN + r"\s*$", re.I)),
        # priority 3: quoted Chinese ratings with a verb or a trailing 評等 (mother rating.quoted_rx)
        (3, "QUOTED_ZH", re.compile(r"(?:維持|調整為|調升為|調降為|調升至|調降至|給予|給與|評等為|建議)[一-鿿\-A-Za-z]{0,6}?\s*[「『“\"]\s*" + ZH
                                    + r"\s*[」』”\"]")),
        (3, "QUOTED_ZH_TAIL", re.compile(r"[「『“\"]\s*" + ZH + r"\s*[」』”\"]\s*(?:之)?(?:投資評等|評等|評級)")),
        # priority 4: English prose (MS 'keep OW', GS 'Maintain Buy', Daiwa 'Reaffirming our Buy (1) call')
        (4, "PROSE_EN", re.compile(r"(?<![A-Za-z])(?:maintain(?:ed|ing)?|reiterat(?:e|ed|ing)|reaffirm(?:ed|ing)?|keep(?:ing)?|"
                                   r"initiat(?:e|ed|ing)\s+(?:coverage\s+)?(?:with|at)|upgrad(?:e|ed|ing)(?:\s+[A-Z][\w\-]*){0,3}\s+to|"
                                   r"downgrad(?:e|ed|ing)(?:\s+[A-Z][\w\-]*){0,3}\s+to)\s+(?:our\s+|an?\s+|the\s+|at\s+)?"
                                   r"(" + _alternation(en) + r"|OW|EW|UW)(?![A-Za-z])", re.I)),
    ]


def filename_rating_slot(filename, rules=None, words=None):
    """Mother 批713 filename_rating_slot: 瑞基(4171,NR_未評等) / 晶心科(6533,N,中立) / 光焱(7728,Note).
    'Note' is a report type, not a rating -> UNRATED_SLOT with the raw slot kept."""
    rx = _rule(rules or {}, "report_field_rules", "filename_rating_slot", "rx",
               default=r"[（(]\s*(?P<code>\d{4,6}[A-Za-z]{0,2})\s*[,，]\s*(?P<slot>[^)）]{1,24})\s*[)）]")
    split = _rule(rules or {}, "report_field_rules", "filename_rating_slot", "slot_split_rx", default=r"[_＿/／,，]")
    m = re.search(rx, filename or "")
    if not m:
        return None
    slot = m.group("slot")
    for token in [t.strip() for t in re.split(split, slot) if t.strip()]:
        canon, fine = normalize_rating_word(token, words)
        if canon:
            return {"value": canon, "fine": fine, "raw": token, "slot": slot, "code": m.group("code"), "state": "HIT"}
    return {"value": None, "fine": None, "raw": slot, "slot": slot, "code": m.group("code"), "state": "UNRATED_SLOT"}


def rating_action(text):
    low = (text or "").lower()
    for action, terms in RATING_ACTION_WORDS.items():
        if any(t in low for t in terms):
            return action
    return ""


def extract_rating(lines, filename="", rules=None, words=None):
    """The report's own current rating.  Returns value / fine / raw / source / line / action / state and
    every candidate seen (append-only evidence)."""
    words = words or merged_rating_words(rules)
    pats = _rating_patterns(words)
    cands = []
    prev_ctx = False
    for i, line in enumerate(lines):
        skip = prev_ctx or bool(_EXCLUDE_RATING_LINE_RX.search(line)) or bool(RECIPIENT_RX.search(line))
        prev_ctx = bool(_PREV_LINE_RX.search(line))
        if skip:
            continue
        for prio, name, rx in pats:
            if name in ("BOX_EN", "COMPANY_EN", "MAINTAINED_EN", "BOX_ZH") and i >= HEADER_LINES:
                continue
            for m in rx.finditer(line):
                raw = next((g for g in m.groups() if g), "")
                pre = line[max(0, m.start() - 6):m.start()]
                if "ESG" in pre.upper() or "前次" in pre:
                    continue
                canon, fine = normalize_rating_word(raw, words)
                if not canon:
                    continue
                cands.append({"value": canon, "fine": fine, "raw": raw, "source": name, "prio": prio, "line": i,
                              "text": line.strip()[:120], "action": rating_action(line)})
    slot = filename_rating_slot(filename, rules, words)
    if slot and slot.get("value"):
        cands.append({"value": slot["value"], "fine": slot.get("fine"), "raw": slot["raw"], "source": "FILENAME_SLOT",
                      "prio": 5, "line": -1, "text": slot["slot"], "action": ""})
    if not cands:
        state = "UNRATED_SLOT" if slot and slot.get("state") == "UNRATED_SLOT" else "NA_NO_TOKEN"
        return {"value": None, "fine": None, "raw": slot["raw"] if slot else "", "source": "", "line": None,
                "action": "", "state": state, "candidates": []}
    best = sorted(cands, key=lambda c: (c["prio"], c["line"] if c["line"] >= 0 else 10 ** 6))[0]
    out = dict(best)
    out["state"] = "HIT"
    out["candidates"] = cands[:20]
    return out


# =====================================================================
# 7 . TARGET PRICE / CURRENT PRICE
# =====================================================================
_CUR = r"(?P<cur>NT\$|NT\s\$|NTD|TWD|US\$|HK\$|\$)"
_NUM = r"(?P<num>[0-9][0-9,]*(?:\.\d+)?)"
_NOT_PRICE_AFTER = re.compile(r"^\s*(?:%|倍|x(?![a-z])|X(?![a-z])|bn|mn|tr|年|億|萬|百萬|千|m\b|b\b)", re.I)
TP_PATTERNS = [
    ("TP_LABEL_EN", re.compile(r"(?<![A-Za-z])(?P<label>(?:(?:12|6)[- ]?(?:month|m|mth)s?\s+)?(?:price\s*target|target\s*price|TP(?![A-Za-z])))"
                               r"(?:\s*\((?:[A-Za-z]{3}[- ]?\d{2,4}|NT\$|TWD|US\$|\$)\))?\s*[:：]?\s*"
                               r"(?:of\s+|at\s+|to\s+|is\s+|raised\s+to\s+|lifted\s+to\s+|cut\s+to\s+|lowered\s+to\s+)?" + _CUR + r"?\s*" + _NUM, re.I)),
    ("TP_LABEL_ZH", re.compile(r"(?<!前次)(?<!前一次)(?<![原舊])(?P<label>(?:(?:12|十二|6|六)\s*個月\s*)?目標(?:股)?價(?:格)?)"
                               r"(?:\s*[\(（](?:NT\$|元|TWD|新台幣)[\)）])?"
                               r"(?:\s*由\s*[0-9][0-9,.]*\s*元?\s*(?:上調|下調|調升|調降|上修|下修|調整|調高|調低)?\s*(?:至|到|為))?"
                               r"\s*(?:維持|調升至|調降至|上修至|下修至|調整至|調高至|調低至|至|為|是)?\s*[:：]?\s*"
                               r"(?:NT\$|\$|新台幣|TWD)?\s*" + _NUM)),
    ("TP_VALUE_FIRST_EN", re.compile(r"(?P<cur>NT\$|TWD|US\$|HK\$)\s*" + _NUM + r"\s+(?:12[- ]?(?:month|m)\s+)?(?:price\s*target|target\s*price|TP)(?![A-Za-z])", re.I)),
]
_TP_EXCLUDE_BEFORE = re.compile(r"(up/downside to|upside to|downside to|previous|prior|old|from|前次|原先|前一次)\s*$", re.I)
CP_PATTERNS = [
    ("CP_EN", re.compile(r"(?<![A-Za-z])(?P<label>shr\.?\s*price,?\s*close|share\s*price|closing\s*price|last\s*price|current\s*price|price|close)"
                         r"\s*(?:\((?P<date>[^)]{2,28})\))?\s*[:：]?\s*" + r"(?P<cur>NT\$|NT\s\$|NTD|TWD|US\$|HK\$)" + r"\s*" + _NUM, re.I)),
    ("CP_PRICE_DASH", re.compile(r"^\s*(?P<cur>NT\$|TWD|US\$|HK\$)\s*" + _NUM + r"\s*[-–]\s*(?:high-conviction\s+)?"
                                 r"(?:outperform|buy|hold|underperform|sell|neutral)", re.I)),
    ("CP_ZH", re.compile(r"(?<!目標)(?P<label>前日收盤價|前一日收盤價|最新收盤價|收盤價|目前股價|最新股價|現價)"
                         r"(?!淨值|表現|走勢|區間|漲跌|報酬)\s*(?:[\(（][^)）]{1,14}[\)）])?\s*(?:[A-Za-z]{3,9}\.?\s*\d{1,2}\s*)?"
                         r"(?:[\(（](?:NT\$|元)[\)）])?\s*[:：]?\s*(?:NT\$|\$)?\s*" + _NUM)),
    ("CP_PAREN_SLOT", re.compile(r"[（(]\s*[1-9]\d{3}\s*TT\s*[,，]\s*(?:NT\$|\$)\s*" + _NUM + r"\s*[,，][^)）]{1,12}[)）]")),
]


def _to_number(s):
    try:
        v = float(str(s).replace(",", ""))
    except (TypeError, ValueError):
        return None
    return v if PRICE_SANE[0] <= v <= PRICE_SANE[1] else None


def _currency(cur, line=""):
    c = (cur or "").upper().replace(" ", "")
    if c in ("NT$", "NTD", "TWD"):
        return "TWD"
    if c in ("US$",):
        return "USD"
    if c in ("HK$",):
        return "HKD"
    if c == "$":
        return "USD" if re.search(r"\bUS\b|AAPL|Nasdaq", line) else "TWD?"
    return "TWD" if re.search(r"元|NT\$", line) else ""


def extract_target_price(lines):
    """Current 12-month target for the covered company.  Previous targets (前次 / from / prior),
    up/downside percentages, multiples and label rows are handled; the maximum is never taken."""
    cands = []
    prev_ctx = False
    for i, line in enumerate(lines):
        skip = prev_ctx or bool(RECIPIENT_RX.search(line))
        prev_ctx = bool(_PREV_LINE_RX.search(line))
        if skip:
            continue
        for name, rx in TP_PATTERNS:
            for m in rx.finditer(line):
                if _TP_EXCLUDE_BEFORE.search(line[max(0, m.start() - 22):m.start()]):
                    continue
                if _NOT_PRICE_AFTER.match(line[m.end():m.end() + 4]):
                    continue
                v = _to_number(m.group("num"))
                if v is None:
                    continue
                cur = m.groupdict().get("cur") or ""
                cands.append({"value": v, "raw": line[m.start():m.end()].strip(), "source": name, "line": i,
                              "currency": _currency(cur, line)})
        # label row: '投資評等 目標價' then the values one line below (兆豐 header box)
        if i + 1 < len(lines) and i < HEADER_LINES and re.search(r"目標價", line) and not re.search(r"\d", line):
            nxt = lines[i + 1]
            m = re.search(r"(?:\$|NT\$)\s*" + _NUM + r"(?![0-9])|" + r"(?<![0-9(（])(?P<n2>[0-9][0-9,]*(?:\.\d+)?)\s*元", nxt)
            if m:
                v = _to_number(m.group("num") or m.group("n2"))
                if v is not None:
                    cands.append({"value": v, "raw": nxt.strip()[:60], "source": "TP_LABEL_ROW", "line": i + 1,
                                  "currency": "TWD"})
    if not cands:
        return {"value": None, "raw": "", "source": "", "line": None, "currency": "", "state": "ABSENT_IN_TEXT",
                "candidates": []}
    best = sorted(cands, key=lambda c: (c["line"], {"TP_LABEL_EN": 0, "TP_LABEL_ZH": 0, "TP_LABEL_ROW": 1,
                                                     "TP_VALUE_FIRST_EN": 2}.get(c["source"], 3)))[0]
    out = dict(best)
    out["state"] = "HIT"
    out["distinct_values"] = sorted({c["value"] for c in cands})
    out["candidates"] = cands[:20]
    return out


def extract_current_price(lines):
    cands = []
    for i, line in enumerate(lines):
        if RECIPIENT_RX.search(line):
            continue
        for name, rx in CP_PATTERNS:
            for m in rx.finditer(line):
                before = line[max(0, m.start() - 16):m.start()].lower()
                label = (m.groupdict().get("label") or "").lower()
                if name == "CP_EN" and ("target" in before or "expected" in before or "52-wk" in before or "range" in before):
                    continue
                if name == "CP_EN" and label in ("price", "close") and re.search(r"target\s*$", before):
                    continue
                if _NOT_PRICE_AFTER.match(line[m.end():m.end() + 4]):
                    continue
                v = _to_number(m.group("num"))
                if v is None:
                    continue
                cands.append({"value": v, "raw": line[m.start():m.end()].strip(), "source": name, "line": i,
                              "currency": _currency(m.groupdict().get("cur") or "", line),
                              "price_date": (m.groupdict().get("date") or "").strip()})
    if not cands:
        return {"value": None, "raw": "", "source": "", "line": None, "state": "ABSENT_IN_TEXT", "candidates": []}
    best = sorted(cands, key=lambda c: c["line"])[0]
    out = dict(best)
    out["state"] = "HIT"
    out["candidates"] = cands[:20]
    return out


# =====================================================================
# 8 . DATES (mother rules.date patterns + English / 五月 19, 2026 forms)
# =====================================================================
_DATE_EXCLUDE_BEFORE = re.compile(r"(前次投資建議|前次|previous|prior|as of|since|截至|price[^)]{0,12}\(|close\s*\(|期間|week of)\s*[:：]?\s*$", re.I)
_DATE_RANGE_AFTER = re.compile(r"^\s*[–\-~～至]\s*\d")
_DATE_RANGE_BEFORE = re.compile(r"[–\-~～至]\s*$")
_COVER_YEAR_RX = re.compile(r"^\s*(2\s*0\s*[0-9]\s*[0-9])\s*$")
_COVER_MD_RX = re.compile(r"^\s*(1[0-2]|0?[1-9])\s*/\s*([12]\d|3[01]|0?[1-9])\s*$")
DATE_PATTERNS = [
    ("AD_SEP", re.compile(r"(?<!\d)(20\d{2})\s*[/\-.年]\s*(0?[1-9]|1[0-2])\s*[/\-.月]\s*(0?[1-9]|[12]\d|3[01])\s*日?(?!\d)"), "AD"),
    ("ROC_SEP", re.compile(r"(?:民國|中華民國)?(?<!\d)(1[0-4]\d)\s*[/\-.年]\s*(0?[1-9]|1[0-2])\s*[/\-.月]\s*(0?[1-9]|[12]\d|3[01])\s*日?(?!\d)"), "ROC"),
    ("AD8", re.compile(r"(?<!\d)(20\d{2})(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])(?!\d)"), "AD"),
    ("EN_DMY", re.compile(r"(?<![0-9A-Za-z])(\d{1,2})\s*" + _MON + r"\.?,?\s*(20\d{2})(?=\D|\d{2}:\d{2}|$)", re.I), "EN_DMY"),
    ("EN_MDY", re.compile(r"(?<![A-Za-z])" + _MON + r"\.?\s*(\d{1,2})(?:st|nd|rd|th)?,?\s*(20\d{2})(?!\d)", re.I), "EN_MDY"),
    ("ZH_MONTH", re.compile(r"(十二|十一|十|[一二三四五六七八九])月\s*(\d{1,2})\s*[,，]\s*(20\d{2})(?!\d)"), "ZH_MDY"),
]


def _mk_date(y, mo, d):
    try:
        dt = datetime.date(int(y), int(mo), int(d))
    except (TypeError, ValueError):
        return None
    if not (DATE_SANE_YEARS[0] <= dt.year <= DATE_SANE_YEARS[1]):
        return None
    return dt.isoformat()


def page_dates(lines):
    out = []
    for i, line in enumerate(lines):
        spans = []
        for name, rx, era in DATE_PATTERNS:
            for m in rx.finditer(line):
                if any(not (m.end() <= a or m.start() >= b) for a, b in spans):
                    continue
                if _DATE_EXCLUDE_BEFORE.search(line[max(0, m.start() - 24):m.start()]):
                    continue
                if _DATE_RANGE_AFTER.match(line[m.end():m.end() + 3]):
                    continue
                if _DATE_RANGE_BEFORE.search(line[max(0, m.start() - 3):m.start()]):
                    continue
                g = m.groups()
                if era == "AD":
                    iso = _mk_date(g[0], g[1], g[2])
                elif era == "ROC":
                    iso = _mk_date(int(g[0]) + 1911, g[1], g[2])
                elif era == "EN_DMY":
                    iso = _mk_date(g[2], EN_MONTHS.get(g[1].lower().rstrip("."), 0), g[0])
                elif era == "EN_MDY":
                    iso = _mk_date(g[2], EN_MONTHS.get(g[0].lower().rstrip("."), 0), g[1])
                else:
                    iso = _mk_date(g[2], ZH_MONTHS.get(g[0], 0), g[1])
                if iso:
                    spans.append((m.start(), m.end()))
                    out.append({"iso": iso, "raw": m.group(0).strip(), "pattern": name, "line": i})
        # cover layout: a spaced year on its own line, month/day a few lines below ('2 0 2 6' ... '9/16')
        ym = _COVER_YEAR_RX.match(line)
        if ym and i < HEADER_LINES:
            for j in range(i + 1, min(i + 4, len(lines))):
                md = _COVER_MD_RX.match(lines[j])
                if md:
                    iso = _mk_date(re.sub(r"\s", "", ym.group(1)), md.group(1), md.group(2))
                    if iso:
                        out.append({"iso": iso, "raw": ym.group(1) + " " + lines[j].strip(), "pattern": "COVER_Y_MD", "line": i})
                    break
    return out


def choose_page_date(dates, filename_date=None):
    """Filename date when the page prints it; else the first header date; else the first date."""
    if not dates:
        return {"iso": None, "source": "", "state": "ABSENT"}
    if filename_date and any(d["iso"] == filename_date for d in dates):
        d = next(d for d in dates if d["iso"] == filename_date)
        return {"iso": d["iso"], "raw": d["raw"], "line": d["line"], "source": "PAGE=FILENAME", "state": "HIT"}
    header = [d for d in dates if d["line"] < HEADER_LINES]
    d = (header or dates)[0]
    return {"iso": d["iso"], "raw": d["raw"], "line": d["line"], "source": "PAGE_HEADER" if header else "PAGE",
            "state": "HIT" if not filename_date else "CONFLICT"}


# =====================================================================
# 9 . ANALYSTS (mother contact rules: e-mail anchors, @-left reverse name, proximity lock)
# =====================================================================
_EN_NAME_RX = re.compile(r"(?<![A-Za-z])((?:[A-Z][a-z]+|[A-Z]{2,}|Mc[A-Z][a-z]+|O'[A-Z][a-z]+)(?:[ -](?:[A-Z][a-z]+|[A-Z]{2,}|[A-Z]\.?)){1,3})(?![A-Za-z])")
_EN_NAME_STOP = {"Equity Analyst", "Research Associate", "Morgan Stanley", "Goldman Sachs", "Asia Pacific", "Taiwan Limited",
                 "Asia Limited", "Global Research", "Equity Research", "Research Analyst", "Head", "Securities"}


def analyst_titles(rules=None):
    zh = _rule(rules or {}, "report_field_rules", "analyst_titles", "zh", default=None) or \
        ["資深分析師", "產業分析師", "策略分析師", "分析師", "研究員", "協理", "經理"]
    en = _rule(rules or {}, "report_field_rules", "analyst_titles", "en", default=None) or \
        ["Equity Analyst", "Research Analyst", "Analyst", "Research Associate", "Strategist", "Economist"]
    return sorted(zh, key=len, reverse=True), sorted(en, key=len, reverse=True)


def reverse_name(local):
    """@-left reverse (mother 批713): only '.' and '_' split, '-' is kept: mei-ling.chen -> Mei-Ling Chen."""
    parts = [p for p in re.split(r"[._]", local or "") if p and not p.isdigit()]
    return " ".join("-".join(x[:1].upper() + x[1:] for x in p.split("-")) for p in parts)


def _name_key(name):
    return re.sub(r"[^a-z]", "", (name or "").lower())


def _same_person(a, b):
    """'John Q Public' ~ 'John Public' ~ 'john.q.public': first and last name tokens agree."""
    ta = [t for t in re.split(r"[\s._]+", (a or "").lower()) if t.isalpha()]
    tb = [t for t in re.split(r"[\s._]+", (b or "").lower()) if t.isalpha()]
    if not ta or not tb:
        return False
    if "".join(ta) == "".join(tb):
        return True
    return len(ta) >= 2 and len(tb) >= 2 and ta[0] == tb[0] and ta[-1] == tb[-1]


def _surnames(rules=None):
    rules = rules if rules is not None else load_rules()
    return str(_rule(rules, "report_field_rules", "name", "surname", "list", default="") or "")


def _chinese_name_above(lines, line_idx, surnames, span=6):
    """A 2-3 character Chinese name standing alone in a row segment within `span` lines above an
    e-mail (KGI/CTBC print '陳大文' beside table rows; the e-mail window of 80 characters misses it)."""
    if not surnames:
        return None
    for j in range(line_idx, max(-1, line_idx - span), -1):
        for seg in re.split(r"\s+|,|，", lines[j]):
            if re.fullmatch(r"[一-鿿]{2,3}", seg) and seg[0] in surnames:
                return seg
    return None


def filename_analysts(filename, lines=None):
    """KGI-style names in the file name ('凱基投顧_1476 儒鴻_陳大文_20260519'): 2-3 CJK characters in a
    token of their own whose first character is a book surname.  PAGE_CONFIRMED when the page prints it."""
    stem = os.path.splitext(os.path.basename(filename or ""))[0]
    sur = _surnames()
    page = re.sub(r"\s", "", "".join(lines or []))
    out = []
    for tok in re.split(r"[_\-\s()（）]+", stem):
        if re.fullmatch(r"[一-鿿]{2,3}", tok) and sur and tok[0] in sur:
            out.append({"name": tok, "source": "FILENAME" + ("+PAGE_CONFIRMED" if tok in page else "")})
    return out


def analysts_from_rulers(lines):
    """Mother CGC_MDL182 contacts_of(): one window per e-mail, starting after the previous anchor (批713).
    Chinese name first (name_cn passed the institution-word blocker), else the English name that agrees
    with the @-left name, else the @-left name itself.  None when the tool is absent."""
    rulers = vcgc("rulers")
    if not rulers:
        return None
    try:
        rows = rulers.contacts_of("\n".join(lines))
    except Exception:
        return None
    out, used = [], set()
    for r in rows or []:
        email = r.get("email") or ""
        cn = [n for n in (r.get("name_cn") or []) if n]
        en = [n for n in (r.get("name_en") or []) if n and n not in _EN_NAME_STOP]
        at_left = r.get("at_left") or ""
        rk = _name_key(at_left)
        agree = [n for n in en if _name_key(n) and (_same_person(n, at_left) or _name_key(n) == rk)]
        if cn:
            name, how = cn[0], "CGC_MDL182:name_cn"
        elif agree:
            name, how = agree[0], "CGC_MDL182:name_en=@left"
        elif r.get("at_left_kind") == "NAME" and at_left:
            name, how = at_left, "CGC_MDL182:@left"
        elif en:
            name, how = en[-1], "CGC_MDL182:name_en"
        else:
            continue
        alias = None
        if not re.search(r"[一-鿿]", name):
            line_idx = next((i for i, l in enumerate(lines) if email and email.split("@")[0] in l), None)
            if line_idx is not None:
                zh = _chinese_name_above(lines, line_idx, _surnames())
                if zh:
                    alias, name, how = name, zh, how + "+ZH_ABOVE"
        key = _name_key(name) or name
        if key in used:
            continue
        used.add(key)
        out.append({"name": name, "alias": alias, "email": email, "source": how, "state": r.get("state"),
                    "broker_domain": r.get("broker"), "phone": (r.get("phone") or [])[:1]})
    return out


def extract_analysts(lines, rules=None):
    """Mother contact rules through CGC_MDL182 when mirrored; the local anchor reader is the fallback.
    Names anchored on e-mails (name on the same line before the e-mail, or up to two lines above),
    cross-checked with the @-left reverse name; Chinese names only next to a title word."""
    via_rulers = analysts_from_rulers(lines)
    if via_rulers:
        return via_rulers
    zh_titles, en_titles = analyst_titles(rules)
    emails = find_emails(lines)
    out = []
    used = set()
    for e in emails:
        cands = []
        same = lines[e["line"]][:e["start"]]
        cands += [(n, "SAME_LINE") for n in _EN_NAME_RX.findall(same)[-1:]]
        for k in (1, 2, 3):
            j = e["line"] - k
            if j < 0:
                break
            for n in _EN_NAME_RX.findall(lines[j]):
                cands.append((n, "ABOVE_%d" % k))
        rev = reverse_name(e["local"])
        pick = None
        for n, src in cands:
            n = re.sub(r",?\s*CFA$", "", n).strip()
            if n in _EN_NAME_STOP or any(n.endswith(t) for t in en_titles):
                continue
            nk = _name_key(n)
            rk = _name_key(rev)
            if nk and rk and (nk == rk or nk in rk or rk in nk or nk.split()[0:1] == rk.split()[0:1]):
                pick = (n, src + "+EMAIL_MATCH")
                break
            first_last = n.split()
            if len(first_last) >= 2 and _name_key(first_last[-1]) and _name_key(first_last[-1]) in rk:
                pick = (n, src + "+SURNAME_MATCH")
                break
        if not pick:
            # Chinese name in the proximity of the e-mail (same / previous line)
            window = " ".join(lines[max(0, e["line"] - 2):e["line"] + 1])
            zh = re.findall(r"(?:" + "|".join(map(re.escape, zh_titles)) + r")\s*[:：]?\s*([一-鿿]{2,4})", window)
            if zh:
                pick = (zh[-1], "ZH_TITLE_NEAR_EMAIL")
            else:
                pick = (rev, "EMAIL_REVERSE")
        key = _name_key(pick[0])
        if key and key not in used:
            used.add(key)
            out.append({"name": pick[0], "email": e["email"], "source": pick[1], "line": e["line"]})
    if not out:
        # measured local signature forms (additive): 台新 '報告人：', 兆豐 '科技產業組張志明', 凱基 '凱基投顧 林小華 (Mia)'
        local = [re.compile(r"報告人\s*[:：]\s*([一-鿿]{2,4})"),
                 re.compile(r"[一-鿿]{1,4}組\s*([一-鿿]{2,3})(?![一-鿿])"),
                 re.compile(r"(?:投顧|證券|研究部)\s*([一-鿿]{2,3})\s*[（(][A-Za-z][A-Za-z .]{1,20}[)）]")]
        for i, raw in enumerate(lines[:HEADER_LINES + 6]):
            line = re.sub(r"(?<=[一-鿿])\s(?=[一-鿿])", "", raw)
            for rx_local in local:
                for m in rx_local.finditer(line):
                    name = m.group(1)
                    if name not in used and (not _surnames() or name[0] in _surnames()):
                        used.add(name)
                        out.append({"name": name, "email": "", "source": "ZH_LOCAL_SIGNATURE", "line": i})
    if not out:
        rx = re.compile(r"(?:" + "|".join(map(re.escape, zh_titles)) + r")\s*[:：]?\s*([一-鿿](?:\s?[一-鿿]){1,3})")
        for i, raw in enumerate(lines):
            line = re.sub(r"(?<=[一-鿿])\s(?=[一-鿿])", "", raw)       # '研 究 員：王小明' -> '研究員：王小明'
            for m in rx.finditer(line):
                name = re.sub(r"\s", "", m.group(1))
                if name not in used:
                    used.add(name)
                    out.append({"name": name, "email": "", "source": "ZH_TITLE", "line": i})
    return out


# =====================================================================
# 10 . COMPANY NAME ON THE PAGE
# =====================================================================
_NAME_PREFIX_RX = re.compile(r"^(?:個股報告|新股介紹|訪談報告|公司訪談|動態更新|研究報告|首次評等|初次評等|Equity Research|Flash|Update|Idea|First Read)\s*")
_NAME_LABEL_RX = re.compile(r"^(?:reuters|bloomberg|ric|bbg|ticker|code|stock code|代號|股票代號|證券代號)\s*[:：]?$", re.I)
_NAME_TAIL_RX = re.compile(r"\s*(?:\|\s*asia pacific|equities|overweight|equal-weight|underweight|neutral|buy|sell|hold|outperform|underperform)\s*$", re.I)
_HEADER_WORDS = re.compile(r"^(?:equity research|global research|asia pacific equity research|research|flash|update|idea|first read|company report|個股報告|新股介紹|產業報告|訪談報告|動態更新)\b", re.I)
_FOREIGN_CODE_RX = re.compile(r"^(.{1,40}?)\s*[（(]\s*[A-Z]{1,5}\s+(?:US|HK|JP|KS|CH|LN|GR|FP|SP)\s*[)）]")


def _clean_name(name):
    name = (name or "").strip(" -–:：|,，")
    name = _NAME_PREFIX_RX.sub("", name).strip()
    name = re.sub(r"^\d{4}[/.\-]\d{1,2}[/.\-]\d{1,2}\s*", "", name)
    prev = None
    while prev != name:
        prev = name
        name = _NAME_TAIL_RX.sub("", name).strip(" -–:：|,，")
    if not name or re.fullmatch(r"[\d\s.,/]+", name) or len(name) > 40 or _NAME_LABEL_RX.match(name):
        return None
    return name


def company_name_on_page(lines, codes, ticker, stock=True):
    """Printed name of the covered company: text before its code ('Hon Hai (2317.TW)' -> Hon Hai),
    the line above a code-only line (KGI '儒鴻' / '(1476.TW/1476 TT)'; JPM 'TSMC Overweight' / '2330.TW'),
    MS 'Name | Asia Pacific', UBS 'Name Equities', a foreign code 'Apple (AAPL US)', CLSA's first line."""
    if not stock:
        return None
    for c in codes or []:
        if ticker and c["core"] != ticker:
            continue
        name = _clean_name(c.get("name_before"))
        if name:
            return name
        line = lines[c["line"]] if 0 <= c["line"] < len(lines) else ""
        if c["line"] > 0 and not line[:line.find(c["raw"])].strip(" (（,") if c["raw"] in line else False:
            name = _clean_name(lines[c["line"] - 1])
            if name and not _HEADER_WORDS.match(name):
                return name
    for i, line in enumerate(lines[:HEADER_LINES]):
        m = re.match(r"^(.{2,50}?)\s*(?:\|\s*Asia Pacific|Equities)\s*$", line)
        if m:
            name = _clean_name(m.group(1))
            if name and not _HEADER_WORDS.match(name):
                return name
        m = _FOREIGN_CODE_RX.match(line)
        if m:
            name = _clean_name(m.group(1))
            if name:
                return name
    first = _clean_name(lines[0]) if lines else None
    if first and not _HEADER_WORDS.match(first) and re.fullmatch(r"[A-Za-z][A-Za-z&.'’\- ]{1,30}", first):
        return first                                  # CLSA prints the company alone on the first line
    return None


# =====================================================================
# 11 . ONE CALL FOR BOTH ENGINES
# =====================================================================
def analyze(filename, lines, alias_table=None, filename_ticker=None, filename_date=None, filename_broker=None,
            rules=None, rating_words=None, official_market=None, text_layer=None, cross_check=False):
    """Everything the first page proves, with the evidence and the state of each field."""
    rules = rules if rules is not None else load_rules()
    lines = [l if isinstance(l, str) else str(l.get("text", "")) for l in (lines or [])]
    words = rating_words or merged_rating_words(rules)
    codes = page_codes(lines)
    rating = extract_rating(lines, filename, rules, words)
    tp = extract_target_price(lines)
    cp = extract_current_price(lines)
    signals = {
        "labeled_rating": rating.get("value") is not None and rating.get("source") not in ("FILENAME_SLOT", "PROSE_EN"),
        "labeled_tp": tp.get("value") is not None,
        "labeled_cp": cp.get("value") is not None,
        "single_company": len({c["core"] for c in codes}) <= 1,
    }
    rtype = detect_report_type(filename, lines, filename_ticker, codes, signals)
    prim = primary_ticker(codes, filename_ticker, rtype["type"])
    tri = tri_code(filename_ticker, codes, prim.get("ticker"), official_market)
    broker = broker_evidence(lines, alias_table or {}, filename_broker, rules)
    dates = page_dates(lines)
    date = choose_page_date(dates, filename_date)
    analysts = extract_analysts(lines, rules)
    fn_analysts = filename_analysts(filename, lines)
    if fn_analysts and not any(re.search(r"[一-鿿]", a.get("name", "")) for a in analysts):
        # a Chinese name in the file name leads when the page gave only e-mail aliases (source priority:
        # page > info zone > filename; the page names stay as aliases, nothing is dropped)
        first = fn_analysts[0]
        analysts = [{"name": first["name"], "email": (analysts[0]["email"] if analysts else ""),
                     "alias": (analysts[0]["name"] if analysts else None), "source": first["source"]}] + analysts[1:]
    stock = rtype["type"] in STOCK_TYPES
    inst = institution_ssot()
    key, key_how = canonical_rating_key(rating.get("raw"), rating.get("value"), rating.get("fine"), inst)
    rating = dict(rating, canonical_key=key, canonical_key_how=key_how)
    has_text = bool(lines) if text_layer is None else bool(text_layer)
    out = {
        "core": "VRN_Evidence_Core " + CORE_VERSION,
        "rule_books": sorted(rules.get("paths", {}).keys()) if isinstance(rules, dict) else [],
        "text_layer": has_text,
        "report_type": rtype,
        "ticker": prim,
        "tri_code": tri,
        "page_codes": codes[:40],
        "company_name": company_name_on_page(lines, codes, prim.get("ticker"), rtype["type"] in STOCK_TYPES),
        "broker": broker,
        "broker_vs_filename": compare_broker(filename_broker, broker),
        "page_date": date,
        "dates": dates[:12],
        "rating": rating if stock else dict(rating, value=None, state="NOT_STOCK", kept_value=rating.get("value")),
        "target_price": tp if stock else dict(tp, value=None, state="NOT_STOCK", kept_value=tp.get("value")),
        "current_price": cp if stock else dict(cp, value=None, state="NOT_STOCK", kept_value=cp.get("value")),
        "analysts": analysts,
        "vcgc": vcgc_status(),
    }
    if cross_check:
        out["canonical_engine"] = canonical_cross_check(lines, filename, prim.get("ticker"))
    if not has_text:
        for key in ("rating", "target_price", "current_price"):
            out[key] = dict(out[key], state="OCR_REQUIRED")
        out["page_date"] = dict(out["page_date"], state="OCR_REQUIRED")
    return out


def canonical_cross_check(lines, filename, ticker=None):
    """The mother's canonical answer (SUP_MDL749 hub -> ENG086) for the same page, kept as evidence.
    Never overrides this core; the loop reports agreement so drift on either side is visible."""
    hub = vcgc("hub")
    if not hub:
        return {"state": "ABSENT"}
    text = "\n".join(lines)
    out = {"state": "OK", "engine": os.path.basename(getattr(hub, "__vcgc_path__", "hub"))}
    try:
        r = hub.rating_of(text)
        out["rating"] = (r or {}).get("canonical") if isinstance(r, dict) else None
    except Exception as exc:
        out["rating_error"] = type(exc).__name__
    try:
        tp = hub.tp_of(text, ticker)
        out["target_price"] = tp[0] if isinstance(tp, tuple) else tp
    except Exception as exc:
        out["tp_error"] = type(exc).__name__
    try:
        b = hub.broker_of(text, filename, ticker)
        out["broker"] = b[0] if isinstance(b, tuple) else b
    except Exception as exc:
        out["broker_error"] = type(exc).__name__
    return out


# =====================================================================
# 12 . COLUMN-ALIGNED TABLES (unruled financial summaries; fallback for the database engine)
# =====================================================================
PERIOD_TOKEN_RX = re.compile(
    r"^(?:FY)?(?:(?:19|20)\d{2}|\d{2})(?:A|E|F|\(F\)|\(E\)|\(A\))?$|^\d{1,2}/\d{2}(?:A|E|F)?$|^[1-4]Q\d{2}(?:A|E|F)?$|"
    r"^(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[-/ ]?\d{2}(?:A|E|F)?$", re.I)
_MONTH_YEAR_RX = re.compile(r"^(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[-/ ]?(\d{2})(A|E|F)?$", re.I)
_GROWTH_LABEL_RX = re.compile(r"yoy|成長|growth|增減|變動|chg|change", re.I)


def normalize_period_label(label):
    """'Dec-24A' -> '2024A' (fiscal year ending in December); other labels unchanged."""
    m = _MONTH_YEAR_RX.match(label or "")
    if m:
        return "20" + m.group(1) + (m.group(2) or "").upper()
    return label
NUMERIC_TOKEN_RX = re.compile(r"^\(?[-−–]?[\d,]*\d(?:\.\d+)?\)?%?$|^(?:n\.?a\.?|nm|n/a|-|--|—)$", re.I)
TABLE_MAX_ROWS = 28
LABEL_GAP_PT = 10.0          # a row label is the run of words right before the first value column


def words_from_chars(chars):
    """Rows of words with x-extents (same gap rule as join_row); used to align table columns."""
    rows = []
    for cs in group_rows(chars):
        words, cur, prev = [], None, None
        for c in sorted(cs, key=lambda g: float(g.get("x0", 0.0))):
            t = c.get("text", "")
            if not t or not t.strip():
                continue
            if prev is not None:
                size = max(float(prev.get("size") or 0.0), float(c.get("size") or 0.0), 1.0)
                gap = (float(c["x0"]) - _glyph_right(prev)) / size
                pair = char_class(prev["text"][-1]) + char_class(t[0])
                if gap > GAP_EM.get(pair, GAP_EM_DEFAULT):
                    words.append(cur)
                    cur = None
            if cur is None:
                cur = {"text": t, "x0": float(c["x0"]), "x1": _glyph_right(c)}
            else:
                cur["text"] += t
                cur["x1"] = max(cur["x1"], _glyph_right(c))
            prev = c
        if cur:
            words.append(cur)
        if words:
            rows.append(words)
    return rows


def _period_columns(words):
    cols = [w for w in words if PERIOD_TOKEN_RX.match(w["text"]) and not re.fullmatch(r"\d{2}", w["text"])]
    cols = [w for w in cols if not (w["text"].isdigit() and len(w["text"]) == 4 and not (2000 <= int(w["text"]) <= 2040))]
    return cols if len(cols) >= 2 else []


def column_tables(chars, page_number=1):
    """Unruled tables whose header row carries two or more period labels ('2024A 2025F 2026F',
    '12/25e 12/26e', 'FY25E').  Values are assigned to the nearest period column; the label is the
    text left of the first column.  Returns [{"PageNumber", "TableID", "Rows"}] like extract_pdf_tables."""
    rows = words_from_chars(chars)
    tables = []
    i = 0
    while i < len(rows):
        cols = _period_columns(rows[i])
        if not cols:
            i += 1
            continue
        centers = [(w["x0"] + w["x1"]) / 2 for w in cols]
        spacing = min((b - a) for a, b in zip(centers, centers[1:])) if len(centers) > 1 else 40.0
        tol = max(6.0, spacing * 0.55)
        left_edge = centers[0] - spacing * 0.6
        right_edge = centers[-1] + spacing * 0.6
        header = [""] + [normalize_period_label(w["text"]) for w in cols]
        body = []
        j = i + 1
        while j < len(rows) and j <= i + TABLE_MAX_ROWS:
            if _period_columns(rows[j]):
                break
            vals = [""] * len(cols)
            label_words = []
            hits = 0
            for w in rows[j]:
                cx = (w["x0"] + w["x1"]) / 2
                if w["x1"] <= left_edge:
                    if label_words and w["x0"] - label_words[-1]["x1"] > LABEL_GAP_PT:
                        label_words = []          # a wide gap: the words before it belong to the prose column
                    label_words.append(w)
                    continue
                if cx > right_edge:
                    continue
                if NUMERIC_TOKEN_RX.match(w["text"]):
                    k = min(range(len(centers)), key=lambda n: abs(centers[n] - cx))
                    if abs(centers[k] - cx) <= tol and not vals[k]:
                        vals[k] = w["text"]
                        hits += 1
            label = " ".join(w["text"] for w in label_words).strip()
            if hits >= 2 and label and not _GROWTH_LABEL_RX.search(label):
                body.append([label] + vals)
            j += 1
        if len(body) >= 2:
            tables.append({"PageNumber": page_number, "TableID": "P%03d_C%03d" % (page_number, len(tables) + 1),
                           "Rows": [header] + body, "Strategy": "COLUMN_ALIGNED"})
        i = j if j > i + 1 else i + 1
    return tables


def transposed_tables(chars, page_number=1):
    """Tables with one row per year and one column per item ('會計年度 營收 … 每股盈餘' / 'Year to Net Profit
    Diluted EPS …' / 華南 '年度EPS預估').  Numeric columns are clustered from the year rows; the one to three
    header rows above are mapped onto them by x-overlap.  Growth columns (YoY / 成長 / growth) and pre-tax EPS
    are left out; the result is turned into the usual orientation (items as rows, years as columns)."""
    rows = words_from_chars(chars)
    tables = []
    i = 0
    while i < len(rows):
        first = rows[i][0] if rows[i] else None
        cand = [w for w in rows[i] if PERIOD_TOKEN_RX.match(w["text"]) and not re.fullmatch(r"\d{2}", w["text"])]
        if not cand or sum(1 for w in rows[i] if NUMERIC_TOKEN_RX.match(w["text"])) < 3:
            i += 1
            continue
        year_x = cand[0]["x0"]
        block, j = [], i
        misses = 0
        while j < len(rows) and misses <= 2:
            ys = [w for w in rows[j] if PERIOD_TOKEN_RX.match(w["text"]) and abs(w["x0"] - year_x) <= 15]
            nums = [w for w in rows[j] if NUMERIC_TOKEN_RX.match(w["text"]) and ys and w["x0"] > ys[0]["x1"]]
            if ys and len(nums) >= 3:
                block.append((ys[0], nums))
                misses = 0
            else:
                misses += 1
            j += 1
        if len(block) < 2:
            i += 1
            continue
        centers = []
        for _y, nums in block:
            for w in nums:
                cx = (w["x0"] + w["x1"]) / 2
                hit = next((c for c in centers if abs(c["x"] - cx) <= 8.0), None)
                if hit:
                    hit["n"] += 1
                    hit["x"] = (hit["x"] * (hit["n"] - 1) + cx) / hit["n"]
                else:
                    centers.append({"x": cx, "n": 1})
        centers = sorted((c for c in centers if c["n"] >= max(2, len(block) // 2)), key=lambda c: c["x"])
        if len(centers) < 2:
            i = j
            continue
        xs = [c["x"] for c in centers]
        half = [((xs[k] - xs[k - 1]) / 2 if k else (xs[1] - xs[0]) / 2) for k in range(len(xs))]
        labels = [""] * len(xs)
        top_idx = next(k for k in range(i, len(rows)) if block[0][0] in rows[k])
        for h in rows[max(0, top_idx - 3):top_idx]:
            for w in h:
                cx = (w["x0"] + w["x1"]) / 2
                for k, x in enumerate(xs):
                    if w["x0"] - 2 <= x <= w["x1"] + 2 or abs(cx - x) <= half[k]:
                        labels[k] = (labels[k] + " " + w["text"]).strip()
                        break
        keep = [k for k, lab in enumerate(labels) if lab and not _GROWTH_LABEL_RX.search(lab)
                and not (re.search(r"稅前", lab) and re.search(r"EPS|每股", lab, re.I))]
        if not keep:
            i = j
            continue
        periods = [normalize_period_label(y["text"].replace("(F)", "F").replace("(E)", "E").replace("(A)", "A")) for y, _n in block]
        out_rows = [[""] + periods]
        for k in keep:
            vals = []
            for _y, nums in block:
                v = next((w["text"] for w in nums if abs((w["x0"] + w["x1"]) / 2 - xs[k]) <= 8.0), "")
                vals.append(v)
            if sum(1 for v in vals if v) >= 2:
                out_rows.append([labels[k]] + vals)
        if len(out_rows) >= 2:
            tables.append({"PageNumber": page_number, "TableID": "P%03d_X%03d" % (page_number, len(tables) + 1),
                           "Rows": out_rows, "Strategy": "TRANSPOSED"})
        i = j
    return tables


def pdf_column_tables(path, pages=3):
    """column_tables() + transposed_tables() for the first `pages` pages of a PDF ([] without a PDF library)."""
    out = []
    for idx in range(pages):
        got = pdf_page_chars(path, idx)
        if not got:
            break
        chars, _size = got
        if not chars:
            if idx == 0:
                continue
            break
        out += column_tables(chars, idx + 1)
        out += transposed_tables(chars, idx + 1)
    return out


# =====================================================================
# CLI (v0102): appendix harvest.  Reads only; writes a report under VIA_Reports/vrn_appendix, never a book.
# =====================================================================
def _appendix_report_md(rep):
    out = ["# VRN 附錄小字收割(VRN_Evidence_Core " + CORE_VERSION + ")", "",
           f"- 產生:{rep['generated']} · 檔數 {rep['n_files']} · 讀到附錄小字 {rep['n_with_small_print']} · "
           f"附錄強證定出券商 {rep['n_strong_broker']} · 評等定義 {rep['n_scale_lines']} 行 · 新詞 {len(rep['new_words'])} 個", "",
           "## 新評等詞(候選;只增不減要經樞紐,這裡不寫任何冊)", "",
           "| 詞 | 配對詞 | 建議鍵 | 出現 | 例 |", "|---|---|---|---|---|"]
    for w in rep["new_words"]:
        out.append(f"| {w['word']} | {' / '.join(w['alts']) or '-'} | {' / '.join(w['suggested']) or '待裁定'} | {w['count']} | "
                   f"{(w['examples'][0] if w['examples'] else '')[:60]} |")
    out += ["", "## 逐檔", "", "| 檔 | 頁 | 附錄小字行 | 附錄券商 | 層級 | 對檔名 | 評等定義 |", "|---|---|---|---|---|---|---|"]
    for f in rep["files"]:
        out.append(f"| {f['file']} | {f['n_pages']} | {f['n_lines']} | {f['broker'] or '-'} | {f['tier']} | {f['vs_filename']} | {f['n_scale']} |")
    return "\n".join(out) + "\n"


def _adj_cli(a):
    """adj <ticker> <report date> [target] [--db path]: one ADJ CLOSE answer, printed (read-only)."""
    db = None
    if "--db" in a:
        k = a.index("--db")
        db = a[k + 1] if k + 1 < len(a) else None
        a = a[:k] + a[k + 2:]
    page = None
    if "--page" in a:
        k = a.index("--page")
        page = a[k + 1] if k + 1 < len(a) else None
        a = a[:k] + a[k + 2:]
    if len(a) < 2:                           # 批730 review: counted after both options are taken off
        print("[ADJ] 用法:python VRN_Evidence_Core.py adj <代號> <報告日> [目標價] [--db 庫] [--page 頁面現價]")
        return 2
    q = adj_basis(a[0], a[1], a[2] if len(a) > 2 else None, db=db, page_price=page)
    print("[ADJ] %s %s 目標價 %s → %s" % (q["ticker"], q["report_date"], q["target_price"], q["state"]))
    for key, label in (("target_price_adj", "目標價(ADJ)"), ("adj_factor", "因子"), ("adj_factor_date", "因子日"),
                       ("adj_factor_basis", "因子分母"), ("price_prev_raw", "前一日交易所收盤"), ("adj_retro_ratio", "Yahoo回調比"),
                       ("adj_event", "未調整事件"),
                       ("price_prev_close", "前一日收盤"), ("price_prev_adj", "前一日 ADJ CLOSE"), ("price_prev_date", "前一日"),
                       ("price_latest_adj", "最新 ADJ CLOSE"), ("price_latest_date", "最新日"), ("upside_adj", "上漲空間 %")):
        if q.get(key) not in (None, ""):
            print("  %-16s %s" % (label, q[key]))
    if q.get("why"):
        print("  因由 " + q["why"])
    print("  引擎 %s · 庫 %s" % (q.get("engine"), q.get("db")))
    return 0 if q["state"] and q["state"] not in ADJ_MISS_STATES else 2


def main(argv=None):
    a = list(sys.argv[1:] if argv is None else argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    if a and a[0] == "adj":
        return _adj_cli(a[1:])
    if not a or a[0] != "appendix":
        print("VRN_Evidence_Core " + CORE_VERSION + " · 用法:python VRN_Evidence_Core.py appendix <pdf|資料夾>... [--out 夾] [--last N]"
              " | adj <代號> <報告日> [目標價] [--db 庫] [--page 頁面現價]")
        return 0 if not a else 2
    out_dir, last, targets, i = None, APPENDIX_PAGES, [], 1
    while i < len(a):
        if a[i] == "--out" and i + 1 < len(a):
            out_dir, i = a[i + 1], i + 2
        elif a[i] == "--last" and i + 1 < len(a):
            last, i = max(1, int(a[i + 1])), i + 2
        else:
            targets.append(a[i])
            i += 1
    files = []
    for t in targets:
        p = os.path.abspath(t)
        if os.path.isdir(p):
            for root, _dirs, names in os.walk(p):
                files += [os.path.join(root, n) for n in sorted(names) if n.lower().endswith(".pdf")]
        elif p.lower().endswith(".pdf") and os.path.isfile(p):
            files.append(p)
    if not files:
        print("[附錄] NODATA · 沒有 PDF(給檔或資料夾;附錄只讀 PDF)")
        return 2
    rules = load_rules()
    table = ssot_alias_table()
    rows, new = [], {}
    for k, f in enumerate(files, 1):
        print(f"[進度] {k}/{len(files)} 附錄 {os.path.basename(f)}", flush=True)
        try:
            ev = appendix_evidence(f, alias_table=table, rules=rules, last_pages=last)
        except Exception as exc:
            rows.append({"file": os.path.basename(f), "n_pages": 0, "n_lines": 0, "broker": None, "tier": "ERROR",
                         "strong": False, "vs_filename": "N/A", "n_scale": 0, "error": f"{type(exc).__name__}: {exc}"[:160]})
            continue
        b = ev["broker"]
        rows.append({"file": os.path.basename(f), "n_pages": ev["info"].get("n_pages"), "n_lines": ev["n_lines"],
                     "broker": b.get("broker"), "tier": b.get("tier"), "strong": bool(b.get("strong")),
                     "vs_filename": ev["vs_filename"], "n_scale": len(ev["rating_scale"]), "state": ev["state"],
                     "why": ev["info"].get("why", "")})
        for r in ev["rating_scale"]:
            for word, known, other in ((r["word"], r["word_known"], r["alt"]), (r["alt"], r["alt_known"], r["word"])):
                if not word or known:
                    continue
                w = new.setdefault(word, {"word": word, "count": 0, "alts": set(), "suggested": set(), "examples": [], "files": set()})
                w["count"] += 1
                if other:
                    w["alts"].add(other)
                if r["suggested"]:
                    w["suggested"].add(r["suggested"])
                if len(w["examples"]) < 2:
                    w["examples"].append(r["definition"])
                w["files"].add(os.path.basename(f))
    new_words = sorted(({**w, "alts": sorted(w["alts"]), "suggested": sorted(w["suggested"]), "files": sorted(w["files"])[:5]}
                        for w in new.values()), key=lambda x: (-x["count"], x["word"]))
    rep = {"engine": "VRN_Evidence_Core " + CORE_VERSION, "generated": __import__("time").strftime("%Y-%m-%dT%H:%M:%S"),
           "n_files": len(rows), "n_with_small_print": sum(1 for r in rows if r.get("n_lines")),
           "n_strong_broker": sum(1 for r in rows if r.get("strong")), "n_scale_lines": sum(r.get("n_scale", 0) for r in rows),
           "new_words": new_words, "files": rows, "last_pages": last}
    root = repo_root() or os.getcwd()
    target = out_dir or os.path.join(str(root), "VIA_Reports", "vrn_appendix")
    os.makedirs(target, exist_ok=True)
    stamp = __import__("time").strftime("%Y%m%d_%H%M%S")
    for name in ("APPENDIX_" + stamp + ".json", "APPENDIX_latest.json"):
        with open(os.path.join(target, name), "w", encoding="utf-8") as h:
            h.write(json.dumps(rep, ensure_ascii=False, indent=1) + "\n")
    with open(os.path.join(target, "APPENDIX_latest.md"), "w", encoding="utf-8") as h:
        h.write(_appendix_report_md(rep))
    print(f"[附錄] 檔 {rep['n_files']} · 讀到附錄小字 {rep['n_with_small_print']} · 附錄強證定出券商 {rep['n_strong_broker']} · "
          f"評等定義 {rep['n_scale_lines']} 行 · 新詞 {len(new_words)} 個 → {os.path.join(target, 'APPENDIX_latest.md')}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
