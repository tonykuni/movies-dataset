# -*- coding: utf-8 -*-
"""
VRN Integrated Report Database Engine
=====================================
單檔整合版：Filename 擷取 + First Page Anchor + TWSE/TPEX Name 驗證 + BasicInfo/FinancialData 入庫。

核心輸出：
    - StockReportBasicInfo.parquet           必選，累積式更新
    - StockReportFinancialData.parquet       必選，累積式更新

可選輸出：
    - CSV：UTF-8 BOM，供 Excel 人工檢查
    - JSON：UTF-8，供 UI / Evidence / SSOT
    - DuckDB：供 SQL / Dashboard
    - Google Sheet：保留接口，需自行提供 gspread 憑證

設計原則：
    1. Parquet 是唯一主資料庫。
    2. CSV / JSON / DuckDB / Google Sheet 全部從 Parquet 派生。
    3. Filename 粗抓候選；首頁資訊區精抓；TWSE/TPEX/SSOT 回填正式 Name。
    4. 4 碼數字優先作 TW_TICKER；6 碼以上優先判斷日期。
    5. 支援 (20)260131、260131、20260131、1140822、114/08/22 等日期。
    6. StockReportBasicInfo 一份報告一列；StockReportFinancialData 一份報告多列 long format。

建議執行：
    python VRN_Integrated_ReportDatabase_Engine.py --input "C:/Reports/incoming" --output "C:/VRN_DB" --csv --json --duckdb
    python VRN_Integrated_ReportDatabase_Engine.py --input "C:\\測試樣本報告" --output "C:\\VRN_DB" --csv --json
    python VRN_Integrated_ReportDatabase_Engine.py --self-test

v0101（自動測試迴圈 2026-09-21）：PDF／DOCX／TXT 一起收（影像檔標 OCR_REQUIRED 略過）；券商／評等／目標價
比對加字母邊界與結構限制；四碼年份帶與 ETF 代碼不當股票也不當日期；日期做日曆驗證；MM/YY 欄頭；Parquet
欄位定型；JSON 無 NaN；線上名稱表需 VIA_NET=1（本引擎絕不自行設定）；可接 VIA_VRN_FirstPageEngine。
"""

# ============================================================
# def 00_GLOBAL_PARAMETERS：所有參數集中於此區
# ============================================================

from __future__ import annotations

# ===== [VIA:ACCEL-BRIDGE:v0100] SuperAccel 加速器橋(批102 全樹導入令;graceful 零行為變更) =====
try:
    import sys as _sa_sys
    from pathlib import Path as _sa_Path
    _sa_p = _sa_Path(__file__).resolve()
    while _sa_p.parent != _sa_p:
        if (_sa_p / "supportive modules" / "VIA_SuperAccel_Module.py").exists():
            _sa_sys.path.insert(0, str(_sa_p / "supportive modules"))
            break
        _sa_p = _sa_p.parent
    import VIA_SuperAccel_Module as VIA_ACCEL  # noqa: N816
except Exception:
    VIA_ACCEL = None  # graceful:加速器缺席零影響
# ===== [VIA:ACCEL-BRIDGE:END] =====
# ===== [VIA:NET-BRIDGE:v0100] 統包網路工具橋(批115 VDF 全導入令;graceful 零行為變更) =====
VIA_NET_TOOL_PATH = None
try:
    from pathlib import Path as _nb_Path
    _nb_p = _nb_Path(__file__).resolve()
    while _nb_p.parent != _nb_p:
        _nb_dir = _nb_p / "supportive modules" / "network"
        if _nb_dir.exists():
            _nb_hits = sorted(_nb_dir.glob("via_net_unified_v*.py"))
            if _nb_hits:
                VIA_NET_TOOL_PATH = str(_nb_hits[-1])
            break
        _nb_p = _nb_p.parent
except Exception:
    VIA_NET_TOOL_PATH = None


def _via_net():
    """統包唯一網路工具惰性載入(法遵雙閘 VIA_NET_CONSENT);缺席回 None(誠實)"""
    if VIA_NET_TOOL_PATH is None:
        return None
    try:
        import importlib.util as _nb_ilu
        _nb_spec = _nb_ilu.spec_from_file_location("VIA_NET_UNIFIED", VIA_NET_TOOL_PATH)
        _nb_mod = _nb_ilu.module_from_spec(_nb_spec)
        _nb_spec.loader.exec_module(_nb_mod)
        return _nb_mod
    except Exception:
        return None
# ===== [VIA:NET-BRIDGE:END] =====

import argparse
import csv
import hashlib
import html
import json
import math
import os
import re
import sys
import time
import traceback
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple
from urllib.request import Request, urlopen

SCHEMA_VERSION = "VRN_SCHEMA_20260924_004"
ENGINE_VERSION = "v0106"
# v0106 (批733; 掉球 Z195 -- the workstation run of 2026-09-24: 7 of 105 PDFs gave "no financial rows"):
#   tables printed without ruling lines.  (1) the column reader (evidence core) now knows the year-first quarter
#   headers of the screenshot table (25Q1 / 25Q4(F)), half years and 2024年; before, it saw only 2024 and 2025(F) in
#   that header and read nothing -- in other layouts it would have put quarter numbers under the year columns.
#   (2) when the ruled tables and the first four pages still give no income-statement row, later pages are read
#   too, but only pages whose text looks like an income statement (long initiation reports carry their statements
#   at the back).  (3) a table read without ruling lines is stored only if it adds up (gate_fallback_tables:
#   the same identities as v0105); one that does not is left out and named in the report's ValidationIssues as
#   FIN_TABLE_REJECTED(...).  Ruled tables are unchanged (their failures stay warnings, as in v0105).
# v0105 (mother 批732; operator's screenshot 2026-09-24 of 20251128兆豐訪談速報-神達(3706) p.4 季度損益表):
#   that table came out wrong three ways -- (1) quarter headers "25Q1".."26Q1(F)" lost their year (FiscalYear None,
#   PERIOD_UNCLEAR), (2) "(F)" was not read as a forecast, (3) 營業成本 / 營業費用 had no metric, so both rows were
#   dropped.  Period headers now go to the mother's parser first (VRN_ENG074 period_parts, the one that already
#   reads 24Q1 / 25Q1(F) / 2025(F) / 1H25 / FY25); what it cannot parse keeps the rules below.  Metrics: COGS,
#   OperatingExpense (+ their ratios, so a percentage never becomes an amount), NonOperatingIncome, InterestIncome.
#   financial_identity_checks(): the table must add up (gross = revenue - cost, operating = gross - expense, four
#   quarters = the year) within the rounding the printed digits carry -- a read table that does not add up is
#   named, not trusted.
# v0104 (mother 批731; operator 2026-09-24「剛剛跑好慢」): process_batch takes an optional args.only_files (file names)
#   -- the self-test loop's idempotency pass re-ingests a named sample instead of the whole folder; without it nothing
#   changes.  The evidence core behind the first-page reader now parses each page / appendix once per file version.
# v0103 (mother 批729; operator 2026-09-24 "所有目標價及各前一日的價格都要換成ADJ CLOSE 上漲空間都要用最新的ADJ CLOSE"):
# UpsideDownsidePct is the upside over the LATEST adj close (the mother's L99 through VRN_Evidence_Core.adj_basis ->
# ENG073 adj_quote); the target and the day-before prices in ADJ CLOSE terms get their own columns (appended only);
# the page arithmetic (target / printed price) is kept as UpsidePagePct and never promoted when ADJ is missing.
# The first-page engine gets the file path, so the appendix small print can name the issuer (broker_tier APPENDIX_*).
# v0101 (auto-test loop 2026-09-21): DOCX/TXT intake, boundary-safe broker/rating/target-price matching,
# year-band and ETF aware tickers, calendar-validated dates, MM/YY period headers, typed Parquet columns,
# NaN-safe JSON, VIA_NET consent gate for the optional online name table, first-page engine hook.
DEFAULT_INPUT_DIR_WINDOWS = r"C:\測試樣本報告"
SUPPORTED_SUFFIXES = (".pdf", ".docx", ".txt")
OCR_SUFFIXES = (".jpg", ".jpeg", ".png", ".tif", ".tiff")
PDF_ENGINE_ORDER = ["PyMuPDF", "pdfplumber"]     # runtime-switchable by the auto-test loop
NETWORK_CONSENT_ENV = "VIA_NET"                   # read only; this engine never sets consent variables
YEAR_BAND = (2021, 2030)
ETF_CODE_REGEX = re.compile(r"^(?:00\d{2,4}|00\d{3}[A-Z]|00\d[A-Z]\d{2})$")
DEFAULT_RUN_ID_PREFIX = "VRN_RUN"
DEFAULT_OUTPUT_DIR = "./vrn_output"
DEFAULT_INPUT_DIR = "./incoming"

BASICINFO_PARQUET_NAME = "StockReportBasicInfo.parquet"
FINANCIALDATA_PARQUET_NAME = "StockReportFinancialData.parquet"
BASICINFO_CSV_NAME = "StockReportBasicInfo.csv"
FINANCIALDATA_CSV_NAME = "StockReportFinancialData.csv"
BASICINFO_JSON_NAME = "StockReportBasicInfo.json"
FINANCIALDATA_JSON_NAME = "StockReportFinancialData.json"
DUCKDB_NAME = "StockReport.duckdb"
HTML_MATRIX_NAME = "VRN_BatchReport.html"
BATCH_SUMMARY_JSON_NAME = "VRN_BatchSummary.json"
BATCH_MATRIX_CSV_NAME = "VRN_BatchMatrix.csv"
BATCH_CHECKPOINT_JSON_NAME = "VRN_BatchCheckpoint.json"

# TWSE/TPEX 公司基本資料 OpenAPI。若未連網或端點變更，仍可使用本地 SSOT CSV。
TWSE_PROFILE_API_URL = "https://openapi.twse.com.tw/v1/opendata/t187ap03_L"
TPEX_PROFILE_API_URL = "https://www.tpex.org.tw/openapi/v1/mopsfin_t187ap03_O"
HTTP_TIMEOUT_SECONDS = 30
HTTP_USER_AGENT = "Mozilla/5.0 VRN-Integrated-Engine/1.0"
ENABLE_ONLINE_NAME_TABLE_BY_DEFAULT = False

# 使用者常用本地 SSOT 路徑，可在 CLI 覆蓋。
DEFAULT_TICKER_SSOT_CSV = ""
DEFAULT_BROKER_SSOT_CSV = ""
DEFAULT_RATING_SSOT_CSV = ""

# Filename 與文字切割符號。
TOKEN_SEPARATOR_REGEX = re.compile(
    r"[\s\t\r\n\-_–—+,.，。、:：;；/\\|()（）\[\]【】{}<>《》〈〉\"'`~!！?？@#$%^&*=]+"
)
MIXED_TOKEN_REGEX = re.compile(r"\d+|[A-Za-z]+|[\u3400-\u4DBF\u4E00-\u9FFF]+")
CHINESE_REGEX = re.compile(r"^[\u3400-\u4DBF\u4E00-\u9FFF]+$")
ENGLISH_REGEX = re.compile(r"^[A-Za-z]+$")
NUMBER_REGEX = re.compile(r"^\d+$")

# Ticker Regex：一般公司 / yfinance / Bloomberg / 報告文字。
TW_STOCK_TICKER_REGEX = re.compile(r"^([1-9]\d{3})$")
TW_YFINANCE_TICKER_REGEX = re.compile(r"\b([1-9]\d{3})\.(TW|TWO)\b", re.IGNORECASE)
TW_BBG_TICKER_REGEX = re.compile(r"\b([1-9]\d{3})\s*(TT)\b", re.IGNORECASE)
TW_TEXT_TICKER_REGEX = re.compile(r"(?:代號|股票代號|證券代號|Ticker|Code|Stock\s*Code)\s*[:：]?\s*(?<![0-9A-Za-z])([1-9]\d{3})(?![0-9A-Za-z])", re.IGNORECASE)
TW_PAREN_TICKER_REGEX = re.compile(r"[\u3400-\u4DBF\u4E00-\u9FFF]{1,12}(?:-KY)?\s*[（(]\s*([1-9]\d{3})(?:\s*TT|\.TWO?)?\s*[,，)）]")

# 日期 Regex。
DATE_YYYYMMDD_REGEX = re.compile(r"(?<!\d)((?:19|20)\d{2})(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])(?!\d)")
DATE_YYMMDD_REGEX = re.compile(r"(?<!\d)(\d{2})(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])(?!\d)")
DATE_ROC_YYYMMDD_REGEX = re.compile(r"(?<!\d)(1\d{2})(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])(?!\d)")
DATE_SPLIT_REGEX = re.compile(
    r"(?<!\d)((?:19|20)\d{2}|1\d{2})[./\-年](0?[1-9]|1[0-2])[./\-月](0?[1-9]|[12]\d|3[01])日?(?!\d)"
)
DATE_SENTENCE_END_REGEX = re.compile(r"[。．\.！!？?]$")

# 首頁區域切割比例。
FIRST_PAGE_TOP_RATIO = 0.42
FIRST_PAGE_LEFT_RATIO = 0.45
FIRST_PAGE_RIGHT_RATIO = 0.55
FIRST_PAGE_MAX_PAGES_FOR_TEXT = 3
PDF_TABLE_MAX_PAGES = 80

# BasicInfo 欄位。
BASICINFO_COLUMNS = [
    "FileID",
    "ReportID",
    "SourceFileName",
    "SourcePath",
    "FileHash",
    "TW_TICKER",
    "YF_TICKER",
    "Market",
    "Name",
    "CompanyNameCandidate",
    "Broker",
    "BrokerCandidate",
    "ReportDate",
    "ReportTitle",
    "ReportType",
    "Language",
    "Analyst",
    "Department",
    "Rating",
    "RatingChange",
    "TargetPrice",
    "CurrentPrice",
    "UpsideDownsidePct",
    "InvestmentThesis",
    "Industry",
    "SubIndustry",
    "MarketCap",
    "ShareCapital",
    "Chairman",
    "GeneralManager",
    "Spokesperson",
    "ConfidenceScore",
    "ValidationStatus",
    "ValidationIssues",
    "EvidenceJson",
    "SourceFormat",
    "ReportDateDisplay",
    "TriCodeVerdict",
    "RatingFine",
    "RatingAction",
    "BrokerDisplay",
    "TargetPriceScenarios",
    "PartialDates",
    "FirstPageEngine",
    "CreatedAt",
    "UpdatedAt",
    "RunID",
    "SchemaVersion",
    # v0102 (2026-09-23, real-report loop + VCGC): appended only, nothing above moved or dropped
    "RatingKey",
    "BrokerTier",
    "BrokerDenied",
    "PageDate",
    "FilenameDate",
    "DateConflict",
    "AnalystEmail",
    "CompanyNamePage",
    "PageCodes",
    "EvidenceCore",
    # v0103 (批729, ADJ CLOSE basis): appended only, nothing above moved or dropped
    "UpsideBasis",
    "UpsideState",
    "UpsidePagePct",
    "TargetPriceAdj",
    "AdjFactor",
    "AdjFactorDate",
    "PrevClose",
    "PrevCloseAdj",
    "PrevCloseDate",
    "CurrentPriceAdj",
    "LatestAdjClose",
    "LatestAdjDate",
]

# FinancialData 欄位，採 long format。
FINANCIALDATA_COLUMNS = [
    "FinancialRowID",
    "FileID",
    "ReportID",
    "TW_TICKER",
    "YF_TICKER",
    "Name",
    "ReportDate",
    "Broker",
    "PeriodType",
    "FiscalYear",
    "FiscalQuarter",
    "PeriodLabelRaw",
    "EstimateFlag",
    "MetricCategory",
    "MetricName",
    "MetricNameRaw",
    "Value",
    "ValueRaw",
    "Unit",
    "Currency",
    "Scale",
    "YoY",
    "QoQ",
    "PER",
    "PBR",
    "EV_EBITDA",
    "DividendYield",
    "ROE",
    "ROA",
    "Revenue",
    "GrossProfit",
    "GrossMargin",
    "OperatingProfit",
    "OperatingMargin",
    "NetIncome",
    "EPS",
    "TotalAssets",
    "TotalLiabilities",
    "Equity",
    "BookValuePerShare",
    "OperatingCashFlow",
    "Capex",
    "FreeCashFlow",
    "PageNumber",
    "TableID",
    "RowIndex",
    "ColumnIndex",
    "SourceText",
    "ConfidenceScore",
    "ValidationStatus",
    "ValidationIssues",
    "CreatedAt",
    "UpdatedAt",
    "RunID",
    "SchemaVersion",
]

# 去重鍵。
BASICINFO_UPSERT_KEYS = ["ReportID"]
FINANCIALDATA_UPSERT_KEYS = ["FinancialRowID"]

# Parquet needs one type per column: these are numeric (NaN when missing), everything else is text.
BASICINFO_NUMERIC_COLUMNS = ("TargetPrice", "CurrentPrice", "UpsideDownsidePct", "ConfidenceScore", "MarketCap", "ShareCapital",
                             "UpsidePagePct", "TargetPriceAdj", "AdjFactor", "PrevClose", "PrevCloseAdj", "CurrentPriceAdj",
                             "LatestAdjClose")
FINANCIALDATA_NUMERIC_COLUMNS = (
    "Value", "Scale", "YoY", "QoQ", "PER", "PBR", "EV_EBITDA", "DividendYield", "ROE", "ROA", "Revenue",
    "GrossProfit", "GrossMargin", "OperatingProfit", "OperatingMargin", "NetIncome", "EPS", "TotalAssets",
    "TotalLiabilities", "Equity", "BookValuePerShare", "OperatingCashFlow", "Capex", "FreeCashFlow",
    "PageNumber", "RowIndex", "ColumnIndex", "ConfidenceScore", "FiscalYear",
)

# 評等同義字，可再由 SSOT CSV 擴充。
DEFAULT_RATING_ALIASES = {
    "BUY": ["BUY", "買進", "買入", "強力買進", "OUTPERFORM", "增持", "ADD", "OVERWEIGHT", "OW", "買進評等"],
    "HOLD": ["HOLD", "中立", "持有", "NEUTRAL", "EQUALWEIGHT", "EW", "MARKET PERFORM", "觀望"],
    "SELL": ["SELL", "賣出", "減持", "UNDERPERFORM", "UNDERWEIGHT", "UW", "降低持股"],
    "NOT_RATED": ["未評等", "NOT RATED", "NR", "UNRATED", "NOT COVERED", "NO RECOMMENDATION", "不予評等", "尚未評級", "未覆蓋", "停止覆蓋"],
}
# short codes count only inside a rating structure (spec section 4); "Note" is a report label, never NR
RATING_SHORT_CODES = {"B": "BUY", "OW": "BUY", "OP": "BUY", "SB": "BUY", "H": "HOLD", "N": "HOLD", "EW": "HOLD", "MP": "HOLD",
                      "S": "SELL", "UW": "SELL", "UP": "SELL", "SS": "SELL", "NR": "NOT_RATED", "CD": "NOT_RATED"}
RATING_FINE = {"strong buy": "STRONG_BUY", "強力買進": "STRONG_BUY", "積極買進": "STRONG_BUY", "conviction buy": "STRONG_BUY",
               "top pick": "STRONG_BUY", "strong sell": "STRONG_SELL", "強力賣出": "STRONG_SELL", "強烈賣出": "STRONG_SELL"}
RATING_ACTION_WORDS = ("維持", "重申", "調升", "調降", "升評", "降評", "初評", "初次評等", "maintain", "reiterate", "upgrade", "downgrade", "initiate")
TARGET_PRICE_TERMS = ["12-month price target", "12-month target price", "12m target price", "fair value estimate", "target price",
                      "price target", "fair value", "tp", "pt", "十二個月目標價", "12個月目標價", "目標價格", "股價目標", "目標股價",
                      "合理價值", "合理股價", "目標價", "合理價"]

# 報告類型同義字。
REPORT_TYPE_ALIASES = {
    "INITIATION": ["初次評等", "初評", "Initiation", "Initiating Coverage"],
    "UPDATE": ["更新", "Update", "Company Update", "研究報告"],
    "FLASH": ["速報", "Flash", "快報", "訪談", "法說", "Earnings Flash"],
    "INDUSTRY": ["產業", "Industry", "Sector"],
}

# 財務指標標準化。
METRIC_ALIASES = {
    "Revenue": ["Revenue", "Sales", "Net Sales", "營收", "營業收入", "銷貨收入", "合併營收", "Net Revenue"],
    "GrossProfit": ["Gross Profit", "毛利", "營業毛利"],
    "GrossMargin": ["Gross Margin", "GPM", "毛利率"],
    "OperatingProfit": ["Operating Profit", "OP", "Operating Income", "營業利益", "營業淨利"],
    "OperatingMargin": ["Operating Margin", "OPM", "營益率", "營業利益率"],
    "PretaxIncome": ["Pretax Income", "PBT", "稅前淨利", "稅前利益"],
    "NetIncome": ["Net Income", "NI", "PAT", "稅後淨利", "淨利", "歸屬母公司淨利"],
    "EPS": ["EPS", "每股盈餘", "每股稅後盈餘", "每股獲利"],
    "PER": ["PER", "P/E", "PE", "本益比"],
    "PBR": ["PBR", "P/B", "PB", "股價淨值比"],
    "EV_EBITDA": ["EV/EBITDA", "EV EBITDA"],
    "DividendYield": ["Dividend Yield", "殖利率", "股息殖利率"],
    "ROE": ["ROE", "股東權益報酬率"],
    "ROA": ["ROA", "資產報酬率"],
    "TotalAssets": ["Total Assets", "總資產"],
    "TotalLiabilities": ["Total Liabilities", "總負債"],
    "Equity": ["Equity", "Shareholders Equity", "股東權益", "淨值"],
    "BookValuePerShare": ["BVPS", "Book Value Per Share", "每股淨值"],
    "OperatingCashFlow": ["Operating Cash Flow", "CFO", "營業現金流", "營運現金流"],
    "Capex": ["Capex", "Capital Expenditure", "資本支出"],
    "FreeCashFlow": ["Free Cash Flow", "FCF", "自由現金流"],
    "MarketCap": ["Market Cap", "市值"],
    "TargetPrice": ["Target Price", "TP", "目標價"],
    "CurrentPrice": ["Current Price", "現價", "股價"],
    # v0105 (批732): cost and expense rows were dropped whole (no metric).  Only added; the ratios are here so that
    # 營業費用率 / 營業成本率 (a percentage) never lands in the amount (the longest alias wins).
    "COGS": ["COGS", "Cost of Revenue", "Cost of Sales", "Cost of Goods Sold", "營業成本", "銷貨成本"],
    "COGSRatio": ["營業成本率", "銷貨成本率"],
    "OperatingExpense": ["Operating Expense", "Operating Expenses", "OPEX", "營業費用"],
    "OperatingExpenseRatio": ["營業費用率", "費用率", "OPEX Ratio"],
    "NonOperatingIncome": ["營業外收支", "營業外收入及支出", "營業外收入", "業外收支", "業外損益", "Non-operating Income",
                           "Non-Operating Income"],
    "InterestIncome": ["利息收入", "Interest Income"],
}

METRIC_CATEGORY_MAP = {
    "Revenue": "Income Statement",
    "GrossProfit": "Income Statement",
    "GrossMargin": "Income Statement",
    "OperatingProfit": "Income Statement",
    "OperatingMargin": "Income Statement",
    "PretaxIncome": "Income Statement",
    "NetIncome": "Income Statement",
    "EPS": "Income Statement",
    "PER": "Valuation",
    "PBR": "Valuation",
    "EV_EBITDA": "Valuation",
    "DividendYield": "Valuation",
    "ROE": "Valuation",
    "ROA": "Valuation",
    "TotalAssets": "Balance Sheet",
    "TotalLiabilities": "Balance Sheet",
    "Equity": "Balance Sheet",
    "BookValuePerShare": "Balance Sheet",
    "OperatingCashFlow": "Cash Flow",
    "Capex": "Cash Flow",
    "FreeCashFlow": "Cash Flow",
    "MarketCap": "Basic Info",
    "TargetPrice": "Valuation",
    "CurrentPrice": "Valuation",
    "COGS": "Income Statement",
    "COGSRatio": "Income Statement",
    "OperatingExpense": "Income Statement",
    "OperatingExpenseRatio": "Income Statement",
    "NonOperatingIncome": "Income Statement",
    "InterestIncome": "Income Statement",
}

# CSV 編碼。
CSV_ENCODING_FOR_EXCEL = "utf-8-sig"
JSON_ENCODING = "utf-8"

# 信心分數權重。
CONFIDENCE_WEIGHTS = {
    "filename_ticker": 0.18,
    "first_page_ticker": 0.22,
    "exchange_name": 0.24,
    "report_date": 0.12,
    "broker": 0.10,
    "rating_or_target": 0.08,
    "title_or_text": 0.06,
}

# ============================================================
# def 01_OPTIONAL_IMPORTS
# ============================================================

def require_pandas():
    try:
        import pandas as pd  # type: ignore
        return pd
    except Exception as exc:
        raise RuntimeError(
            f"pandas 匯入失敗（{exc.__class__.__name__}: {exc}）。直譯器 {sys.executable}。"
            "請先安裝：pip install pandas pyarrow"
        ) from exc


def optional_import_pyarrow_available() -> bool:
    try:
        import pyarrow  # noqa: F401
        return True
    except Exception:
        return False


def optional_import_duckdb():
    try:
        import duckdb  # type: ignore
        return duckdb
    except Exception:
        return None


def optional_import_pdfplumber():
    try:
        import pdfplumber  # type: ignore
        return pdfplumber
    except Exception:
        return None


def optional_import_fitz():
    try:
        import fitz  # type: ignore
        return fitz
    except Exception:
        return None

# ============================================================
# def 02_BASIC_UTILS
# ============================================================

def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def make_run_id(prefix: str = DEFAULT_RUN_ID_PREFIX) -> str:
    return f"{prefix}_{datetime.now().strftime('%Y%m%dT%H%M%S')}_{os.getpid()}"


def ensure_dir(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path


def normalize_unicode_text(text: Any) -> str:
    if text is None:
        return ""
    value = str(text)
    value = unicodedata.normalize("NFKC", value)
    value = value.replace("\u3000", " ")
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def file_sha256(path: Path, chunk_size: int = 1024 * 1024) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(chunk_size), b""):
            h.update(chunk)
    return h.hexdigest()


def stable_hash(parts: Sequence[Any], size: int = 16) -> str:
    raw = "||".join(normalize_unicode_text(p) for p in parts)
    return hashlib.sha256(raw.encode("utf-8", errors="ignore")).hexdigest()[:size]


def safe_json_dumps(obj: Any) -> str:
    try:
        return json.dumps(obj, ensure_ascii=False, sort_keys=True, default=str)
    except Exception:
        return json.dumps({"error": "json_dump_failed", "repr": repr(obj)}, ensure_ascii=False)


def read_json_file(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def write_json_file(path: Path, data: Any) -> None:
    with path.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2, default=str)


def safe_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    text = normalize_unicode_text(value)
    if text == "" or text in {"-", "--", "NA", "N/A", "nan", "None"}:
        return None
    text = text.replace(",", "")
    negative = False
    if text.startswith("(") and text.endswith(")"):
        negative = True
        text = text[1:-1]
    text = text.replace("%", "").replace("x", "").replace("X", "")
    m = re.search(r"[-+]?\d+(?:\.\d+)?", text)
    if not m:
        return None
    try:
        num = float(m.group(0))
        return -num if negative else num
    except Exception:
        return None


def safe_int(value: Any) -> Optional[int]:
    num = safe_float(value)
    if num is None or math.isnan(num):
        return None
    return int(num)


def first_non_empty(values: Iterable[Any]) -> str:
    for value in values:
        text = normalize_unicode_text(value)
        if text:
            return text
    return ""


def scan_input_files(input_dir: Path) -> List[Path]:
    if input_dir.is_file():
        return [input_dir]
    files: List[Path] = []
    for path in input_dir.rglob("*"):
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES + OCR_SUFFIXES:
            files.append(path)
    unique = sorted(set(files), key=lambda p: str(p).lower())
    return unique


def canonical_broker_key(key: str) -> str:
    """Bridge every spelling of a broker key to the SYNONYM_LIBRARY canonical (uppercase)."""
    raw = normalize_unicode_text(key)
    compact = re.sub(r"[^A-Za-z0-9]", "", raw).upper()
    return BROKER_KEY_BRIDGE.get(compact, compact or raw.upper())


BROKER_KEY_BRIDGE = {
    "YUANTA": "YUANTA", "YT": "YUANTA", "KGI": "KGI", "FUBON": "FUBON", "FB": "FUBON", "CATHAY": "CATHAY", "CT": "CATHAY",
    "SINOPAC": "SINOPAC", "SP": "SINOPAC", "CAPITAL": "CAPITAL", "CAP": "CAPITAL", "MASTERLINK": "MASTERLINK",
    "JIHSUN": "JIHSUN", "MEGA": "MEGA", "MEGABANK": "MEGA", "CTBC": "CTBC", "NOMURA": "NOMURA", "MORGANSTANLEY": "MS",
    "MS": "MS", "JPMORGAN": "JPM", "JPM": "JPM", "JP": "JPM", "GOLDMANSACHS": "GS", "GOLDMAN": "GS", "GS": "GS",
    "CITI": "CITI", "CITIGROUP": "CITI", "UBS": "UBS", "CLSA": "CLSA", "CLST": "CLSA", "MACQUARIE": "MACQUARIE",
    "MQ": "MACQUARIE", "MCQ": "MACQUARIE", "HSBC": "HSBC", "HUANAN": "HUANAN", "HN": "HUANAN", "TAISHIN": "TAISHIN",
    "TSC": "TAISHIN", "PRESIDENT": "PRESIDENT", "PSC": "PRESIDENT", "DAIWA": "DAIWA", "BOFA": "BOFA", "BOA": "BOFA",
     "GFHK": "GF", "ESUN": "ESUN", "FIRST": "FIRST", "CS": "CS", "CREDITSUISSE": "CS",
}


def alias_hits(alias: str, hay: str) -> bool:
    """CJK alias: substring. ASCII alias: letter/digit boundaries on both sides (MS never hits 'systems')."""
    a = normalize_unicode_text(alias).lower()
    h = normalize_unicode_text(hay).lower()
    if not a or not h:
        return False
    if re.search(r"[\u3400-\u4DBF\u4E00-\u9FFF]", a):
        return a in h
    return re.search(r"(?<![A-Za-z0-9])" + re.escape(a) + r"(?![A-Za-z0-9])", h) is not None


def company_fragments_after_ticker(tokens: Sequence[str]) -> List[str]:
    """CJK tokens directly after a 4-digit ticker token are the company (凱基投顧_2891 中信金 -> 中信金)."""
    out: List[str] = []
    after = 0
    for tok in tokens:
        tok = normalize_unicode_text(tok)
        if TW_STOCK_TICKER_REGEX.match(tok) and not (YEAR_BAND[0] <= int(tok) <= YEAR_BAND[1]):
            after = 1
            continue
        if after and CHINESE_REGEX.match(tok) and after <= 2:
            out.append(tok)
            after += 1
        else:
            after = 0
    return out


def find_knowledge_root(start: Optional[Path] = None) -> Dict[str, Optional[Path]]:
    """Locate optional knowledge files (attachments dicts, SYNONYM_LIBRARY, first-page engine)."""
    here = (start or Path(__file__)).resolve().parent
    found: Dict[str, Optional[Path]] = {"broker_dict": None, "rating_dict": None, "synonym_library": None, "first_page_engine": None}
    probe = here
    for _ in range(8):
        candidates = {
            "broker_dict": [probe / "attachments" / "VRN_Broker_Dict_v0100.json"],
            "rating_dict": [probe / "attachments" / "VRN_Rating_Dict_v0100.json"],
            "synonym_library": [probe / "knowledge" / "SYNONYM_LIBRARY_v4.json",
                                probe / "functional modules" / "VRN" / "knowledge" / "SYNONYM_LIBRARY_v4.json"],
            "first_page_engine": [probe / "VIA_VRN_FirstPageEngine.py", probe / "engine" / "VIA_VRN_FirstPageEngine.py"],
        }
        for key, paths in candidates.items():
            if found[key] is None:
                for path in paths:
                    if path.is_file():
                        found[key] = path
                        break
        if probe.parent == probe:
            break
        probe = probe.parent
    return found


def load_synonym_scope(path: Optional[Path], scope: str) -> Dict[str, List[str]]:
    """SYNONYM_LIBRARY scope -> {term: [canonical, ...]} keeping every source's canonical."""
    if not path or not path.exists():
        return {}
    try:
        payload = read_json_file(path)
    except Exception:
        return {}
    table: Dict[str, List[str]] = {}
    for term, records in ((payload.get("scopes") or {}).get(scope) or {}).items():
        canon: List[str] = []
        for rec in records or []:
            c = rec.get("canonical")
            if c and c not in canon:
                canon.append(str(c))
        if canon:
            table[normalize_unicode_text(term).lower()] = canon
    return table


_FIRST_PAGE_MODULE_CACHE: Dict[str, Any] = {}


def optional_import_first_page_engine():
    """Load VIA_VRN_FirstPageEngine once per process (v0101 re-executed it for every report)."""
    path = find_knowledge_root().get("first_page_engine")
    if not path:
        return None
    key = str(path)
    if key in _FIRST_PAGE_MODULE_CACHE:
        return _FIRST_PAGE_MODULE_CACHE[key]
    try:
        import importlib.util
        spec = importlib.util.spec_from_file_location("via_vrn_first_page_engine", str(path))
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)  # type: ignore[union-attr]
    except Exception:
        module = None
    _FIRST_PAGE_MODULE_CACHE[key] = module
    return module


_EVIDENCE_CORE_CACHE: Dict[str, Any] = {}


def optional_import_evidence_core():
    """VRN_Evidence_Core (same folder as this engine); None when absent."""
    if "core" in _EVIDENCE_CORE_CACHE:
        return _EVIDENCE_CORE_CACHE["core"]
    module = None
    try:
        import importlib.util
        # mother 批728: unique module name, and only the sibling file (an older intake VRN_Evidence_Core.py exists)
        modname = "vrn_engine_evidence_core"
        path = Path(__file__).resolve().parent / "VRN_Evidence_Core.py"
        cached = sys.modules.get(modname)
        if cached is not None and Path(getattr(cached, "__file__", "") or ".").resolve() == path:
            module = cached
        elif path.is_file():
            spec = importlib.util.spec_from_file_location(modname, str(path))
            module = importlib.util.module_from_spec(spec)
            sys.modules[modname] = module
            spec.loader.exec_module(module)  # type: ignore[union-attr]
    except Exception:
        module = None
    _EVIDENCE_CORE_CACHE["core"] = module
    return module

# ============================================================
# def 03_FILENAME_TOKENIZATION_AND_DATE
# ============================================================

def split_mixed_token(token: str) -> List[str]:
    token = normalize_unicode_text(token)
    if not token:
        return []
    return [x for x in MIXED_TOKEN_REGEX.findall(token) if x]


def tokenize_filename(filename: str) -> Dict[str, Any]:
    stem = Path(filename).stem
    normalized = normalize_unicode_text(stem)
    rough = TOKEN_SEPARATOR_REGEX.split(normalized)
    tokens: List[str] = []
    for item in rough:
        item = normalize_unicode_text(item)
        if not item:
            continue
        parts = split_mixed_token(item)
        if parts:
            tokens.extend(parts)
        else:
            tokens.append(item)
    tokens = [normalize_unicode_text(x) for x in tokens if normalize_unicode_text(x)]
    number_tokens = [x for x in tokens if NUMBER_REGEX.match(x)]
    chinese_tokens = [x for x in tokens if CHINESE_REGEX.match(x)]
    english_tokens = [x for x in tokens if ENGLISH_REGEX.match(x)]
    mixed_tokens = [x for x in tokens if x not in number_tokens and x not in chinese_tokens and x not in english_tokens]
    return {
        "source_stem": stem,
        "normalized": normalized,
        "tokens": tokens,
        "number_tokens": number_tokens,
        "chinese_tokens": chinese_tokens,
        "english_tokens": english_tokens,
        "mixed_tokens": mixed_tokens,
    }


def parse_date_from_yyyymmdd(raw: str) -> Optional[str]:
    text = normalize_unicode_text(raw)
    m = DATE_YYYYMMDD_REGEX.search(text)
    if not m:
        return None
    year = int(m.group(1))
    month = int(m.group(2))
    day = int(m.group(3))
    return validate_date_parts(year, month, day)


def parse_date_from_yymmdd(raw: str) -> Optional[str]:
    text = normalize_unicode_text(raw)
    m = DATE_YYMMDD_REGEX.search(text)
    if not m:
        return None
    yy = int(m.group(1))
    year = 2000 + yy
    month = int(m.group(2))
    day = int(m.group(3))
    return validate_date_parts(year, month, day)


def parse_date_from_roc_yyyymmdd(raw: str) -> Optional[str]:
    text = normalize_unicode_text(raw)
    m = DATE_ROC_YYYMMDD_REGEX.search(text)
    if not m:
        return None
    roc_year = int(m.group(1))
    if not (80 <= roc_year <= 150):
        return None
    year = roc_year + 1911
    month = int(m.group(2))
    day = int(m.group(3))
    return validate_date_parts(year, month, day)


def parse_date_from_split_text(raw: str) -> Optional[str]:
    text = normalize_unicode_text(raw)
    m = DATE_SPLIT_REGEX.search(text)
    if not m:
        return None
    year_value = int(m.group(1))
    if 80 <= year_value <= 150:
        year = year_value + 1911
    else:
        year = year_value
    month = int(m.group(2))
    day = int(m.group(3))
    return validate_date_parts(year, month, day)


def validate_date_parts(year: int, month: int, day: int) -> Optional[str]:
    try:
        dt = datetime(year, month, day)
        if not (1990 <= dt.year <= 2099):
            return None
        return dt.strftime("%Y-%m-%d")
    except Exception:
        return None


def parse_date_any(raw: str) -> Optional[str]:
    text = normalize_unicode_text(raw)
    if not text:
        return None
    for parser in [parse_date_from_split_text, parse_date_from_yyyymmdd, parse_date_from_roc_yyyymmdd, parse_date_from_yymmdd]:
        parsed = parser(text)
        if parsed:
            return parsed
    return None


def parse_date_candidates_from_number_tokens(number_tokens: Sequence[str]) -> List[Dict[str, Any]]:
    candidates: List[Dict[str, Any]] = []
    for idx, token in enumerate(number_tokens):
        token = normalize_unicode_text(token)
        parsed = None
        kind = ""
        if ETF_CODE_REGEX.match(token):
            continue
        if len(token) == 8:
            parsed = parse_date_from_yyyymmdd(token)
            kind = "YYYYMMDD"
        elif len(token) == 7:
            parsed = parse_date_from_roc_yyyymmdd(token)
            kind = "ROC_YYYMMDD"
        elif len(token) == 6:
            parsed = parse_date_from_yymmdd(token)
            kind = "YYMMDD"
        if parsed:
            candidates.append({"raw": token, "date": parsed, "kind": kind, "index": idx})

    # 支援 (20)260131 被切成 20 + 260131 的情況。
    for idx in range(len(number_tokens) - 1):
        left = normalize_unicode_text(number_tokens[idx])
        right = normalize_unicode_text(number_tokens[idx + 1])
        if left in {"19", "20"} and len(right) == 6:
            merged = left + right
            parsed = parse_date_from_yyyymmdd(merged)
            if parsed:
                candidates.append({"raw": f"{left}+{right}", "date": parsed, "kind": "PREFIX_YYYYMMDD", "index": idx})
    return candidates


def choose_best_report_date(filename_candidates: List[Dict[str, Any]], first_page_text: str) -> Tuple[str, List[Dict[str, Any]]]:
    all_candidates = list(filename_candidates)
    for m in DATE_SPLIT_REGEX.finditer(first_page_text or ""):
        parsed = parse_date_any(m.group(0))
        if parsed:
            all_candidates.append({"raw": m.group(0), "date": parsed, "kind": "FIRST_PAGE_SPLIT", "index": -1})
    for m in DATE_YYYYMMDD_REGEX.finditer(first_page_text or ""):
        parsed = parse_date_any(m.group(0))
        if parsed:
            all_candidates.append({"raw": m.group(0), "date": parsed, "kind": "FIRST_PAGE_YYYYMMDD", "index": -1})
    # filename 優先，其次 first page；取第一個合理值。
    return (all_candidates[0]["date"] if all_candidates else "", all_candidates)

# ============================================================
# def 04_SSOT_LOADERS_AND_NAME_TABLE
# ============================================================

def read_csv_rows_flexible(path: Path) -> List[Dict[str, str]]:
    if not path or not path.exists():
        return []
    encodings = ["utf-8-sig", "utf-8", "big5", "cp950"]
    last_exc: Optional[Exception] = None
    for enc in encodings:
        try:
            with path.open("r", encoding=enc, newline="") as f:
                reader = csv.DictReader(f)
                return [{normalize_unicode_text(k): normalize_unicode_text(v) for k, v in row.items()} for row in reader]
        except Exception as exc:
            last_exc = exc
            continue
    raise RuntimeError(f"CSV 讀取失敗：{path}，最後錯誤：{last_exc}")


def pick_row_value(row: Dict[str, Any], candidates: Sequence[str]) -> str:
    lowered = {normalize_unicode_text(k).lower(): v for k, v in row.items()}
    for key in candidates:
        if key in row and normalize_unicode_text(row.get(key)):
            return normalize_unicode_text(row.get(key))
        lk = key.lower()
        if lk in lowered and normalize_unicode_text(lowered.get(lk)):
            return normalize_unicode_text(lowered.get(lk))
    return ""


def load_local_ticker_ssot(path: Optional[Path]) -> Dict[str, Dict[str, Any]]:
    lookup: Dict[str, Dict[str, Any]] = {}
    if not path or not path.exists():
        return lookup
    rows = read_csv_rows_flexible(path)
    for row in rows:
        code = pick_row_value(row, ["TW_TICKER", "Ticker", "Code", "證券代號", "股票代號", "公司代號", "代號"])
        code_match = re.search(r"([1-9]\d{3})", code)
        if not code_match:
            continue
        ticker = code_match.group(1)
        name = pick_row_value(row, ["Name", "證券名稱", "公司名稱", "公司簡稱", "簡稱", "中文簡稱"])
        market = pick_row_value(row, ["Market", "市場", "上市櫃", "Exchange"])
        yf = pick_row_value(row, ["YF_TICKER", "YFinance", "Yahoo", "yfinance"])
        english = pick_row_value(row, ["EnglishName", "English", "英文名稱", "英文簡稱"])
        industry = pick_row_value(row, ["Industry", "產業", "產業別"])
        if not yf and market.upper() in {"TWSE", "上市"}:
            yf = f"{ticker}.TW"
        if not yf and market.upper() in {"TPEX", "OTC", "上櫃"}:
            yf = f"{ticker}.TWO"
        lookup[ticker] = {
            "TW_TICKER": ticker,
            "YF_TICKER": yf,
            "Market": normalize_market_name(market),
            "Name": name,
            "EnglishName": english,
            "Industry": industry,
            "source": str(path),
        }
    return lookup


def normalize_market_name(value: str) -> str:
    text = normalize_unicode_text(value).upper()
    if text in {"TWSE", "上市", "上市公司", "TSE"}:
        return "TWSE"
    if text in {"TPEX", "OTC", "上櫃", "櫃買", "上櫃公司"}:
        return "TPEX"
    return normalize_unicode_text(value)


def http_get_json(url: str, timeout: int = HTTP_TIMEOUT_SECONDS) -> Any:
    req = Request(url, headers={"User-Agent": HTTP_USER_AGENT, "Accept": "application/json,text/plain,*/*"})
    with urlopen(req, timeout=timeout) as resp:
        raw = resp.read()
    text = raw.decode("utf-8", errors="replace")
    return json.loads(text)


def normalize_exchange_profile_row(row: Dict[str, Any], market: str, source_url: str) -> Optional[Dict[str, Any]]:
    code = pick_row_value(row, ["公司代號", "Code", "code", "SecuritiesCode", "股票代號", "有價證券代號"])
    m = re.search(r"([1-9]\d{3})", code)
    if not m:
        return None
    ticker = m.group(1)
    name = pick_row_value(row, ["公司簡稱", "公司名稱", "證券名稱", "Name", "CompanyName", "簡稱"])
    full_name = pick_row_value(row, ["公司名稱", "FullName", "Full Name", "發行公司名稱"])
    english = pick_row_value(row, ["英文簡稱", "英文名稱", "EnglishName", "English Abbreviation"])
    industry = pick_row_value(row, ["產業別", "Industry", "IndustryCode"])
    chairman = pick_row_value(row, ["董事長", "Chairman"])
    general_manager = pick_row_value(row, ["總經理", "General Manager", "GeneralManager"])
    spokesperson = pick_row_value(row, ["發言人", "Spokesperson"])
    yf = f"{ticker}.TW" if market == "TWSE" else f"{ticker}.TWO"
    return {
        "TW_TICKER": ticker,
        "YF_TICKER": yf,
        "Market": market,
        "Name": name or full_name,
        "FullName": full_name,
        "EnglishName": english,
        "Industry": industry,
        "Chairman": chairman,
        "GeneralManager": general_manager,
        "Spokesperson": spokesperson,
        "source": source_url,
    }


def fetch_exchange_name_table(enable_online: bool) -> Dict[str, Dict[str, Any]]:
    if not enable_online:
        return {}
    if os.environ.get(NETWORK_CONSENT_ENV, "0") != "1":
        print(f"[WARN] 線上名稱表需要 {NETWORK_CONSENT_ENV}=1（LIVE 預設關閉），本次略過。", file=sys.stderr)
        return {}
    lookup: Dict[str, Dict[str, Any]] = {}
    endpoints = [(TWSE_PROFILE_API_URL, "TWSE"), (TPEX_PROFILE_API_URL, "TPEX")]
    for url, market in endpoints:
        try:
            data = http_get_json(url)
            if isinstance(data, dict):
                rows = data.get("data") or data.get("items") or data.get("result") or []
            else:
                rows = data
            if not isinstance(rows, list):
                continue
            for row in rows:
                if not isinstance(row, dict):
                    continue
                item = normalize_exchange_profile_row(row, market, url)
                if item:
                    lookup[item["TW_TICKER"]] = item
        except Exception as exc:
            print(f"[WARN] 線上名稱表抓取失敗：{market} {url} -> {exc}", file=sys.stderr)
    return lookup


def merge_ticker_lookups(*lookups: Dict[str, Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    merged: Dict[str, Dict[str, Any]] = {}
    for lookup in lookups:
        for ticker, row in lookup.items():
            if ticker not in merged:
                merged[ticker] = dict(row)
            else:
                for key, value in row.items():
                    if normalize_unicode_text(value):
                        merged[ticker][key] = value
    return merged


def load_broker_alias_ssot(path: Optional[Path]) -> Dict[str, List[str]]:
    alias: Dict[str, List[str]] = {}
    if path and path.exists():
        rows = read_csv_rows_flexible(path)
        for row in rows:
            broker = pick_row_value(row, ["Broker", "broker", "券商", "標準券商", "Canonical"])
            names = pick_row_value(row, ["Alias", "alias", "同義字", "別名", "Aliases"])
            if not broker:
                continue
            parts = [broker] + [normalize_unicode_text(x) for x in re.split(r"[|,，;；/]+", names) if normalize_unicode_text(x)]
            alias[broker] = sorted(set(parts), key=len, reverse=True)
    if not alias:
        alias = default_broker_aliases()
    alias = {canonical_broker_key(k): list(v) for k, v in alias.items()}
    # append-only knowledge, one owner per alias: built-in first, SYNONYM_LIBRARY second, attachments dict last
    owner: Dict[str, str] = {}
    for key, words in alias.items():
        for word in words:
            owner.setdefault(normalize_unicode_text(word).lower(), key)

    def add(key: str, word: str) -> None:
        clean = normalize_unicode_text(word)
        low = clean.lower()
        if not clean or (low in owner and owner[low] != key):
            return
        alias.setdefault(key, []).append(clean)
        owner[low] = key

    roots = find_knowledge_root()
    for term, canon in load_synonym_scope(roots.get("synonym_library"), "broker").items():
        add(canonical_broker_key(canon[0]), term)
    try:
        payload = read_json_file(roots["broker_dict"]) if roots.get("broker_dict") else None
    except Exception:
        payload = None
    if isinstance(payload, dict):
        for canon_zh, spec in (payload.get("brokers") or {}).items():
            words = [str(a) for a in (spec.get("aliases") or [])] + [str(canon_zh)]
            owners = {owner[normalize_unicode_text(w).lower()] for w in words if normalize_unicode_text(w).lower() in owner}
            if len(owners) > 1:
                continue        # a lumped entry (MS + JPM in one row) never rewrites an identity
            key = owners.pop() if owners else canonical_broker_key(str(spec.get("abbr") or canon_zh))
            for word in words:
                add(key, word)
            BROKER_DISPLAY.setdefault(key, str(canon_zh))
    return {key: sorted(set(words), key=len, reverse=True) for key, words in alias.items()}


BROKER_DISPLAY: Dict[str, str] = {}


def default_broker_aliases() -> Dict[str, List[str]]:
    return {
        "KGI": ["KGI", "凱基", "凱基證券", "凱基投顧"],
        "Yuanta": ["Yuanta", "元大", "元大投顧", "元大證券"],
        "Fubon": ["Fubon", "富邦", "富邦投顧", "富邦證券"],
        "Cathay": ["Cathay", "國泰", "國泰證券", "國泰投顧"],
        "SinoPac": ["SinoPac", "永豐", "永豐金", "永豐投顧", "永豐金證券"],
        "Capital": ["Capital", "群益", "群益投顧", "群益證券"],
        "MasterLink": ["MasterLink", "元富", "元富投顧", "元富證券"],
        "Jih Sun": ["Jih Sun", "日盛", "日盛投顧", "日盛證券"],
        "Mega": ["Mega", "兆豐", "兆豐投顧", "兆豐證券"],
        "CTBC": ["CTBC", "中信", "中信投顧", "中國信託", ],
        "Nomura": ["Nomura", "野村"],
        "Morgan Stanley": ["Morgan Stanley", "MS", "摩根士丹利"],
        "JPMorgan": ["JPMorgan", "J.P. Morgan", "JP Morgan", "摩根大通"],
        "Goldman Sachs": ["Goldman Sachs", "GS", "高盛"],
        "Citi": ["Citi", "Citigroup", "花旗"],
        "UBS": ["UBS", "瑞銀"],
        "CLSA": ["CLSA", "里昂"],
        "Macquarie": ["Macquarie", "麥格理"],
        "HSBC": ["HSBC", "匯豐"],
    }


def load_rating_alias_ssot(path: Optional[Path]) -> Dict[str, List[str]]:
    alias = {k: list(v) for k, v in DEFAULT_RATING_ALIASES.items()}
    if path and path.exists():
        rows = read_csv_rows_flexible(path)
        for row in rows:
            canonical = pick_row_value(row, ["Rating", "Canonical", "標準評等"])
            names = pick_row_value(row, ["Alias", "Aliases", "同義字", "別名"])
            if not canonical:
                continue
            parts = [canonical] + [normalize_unicode_text(x) for x in re.split(r"[|,，;；/]+", names) if normalize_unicode_text(x)]
            alias[canonical] = sorted(set(alias.get(canonical, []) + parts), key=len, reverse=True)
    return alias

# ============================================================
# def 05_PDF_FIRST_PAGE_AND_LAYOUT
# ============================================================

def extract_pdf_text_and_first_page_zones(pdf_path: Path) -> Dict[str, Any]:
    result = {
        "full_text": "",
        "first_page_text": "",
        "top_left_text": "",
        "top_center_text": "",
        "top_right_text": "",
        "first_page_repaired_sentences": [],
        "engine": "",
        "issues": [],
    }
    engines = [str(x) for x in PDF_ENGINE_ORDER]
    if engines and engines[0] == "pdfplumber":
        plumber_first = extract_pdf_text_with_pdfplumber(pdf_path, dict(result))
        if plumber_first.get("engine"):
            return plumber_first
    fitz = optional_import_fitz()
    if fitz is not None:
        try:
            doc = fitz.open(str(pdf_path))
            texts = []
            for page_idx in range(min(len(doc), FIRST_PAGE_MAX_PAGES_FOR_TEXT)):
                texts.append(doc[page_idx].get_text("text") or "")
            result["full_text"] = "\n".join(texts)
            if len(doc) > 0:
                page = doc[0]
                result["first_page_text"] = page.get_text("text") or ""
                blocks = page.get_text("blocks") or []
                width = float(page.rect.width)
                height = float(page.rect.height)
                zones = split_blocks_into_top_zones(blocks, width, height)
                result.update(zones)
                zone_lines = []
                for key in ["top_left_text", "top_center_text", "top_right_text"]:
                    zone_lines.extend(text_to_lines(result.get(key, "")))
                result["first_page_repaired_sentences"] = repair_sentences_until_period(zone_lines)
            result["engine"] = "PyMuPDF"
            doc.close()
            return result
        except Exception as exc:
            result["issues"].append(f"PyMuPDF failed: {exc}")

    plumber = extract_pdf_text_with_pdfplumber(pdf_path, result)
    if plumber.get("engine"):
        return plumber
    result["issues"].extend(x for x in plumber.get("issues", []) if x not in result["issues"])
    result["issues"].append("No PDF extraction engine available. Install PyMuPDF or pdfplumber.")
    return result


def extract_pdf_text_with_pdfplumber(pdf_path: Path, result: Dict[str, Any]) -> Dict[str, Any]:
    pdfplumber = optional_import_pdfplumber()
    if pdfplumber is None:
        return result
    try:
        with pdfplumber.open(str(pdf_path)) as pdf:
            texts = []
            for page_idx, page in enumerate(pdf.pages[:FIRST_PAGE_MAX_PAGES_FOR_TEXT]):
                page_text = page.extract_text() or ""
                texts.append(page_text)
                if page_idx == 0:
                    result["first_page_text"] = page_text
                    width, height = float(page.width), float(page.height)
                    blocks = []
                    for word in page.extract_words() or []:
                        blocks.append((float(word["x0"]), float(word["top"]), float(word["x1"]), float(word["bottom"]), str(word["text"])))
                    result.update(split_blocks_into_top_zones(blocks, width, height))
            result["full_text"] = "\n".join(texts)
            zone_lines = []
            for key in ["top_left_text", "top_center_text", "top_right_text"]:
                zone_lines.extend(text_to_lines(result.get(key, "")))
            result["first_page_repaired_sentences"] = repair_sentences_until_period(zone_lines or text_to_lines(result["first_page_text"]))
        result["engine"] = "pdfplumber"
        return result
    except Exception as exc:
        result["issues"].append(f"pdfplumber failed: {exc}")
        return result


def extract_docx_text_and_zones(docx_path: Path) -> Dict[str, Any]:
    """DOCX intake: paragraphs are the page text, the first 40 paragraphs form the first page."""
    result = {"full_text": "", "first_page_text": "", "top_left_text": "", "top_center_text": "", "top_right_text": "",
              "first_page_repaired_sentences": [], "engine": "", "issues": [], "tables": []}
    try:
        import docx  # type: ignore
    except Exception:
        result["issues"].append("python-docx not installed: pip install python-docx")
        return result
    try:
        document = docx.Document(str(docx_path))
        paragraphs = [normalize_unicode_text(p.text) for p in document.paragraphs if normalize_unicode_text(p.text)]
        result["full_text"] = "\n".join(paragraphs)
        result["first_page_text"] = "\n".join(paragraphs[:40])
        result["top_left_text"] = "\n".join(paragraphs[:12])
        result["first_page_repaired_sentences"] = repair_sentences_until_period(paragraphs[:12])
        for t_idx, table in enumerate(document.tables, start=1):
            rows = [[normalize_unicode_text(c.text) for c in row.cells] for row in table.rows]
            cleaned = clean_raw_table(rows)
            if cleaned:
                result["tables"].append({"PageNumber": 1, "TableID": f"P001_T{t_idx:03d}", "Rows": cleaned})
        result["engine"] = "python-docx"
    except Exception as exc:
        result["issues"].append(f"python-docx failed: {exc}")
    return result


def extract_txt_text_and_zones(txt_path: Path) -> Dict[str, Any]:
    result = {"full_text": "", "first_page_text": "", "top_left_text": "", "top_center_text": "", "top_right_text": "",
              "first_page_repaired_sentences": [], "engine": "", "issues": [], "tables": []}
    raw = b""
    try:
        raw = txt_path.read_bytes()
    except Exception as exc:
        result["issues"].append(f"read failed: {exc}")
        return result
    text = ""
    for enc in ("utf-8-sig", "utf-8", "cp950", "big5", "utf-16"):
        try:
            text = raw.decode(enc)
            break
        except Exception:
            continue
    lines = text_to_lines(text)
    result["full_text"] = "\n".join(lines)
    result["first_page_text"] = "\n".join(lines[:60])
    result["top_left_text"] = "\n".join(lines[:12])
    result["first_page_repaired_sentences"] = repair_sentences_until_period(lines[:12])
    result["engine"] = "text"
    return result


def extract_document_text_and_zones(path: Path) -> Dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_text_and_first_page_zones(path)
    if suffix == ".docx":
        return extract_docx_text_and_zones(path)
    if suffix == ".txt":
        return extract_txt_text_and_zones(path)
    result = {"full_text": "", "first_page_text": "", "top_left_text": "", "top_center_text": "", "top_right_text": "",
              "first_page_repaired_sentences": [], "engine": "", "issues": [], "tables": []}
    if suffix in OCR_SUFFIXES:
        result["issues"].append("OCR_REQUIRED: image input, OCR is not wired in this engine")
    else:
        result["issues"].append(f"UNSUPPORTED_SUFFIX:{suffix}")
    return result


COLUMN_TABLE_PAGES = 4      # unruled summary tables are read from the first pages only
PDF_COLUMN_LATER_MAX = 40   # v0106: ... and, when those give no income-statement row, up to this page (keyword-gated)
INCOME_METRICS = ("Revenue", "GrossProfit", "OperatingProfit", "NetIncome", "EPS", "PretaxIncome")
FALLBACK_STRATEGIES = ("COLUMN_ALIGNED", "TRANSPOSED")


def _number_cells(table: Dict[str, Any]) -> set:
    """The numbers printed in a table's body (label column and header left out), as printed minus commas."""
    rows = table.get("Rows") or []
    return {str(v).replace(",", "").strip() for row in rows[1:] for v in list(row)[1:]
            if v and safe_float(str(v)) is not None}


def drop_reread_tables(tables: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """v0106 (批733): the column reader also sees the text of a ruled table -- the fallback runs on any page when the
    ruled tables give no EPS row (the screenshot table has none), and with the quarter headers now read it read that
    table a second time (72 rows for 36 cells).  A fallback table whose numbers are mostly (>= 60 %) numbers of a
    ruled table on the same page is the same table read twice: left out."""
    ruled: Dict[Any, set] = {}
    for t in tables:
        if t.get("Strategy") not in FALLBACK_STRATEGIES:
            ruled.setdefault(t.get("PageNumber"), set()).update(_number_cells(t))
    out = []
    for t in tables:
        if t.get("Strategy") in FALLBACK_STRATEGIES:
            nums = _number_cells(t)
            if nums and len(nums & ruled.get(t.get("PageNumber"), set())) >= 0.6 * len(nums):
                continue
        out.append(t)
    return out


def _has_income_rows(tables: List[Dict[str, Any]]) -> bool:
    """v0106: does any table already give an income-statement row (the trigger for reading later pages)?"""
    try:
        probe = parse_financial_tables_to_records(tables, {"ReportID": "probe", "FileID": "probe"}, "probe")
    except Exception:  # noqa: BLE001 -- an unparseable probe counts as "no rows": the later pages are read
        return False
    return any(r.get("MetricName") in INCOME_METRICS for r in probe)


def extract_document_tables(path: Path, extracted: Optional[Dict[str, Any]] = None) -> List[Dict[str, Any]]:
    if path.suffix.lower() == ".pdf":
        tables = extract_pdf_tables(path)
        # v0102: sell-side summary tables rarely carry ruling lines.  When the ruled tables yield no EPS row,
        # the evidence core reads column-aligned and transposed tables ('12/25E 12/26E', '會計年度 … 每股盈餘').
        try:
            probe = parse_financial_tables_to_records(tables, {"ReportID": "probe", "FileID": "probe"}, "probe")
        except Exception:
            probe = []
        if not any(r.get("MetricName") == "EPS" for r in probe):
            core = optional_import_evidence_core()
            if core is not None:
                try:
                    tables += core.pdf_column_tables(str(path), COLUMN_TABLE_PAGES)
                except Exception as exc:  # graceful: ruled tables stay as they are
                    print(f"[WARN] column tables failed: {path.name} -> {exc}", file=sys.stderr)
                # v0106 (批733 Z195): still no income-statement row -> later pages, the ones whose text looks like one
                if callable(getattr(core, "pdf_column_tables_later", None)) and not _has_income_rows(tables):
                    try:
                        tables += core.pdf_column_tables_later(str(path), COLUMN_TABLE_PAGES, PDF_COLUMN_LATER_MAX)
                    except Exception as exc:  # graceful, as above
                        print(f"[WARN] later-page column tables failed: {path.name} -> {exc}", file=sys.stderr)
        return drop_reread_tables(tables)
    if extracted and extracted.get("tables"):
        return list(extracted["tables"])
    return []


def split_blocks_into_top_zones(blocks: Sequence[Any], width: float, height: float) -> Dict[str, str]:
    left: List[Tuple[float, float, str]] = []
    center: List[Tuple[float, float, str]] = []
    right: List[Tuple[float, float, str]] = []
    top_limit = height * FIRST_PAGE_TOP_RATIO
    for block in blocks:
        try:
            x0, y0, x1, y1, text = float(block[0]), float(block[1]), float(block[2]), float(block[3]), str(block[4])
        except Exception:
            continue
        if y0 > top_limit:
            continue
        cleaned = normalize_unicode_text(text)
        if not cleaned:
            continue
        mid = (x0 + x1) / 2.0
        item = (y0, x0, cleaned)
        if mid <= width * FIRST_PAGE_LEFT_RATIO:
            left.append(item)
        elif mid >= width * FIRST_PAGE_RIGHT_RATIO:
            right.append(item)
        else:
            center.append(item)
    def join(items: List[Tuple[float, float, str]]) -> str:
        items = sorted(items, key=lambda x: (x[0], x[1]))
        return "\n".join(x[2] for x in items)
    return {
        "top_left_text": join(left),
        "top_center_text": join(center),
        "top_right_text": join(right),
    }


def text_to_lines(text: str) -> List[str]:
    raw_lines = str(text or "").splitlines()
    return [normalize_unicode_text(line) for line in raw_lines if normalize_unicode_text(line)]


def repair_sentences_until_period(lines: Sequence[str]) -> List[str]:
    repaired: List[str] = []
    buffer = ""
    for line in lines:
        line = normalize_unicode_text(line)
        if not line:
            continue
        if buffer:
            if should_join_without_space(buffer, line):
                buffer += line
            else:
                buffer += " " + line
        else:
            buffer = line
        if DATE_SENTENCE_END_REGEX.search(buffer) or len(buffer) >= 240:
            repaired.append(buffer.strip())
            buffer = ""
    if buffer:
        repaired.append(buffer.strip())
    return repaired


def should_join_without_space(left: str, right: str) -> bool:
    if not left or not right:
        return False
    return bool(CHINESE_REGEX.search(left[-1]) and CHINESE_REGEX.search(right[0]))

# ============================================================
# def 06_TICKER_BROKER_COMPANY_EXTRACTION
# ============================================================

def extract_filename_ticker_candidates(token_info: Dict[str, Any], ticker_lookup: Dict[str, Dict[str, Any]]) -> List[Dict[str, Any]]:
    candidates: List[Dict[str, Any]] = []
    for token in token_info.get("number_tokens", []):
        if TW_STOCK_TICKER_REGEX.match(token):
            exists = token in ticker_lookup if ticker_lookup else True
            if YEAR_BAND[0] <= int(token) <= YEAR_BAND[1] and not (ticker_lookup and token in ticker_lookup):
                candidates.append({"ticker": "", "raw": token, "source": "filename", "exists_in_ssot": False, "status": "YEAR_BAND"})
                continue
            candidates.append({"ticker": token, "raw": token, "source": "filename", "exists_in_ssot": exists})
        elif ETF_CODE_REGEX.match(token):
            candidates.append({"ticker": "", "raw": token, "source": "filename", "exists_in_ssot": False, "status": "ETF_CANDIDATE"})
    return candidates


def parse_partial_dates_from_filename(stem: str) -> List[Dict[str, Any]]:
    """CTBC0915 -> month/day with YEAR_MISSING; the year is never invented (spec section 2)."""
    out: List[Dict[str, Any]] = []
    for m in re.finditer(r"(?<![0-9])([A-Za-z]{2,6})(0[1-9]|1[0-2])(0[1-9]|[12][0-9]|3[01])(?![0-9])", normalize_unicode_text(stem)):
        out.append({"raw": m.group(0), "month": int(m.group(2)), "day": int(m.group(3)), "status": "YEAR_MISSING"})
    return out


def extract_text_ticker_candidates(text: str, source: str, ticker_lookup: Dict[str, Dict[str, Any]]) -> List[Dict[str, Any]]:
    candidates: List[Dict[str, Any]] = []
    seen = set()
    text = normalize_unicode_text(text)
    for regex, kind in [(TW_YFINANCE_TICKER_REGEX, "yfinance"), (TW_BBG_TICKER_REGEX, "bloomberg"),
                        (TW_PAREN_TICKER_REGEX, "paren"), (TW_TEXT_TICKER_REGEX, "text")]:
        for m in regex.finditer(text):
            ticker = m.group(1)
            if kind in {"text", "paren"} and YEAR_BAND[0] <= int(ticker) <= YEAR_BAND[1] and not (ticker_lookup and ticker in ticker_lookup):
                continue
            key = (ticker, kind, m.group(0))
            if key in seen:
                continue
            seen.add(key)
            exists = ticker in ticker_lookup if ticker_lookup else True
            candidates.append({
                "ticker": ticker,
                "raw": m.group(0),
                "source": source,
                "kind": kind,
                "exists_in_ssot": exists,
            })
    return candidates


def choose_canonical_ticker(filename_candidates: List[Dict[str, Any]], page_candidates: List[Dict[str, Any]], ticker_lookup: Dict[str, Dict[str, Any]]) -> Tuple[str, str, List[str]]:
    issues: List[str] = []
    score: Dict[str, float] = {}
    for c in filename_candidates:
        ticker = c.get("ticker", "")
        if not ticker:
            continue
        score[ticker] = score.get(ticker, 0) + 2.0
        if c.get("exists_in_ssot"):
            score[ticker] += 1.0
    for c in page_candidates:
        ticker = c.get("ticker", "")
        if not ticker:
            continue
        score[ticker] = score.get(ticker, 0) + 2.3
        if c.get("kind") in {"yfinance", "bloomberg"}:
            score[ticker] += 0.4
        if c.get("exists_in_ssot"):
            score[ticker] += 1.0
    if not score:
        return "", "", ["NO_TICKER_CANDIDATE"]
    sorted_scores = sorted(score.items(), key=lambda x: (-x[1], x[0]))
    ticker = sorted_scores[0][0]
    if len(sorted_scores) > 1 and sorted_scores[0][1] - sorted_scores[1][1] < 0.5:
        issues.append(f"TICKER_AMBIGUOUS:{sorted_scores[:3]}")
    if ticker_lookup and ticker not in ticker_lookup:
        issues.append(f"TICKER_NOT_IN_SSOT:{ticker}")
    row = ticker_lookup.get(ticker, {})
    yf = normalize_unicode_text(row.get("YF_TICKER", ""))
    if not yf:
        market = normalize_unicode_text(row.get("Market", ""))
        if market == "TWSE":
            yf = f"{ticker}.TW"
        elif market == "TPEX":
            yf = f"{ticker}.TWO"
    return ticker, yf, issues


def match_broker_from_tokens_and_text(chinese_tokens: Sequence[str], english_tokens: Sequence[str], text: str, broker_alias: Dict[str, List[str]], company_fragments: Sequence[str] = ()) -> Tuple[str, List[str], List[Dict[str, str]]]:
    haystacks = [normalize_unicode_text(x) for x in list(chinese_tokens) + list(english_tokens)]
    long_text = normalize_unicode_text(text)
    fragments = [normalize_unicode_text(f) for f in (company_fragments or [])]
    matches: List[Dict[str, str]] = []
    core = optional_import_evidence_core()      # mother 批728: the one deny gate (overlay deny_keys + CGC_MDL177)
    for broker, aliases in broker_alias.items():
        ordered = [normalize_unicode_text(a) for a in sorted(aliases, key=len, reverse=True) if normalize_unicode_text(a)]
        token_match = None
        for alias in ordered:                       # pass 1: filename tokens, longest alias first
            for idx, h in enumerate(haystacks):
                if h in fragments and h.lower().startswith(alias.lower()):
                    continue        # company fragment after the code: 中信 inside 中信金 is not the broker
                if core is not None and core.deny_shadowed(alias, h):
                    continue        # denied name, or inside a longer denied name (中信 in 中信證券): never a broker
                if alias_hits(alias, h):
                    token_match = {"broker": broker, "alias": alias, "hit": "token", "index": str(idx)}
                    break
            if token_match:
                break
        if token_match:
            matches.append(token_match)
            continue
        for alias in ordered:                       # pass 2: page text, longest alias first
            if core is not None and core.deny_shadowed(alias, long_text):
                continue
            if alias_hits(alias, long_text):
                matches.append({"broker": broker, "alias": alias, "hit": "text", "index": "999"})
                break
    if not matches:
        return "", [], []
    # token hit first, then the longest alias (spec: 優先長別名), then the earliest token
    matches = sorted(matches, key=lambda x: (x["hit"] != "token", -len(x["alias"]), int(x.get("index", 999)), x["broker"]))
    broker = matches[0]["broker"]
    candidates = [m["broker"] for m in matches]
    return broker, sorted(set(candidates)), matches


def find_company_name_candidates(token_info: Dict[str, Any], ticker_row: Dict[str, Any], broker_candidates: Sequence[str]) -> List[str]:
    official_names = [
        normalize_unicode_text(ticker_row.get("Name", "")),
        normalize_unicode_text(ticker_row.get("FullName", "")),
        normalize_unicode_text(ticker_row.get("EnglishName", "")),
    ]
    broker_set = {normalize_unicode_text(x).lower() for x in broker_candidates}
    candidates: List[str] = []
    for token in token_info.get("chinese_tokens", []) + token_info.get("english_tokens", []):
        token_clean = normalize_unicode_text(token)
        if not token_clean:
            continue
        if token_clean.lower() in broker_set:
            continue
        if any(token_clean and (token_clean in n or n in token_clean) for n in official_names if n):
            candidates.append(token_clean)
        elif len(token_clean) >= 2:
            candidates.append(token_clean)
    return sorted(set(candidates), key=lambda x: (-len(x), x))[:8]


def extract_report_title(first_page_text: str, repaired_sentences: Sequence[str]) -> str:
    lines = text_to_lines(first_page_text)
    # 優先取有中文/英文且不是純日期/頁碼的前幾行。
    for line in lines[:12]:
        if len(line) < 4:
            continue
        if re.fullmatch(r"[\d./\- ]+", line):
            continue
        if "http" in line.lower():
            continue
        return line[:240]
    return first_non_empty(repaired_sentences)[:240]


def detect_report_type(filename_text: str, page_text: str) -> str:
    joined = f"{filename_text} {page_text}"
    for canonical, aliases in REPORT_TYPE_ALIASES.items():
        for alias in aliases:
            if normalize_unicode_text(alias).lower() in joined.lower():
                return canonical
    return ""


def detect_language(text: str) -> str:
    text = normalize_unicode_text(text)
    cn = len(re.findall(r"[\u3400-\u4DBF\u4E00-\u9FFF]", text))
    en = len(re.findall(r"[A-Za-z]", text))
    if cn > 0 and en > 0:
        return "mixed"
    if cn > 0:
        return "zh"
    if en > 0:
        return "en"
    return "unknown"

# ============================================================
# def 07_RATING_TARGET_PRICE_EXTRACTION
# ============================================================

def extract_rating_structures(text: str) -> List[Dict[str, str]]:
    """Structured rating cells: (4171,NR_未評等), (6533,N,中立), 投資評等: 買進, Rating: OW."""
    hay = normalize_unicode_text(text)
    found: List[Dict[str, str]] = []
    for m in re.finditer(r"\(\s*([1-9]\d{3})\s*,\s*([A-Za-z]{1,2})\s*[_,]\s*([^)]{1,12})\)", hay):
        code = m.group(2).upper()
        rating = RATING_SHORT_CODES.get(code, "")
        found.append({"ticker": m.group(1), "code": code, "label": m.group(3), "rating": rating, "source": "FILENAME_STRUCTURE"})
    for m in re.finditer(r"(?:投資評等|投資評級|投資建議|評等|評級|rating|recommendation)\s*[:：]\s*([A-Za-z]{1,2}(?![A-Za-z])|[^\s,，;；)]{2,12})", hay, re.IGNORECASE):
        token = m.group(1)
        rating = RATING_SHORT_CODES.get(token.upper(), "") if len(token) <= 2 else ""
        found.append({"code": token, "rating": rating, "source": "RATING_FIELD"})
    return found


def extract_rating(text: str, rating_alias: Dict[str, List[str]]) -> Tuple[str, List[Dict[str, str]]]:
    """Coarse rating (BUY/HOLD/SELL/NOT_RATED). Short codes only in structures; word boundaries for
    ASCII aliases; action words (維持/調升/...) are not ratings; 'Note' is not NOT_RATED."""
    hay = normalize_unicode_text(text)
    structures = [s for s in extract_rating_structures(hay) if s.get("rating")]
    if structures:
        s = structures[0]
        return s["rating"], [{"rating": s["rating"], "alias": s.get("code", ""), "source": s["source"]}]
    matches: List[Dict[str, str]] = []
    for canonical, aliases in rating_alias.items():
        for alias in sorted(aliases, key=len, reverse=True):
            a = normalize_unicode_text(alias)
            if not a or a.upper() in RATING_SHORT_CODES or a in RATING_ACTION_WORDS:
                continue
            if alias_hits(a, hay):
                matches.append({"rating": canonical, "alias": a, "fine": RATING_FINE.get(a.lower(), "")})
                break
    for term, fine in RATING_FINE.items():
        if alias_hits(term, hay) and not any(m["alias"].lower() == term for m in matches):
            base = "BUY" if fine == "STRONG_BUY" else "SELL"
            matches.append({"rating": base, "alias": term, "fine": fine})
    if not matches:
        return "", []
    matches = sorted(matches, key=lambda x: (-len(x["alias"]), x["rating"]))
    return matches[0]["rating"], matches


def extract_rating_change(text: str) -> str:
    hay = normalize_unicode_text(text)
    patterns = {
        "UPGRADE": ["升評", "調升", "Upgrade", "Upgraded"],
        "DOWNGRADE": ["降評", "調降", "Downgrade", "Downgraded"],
        "MAINTAIN": ["維持", "Maintain", "Maintained", "重申", "Reiterate"],
        "INITIATE": ["初評", "初次", "Initiate", "Initiating"],
    }
    for canonical, aliases in patterns.items():
        for alias in aliases:
            if alias.lower() in hay.lower():
                return canonical
    return ""


_TARGET_PRICE_REGEX: Optional[re.Pattern] = None


def target_price_regex() -> re.Pattern:
    """Target-price terms (built-in + SYNONYM_LIBRARY scope) followed by a price with currency/unit context;
    percentages and multiples never count."""
    global _TARGET_PRICE_REGEX
    if _TARGET_PRICE_REGEX is None:
        terms = list(TARGET_PRICE_TERMS)
        for term in load_synonym_scope(find_knowledge_root().get("synonym_library"), "target_price"):
            if term not in terms and term not in ("target", "tgt", "valuation", "目標", "預期價", "目標區間", "fv",
                                                    "base case", "bull case", "bear case"):
                terms.append(term)
        terms.sort(key=len, reverse=True)
        alts = []
        for t in terms:
            if re.search(r"[\u3400-\u4DBF\u4E00-\u9FFF]", t):
                alts.append(re.escape(t))
            else:
                alts.append(r"(?<![A-Za-z])" + re.escape(t) + r"(?![A-Za-z])")
        _TARGET_PRICE_REGEX = re.compile(
            r"(?:" + "|".join(alts) + r")\s*[:：]?\s*(?:\((?:NT\$|TWD)\))?\s*(?:NT\$|NT\s?\$|新台幣|TWD|NTD|US\$|USD)?\s*"
            r"([0-9][0-9,]*(?:\.\d+)?)\s*(?:元|美元)?(?!\s*[%倍xX])", re.IGNORECASE)
    return _TARGET_PRICE_REGEX


def extract_target_price(text: str) -> Optional[float]:
    hay = normalize_unicode_text(text)
    m = target_price_regex().search(hay)
    if m:
        return safe_float(m.group(1))
    return None


def extract_target_price_scenarios(text: str) -> Dict[str, float]:
    """Base/Bull/Bear case prices kept apart; the engine never takes the maximum."""
    hay = normalize_unicode_text(text)
    out: Dict[str, float] = {}
    for m in re.finditer(r"(?<![A-Za-z])(base|bull|bear)\s*case(?![A-Za-z])[^0-9]{0,40}([0-9][0-9,]*(?:\.\d+)?)", hay, re.IGNORECASE):
        value = safe_float(m.group(2))
        if value is not None:
            out[m.group(1).upper()] = value
    return out


def extract_current_price(text: str) -> Optional[float]:
    hay = normalize_unicode_text(text)
    patterns = [
        r"(?<!目標)(?<!Target )(?<!target )(?:最新收盤價|收盤價|最新股價|現價|股價|Current\s*Price|Share\s*Price|Last\s*Price|Close)"
        r"\s*[:：]?\s*(?:\((?:NT\$|TWD)\))?\s*(?:NT\$|NT\s?\$|新台幣|TWD|NTD)?\s*([0-9][0-9,]*(?:\.\d+)?)\s*(?:元)?(?!\s*[%倍xX])",
    ]
    for pat in patterns:
        m = re.search(pat, hay, re.IGNORECASE)
        if m:
            return safe_float(m.group(1))
    return None


def calculate_upside_downside(target: Optional[float], current: Optional[float]) -> Optional[float]:
    """Page arithmetic: target / printed price (v0103 keeps it as UpsidePagePct; the upside itself is ADJ)."""
    if target is None or current is None or current == 0:
        return None
    return round((target / current - 1.0) * 100.0, 4)


def adj_upside_columns(tw_ticker: str, report_date: str, target_price: Optional[float],
                       current_price: Optional[float], db: Optional[str] = None) -> Dict[str, Any]:
    """v0103 (批729): the ADJ CLOSE columns of one report.  UpsideDownsidePct = the upside over the latest adj close
    (VRN_Evidence_Core.adj_basis -> ENG073 adj_quote); CurrentPriceAdj = the printed price x the same adj factor;
    UpsidePagePct = the page arithmetic.  When ADJ is missing the cells stay empty and UpsideState says why."""
    page = calculate_upside_downside(target_price, current_price)
    cols: Dict[str, Any] = {"UpsideDownsidePct": "", "UpsideBasis": "ADJ_LATEST", "UpsideState": "",
                            "UpsidePagePct": page if page is not None else ""}
    for col in ("TargetPriceAdj", "AdjFactor", "AdjFactorDate", "PrevClose", "PrevCloseAdj", "PrevCloseDate",
                "CurrentPriceAdj", "LatestAdjClose", "LatestAdjDate"):
        cols[col] = ""
    core = optional_import_evidence_core()
    if core is None or not hasattr(core, "adj_basis"):
        cols["UpsideState"] = "ADJ_NO_CORE"
        return cols
    try:
        q = core.adj_basis(tw_ticker or None, report_date or None, target_price, db=db, page_price=current_price)
    except Exception as exc:  # graceful by design
        cols["UpsideState"] = f"ADJ_ERROR({exc.__class__.__name__})"
        return cols
    fac = q.get("adj_factor")

    def cell(value: Any) -> Any:
        return "" if value is None else value

    cols.update({
        "UpsideDownsidePct": cell(q.get("upside_adj")),
        "UpsideState": q.get("state") or "",
        "TargetPriceAdj": cell(q.get("target_price_adj")),
        "AdjFactor": cell(round(fac, 6) if fac is not None else None),
        "AdjFactorDate": q.get("adj_factor_date") or "",
        "PrevClose": cell(q.get("price_prev_close")),
        "PrevCloseAdj": cell(q.get("price_prev_adj")),
        "PrevCloseDate": q.get("price_prev_date") or "",
        "CurrentPriceAdj": cell(round(current_price * fac, 4) if (current_price and fac) else None),
        "LatestAdjClose": cell(q.get("price_latest_adj")),
        "LatestAdjDate": q.get("price_latest_date") or "",
    })
    return cols

# ============================================================
# def 08_FINANCIAL_TABLE_EXTRACTION
# ============================================================

def extract_pdf_tables_with_fitz(pdf_path: Path, max_pages: int = PDF_TABLE_MAX_PAGES) -> List[Dict[str, Any]]:
    """PyMuPDF table finder (ruled tables); used when pdfplumber is not installed."""
    tables: List[Dict[str, Any]] = []
    fitz = optional_import_fitz()
    if fitz is None:
        return tables
    try:
        doc = fitz.open(str(pdf_path))
        for page_idx, page in enumerate(doc, start=1):
            if page_idx > max_pages:
                break
            try:
                found = page.find_tables()
            except Exception:
                continue
            for table_idx, table in enumerate(getattr(found, "tables", []) or [], start=1):
                try:
                    cleaned = clean_raw_table(table.extract())
                except Exception:
                    cleaned = []
                if cleaned:
                    tables.append({"PageNumber": page_idx, "TableID": f"P{page_idx:03d}_T{table_idx:03d}", "Rows": cleaned})
        doc.close()
    except Exception as exc:
        print(f"[WARN] PyMuPDF 表格抽取失敗：{pdf_path} -> {exc}", file=sys.stderr)
    return tables


def extract_pdf_tables(pdf_path: Path, max_pages: int = PDF_TABLE_MAX_PAGES) -> List[Dict[str, Any]]:
    tables: List[Dict[str, Any]] = []
    pdfplumber = optional_import_pdfplumber()
    if pdfplumber is None:
        return extract_pdf_tables_with_fitz(pdf_path, max_pages)
    try:
        with pdfplumber.open(str(pdf_path)) as pdf:
            for page_idx, page in enumerate(pdf.pages[:max_pages], start=1):
                try:
                    page_tables = page.extract_tables() or []
                except Exception:
                    page_tables = []
                for table_idx, table in enumerate(page_tables, start=1):
                    cleaned = clean_raw_table(table)
                    if cleaned:
                        tables.append({
                            "PageNumber": page_idx,
                            "TableID": f"P{page_idx:03d}_T{table_idx:03d}",
                            "Rows": cleaned,
                        })
    except Exception as exc:
        print(f"[WARN] 表格抽取失敗：{pdf_path} -> {exc}", file=sys.stderr)
    return tables


def clean_raw_table(table: Any) -> List[List[str]]:
    if not table:
        return []
    rows: List[List[str]] = []
    for row in table:
        if row is None:
            continue
        cleaned = [normalize_unicode_text(cell) for cell in row]
        if any(cell for cell in cleaned):
            rows.append(cleaned)
    return rows


def normalize_metric_name(raw: str) -> str:
    raw_clean = normalize_unicode_text(raw)
    if not raw_clean:
        return ""
    raw_lower = raw_clean.lower().replace(" ", "")
    best = ""
    best_len = 0
    for metric, aliases in METRIC_ALIASES.items():
        for alias in aliases:
            alias_clean = normalize_unicode_text(alias)
            alias_lower = alias_clean.lower().replace(" ", "")
            if alias_lower and alias_lower in raw_lower and len(alias_lower) > best_len:
                best = metric
                best_len = len(alias_lower)
    return best


def infer_unit_and_scale(metric_raw: str, table_context: str) -> Tuple[str, str, float]:
    context = normalize_unicode_text(f"{metric_raw} {table_context}")
    unit = ""
    currency = "TWD"
    scale = 1.0
    if "%" in context or any(x in context for x in ["率", "Margin", "Yield", "ROE", "ROA"]):
        unit = "%"
    elif re.search(r"\b(x|X|倍)\b", context):
        unit = "x"
    elif any(x in context for x in ["百萬", "NT$m", "NTD m", "新台幣百萬元", "台幣百萬元"]):
        unit = "NT$m"
        scale = 1_000_000.0
    elif any(x in context for x in ["十億", "bn", "billion", "NT$bn"]):
        unit = "NT$bn"
        scale = 1_000_000_000.0
    elif any(x in context for x in ["千元", "仟元"]):
        unit = "NT$000"
        scale = 1_000.0
    elif any(x in context for x in ["美元", "USD", "US$"]):
        currency = "USD"
    if not unit:
        unit = "raw"
    return unit, currency, scale


def _period_from_eng074(raw: str) -> Optional[Dict[str, Any]]:
    """v0105 (批732): the mother's reading of one header cell (VRN_ENG074 period_parts, through the evidence core --
    one period parser for the VRN), in this engine's columns; None = the local rules below decide."""
    core = optional_import_evidence_core()
    p = core.period_parts(raw) if core is not None and callable(getattr(core, "period_parts", None)) else None
    if not p:
        return None
    quarter = f"Q{p['fiscal_quarter']}" if p.get("fiscal_quarter") else (f"H{p['half']}" if p.get("half") else "")
    return {"PeriodType": p["period_type"], "FiscalYear": p.get("fiscal_year"), "FiscalQuarter": quarter,
            "EstimateFlag": {"E": "Estimate", "F": "Estimate", "A": "Actual"}.get(str(p.get("estimate") or "").upper(), "")}


def parse_period_label(label: str) -> Dict[str, Any]:
    raw = normalize_unicode_text(label)
    result = {"PeriodType": "", "FiscalYear": None, "FiscalQuarter": "", "EstimateFlag": "", "PeriodLabelRaw": raw}
    if not raw:
        return result
    canon = _period_from_eng074(raw)
    if canon is not None:
        if not canon["EstimateFlag"] and ("預估" in raw or "估" in raw or "FORECAST" in raw.upper() or "ESTIMATE" in raw.upper()):
            canon["EstimateFlag"] = "Estimate"
        result.update(canon)
        return result
    upper = raw.upper()
    if re.search(r"(?<![A-Z])(TTM|LTM)(?![A-Z])", upper):
        result["PeriodType"] = "TTM"
    # estimate / actual flags come from a suffix letter after digits or an explicit word, never from any E/F in the label
    if re.search(r"\d\s*[EF](?![A-Z])", upper) or "預估" in raw or "估" in raw or "FORECAST" in upper or "ESTIMATE" in upper:
        result["EstimateFlag"] = "Estimate"
    elif re.search(r"\d\s*A(?![A-Z])", upper) or "實際" in raw or "ACTUAL" in upper:
        result["EstimateFlag"] = "Actual"
    else:
        result["EstimateFlag"] = ""
    year: Optional[int] = None
    q_match = re.search(r"(?:([1-4])Q(\d{2,4})|Q([1-4])\s?(\d{2,4})?|(\d{4})[.\s]?Q([1-4])|(\d{2,4})年?第?([1-4])季|([1-4])季)", upper)
    if q_match:
        g = q_match.groups()
        if g[0]:
            q, y = g[0], g[1]
        elif g[2]:
            q, y = g[2], g[3]
        elif g[4]:
            y, q = g[4], g[5]
        elif g[6]:
            y, q = g[6], g[7]
        else:
            q, y = g[8], None
        result["FiscalQuarter"] = f"Q{q}"
        result["PeriodType"] = "FQ"
        if y:
            year = int(y)
    mm_yy = re.match(r"^(\d{1,2})/(\d{2})\s*[AEF]?$", upper)
    if year is None and mm_yy:
        year = 2000 + int(mm_yy.group(2))      # 12/24A -> FY2024 (month/year column header)
    if year is None:
        year_match = re.search(r"(?<!\d)((?:19|20)\d{2})(?!\d)", upper) or re.search(r"(?<![\d/])(\d{2})\s*[AEF](?![A-Z])", upper) \
            or re.search(r"^FY\s?(\d{2})(?!\d)", upper) or (re.search(r"^(\d{2})$", upper) if len(upper) == 2 else None)
        if year_match:
            year = int(year_match.group(1))
    if year is not None:
        if year < 100:
            year += 2000
        if 1990 <= year <= 2099:
            result["FiscalYear"] = year
            if not result["PeriodType"]:
                result["PeriodType"] = "FY"
    return result


def parse_financial_tables_to_records(tables: List[Dict[str, Any]], basic_row: Dict[str, Any], run_id: str) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    for table in tables:
        rows = table.get("Rows") or []
        if len(rows) < 2:
            continue
        header_idx = find_best_header_row(rows)
        if header_idx is None:
            continue
        header = rows[header_idx]
        data_rows = rows[header_idx + 1:]
        table_context = " ".join(" ".join(r) for r in rows[: min(len(rows), 3)])
        for r_idx, row in enumerate(data_rows, start=header_idx + 1):
            if not row:
                continue
            metric_raw = normalize_unicode_text(row[0] if len(row) > 0 else "")
            metric = normalize_metric_name(metric_raw)
            if not metric:
                continue
            category = METRIC_CATEGORY_MAP.get(metric, "Other")
            unit, currency, scale = infer_unit_and_scale(metric_raw, table_context)
            max_cols = max(len(row), len(header))
            for c_idx in range(1, max_cols):
                value_raw = normalize_unicode_text(row[c_idx] if c_idx < len(row) else "")
                if not value_raw:
                    continue
                value = safe_float(value_raw)
                if value is None:
                    continue
                period_label = normalize_unicode_text(header[c_idx] if c_idx < len(header) else "")
                period = parse_period_label(period_label)
                source_text = f"{metric_raw} | {period_label} | {value_raw}"
                fin_id = stable_hash([
                    basic_row.get("ReportID", ""),
                    table.get("TableID", ""),
                    r_idx,
                    c_idx,
                    metric,
                    period_label,
                    value_raw,
                ], size=24)
                record = empty_financial_row()
                record.update({
                    "FinancialRowID": fin_id,
                    "FileID": basic_row.get("FileID", ""),
                    "ReportID": basic_row.get("ReportID", ""),
                    "TW_TICKER": basic_row.get("TW_TICKER", ""),
                    "YF_TICKER": basic_row.get("YF_TICKER", ""),
                    "Name": basic_row.get("Name", ""),
                    "ReportDate": basic_row.get("ReportDate", ""),
                    "Broker": basic_row.get("Broker", ""),
                    "PeriodType": period.get("PeriodType", ""),
                    "FiscalYear": period.get("FiscalYear"),
                    "FiscalQuarter": period.get("FiscalQuarter", ""),
                    "PeriodLabelRaw": period.get("PeriodLabelRaw", ""),
                    "EstimateFlag": period.get("EstimateFlag", ""),
                    "MetricCategory": category,
                    "MetricName": metric,
                    "MetricNameRaw": metric_raw,
                    "Value": value,
                    "ValueRaw": value_raw,
                    "Unit": unit,
                    "Currency": currency,
                    "Scale": scale,
                    "PageNumber": table.get("PageNumber", ""),
                    "TableID": table.get("TableID", ""),
                    "RowIndex": r_idx,
                    "ColumnIndex": c_idx,
                    "SourceText": source_text,
                    "ConfidenceScore": 0.72 if period.get("FiscalYear") else 0.62,
                    "ValidationStatus": "PASS" if period.get("FiscalYear") else "REVIEW",
                    "ValidationIssues": "" if period.get("FiscalYear") else "PERIOD_UNCLEAR",
                    "CreatedAt": now_iso(),
                    "UpdatedAt": now_iso(),
                    "RunID": run_id,
                    "SchemaVersion": SCHEMA_VERSION,
                })
                # 對常用寬欄也同步寫值，方便 dashboard 使用；long 欄仍以 MetricName/Value 為主。
                if metric in record:
                    record[metric] = value
                records.append(record)
    return records


IDENTITY_RULES = (
    ("GrossProfit", "Revenue", "COGS", "毛利 = 營收 − 營業成本"),
    ("OperatingProfit", "GrossProfit", "OperatingExpense", "營業利益 = 毛利 − 營業費用"),
)
QUARTER_SUM_METRICS = ("Revenue", "COGS", "GrossProfit", "OperatingExpense", "OperatingProfit", "PretaxIncome",
                       "NetIncome", "InterestIncome")


def _printed_step(value_raw: Any) -> float:
    """The last printed digit of a cell ("61,360" -> 1, "5.12" -> 0.01): the rounding the table itself carries."""
    m = re.search(r"\.(\d+)", str(value_raw or ""))
    return 10.0 ** -len(m.group(1)) if m else 1.0


def financial_identity_checks(records: List[Dict[str, Any]]) -> Dict[str, Any]:
    """v0105 (批732): does one report's extracted table add up?  In each table column: gross profit = revenue - cost
    and operating profit = gross profit - operating expense; in each table: four quarters = the year, for the flow
    items.  Tolerance = half a last printed digit per number involved (the table's own rounding) -- nothing looser.
    A cost printed as a negative number counts by its size.  A check runs only when every number it needs was
    read; nothing is filled in.  Returns counts, the failures in words, and two plain counts that say whether the
    period headers were read: quarter/half rows without a year, and rows marked as a forecast."""
    checks: List[Dict[str, Any]] = []
    columns: Dict[Tuple[Any, Any], Dict[str, Dict[str, Any]]] = {}
    for r in records:
        if r.get("MetricName") and isinstance(r.get("Value"), (int, float)):
            columns.setdefault((r.get("TableID"), r.get("ColumnIndex")), {}).setdefault(r["MetricName"], r)
    for (table, _col), m in sorted(columns.items(), key=lambda kv: (str(kv[0][0]), str(kv[0][1]))):
        for target, left, right, rule in IDENTITY_RULES:
            if target in m and left in m and right in m:
                got, expected = m[target]["Value"], m[left]["Value"] - abs(m[right]["Value"])
                tol = 0.5 * sum(_printed_step(m[k].get("ValueRaw")) for k in (target, left, right)) + 1e-9
                checks.append({"rule": rule, "table": table, "period": m[target].get("PeriodLabelRaw", ""),
                               "got": got, "expected": round(expected, 6), "tol": tol, "ok": abs(got - expected) <= tol})
    groups: Dict[Tuple[Any, str, Any], Dict[str, Dict[str, Any]]] = {}
    for r in records:
        if r.get("MetricName") in QUARTER_SUM_METRICS and r.get("FiscalYear") and isinstance(r.get("Value"), (int, float)):
            key = "FY" if r.get("PeriodType") == "FY" else (r.get("FiscalQuarter") if r.get("PeriodType") == "FQ" else "")
            if key:
                groups.setdefault((r.get("TableID"), r["MetricName"], r["FiscalYear"]), {}).setdefault(key, r)
    for (table, metric, year), g in sorted(groups.items(), key=lambda kv: (str(kv[0][0]), kv[0][1], str(kv[0][2]))):
        if "FY" in g and all(q in g for q in ("Q1", "Q2", "Q3", "Q4")):
            parts = [g[q] for q in ("Q1", "Q2", "Q3", "Q4")]
            total = sum(p["Value"] for p in parts)
            tol = 0.5 * sum(_printed_step(p.get("ValueRaw")) for p in parts + [g["FY"]]) + 1e-9
            checks.append({"rule": f"{metric} 四季和 = 全年", "table": table, "period": str(int(year)),
                           "got": g["FY"]["Value"], "expected": round(total, 6), "tol": tol,
                           "ok": abs(g["FY"]["Value"] - total) <= tol})
    failed = [c for c in checks if not c["ok"]]
    return {"checked": len(checks), "passed": len(checks) - len(failed), "failed": failed,
            "periods_unclear": sum(1 for r in records if r.get("PeriodType") in ("FQ", "FH") and not r.get("FiscalYear")),
            "estimate_rows": sum(1 for r in records if r.get("EstimateFlag") == "Estimate")}


def _plain(x: float) -> str:
    return f"{int(round(x)):,}" if abs(x - round(x)) < 1e-9 else f"{x:,.4f}".rstrip("0").rstrip(".")


def identity_failure_text(check: Dict[str, Any]) -> str:
    return (f"{check['rule']} · {check['period']} · 表 {check['table']}:印 {_plain(check['got'])} · "
            f"算 {_plain(check['expected'])}(差 {_plain(abs(check['got'] - check['expected']))},容許 {_plain(check['tol'])})")


def find_best_header_row(rows: List[List[str]]) -> Optional[int]:
    best_idx = None
    best_score = -1
    for idx, row in enumerate(rows[:6]):
        score = 0
        joined = " ".join(row)
        if re.search(r"(?:19|20)\d{2}|\d{2}[EF](?![A-Za-z])|Q[1-4]|[1-4]Q|\d{1,2}/\d{2}", joined, re.IGNORECASE):
            score += 2
        period_cells = sum(1 for cell in row[1:] if parse_period_label(cell).get("FiscalYear"))
        score += 3 * period_cells
        score += sum(1 for cell in row if normalize_unicode_text(cell))
        score -= 2 * sum(1 for cell in row[1:] if safe_float(cell) is not None and not parse_period_label(cell).get("FiscalYear"))
        if score > best_score:
            best_score = score
            best_idx = idx
    if best_score <= 1 or best_idx is None:
        return None
    if not any(parse_period_label(cell).get("FiscalYear") for cell in rows[best_idx][1:]):
        return None
    return best_idx

# ============================================================
# def 09_BASICINFO_RECORD_BUILDER
# ============================================================

def empty_basic_row() -> Dict[str, Any]:
    return {col: "" for col in BASICINFO_COLUMNS}


def empty_financial_row() -> Dict[str, Any]:
    return {col: "" for col in FINANCIALDATA_COLUMNS}


def build_basicinfo_record(pdf_path: Path, ticker_lookup: Dict[str, Dict[str, Any]], broker_alias: Dict[str, List[str]], rating_alias: Dict[str, List[str]], run_id: str) -> Dict[str, Any]:
    source_filename = pdf_path.name
    source_path = str(pdf_path)
    file_hash = file_sha256(pdf_path)
    file_id = stable_hash([source_filename, file_hash], size=24)
    token_info = tokenize_filename(source_filename)
    pdf_info = extract_document_text_and_zones(pdf_path)
    first_page_raw = str(pdf_info.get("first_page_text", "") or "")
    first_page_text = normalize_unicode_text(first_page_raw)
    company_fragments = company_fragments_after_ticker(token_info.get("tokens", []))
    partial_dates = parse_partial_dates_from_filename(token_info.get("source_stem", ""))
    top_zone_text = "\n".join([
        pdf_info.get("top_left_text", ""),
        pdf_info.get("top_center_text", ""),
        pdf_info.get("top_right_text", ""),
    ])
    filename_ticker_candidates = extract_filename_ticker_candidates(token_info, ticker_lookup)
    page_ticker_candidates = extract_text_ticker_candidates(first_page_text + "\n" + top_zone_text, "first_page", ticker_lookup)
    tw_ticker, yf_ticker, ticker_issues = choose_canonical_ticker(filename_ticker_candidates, page_ticker_candidates, ticker_lookup)
    ticker_row = ticker_lookup.get(tw_ticker, {}) if tw_ticker else {}
    filename_date_candidates = parse_date_candidates_from_number_tokens(token_info.get("number_tokens", []))
    report_date, all_date_candidates = choose_best_report_date(filename_date_candidates, first_page_text)
    broker, broker_candidates, broker_matches = match_broker_from_tokens_and_text(
        token_info.get("chinese_tokens", []),
        token_info.get("english_tokens", []),
        top_zone_text + "\n" + first_page_text,
        broker_alias,
        company_fragments,
    )
    company_candidates = find_company_name_candidates(token_info, ticker_row, broker_candidates + company_fragments)
    repaired_sentences = pdf_info.get("first_page_repaired_sentences", []) or []
    report_title = extract_report_title(first_page_raw, repaired_sentences)
    report_type = detect_report_type(token_info.get("normalized", ""), first_page_text)
    language = detect_language(first_page_text or token_info.get("normalized", ""))
    rating_text = token_info.get("source_stem", "") + "\n" + first_page_text + "\n" + top_zone_text
    rating, rating_matches = extract_rating(rating_text, rating_alias)
    rating_fine = next((m.get("fine", "") for m in rating_matches if m.get("fine")), "")
    rating_change = extract_rating_change(first_page_text + "\n" + token_info.get("normalized", ""))
    target_price = extract_target_price(first_page_text + "\n" + top_zone_text)
    target_scenarios = extract_target_price_scenarios(first_page_text + "\n" + top_zone_text)
    current_price = extract_current_price(first_page_text + "\n" + top_zone_text)
    tri_code = tri_code_verdict(filename_ticker_candidates, page_ticker_candidates)
    first_page_engine = run_first_page_engine(pdf_path, source_filename)
    fn_ticker = next((c.get("ticker", "") for c in filename_ticker_candidates if c.get("ticker")), "")
    fn_date = filename_date_candidates[0].get("date", "") if filename_date_candidates else ""
    fn_broker, _fn_cands, _fn_matches = match_broker_from_tokens_and_text(
        token_info.get("chinese_tokens", []), token_info.get("english_tokens", []), "", broker_alias, company_fragments)
    ev = evidence_from_core(pdf_path, source_filename, first_page_engine, fn_ticker, fn_date, fn_broker, broker_alias)
    if ev and not ev.get("error"):
        # v0102: one rule set with the first-page engine; the v0101 values stay in EvidenceJson (append-only)
        legacy = {"tw_ticker": tw_ticker, "report_date": report_date, "broker": broker, "rating": rating,
                  "target_price": target_price, "current_price": current_price, "report_type": report_type, "tri_code": tri_code}
        report_type = ev.get("type") or report_type
        if report_type == "STOCK":
            if ev.get("ticker"):
                tw_ticker = ev["ticker"]
        else:
            tw_ticker = ""                          # NOT_APPLICABLE: industry / market / macro / ETF / event report
        ticker_row = ticker_lookup.get(tw_ticker, {}) if tw_ticker else {}
        page_yf = next((c.get("raw", "") for c in ev.get("page_codes", []) if c.get("core") == tw_ticker and c.get("kind") == "yfinance"), "")
        yf_ticker = (page_yf.upper() if page_yf else "") or ticker_row.get("YF_TICKER", "") if tw_ticker else ""
        broker = ev.get("broker") or ""
        rating = ev.get("rating") or ""
        target_price = ev.get("tp")
        current_price = ev.get("cp")
        report_date = ev.get("page_date") or report_date
        tri_code = ev.get("tri") or tri_code
    else:
        legacy = {}
        if not broker and first_page_engine.get("broker"):
            broker = canonical_broker_key(first_page_engine["broker"])
        if not rating and (first_page_engine.get("rating") or {}).get("canonical"):
            rating = first_page_engine["rating"]["canonical"]
        if target_price is None and (first_page_engine.get("target_prices") or {}).get("primary") is not None:
            target_price = first_page_engine["target_prices"]["primary"]
    confidence, status, validation_issues = score_basicinfo_confidence(
        filename_ticker_candidates,
        page_ticker_candidates,
        ticker_row,
        report_date,
        broker,
        rating,
        target_price,
        report_title,
        ticker_issues,
        pdf_info.get("issues", []),
    )
    report_id = stable_hash([
        tw_ticker,
        report_date,
        broker,
        source_filename,
        file_hash[:16],
    ], size=24)
    evidence = {
        "filename_tokens": token_info,
        "filename_ticker_candidates": filename_ticker_candidates,
        "first_page_ticker_candidates": page_ticker_candidates,
        "date_candidates": all_date_candidates,
        "broker_matches": broker_matches,
        "rating_matches": rating_matches,
        "company_candidates": company_candidates,
        "first_page_left_block": pdf_info.get("top_left_text", ""),
        "first_page_center_block": pdf_info.get("top_center_text", ""),
        "first_page_right_block": pdf_info.get("top_right_text", ""),
        "first_page_repaired_sentences": repaired_sentences,
        "twse_tpex_match": ticker_row,
        "pdf_engine": pdf_info.get("engine", ""),
        "pdf_issues": pdf_info.get("issues", []),
    }
    row = empty_basic_row()
    row.update({
        "FileID": file_id,
        "ReportID": report_id,
        "SourceFileName": source_filename,
        "SourcePath": source_path,
        "FileHash": file_hash,
        "TW_TICKER": tw_ticker,
        "YF_TICKER": yf_ticker or ticker_row.get("YF_TICKER", ""),
        "Market": ticker_row.get("Market", ""),
        "Name": ticker_row.get("Name", ""),
        "CompanyNameCandidate": " | ".join(company_candidates),
        "Broker": broker,
        "BrokerCandidate": " | ".join(broker_candidates),
        "ReportDate": report_date,
        "ReportTitle": report_title,
        "ReportType": report_type,
        "Language": language,
        "Rating": rating,
        "RatingChange": rating_change,
        "TargetPrice": target_price if target_price is not None else "",
        "CurrentPrice": current_price if current_price is not None else "",
        "UpsideDownsidePct": "",          # v0103: the ADJ upside, filled by adj_upside_columns() below
        "InvestmentThesis": first_non_empty(repaired_sentences[:2])[:800],
        "SourceFormat": pdf_path.suffix.lower().lstrip("."),
        "ReportDateDisplay": report_date.replace("-", "/") if report_date else "",
        "TriCodeVerdict": tri_code,
        "RatingFine": rating_fine,
        "RatingAction": rating_change,
        "BrokerDisplay": BROKER_DISPLAY.get(broker, ""),
        "TargetPriceScenarios": safe_json_dumps(target_scenarios) if target_scenarios else "",
        "PartialDates": safe_json_dumps(partial_dates) if partial_dates else "",
        "FirstPageEngine": safe_json_dumps(first_page_engine) if first_page_engine else "",
        "Industry": ticker_row.get("Industry", ""),
        "Chairman": ticker_row.get("Chairman", ""),
        "GeneralManager": ticker_row.get("GeneralManager", ""),
        "Spokesperson": ticker_row.get("Spokesperson", ""),
        "ConfidenceScore": confidence,
        "ValidationStatus": status,
        "ValidationIssues": " | ".join(validation_issues),
        "EvidenceJson": safe_json_dumps(evidence),
        "CreatedAt": now_iso(),
        "UpdatedAt": now_iso(),
        "RunID": run_id,
        "SchemaVersion": SCHEMA_VERSION,
    })
    # v0103 (批729): the upside is on the latest ADJ close; the page arithmetic stays as UpsidePagePct
    row.update(adj_upside_columns(tw_ticker, report_date, target_price, current_price))
    if ev and not ev.get("error"):
        names = [a.get("name", "") for a in ev.get("analysts", []) if a.get("name")]
        row.update({
            "Analyst": " | ".join(names),
            "AnalystEmail": " | ".join(a.get("email", "") for a in ev.get("analysts", []) if a.get("email")),
            "RatingFine": ev.get("rating_fine") or row.get("RatingFine", ""),
            "RatingAction": ev.get("rating_action") or row.get("RatingAction", ""),
            "RatingKey": ev.get("rating_key") or "",
            "BrokerTier": ev.get("broker_tier") or "",
            "BrokerDenied": " | ".join(sorted(str(k) for k in (ev.get("denied") or {}) if k)),
            "PageDate": ev.get("page_date") or "",
            "FilenameDate": fn_date or "",
            "DateConflict": "Y" if ev.get("date_conflict") else "",
            "CompanyNamePage": ev.get("company_page") or "",
            "PageCodes": " | ".join(sorted({c.get("core", "") for c in ev.get("page_codes", []) if c.get("core")})),
            "EvidenceCore": safe_json_dumps({"source": ev.get("source"), "rating_raw": ev.get("rating_raw"),
                                             "broker_vs_filename": ev.get("broker_vs_filename"), "legacy_v0101": legacy}),
        })
        if ev.get("company_page") and ev["company_page"] not in row.get("CompanyNameCandidate", ""):
            row["CompanyNameCandidate"] = " | ".join(x for x in [ev["company_page"], row.get("CompanyNameCandidate", "")] if x)
    elif ev:
        row["EvidenceCore"] = safe_json_dumps({"error": ev.get("error")})
    return row


def evidence_from_core(path: Path, source_filename: str, first_page_engine: Dict[str, Any],
                       filename_ticker: str, filename_date: str, filename_broker: str,
                       broker_alias: Dict[str, List[str]]) -> Dict[str, Any]:
    """The shared evidence core's answer for this report: from the first-page hook when it ran (v0103
    engines carry it), else straight from VRN_Evidence_Core (the F2 fix turns the hook off).  {} if absent."""
    fp = first_page_engine or {}
    hook_had_text = bool((fp.get("evidence_core") or {}).get("text_layer"))
    if fp.get("report_type") and not fp.get("error") and hook_had_text:
        rating = fp.get("rating") or {}
        tps = fp.get("target_prices") or {}
        return {
            "source": "FIRST_PAGE_HOOK",
            "type": (fp.get("report_type") or {}).get("type", ""),
            "ticker": (fp.get("ticker_page") or {}).get("ticker"),
            "tri": (fp.get("tri_code") or {}).get("verdict", ""),
            "broker": fp.get("broker") or None,
            "broker_tier": fp.get("broker_tier") or "",
            "broker_vs_filename": fp.get("broker_vs_filename") or "",
            "denied": fp.get("broker_denied") or {},
            "rating": rating.get("canonical"), "rating_fine": rating.get("fine"), "rating_key": fp.get("rating_key"),
            "rating_raw": fp.get("rating_raw"), "rating_action": rating.get("action"),
            "tp": tps.get("primary"), "cp": fp.get("current_price"),
            "page_date": fp.get("page_date"), "date_conflict": bool(fp.get("date_conflict")),
            "analysts": fp.get("analysts") or [], "company_page": fp.get("company_name_page"),
            "page_codes": fp.get("page_codes_core") or [],
        }
    core = optional_import_evidence_core()
    if core is None:
        return {}
    try:
        got = core.document_first_page_lines(str(path))
        a = core.analyze(source_filename, got["lines"], alias_table=broker_alias, filename_ticker=filename_ticker or None,
                         filename_date=filename_date or None, filename_broker=filename_broker or None,
                         text_layer=got["text_layer"])
    except Exception as exc:
        return {"error": f"{exc.__class__.__name__}: {exc}"}
    denied = a.get("broker_vs_filename") == "DENIED"
    return {
        "source": "CORE_DIRECT",
        "type": a["report_type"]["type"], "ticker": a["ticker"].get("ticker"), "tri": a["tri_code"].get("verdict", ""),
        "broker": a["broker"].get("broker") or (None if denied else (filename_broker or None)),
        "broker_tier": a["broker"].get("tier", ""), "broker_vs_filename": a.get("broker_vs_filename", ""),
        "denied": a["broker"].get("denied") or ({filename_broker: 1} if denied else {}),
        "rating": a["rating"].get("value"), "rating_fine": a["rating"].get("fine"), "rating_key": a["rating"].get("canonical_key"),
        "rating_raw": a["rating"].get("raw"), "rating_action": a["rating"].get("action"),
        "tp": a["target_price"].get("value"), "cp": a["current_price"].get("value"),
        "page_date": a["page_date"].get("iso"),
        "date_conflict": bool(a["page_date"].get("iso") and filename_date and a["page_date"].get("iso") != filename_date),
        "analysts": a.get("analysts") or [], "company_page": a.get("company_name"), "page_codes": a.get("page_codes") or [],
    }


def tri_code_verdict(filename_candidates: List[Dict[str, Any]], page_candidates: List[Dict[str, Any]]) -> str:
    """Spec section 2: PASS / MISMATCH / INSUFFICIENT_EVIDENCE / N/A from filename core vs page platform codes."""
    filename_cores = {c["ticker"] for c in filename_candidates if c.get("ticker")}
    page_cores = {c["ticker"] for c in page_candidates if c.get("ticker") and c.get("kind") in {"yfinance", "bloomberg", "paren"}}
    if not filename_cores and not page_cores:
        return "N/A"
    if not filename_cores or not page_cores:
        return "INSUFFICIENT_EVIDENCE"
    if filename_cores == page_cores or (len(filename_cores) == 1 and filename_cores <= page_cores):
        return "PASS"
    return "MISMATCH"


FIRST_PAGE_ENGINE_ENABLED = True


def run_first_page_engine(path: Path, source_filename: str) -> Dict[str, Any]:
    """Optional hook into VIA_VRN_FirstPageEngine (same directory); a compact evidence dict or {}."""
    if not FIRST_PAGE_ENGINE_ENABLED:
        return {}
    module = optional_import_first_page_engine()
    if module is None:
        return {}
    try:
        engine = _FIRST_PAGE_MODULE_CACHE.get("engine")
        if engine is None or _FIRST_PAGE_MODULE_CACHE.get("engine_module") is not module:
            engine = module.FirstPageEngine()
            _FIRST_PAGE_MODULE_CACHE["engine"] = engine
            _FIRST_PAGE_MODULE_CACHE["engine_module"] = module
        chars, size = None, (None, None)
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            got = module._chars_from_pdf(str(path))
        elif suffix == ".docx":
            got = module._chars_from_docx(str(path))
        else:
            got = None
        if got is not None:
            chars, size = got
        # 批729: the file path lets the first-page engine read the appendix small print (issuer when page 1 is weak)
        result = engine.run(source_filename, chars=chars, page_width=size[0], page_height=size[1],
                            source_path=str(path) if suffix == ".pdf" else None)
        return {
            "engine": result.get("engine", ""),
            "ticker": result.get("ticker", {}),
            "broker": result.get("broker", ""),
            "broker_tier": result.get("broker_tier"),
            "broker_conflict": result.get("broker_conflict"),
            "rating": {k: result.get("rating", {}).get(k) for k in ("canonical", "fine", "scope", "action")},
            "target_prices": result.get("target_prices", {}),
            "current_price": result.get("current_price"),
            "upside": result.get("upside"),           # v0103 first-page engine v0104: ADJ basis (latest adj close)
            "upside_page": result.get("upside_page"),
            "tri_code": result.get("tri_code", {}),
            "report_date": result.get("report_date", ""),
            "company_name": (result.get("layout") or {}).get("company_name", ""),
            "xv_filename_vs_page": result.get("xv_filename_vs_page", {}),
            # v0103 first-page engine: the shared evidence core's answers (absent on older engines)
            "report_type": result.get("report_type"),
            "ticker_page": result.get("ticker_page"),
            "broker_tier": result.get("broker_tier"),
            "broker_vs_filename": result.get("broker_vs_filename"),
            "broker_denied": result.get("broker_denied"),
            "rating_key": (result.get("rating") or {}).get("canonical_key"),
            "rating_raw": (result.get("rating") or {}).get("raw"),
            "page_date": result.get("page_date"),
            "date_conflict": result.get("date_conflict"),
            "analysts": result.get("analysts"),
            "company_name_page": result.get("company_name_page"),
            "page_codes_core": result.get("page_codes_core"),
            "evidence_core": result.get("evidence_core"),
            "filename_fields": result.get("filename_fields"),
        }
    except Exception as exc:  # graceful by design
        return {"error": f"{exc.__class__.__name__}: {exc}"}


def score_basicinfo_confidence(filename_tickers: List[Dict[str, Any]], page_tickers: List[Dict[str, Any]], ticker_row: Dict[str, Any], report_date: str, broker: str, rating: str, target_price: Optional[float], report_title: str, ticker_issues: Sequence[str], pdf_issues: Sequence[str]) -> Tuple[float, str, List[str]]:
    score = 0.0
    issues: List[str] = []
    if filename_tickers:
        score += CONFIDENCE_WEIGHTS["filename_ticker"]
    else:
        issues.append("NO_FILENAME_TICKER")
    if page_tickers:
        score += CONFIDENCE_WEIGHTS["first_page_ticker"]
    else:
        issues.append("NO_FIRST_PAGE_TICKER")
    if ticker_row and ticker_row.get("Name"):
        score += CONFIDENCE_WEIGHTS["exchange_name"]
    else:
        issues.append("NO_EXCHANGE_NAME_MATCH")
    if report_date:
        score += CONFIDENCE_WEIGHTS["report_date"]
    else:
        issues.append("NO_REPORT_DATE")
    if broker:
        score += CONFIDENCE_WEIGHTS["broker"]
    else:
        issues.append("NO_BROKER")
    if rating or target_price is not None:
        score += CONFIDENCE_WEIGHTS["rating_or_target"]
    else:
        issues.append("NO_RATING_OR_TARGET_PRICE")
    if report_title:
        score += CONFIDENCE_WEIGHTS["title_or_text"]
    else:
        issues.append("NO_REPORT_TITLE")
    issues.extend(ticker_issues)
    issues.extend(str(x) for x in pdf_issues)
    score = round(min(score, 1.0), 4)
    if score >= 0.78 and not any(x.startswith("TICKER_AMBIGUOUS") for x in issues):
        status = "PASS"
    elif score >= 0.45:
        status = "REVIEW"
    else:
        status = "FAIL"
    return score, status, issues

# ============================================================
# def 10_DATAFRAME_SCHEMA_AND_UPSERT
# ============================================================

def make_dataframe(rows: List[Dict[str, Any]], columns: Sequence[str]):
    pd = require_pandas()
    if not rows:
        return pd.DataFrame(columns=list(columns))
    df = pd.DataFrame(rows)
    for col in columns:
        if col not in df.columns:
            df[col] = ""
    return df[list(columns)]


def enforce_basicinfo_schema(df):
    pd = require_pandas()
    if df is None:
        return pd.DataFrame(columns=BASICINFO_COLUMNS)
    for col in BASICINFO_COLUMNS:
        if col not in df.columns:
            df[col] = ""
    return df[BASICINFO_COLUMNS]


def enforce_financialdata_schema(df):
    pd = require_pandas()
    if df is None:
        return pd.DataFrame(columns=FINANCIALDATA_COLUMNS)
    for col in FINANCIALDATA_COLUMNS:
        if col not in df.columns:
            df[col] = ""
    return df[FINANCIALDATA_COLUMNS]


def read_existing_parquet(path: Path, columns: Sequence[str]):
    pd = require_pandas()
    if not path.exists():
        return pd.DataFrame(columns=list(columns))
    try:
        df = pd.read_parquet(path)
        for col in columns:
            if col not in df.columns:
                df[col] = ""
        return df[list(columns)]
    except Exception as exc:
        backup_path = path.with_suffix(path.suffix + f".read_failed.{int(time.time())}.bak")
        try:
            path.rename(backup_path)
        except Exception:
            pass
        print(f"[WARN] 舊 Parquet 讀取失敗，已嘗試備份：{path} -> {backup_path}，錯誤：{exc}", file=sys.stderr)
        return pd.DataFrame(columns=list(columns))


def cumulative_upsert(existing_df, new_df, key_columns: Sequence[str], columns: Sequence[str]):
    pd = require_pandas()
    existing_df = existing_df.copy() if existing_df is not None else pd.DataFrame(columns=list(columns))
    new_df = new_df.copy() if new_df is not None else pd.DataFrame(columns=list(columns))
    for col in columns:
        if col not in existing_df.columns:
            existing_df[col] = ""
        if col not in new_df.columns:
            new_df[col] = ""
    if new_df.empty:
        return existing_df[list(columns)]
    frames = [frame[list(columns)] for frame in (existing_df, new_df) if len(frame)]
    combined = pd.concat(frames, ignore_index=True) if frames else existing_df[list(columns)]
    for key in key_columns:
        if key not in combined.columns:
            combined[key] = ""
    # 保留最後一次抽取結果，但不覆蓋不同 ReportID。
    combined = combined.drop_duplicates(subset=list(key_columns), keep="last")
    return combined[list(columns)]


def coerce_frame_types(df, numeric_columns: Sequence[str]):
    """One type per column: numeric columns become float (NaN when blank), the rest become text."""
    pd = require_pandas()
    out = df.copy()
    for col in out.columns:
        if col in numeric_columns:
            out[col] = pd.to_numeric(out[col].replace({"": None}), errors="coerce").astype("float64")
        else:
            out[col] = out[col].map(lambda v: "" if v is None or (isinstance(v, float) and math.isnan(v)) else str(v)).astype("string")
    return out


def write_parquet_mandatory(df, path: Path, numeric_columns: Sequence[str] = ()) -> None:
    if not optional_import_pyarrow_available():
        raise RuntimeError("Parquet 是必選輸出，但缺少 pyarrow。請先安裝：pip install pyarrow")
    if not numeric_columns:
        numeric_columns = BASICINFO_NUMERIC_COLUMNS if "ReportTitle" in df.columns else FINANCIALDATA_NUMERIC_COLUMNS
    coerce_frame_types(df, numeric_columns).to_parquet(path, index=False)

# ============================================================
# def 11_OPTIONAL_OUTPUTS
# ============================================================

def write_csv_optional(df, path: Path) -> None:
    df.to_csv(path, index=False, encoding=CSV_ENCODING_FOR_EXCEL)


def write_json_optional(df, path: Path) -> None:
    pd = require_pandas()
    records = df.astype(object).where(pd.notna(df), None).to_dict(orient="records")
    with path.open("w", encoding=JSON_ENCODING) as f:
        json.dump(records, f, ensure_ascii=False, indent=2, default=str)


def write_duckdb_optional(output_dir: Path, basic_parquet: Path, financial_parquet: Path) -> Optional[Path]:
    duckdb = optional_import_duckdb()
    if duckdb is None:
        print("[WARN] 未安裝 duckdb，略過 DuckDB 輸出。pip install duckdb", file=sys.stderr)
        return None
    db_path = output_dir / DUCKDB_NAME
    con = duckdb.connect(str(db_path))
    con.execute("CREATE OR REPLACE TABLE StockReportBasicInfo AS SELECT * FROM read_parquet(?)", [str(basic_parquet)])
    con.execute("CREATE OR REPLACE TABLE StockReportFinancialData AS SELECT * FROM read_parquet(?)", [str(financial_parquet)])
    con.execute("CREATE OR REPLACE VIEW v_report_financial AS SELECT b.*, f.MetricName, f.Value, f.ValueRaw, f.PeriodType, f.FiscalYear, f.FiscalQuarter, f.EstimateFlag, f.PageNumber, f.TableID FROM StockReportBasicInfo b LEFT JOIN StockReportFinancialData f USING (ReportID)")
    con.close()
    return db_path


def write_google_sheet_optional(basic_df, financial_df, spreadsheet_name: str, credentials_json: str) -> bool:
    if not spreadsheet_name or not credentials_json:
        return False
    try:
        import gspread  # type: ignore
        from gspread_dataframe import set_with_dataframe  # type: ignore
        gc = gspread.service_account(filename=credentials_json)
        try:
            sh = gc.open(spreadsheet_name)
        except Exception:
            sh = gc.create(spreadsheet_name)
        def upsert_ws(title: str, df):
            try:
                ws = sh.worksheet(title)
                ws.clear()
            except Exception:
                ws = sh.add_worksheet(title=title, rows=max(len(df) + 10, 100), cols=max(len(df.columns) + 5, 20))
            set_with_dataframe(ws, df)
        upsert_ws("StockReportBasicInfo", basic_df)
        upsert_ws("StockReportFinancialData", financial_df)
        return True
    except Exception as exc:
        print(f"[WARN] Google Sheet 輸出失敗：{exc}", file=sys.stderr)
        return False

# ============================================================
# def 12_HTML_MATRIX_AND_SUMMARY
# ============================================================

def write_batch_matrix_csv(path: Path, rows: List[Dict[str, Any]]) -> None:
    fieldnames = [
        "SourceFileName", "TW_TICKER", "Name", "Broker", "ReportDate", "Rating", "TargetPrice",
        "ConfidenceScore", "ValidationStatus", "ValidationIssues", "FinancialRows",
    ]
    with path.open("w", encoding=CSV_ENCODING_FOR_EXCEL, newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in fieldnames})


def write_html_matrix(path: Path, basic_rows: List[Dict[str, Any]], financial_counts: Dict[str, int], summary: Dict[str, Any]) -> None:
    status_counts: Dict[str, int] = {}
    for row in basic_rows:
        status = normalize_unicode_text(row.get("ValidationStatus", "UNKNOWN")) or "UNKNOWN"
        status_counts[status] = status_counts.get(status, 0) + 1
    body_rows = []
    for row in basic_rows:
        report_id = row.get("ReportID", "")
        fin_count = financial_counts.get(report_id, 0)
        status = row.get("ValidationStatus", "")
        cls = "pass" if status == "PASS" else "review" if status == "REVIEW" else "fail"
        body_rows.append(
            "<tr class='{cls}'>".format(cls=cls)
            + f"<td>{html.escape(str(row.get('SourceFileName', '')))}</td>"
            + f"<td>{html.escape(str(row.get('TW_TICKER', '')))}</td>"
            + f"<td>{html.escape(str(row.get('Name', '')))}</td>"
            + f"<td>{html.escape(str(row.get('Broker', '')))}</td>"
            + f"<td>{html.escape(str(row.get('ReportDate', '')))}</td>"
            + f"<td>{html.escape(str(row.get('Rating', '')))}</td>"
            + f"<td>{html.escape(str(row.get('TargetPrice', '')))}</td>"
            + f"<td>{html.escape(str(row.get('ConfidenceScore', '')))}</td>"
            + f"<td>{html.escape(str(status))}</td>"
            + f"<td>{fin_count}</td>"
            + f"<td>{html.escape(str(row.get('ValidationIssues', ''))[:400])}</td>"
            + "</tr>"
        )
    doc = f"""<!doctype html>
<html lang="zh-Hant">
<head>
<meta charset="utf-8" />
<title>VRN Batch Report</title>
<style>
body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; background:#fbfaf7; color:#1f2937; margin:24px; font-size:13px; }}
h1 {{ margin:0 0 4px 0; font-size:22px; }}
.sub {{ color:#6b7280; margin-bottom:18px; }}
.cards {{ display:grid; grid-template-columns: repeat(5, minmax(120px,1fr)); gap:10px; margin-bottom:16px; }}
.card {{ background:#fff; border:1px solid #e5e7eb; border-radius:12px; padding:12px; }}
.card b {{ display:block; font-size:20px; margin-top:4px; }}
table {{ width:100%; border-collapse:collapse; background:#fff; border:1px solid #e5e7eb; }}
th, td {{ border-bottom:1px solid #e5e7eb; padding:8px; text-align:left; vertical-align:top; }}
th {{ background:#f3f4f6; position:sticky; top:0; }}
tr.pass td:first-child {{ border-left:5px solid #16a34a; }}
tr.review td:first-child {{ border-left:5px solid #f59e0b; }}
tr.fail td:first-child {{ border-left:5px solid #dc2626; }}
.small {{ font-size:12px; color:#6b7280; }}
</style>
</head>
<body>
<h1>VERITAS INTELLIGENCE ANALYTICS</h1>
<div class="sub">VeritasReportNova · Stock Report Database Matrix · {html.escape(str(summary.get('RunID','')))}</div>
<div class="cards">
  <div class="card">PDF processed<b>{summary.get('ProcessedFiles', 0)}</b></div>
  <div class="card">Basic rows<b>{summary.get('BasicRows', 0)}</b></div>
  <div class="card">Financial rows<b>{summary.get('FinancialRows', 0)}</b></div>
  <div class="card">PASS<b>{status_counts.get('PASS', 0)}</b></div>
  <div class="card">REVIEW/FAIL<b>{status_counts.get('REVIEW', 0) + status_counts.get('FAIL', 0)}</b></div>
</div>
<table>
<thead><tr><th>File</th><th>Ticker</th><th>Name</th><th>Broker</th><th>Date</th><th>Rating</th><th>Target</th><th>Score</th><th>Status</th><th>FinRows</th><th>Issues</th></tr></thead>
<tbody>
{''.join(body_rows)}
</tbody>
</table>
<p class="small">Parquet is the mandatory source of truth. CSV / JSON / DuckDB / Google Sheet are optional derived outputs.</p>
</body>
</html>"""
    path.write_text(doc, encoding="utf-8")

# ============================================================
# def 13_PROCESSING_PIPELINE
# ============================================================

def gate_fallback_tables(tables: List[Dict[str, Any]], records: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], List[str]]:
    """v0106 (批733 Z195): a table read without ruling lines (column-aligned / transposed) is stored only if it adds
    up -- the identities of financial_identity_checks, per table.  Returns (the records kept, one line per table
    left out).  A table with nothing to check is kept (nothing says it is wrong); ruled tables are not gated."""
    fallback = sorted({t.get("TableID") for t in tables if t.get("Strategy") in FALLBACK_STRATEGIES})
    drop, rejected = set(), []
    for tid in fallback:
        check = financial_identity_checks([r for r in records if r.get("TableID") == tid])
        if check["failed"]:
            drop.add(tid)
            more = f"(另 {len(check['failed']) - 1} 條)" if len(check["failed"]) > 1 else ""
            rejected.append(f"{tid} " + identity_failure_text(check["failed"][0]) + more)
    return [r for r in records if r.get("TableID") not in drop], rejected


def process_one_pdf(pdf_path: Path, ticker_lookup: Dict[str, Dict[str, Any]], broker_alias: Dict[str, List[str]], rating_alias: Dict[str, List[str]], run_id: str) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    basic = build_basicinfo_record(pdf_path, ticker_lookup, broker_alias, rating_alias, run_id)
    extracted = None
    if pdf_path.suffix.lower() != ".pdf":
        extracted = extract_document_text_and_zones(pdf_path)
    tables = extract_document_tables(pdf_path, extracted)
    financial = parse_financial_tables_to_records(tables, basic, run_id)
    financial, rejected = gate_fallback_tables(tables, financial)      # v0106 (批733 Z195)
    if rejected:
        issues = [x for x in str(basic.get("ValidationIssues") or "").split(" | ") if x]
        basic["ValidationIssues"] = " | ".join(issues + ["FIN_TABLE_REJECTED(" + " ; ".join(rejected) + ")"])
    return basic, financial


process_one_document = process_one_pdf


def batch_pool_results(files: List[Path], ticker_lookup: Dict[str, Dict[str, Any]], broker_alias: Dict[str, List[str]],
                       rating_alias: Dict[str, List[str]], run_id: str) -> Dict[int, Tuple[Any, ...]]:
    """v0104 (批731): the reports of one batch in worker processes (VRN_BatchPool, sized by the accelerator's thread
    budget).  {idx: (status, basic | error, financial)} in the batch numbering, progress lines printed as reports
    finish; {} = everything runs in this process (small batch, VRN_BATCH_WORKERS=1, or the pool could not run --
    then the unfinished files go through the sequential loop and one line says why)."""
    todo = [(idx, p) for idx, p in enumerate(files, start=1) if p.suffix.lower() not in OCR_SUFFIXES]
    here = str(Path(__file__).resolve().parent)
    try:
        if here not in sys.path:
            sys.path.insert(0, here)                 # the workers import the pool by this stable name
        import VRN_BatchPool as pool
    except ImportError:
        return {}
    workers = pool.workers_for(len(todo))
    if workers <= 1:
        return {}
    done: Dict[int, Tuple[Any, ...]] = {}
    count = [0]

    def on_done(idx: int) -> None:
        count[0] += 1
        print(f"[{count[0]}/{len(files)}] VRN processing: {files[idx - 1].name}")

    for idx, p in enumerate(files, start=1):         # image files are skipped at once; number them first
        if p.suffix.lower() in OCR_SUFFIXES:
            done[idx] = ("OCR", None, None)
            on_done(idx)
    why = pool.workers_reason(len(todo)) if callable(getattr(pool, "workers_reason", None)) else "VRN_BATCH_WORKERS=1 可改回序跑"
    print(f"[VRN] 平行池 {workers} 個工作行程({len(todo)} 份;VRN_BatchPool {pool.POOL_VERSION};{why})",
          file=sys.stderr)
    try:
        done.update(pool.map_reports(todo, Path(__file__).resolve(), (ticker_lookup, broker_alias, rating_alias),
                                     run_id, workers, on_done))
    except Exception as exc:  # noqa: BLE001 -- the pool is an accelerator, never a gate
        print(f"[VRN] 平行池中斷,沒跑完的改在本行程序跑:{type(exc).__name__}: {exc!r}"[:220], file=sys.stderr)
        done = {k: v for k, v in done.items() if v[0] != "OCR"}
    return done


def process_batch(args: argparse.Namespace) -> Dict[str, Any]:
    output_dir = ensure_dir(Path(args.output).expanduser().resolve())
    input_path = Path(args.input).expanduser().resolve()
    run_id = args.run_id or make_run_id()

    ticker_ssot_path = Path(args.ticker_ssot).expanduser().resolve() if args.ticker_ssot else None
    broker_ssot_path = Path(args.broker_ssot).expanduser().resolve() if args.broker_ssot else None
    rating_ssot_path = Path(args.rating_ssot).expanduser().resolve() if args.rating_ssot else None

    local_lookup = load_local_ticker_ssot(ticker_ssot_path)
    online_lookup = fetch_exchange_name_table(bool(args.online_name_update))
    ticker_lookup = merge_ticker_lookups(local_lookup, online_lookup)
    broker_alias = load_broker_alias_ssot(broker_ssot_path)
    rating_alias = load_rating_alias_ssot(rating_ssot_path)

    files = scan_input_files(input_path)
    only = getattr(args, "only_files", None)
    if only:                                    # v0104: a named sample (the loop's idempotency pass)
        keep = {str(name) for name in only}
        files = [f for f in files if f.name in keep]
    basic_rows: List[Dict[str, Any]] = []
    financial_rows: List[Dict[str, Any]] = []
    checkpoint = {
        "RunID": run_id,
        "StartedAt": now_iso(),
        "Input": str(input_path),
        "Output": str(output_dir),
        "Files": [],
        "SchemaVersion": SCHEMA_VERSION,
    }

    pooled = batch_pool_results(files, ticker_lookup, broker_alias, rating_alias, run_id)   # v0104; {} = in this process
    for idx, pdf_path in enumerate(files, start=1):
        pre = pooled.get(idx)
        if pre is None:
            print(f"[{idx}/{len(files)}] VRN processing: {pdf_path.name}")
        item = {"path": str(pdf_path), "status": "", "error": ""}
        if pdf_path.suffix.lower() in OCR_SUFFIXES:
            item["status"] = "SKIP_OCR_REQUIRED"
            checkpoint["Files"].append(item)
            write_json_file(output_dir / BATCH_CHECKPOINT_JSON_NAME, checkpoint)
            continue
        if pre is not None and pre[0] == "FAIL":
            item["status"] = "FAIL"
            item["error"] = pre[1]
            print(f"[ERROR] {pdf_path}: {str(pre[1]).splitlines()[0] if pre[1] else ''}", file=sys.stderr)
            checkpoint["Files"].append(item)
            write_json_file(output_dir / BATCH_CHECKPOINT_JSON_NAME, checkpoint)
            continue
        try:
            if pre is not None:
                basic, financial = pre[1], pre[2]
            else:
                basic, financial = process_one_pdf(pdf_path, ticker_lookup, broker_alias, rating_alias, run_id)
            basic_rows.append(basic)
            financial_rows.extend(financial)
            item["status"] = "OK"
            item["ReportID"] = basic.get("ReportID", "")
            item["FinancialRows"] = len(financial)
        except Exception as exc:
            item["status"] = "FAIL"
            item["error"] = f"{exc}\n{traceback.format_exc()}"
            print(f"[ERROR] {pdf_path}: {exc}", file=sys.stderr)
        checkpoint["Files"].append(item)
        write_json_file(output_dir / BATCH_CHECKPOINT_JSON_NAME, checkpoint)

    basic_new = make_dataframe(basic_rows, BASICINFO_COLUMNS)
    financial_new = make_dataframe(financial_rows, FINANCIALDATA_COLUMNS)

    basic_path = output_dir / BASICINFO_PARQUET_NAME
    financial_path = output_dir / FINANCIALDATA_PARQUET_NAME
    basic_existing = read_existing_parquet(basic_path, BASICINFO_COLUMNS)
    financial_existing = read_existing_parquet(financial_path, FINANCIALDATA_COLUMNS)
    basic_final = cumulative_upsert(basic_existing, basic_new, BASICINFO_UPSERT_KEYS, BASICINFO_COLUMNS)
    financial_final = cumulative_upsert(financial_existing, financial_new, FINANCIALDATA_UPSERT_KEYS, FINANCIALDATA_COLUMNS)

    write_parquet_mandatory(basic_final, basic_path, BASICINFO_NUMERIC_COLUMNS)
    write_parquet_mandatory(financial_final, financial_path, FINANCIALDATA_NUMERIC_COLUMNS)

    if args.csv:
        write_csv_optional(basic_final, output_dir / BASICINFO_CSV_NAME)
        write_csv_optional(financial_final, output_dir / FINANCIALDATA_CSV_NAME)
    if args.json:
        write_json_optional(basic_final, output_dir / BASICINFO_JSON_NAME)
        write_json_optional(financial_final, output_dir / FINANCIALDATA_JSON_NAME)
    duckdb_path = None
    if args.duckdb:
        duckdb_path = write_duckdb_optional(output_dir, basic_path, financial_path)
    google_sheet_ok = False
    if args.google_sheet:
        google_sheet_ok = write_google_sheet_optional(basic_final, financial_final, args.google_sheet, args.google_credentials)

    financial_counts: Dict[str, int] = {}
    for row in financial_rows:
        rid = normalize_unicode_text(row.get("ReportID", ""))
        if rid:
            financial_counts[rid] = financial_counts.get(rid, 0) + 1

    matrix_rows = []
    for row in basic_rows:
        rid = normalize_unicode_text(row.get("ReportID", ""))
        copy = dict(row)
        copy["FinancialRows"] = financial_counts.get(rid, 0)
        matrix_rows.append(copy)
    write_batch_matrix_csv(output_dir / BATCH_MATRIX_CSV_NAME, matrix_rows)

    summary = {
        "RunID": run_id,
        "StartedAt": checkpoint.get("StartedAt"),
        "FinishedAt": now_iso(),
        "Input": str(input_path),
        "Output": str(output_dir),
        "ProcessedFiles": len(files),
        "SkippedOcrFiles": sum(1 for f in checkpoint["Files"] if f.get("status") == "SKIP_OCR_REQUIRED"),
        "FailedFiles": sum(1 for f in checkpoint["Files"] if f.get("status") == "FAIL"),
        "EngineVersion": ENGINE_VERSION,
        "BasicRowsNew": len(basic_rows),
        "FinancialRowsNew": len(financial_rows),
        "BasicRows": int(len(basic_final)),
        "FinancialRows": int(len(financial_final)),
        "TickerLookupRows": len(ticker_lookup),
        "OutputFiles": {
            "BasicInfoParquet": str(basic_path),
            "FinancialDataParquet": str(financial_path),
            "BasicInfoCSV": str(output_dir / BASICINFO_CSV_NAME) if args.csv else "",
            "FinancialDataCSV": str(output_dir / FINANCIALDATA_CSV_NAME) if args.csv else "",
            "BasicInfoJSON": str(output_dir / BASICINFO_JSON_NAME) if args.json else "",
            "FinancialDataJSON": str(output_dir / FINANCIALDATA_JSON_NAME) if args.json else "",
            "DuckDB": str(duckdb_path) if duckdb_path else "",
            "GoogleSheet": args.google_sheet if google_sheet_ok else "",
            "HTML": str(output_dir / HTML_MATRIX_NAME),
        },
        "SchemaVersion": SCHEMA_VERSION,
    }
    write_json_file(output_dir / BATCH_SUMMARY_JSON_NAME, summary)
    write_html_matrix(output_dir / HTML_MATRIX_NAME, basic_rows, financial_counts, summary)
    print(safe_json_dumps(summary))
    return summary

# ============================================================
# def 14_SELF_TEST
# ============================================================

def run_self_test() -> None:
    print("[SELF-TEST] Start")
    sample = "凱基_神達3706 TT_初次評等_(20)260131_TargetPrice 168.pdf"
    token_info = tokenize_filename(sample)
    assert "3706" in token_info["number_tokens"], token_info
    assert "20" in token_info["number_tokens"] and "260131" in token_info["number_tokens"], token_info
    dates = parse_date_candidates_from_number_tokens(token_info["number_tokens"])
    assert any(d["date"] == "2026-01-31" for d in dates), dates
    local_lookup = {
        "3706": {"TW_TICKER": "3706", "YF_TICKER": "3706.TW", "Market": "TWSE", "Name": "神達", "Industry": "電腦及週邊"}
    }
    ft = extract_filename_ticker_candidates(token_info, local_lookup)
    pt = extract_text_ticker_candidates("神達 3706 TT Target Price 168 Buy", "first_page", local_lookup)
    ticker, yf, issues = choose_canonical_ticker(ft, pt, local_lookup)
    assert ticker == "3706" and yf == "3706.TW", (ticker, yf, issues)
    broker_alias = default_broker_aliases()
    broker, broker_candidates, _ = match_broker_from_tokens_and_text(token_info["chinese_tokens"], token_info["english_tokens"], "凱基證券", broker_alias)
    assert broker == "KGI", broker
    rating, _ = extract_rating("本報告維持買進評等，目標價168元", DEFAULT_RATING_ALIASES)
    assert rating == "BUY", rating
    tp = extract_target_price("目標價: NT$168")
    assert tp == 168.0, tp
    # v0101: spec section 2 table
    assert parse_date_any("1140122") == "2025-01-22"
    assert parse_date_any("20250122") == "2025-01-22"
    assert parse_date_any("250122") == "2025-01-22"
    assert parse_date_any("20250230") is None, "impossible calendar date must be rejected"
    assert parse_date_any("926708") is None
    assert not parse_date_candidates_from_number_tokens(["006208"]), "ETF code is never a date"
    partial = parse_partial_dates_from_filename("主動式ETF籌碼追蹤-CTBC0915")
    assert partial and partial[0]["status"] == "YEAR_MISSING" and partial[0]["month"] == 9, partial
    ti = tokenize_filename("第一場 2026年投資大趨勢 - 華南投顧 -1141201.pdf")
    assert all(not c["ticker"] for c in extract_filename_ticker_candidates(ti, {})), "2026 is a year, not a ticker"
    ti = tokenize_filename("凱基投顧_2891 中信金_施志鴻_20260519.pdf")
    frags = company_fragments_after_ticker(ti["tokens"])
    b, _, _ = match_broker_from_tokens_and_text(ti["chinese_tokens"], ti["english_tokens"], "", load_broker_alias_ssot(None), frags)
    assert b == "KGI", b
    b, _, _ = match_broker_from_tokens_and_text([], ["ms"], "", load_broker_alias_ssot(None))
    assert b == "MS", b
    b, _, _ = match_broker_from_tokens_and_text([], [], "thermal systems and earnings", load_broker_alias_ssot(None))
    assert b == "", b
    r, _ = extract_rating("光焱(7728,Note)-CTBC260918", DEFAULT_RATING_ALIASES)
    assert r == "", r
    r, _ = extract_rating("瑞基(4171,NR_未評等)-CTBC251208", DEFAULT_RATING_ALIASES)
    assert r == "NOT_RATED", r
    r, _ = extract_rating("維持中立評等", DEFAULT_RATING_ALIASES)
    assert r == "HOLD", r
    assert extract_target_price("上漲空間 30.4% 目標價 NT$168 元") == 168.0
    assert extract_target_price("Target Price (NT$) 1,250") == 1250.0
    assert extract_current_price("Target Price 168 收盤價 129") == 129.0
    period = parse_period_label("12/24A")
    assert period["FiscalYear"] == 2024 and period["EstimateFlag"] == "Actual", period
    period = parse_period_label("FY2025")
    assert period["FiscalYear"] == 2025 and period["EstimateFlag"] == "", period
    period = parse_period_label("2025.Q1")
    assert period["FiscalYear"] == 2025 and period["FiscalQuarter"] == "Q1", period
    period = parse_period_label("3Q26")
    assert period["FiscalYear"] == 2026 and period["FiscalQuarter"] == "Q3", period
    print("[SELF-TEST] PASS")

# ============================================================
# def 15_CLI
# ============================================================

def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="VRN Integrated Report Database Engine")
    parser.add_argument("--input", default=DEFAULT_INPUT_DIR, help="PDF 檔案或資料夾")
    parser.add_argument("--output", default=DEFAULT_OUTPUT_DIR, help="輸出資料夾")
    parser.add_argument("--ticker-ssot", default=DEFAULT_TICKER_SSOT_CSV, help="Ticker / Name SSOT CSV，本地優先")
    parser.add_argument("--broker-ssot", default=DEFAULT_BROKER_SSOT_CSV, help="Broker alias SSOT CSV")
    parser.add_argument("--rating-ssot", default=DEFAULT_RATING_SSOT_CSV, help="Rating alias SSOT CSV")
    parser.add_argument("--online-name-update", action="store_true", default=ENABLE_ONLINE_NAME_TABLE_BY_DEFAULT, help="啟用 TWSE/TPEX 線上名稱表更新")
    parser.add_argument("--csv", action="store_true", help="輸出 UTF-8 BOM CSV")
    parser.add_argument("--json", action="store_true", help="輸出 UTF-8 JSON")
    parser.add_argument("--duckdb", action="store_true", help="輸出 DuckDB")
    parser.add_argument("--google-sheet", default="", help="Google Sheet 名稱，可選")
    parser.add_argument("--google-credentials", default="", help="gspread service account json path")
    parser.add_argument("--run-id", default="", help="指定 RunID")
    parser.add_argument("--self-test", action="store_true", help="只跑內建測試")
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    import logging
    logging.getLogger("pdfminer").setLevel(logging.ERROR)   # 'Could not get FontBBox …' is noise, not an error
    parser = build_arg_parser()
    args = parser.parse_args(argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[attr-defined]
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[attr-defined]
    except Exception:
        pass
    if args.input == DEFAULT_INPUT_DIR and os.name == "nt" and Path(DEFAULT_INPUT_DIR_WINDOWS).exists():
        args.input = DEFAULT_INPUT_DIR_WINDOWS
    if args.self_test:
        run_self_test()
        return 0
    process_batch(args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
