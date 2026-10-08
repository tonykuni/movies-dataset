#!/usr/bin/env python3
PARAMS = {
    'version': 'v0108-handover-v0100',
    'reference_manifest_sha256': '5c6352a2a3165fbb9a55f3777dbb268775f6b65668b7d5409b52f9c7bc40d0c2',
    'reference_files': 199,
    'upstream_locked_files': 26,
    'facts': 4415, 'equations': 409, 'source_cells': 5462,
    'tables': 43, 'pages': 12, 'datasets': 16, 'release_checks': 17,
    'money_facts': 2547, 'policy_tests': 16, 'semantic_checks': 11,
    'fact_key': ['FILENAME_SHA256', 'TABLE_ID', 'SOURCE_ROW_INDEX', 'PERIOD'],
    'default_pack': '../pack',
    'csv_field_size_limit': 16000000,
    'money_rounding': 'ROUND_HALF_UP',
}
import argparse, ast, csv, hashlib, json, sys, zipfile
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path


def def_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def def_hash(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def def_require(condition, message):
    if not condition:
        raise ValueError(message)


def def_record(rows, name, passed, evidence):
    def_require(passed, name + ': ' + evidence)
    rows.append({'check': name, 'status': 'PASS', 'evidence': evidence})


def def_verify_manifest(root, expected_pin=None):
    path = root / 'SHA256_MANIFEST.json'
    if expected_pin:
        def_require(def_hash(path) == expected_pin, 'REFERENCE_MANIFEST_PIN_MISMATCH')
    manifest = def_json(path)
    for name, digest in manifest.items():
        file = (root / name).resolve()
        def_require(file.is_relative_to(root.resolve()), 'UNSAFE_MANIFEST_PATH: ' + name)
        def_require(file.is_file(), 'MISSING_FILE: ' + name)
        def_require(def_hash(file) == digest, 'FILE_SHA256_MISMATCH: ' + name)
    return manifest


def def_integrity(pack):
    checks = []
    manifest = def_verify_manifest(pack, PARAMS['reference_manifest_sha256'])
    def_record(checks, 'REFERENCE_199_FILES', len(manifest) == PARAMS['reference_files'], 'Pinned v0108 reference and every listed file verified')
    lock = def_json(pack / 'engine/UPSTREAM_LOCK.json')
    def_record(checks, 'UPSTREAM_26_FILES', len(lock) == PARAMS['upstream_locked_files'] and all(def_hash(pack / 'engine' / name) == digest for name, digest in lock.items()), 'Original upstream engines/rules/configs unchanged')
    modules = sorted(pack.glob('VRN_Annual*_v0108.py'))
    for module in modules:
        code = module.read_text(encoding='utf-8')
        ast.parse(code)
        compile(code, str(module), 'exec')
    def_record(checks, 'FIVE_MODULES_AST_COMPILE', len(modules) == 5, 'No module execution or engine replacement during integrity-only check')
    return checks


def def_fact_contract(facts, baseline):
    keys = [tuple(row[k] for k in PARAMS['fact_key']) for row in facts]
    def_require(len(facts) == PARAMS['facts'] and len(set(keys)) == len(keys), 'FACT_COUNT_OR_DUPLICATE_KEY')
    def_require(all(x['FREQUENCY'] == 'ANNUAL' and len(x['PERIOD']) == 4 and x['PERIOD'].isdigit() for x in facts), 'QUARTER_OR_INVALID_PERIOD')
    index = {tuple(row[k] for k in PARAMS['fact_key']): row for row in baseline}
    def_require(len(index) == len(facts) and all(index.get(key) == row for key, row in zip(keys, facts)), 'FACTS_DIFFER_FROM_VERIFIED_REFERENCE')
    def_require(all(x['UNIT'] != 'UNSPECIFIED' and x['EXTRACT_CELL_VALIDATION_STATUS'] == 'PASS' for x in facts), 'UNIT_OR_EXTRACTION_REVIEW')
    money = [x for x in facts if x['OUTPUT_DECIMAL_PLACES'] is not None]
    def_require(len(money) == PARAMS['money_facts'], 'MONEY_COUNT')
    for fact in money:
        places = 0 if fact['SOURCE_DECIMAL_PLACES'] == 0 else 1
        rounded = Decimal(fact['VALUE_CHECK_DECIMAL']).quantize(Decimal(1).scaleb(-places), rounding=PARAMS['money_rounding'])
        def_require(fact['OUTPUT_DECIMAL_PLACES'] == places and fact['ROUNDING_STAGE'] == 'AFTER_ALL_ANNUAL_CHECKS' and Decimal(str(fact['VALUE_NUMERIC'])) == rounded, 'ROUNDING_POLICY_MISMATCH')
    missing = [x for x in facts if x['TABLE_ID'] == 'N01-P2-T3' and x['METRIC'] == 'Otheroperatingcashflow' and x['PERIOD'] == '2028']
    def_require(len(missing) == 1 and missing[0]['VALUE_RAW'] == '--' and missing[0]['VALUE_NUMERIC'] is None and missing[0]['VALUE_CHECK_DECIMAL'] is None, 'SOURCE_MISSING_VALUE_WAS_FILLED')


def def_csv_value(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def def_equivalent(left, right):
    if left is None or right is None:
        return left is right
    if isinstance(left, (dict, list)):
        return left == right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return Decimal(str(left)) == Decimal(str(right))
    return str(left) == str(right)


def def_results(pack, results):
    import polars as pl
    import fitz
    from openpyxl import load_workbook
    checks = []
    manifest = def_verify_manifest(results)
    def_record(checks, 'RESULTS_FILE_HASHES', bool(manifest), 'Every generated file listed in result manifest matches')
    facts = def_json(results / 'data/AnnualFinancialData.json')
    def_fact_contract(facts, def_json(pack / 'data/AnnualFinancialData.json'))
    def_record(checks, '4415_FACTS_AND_ROUNDING', True, 'Exact field-by-field reference match; annual scope; unique keys; original values; final rounding; missing value retained')
    for name, count in [('AnnualEquations', 409), ('SemanticChecks', 11), ('AmountPolicyTests', 16), ('ReleaseGate', 17)]:
        records = def_json(results / ('data/' + name + '.json'))
        def_record(checks, name, len(records) == count and all(r['結果'] == 'PASS' for r in records), str(count) + ' PASS')
    equations = def_json(results / 'data/AnnualEquations.json')
    def_record(checks, 'EQUATIONS_REFERENCE_MATCH', equations == def_json(pack / 'data/AnnualEquations.json'), 'All 409 formula inputs/results/precision intervals match the reference')
    cells = def_json(results / 'data/AnnualSourceCells.json')
    tables = def_json(results / 'data/FreshSourceTables.json')
    def_record(checks, 'FRESH_SOURCE_READS', len(cells) == 5462 and all(x['status'] == 'PASS' for x in cells) and len(tables) == 43 and all(x['結果'] == 'PASS' and x['待覆核格數'] == 0 for x in tables), '5462 cells and 43 freshly read annual tables')
    skips = def_json(results / 'data/FormulaSkips.json')
    conflicts = def_json(results / 'data/SourceConflicts.json')
    def_record(checks, 'SOURCE_LIMITATIONS_RETAINED', len(skips) == 1 and skips[0]['結果'] == 'NOT_CALCULABLE' and len(conflicts) == 1 and conflicts[0]['問題狀態'] == 'OPEN_SOURCE_CORRECTION_REQUIRED', 'GS missing value and CLST narrative conflict remain visible')
    files = sorted((results / 'data').glob('*.json'))
    def_require(len(files) == PARAMS['datasets'], 'DATASET_COUNT')
    csv.field_size_limit(PARAMS['csv_field_size_limit'])
    for file in files:
        records = def_json(file)
        with file.with_suffix('.csv').open(encoding='utf-8-sig', newline='') as handle:
            csv_rows = list(csv.DictReader(handle))
        pq_rows = pl.read_parquet(file.with_suffix('.parquet')).to_dicts()
        def_require(len(csv_rows) == len(pq_rows) == len(records), 'ROW_COUNT_MISMATCH: ' + file.stem)
        for expected, csv_row, pq_row in zip(records, csv_rows, pq_rows):
            def_require(all(csv_row[k] == def_csv_value(v) for k, v in expected.items()), 'CSV_VALUE_MISMATCH: ' + file.stem)
            def_require(all(def_equivalent(v, pq_row[k]) for k, v in expected.items()), 'PARQUET_VALUE_MISMATCH: ' + file.stem)
    def_record(checks, '16_DATASETS_THREE_FORMATS', True, 'Every JSON/CSV/Parquet row and value verified; declared string coercions allowed')
    workbook = load_workbook(results / 'VRN_RealTest_v0104.xlsx', read_only=True)
    sheet = workbook['年度財務']
    columns = [c.value for c in sheet[1]]
    def_require(sheet.max_row == 4416, 'EXCEL_ROW_COUNT')
    for row, fact in zip(sheet.iter_rows(min_row=2), facts):
        for key in ['VALUE_RAW', 'VALUE_TRIM', 'VALUE_CHECK_DECIMAL', 'UNIT', 'PERIOD']:
            def_require((row[columns.index(key)].value or '') == (fact[key] or ''), 'EXCEL_VALUE_MISMATCH: ' + key)
        places = fact['OUTPUT_DECIMAL_PLACES']
        if places is not None:
            cell = row[columns.index('VALUE_NUMERIC')]
            def_require(cell.number_format == ('0' if places == 0 else '0.0') and Decimal(str(cell.value)) == Decimal(str(fact['VALUE_NUMERIC'])), 'EXCEL_ROUNDING_FORMAT')
    workbook.close()
    def_record(checks, 'EXCEL_READBACK', True, '4415 rows, raw/check values, units, periods, numeric output and display formats verified')
    with fitz.open(results / 'VRN_AnnualFinancialPages_v0108.pdf') as pdf:
        def_require(len(pdf) == 12, 'ANNUAL_PDF_PAGE_COUNT')
    with fitz.open(results / 'VRN_ProblemsSolved_v0108.pdf') as pdf:
        text = ''.join(page.get_text() for page in pdf)
        def_require('\x00' not in text and 'OPEN_SOURCE_CORRECTION_REQUIRED' in text, 'REPORT_TEXT_OR_LIMITATIONS')
    def_record(checks, 'PDF_READBACK', True, 'Annual PDF 12 pages; readable report retains limitation status; visual review still required')
    for name in ['VRN_AnnualFinancial_v0108.py', 'VRN_AnnualAmounts_v0108.py', 'VRN_AnnualLayouts_v0108.py', 'VRN_AnnualCloseout_v0108.py', 'VRN_AnnualAudit_v0108.py']:
        def_require(def_hash(results / name) == def_hash(pack / name), 'RERUN_MODULE_CHANGED: ' + name)
    engine_files = [p for p in def_json(pack / 'SHA256_MANIFEST.json') if p.startswith('engine/')]
    def_require(all(def_hash(results / name) == def_hash(pack / name) for name in engine_files), 'EXISTING_ENGINE_CHANGED')
    def_record(checks, 'ENGINE_FILES_UNCHANGED', True, 'All 54 engine tree files and 5 orchestration modules match the delivered pack')
    archive = results / 'VRN_AnnualFinancial_v0108.zip'
    # Delivered baseline is expanded in pack/ without a redundant nested ZIP; reruns must emit one.
    if results.resolve() != pack.resolve():
        def_require(archive.is_file(), 'RERUN_ZIP_MISSING')
    if archive.exists():
        with zipfile.ZipFile(archive) as zipped:
            def_require(zipped.testzip() is None, 'ZIP_CRC_FAILED')
            def_require(all(hashlib.sha256(zipped.read(k)).hexdigest() == v for k, v in manifest.items()), 'ZIP_MANIFEST_MISMATCH')
        def_record(checks, 'RERUN_ZIP_INTEGRITY', True, 'ZIP CRC and manifest hashes match')
    return checks


def def_main():
    parser = argparse.ArgumentParser(description='Read-only VRN v0108 pack and output verifier')
    parser.add_argument('--pack', type=Path, default=Path(__file__).resolve().parent / PARAMS['default_pack'])
    parser.add_argument('--results', type=Path)
    parser.add_argument('--receipt', type=Path, required=True)
    parser.add_argument('--integrity-only', action='store_true')
    args = parser.parse_args()
    pack = args.pack.resolve()
    results = args.results.resolve() if args.results else pack
    receipt = args.receipt.resolve()
    def_require(not receipt.is_relative_to(pack) and not receipt.is_relative_to(results), 'RECEIPT_MUST_BE_OUTSIDE_REFERENCE_AND_RESULTS')
    checks = []
    report = {'verifier': PARAMS['version'], 'utc_time': datetime.now(timezone.utc).isoformat(), 'mode': 'INTEGRITY_ONLY' if args.integrity_only else 'FULL_RESULT_CHECK', 'pack': str(pack), 'results': str(results), 'checks': checks}
    code = 0
    try:
        checks.extend(def_integrity(pack))
        if not args.integrity_only:
            checks.extend(def_results(pack, results))
        report['status'] = 'PASS_INTEGRITY_ONLY' if args.integrity_only else 'PASS_WITH_SOURCE_LIMITATIONS'
    except Exception as exc:
        report['status'] = 'FAIL'
        report['error'] = type(exc).__name__ + ': ' + str(exc)
        code = 1
    receipt.parent.mkdir(parents=True, exist_ok=True)
    receipt.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'status': report['status'], 'checks_pass': len(checks), 'receipt': str(receipt), 'error': report.get('error')}, ensure_ascii=False))
    return code


if __name__ == '__main__':
    sys.exit(def_main())
