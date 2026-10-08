#!/usr/bin/env python3
PARAMS={
 'version':'v0108','anchor_match_policy':'PER_COLUMN_SOURCE_ANCHORS',
 'minimum_source_anchor_count':1,
 'jp_sha':'681de7210e2c50df56bfd251457a52ee1111999fcde41a8b1c03d665efa32eda',
 'jp_balance_table':'FIN-P20-T2','jp_source_footer_note_regex':r'NT\$ in millions',
 'footer_clip_padding':0.5,
 'expected_documents':9,'expected_source_pages':12,'expected_tables':43,
 'expected_facts':4415,'expected_equations':409,'expected_semantic_checks':11,
 'expected_source_cells':5462,
 'source_limitations':{'GS_MISSING_SOURCE_VALUE':'CLOSED_WITH_SOURCE_LIMITATION','CLST_NARRATIVE_CONFLICT':'OPEN_SOURCE_CORRECTION_REQUIRED'},
}
import copy,json,hashlib,re
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import fitz,pdfplumber


def def_column_anchor_row(words,edges,anchors,tolerance):
 """Match each column against its own source anchors, preserving font baseline offsets."""
 valid=[a for a in anchors or []if a and len(a)==4]
 if not valid:return None
 centers=[[]for _ in range(len(edges)-1)]
 for anchor in valid:
  cx=(anchor[0]+anchor[2])/2;cy=(anchor[1]+anchor[3])/2
  col=next((i for i in range(len(edges)-1)if edges[i]<=cx<edges[i+1]),None)
  if col is not None:centers[col].append(cy)
 fallback=[(a[1]+a[3])/2 for a in valid]
 buckets=[[]for _ in centers]
 for word in words:
  cx=(word['x0']+word['x1'])/2;cy=(word['top']+word['bottom'])/2
  col=next((i for i in range(len(edges)-1)if edges[i]<=cx<edges[i+1]),None)
  if col is None:continue
  target=centers[col]or fallback
  if min(abs(cy-c)for c in target)<=tolerance:buckets[col].append((word['top'],word['x0'],word['text']))
 # Sort by source x coordinate, preserving existing single-row engine semantics.
 return [' '.join(text for _,_,text in sorted(bucket,key=lambda w:(w[1],w[0])))for bucket in buckets]


def def_require(condition,message):
 if not condition:raise ValueError(message)


def def_reader_cache():
 """Keep one source document open; retain each reader's exact original crop semantics."""
 native_documents={};plumber_documents={}
 def native(path,table):
  key=str(Path(path).resolve())
  if key not in native_documents:native_documents[key]=fitz.open(path)
  page=native_documents[key][table['page']-1]
  return [{'text':w[4],'x0':w[0],'top':w[1],'x1':w[2],'bottom':w[3]}for w in page.get_text('words',clip=fitz.Rect(table['bbox_pt']))]
 def plumber(path,table):
  key=str(Path(path).resolve())
  if key not in plumber_documents:plumber_documents[key]=pdfplumber.open(path)
  page=plumber_documents[key].pages[table['page']-1]
  return page.crop(tuple(table['bbox_pt'])).extract_words(x_tolerance=1,y_tolerance=1)
 def close():
  for doc in native_documents.values():doc.close()
  for doc in plumber_documents.values():doc.close()
  native_documents.clear();plumber_documents.clear()
 return native,plumber,close


def def_complete_source_crops(path,tables,source):
 receipts=[]
 if source['source_sha256']==PARAMS['jp_sha']:
  table=next(t for t in tables if t['id']==PARAMS['jp_balance_table'])
  before=table['bbox_pt'][:]
  with fitz.open(path)as doc:
   words=doc[table['page']-1].get_text('words')
   note=[w for w in words if table['bbox_pt'][0]<=w[0]<table['bbox_pt'][2]and 673<=w[1]<681]
   note_text=' '.join(w[4]for w in sorted(note,key=lambda w:w[0]))
   def_require(bool(re.search(PARAMS['jp_source_footer_note_regex'],note_text)),'JP source unit footer not corroborated')
   table['bbox_pt'][3]=max(w[3]for w in note)+PARAMS['footer_clip_padding']
  receipts.append({'檔案':source['filename'],'來源頁':table['page'],'表格':table['id'],'原裁切':[float(x)for x in before],'完成後裁切':[float(x)for x in table['bbox_pt']],'處理':'末列 y/y Growth 字形與來源／百萬元單位註記完整保留；來源值未改寫','結果':'PASS'})
 for table in tables:
  x0,y0,x1,y1=table['bbox_pt']
  for row in table['records']:
   if row.get('semantic_component',row['component'])!='ROW':continue
   outside=[a for a in row.get('anchors',[])if a[0]<x0-.05 or a[2]>x1+.05 or a[1]<y0-.05 or a[3]>y1+.05]
   def_require(not outside,'SOURCE_GLYPH_CUT: '+source['filename']+' '+table['id']+' '+str(row['row_index']))
 return receipts


def def_fresh_validate(path,tables,source,engine,reader_cache):
 def_require(hashlib.sha256(Path(path).read_bytes()).hexdigest()==source['source_sha256'],'SOURCE_SHA256_MISMATCH: '+source['filename'])
 before=copy.deepcopy(tables);calls=[]
 previous=(engine.def_words_to_row,engine.def_native_table_words,engine.def_plumber_table_words)
 engine.def_words_to_row=def_column_anchor_row
 engine.def_native_table_words,engine.def_plumber_table_words=reader_cache[:2]
 try:engine.def_table_fast_validate(path,{'tables':tables,'profile':'SOURCE_HASH_LOCKED'},engine.PARAMS,calls,None)
 finally:engine.def_words_to_row,engine.def_native_table_words,engine.def_plumber_table_words=previous
 receipts=[]
 for old,table in zip(before,tables):
  unchanged=[r.get('cells')for r in old['records']]==[r.get('cells')for r in table['records']]
  def_require(unchanged,'Fresh source differs from retained values: '+source['filename']+' '+table['id'])
  audit=table['fast_validation']
  def_require(bool(audit['cells'])and all(c['status']=='PASS'for c in audit['cells']),'Fresh source validation requires review: '+source['filename']+' '+table['id']+' '+json.dumps([c for c in audit['cells']if c['status']!='PASS'],ensure_ascii=False))
  # Retain historical JP clip evidence and raw validation snapshot separately.
  audit['previous_review_cells']=old.get('fast_validation',{}).get('previous_review_cells',[])
  audit['previous_validation_snapshot_status']=old.get('fast_validation',{}).get('status')
  audit['fresh_validation_basis']='PDF re-read this run; per-column anchors; same-source native readers; not official economic truth'
  receipts.append({'檔案':source['filename'],'來源頁':table['page'],'表格':table['id'],'比對格數':len(audit['cells']),'待覆核格數':0,'來源原值保留':'PASS'if unchanged else'FAIL','讀取方式':'重新讀 PDF：fitz + pdfplumber；逐欄原始座標','結果':'PASS'})
 return receipts,calls


def def_release_checks(facts,equations,source_cells,semantic,skips,missing,conflicts,table_receipts):
 checks=[]
 def check(name,condition,evidence):
  def_require(condition,'Release gate failed: '+name)
  checks.append({'檢查':name,'結果':'PASS','證據':evidence})
 check('整批來源重讀',len(table_receipts)==PARAMS['expected_tables']and all(r['結果']=='PASS'for r in table_receipts),'43 表每次重跑均重新讀原 PDF，沒有沿用快照 PASS')
 check('年度範圍',len(facts)==PARAMS['expected_facts']and all(f['FREQUENCY']=='ANNUAL'and re.fullmatch(r'\d{4}',f['PERIOD'])for f in facts),'4415 年度格；季度欄 0')
 check('來源文字數字全格一致',len(source_cells)==PARAMS['expected_source_cells']and all(r['status']=='PASS'for r in source_cells),'5462 格，含項目與分段；增加的 18 格為 MS 分段格')
 check('全格單位已校準',all(f['UNIT']!='UNSPECIFIED'for f in facts),'單位未判定 0；SOURCE_UNIT 與判定依據保留')
 check('年度算式及期間',len(equations)==PARAMS['expected_equations']and all(r['結果']=='PASS'and re.fullmatch(r'\d{4}',str(r['年度']))for r in equations),'409 項完整精度、原印精度區間 PASS')
 check('語意勾稽',len(semantic)==PARAMS['expected_semantic_checks']and all(r['結果']=='PASS'for r in semantic),'11 項；Net cash 不轉 0')
 check('來源缺值不回填',len(skips)==len(missing)==1 and skips[0]['結果']=='NOT_CALCULABLE'and any(f['TABLE_ID']=='N01-P2-T3'and f['METRIC']=='Otheroperatingcashflow'and f['PERIOD']=='2028'and f['VALUE_NUMERIC']is None and f['VALUE_CHECK_DECIMAL']is None for f in facts),'GS -- 維持來源空值；推算診斷分欄')
 check('原文矛盾可追溯',len(conflicts)==1 and conflicts[0]['問題狀態']=='OPEN_SOURCE_CORRECTION_REQUIRED','CLST 10%／約 20.0106% 保留；整體狀態含來源限制')
 check('禁止驗算前捨入',all(f['ROUNDING_STAGE']=='PENDING_ALL_ANNUAL_CHECKS'for f in facts),'全批驗算完成前不修改最終顯示值')
 return checks


def def_verify_outputs(out,facts,policy_tests):
 from openpyxl import load_workbook
 import polars as pl,ast,importlib.util,sys
 checks=[]
 def check(name,condition,evidence):
  def_require(condition,'Output verification failed: '+name);checks.append({'檢查':name,'結果':'PASS','證據':evidence})
 check('金額政策測試',len(policy_tests)==16 and all(r['結果']=='PASS'for r in policy_tests),'16 項，含千元合成案例及過早捨入拒絕測試')
 money=[f for f in facts if f['OUTPUT_DECIMAL_PLACES']is not None]
 check('全精度後捨入',len(money)==2547 and all(f['ROUNDING_STAGE']=='AFTER_ALL_ANNUAL_CHECKS'and Decimal(str(f['VALUE_NUMERIC']))==Decimal(f['VALUE_CHECK_DECIMAL']).quantize(Decimal(1).scaleb(-f['OUTPUT_DECIMAL_PLACES']),rounding='ROUND_HALF_UP')for f in money),'2547 金額格；來源整數→整數，有小數→一位')
 wb=load_workbook(out/'VRN_RealTest_v0104.xlsx',read_only=True);sheet=wb['年度財務'];fields=[c.value for c in sheet[1]];rawcol=fields.index('VALUE_RAW');valuecol=fields.index('VALUE_NUMERIC');valid=True
 for row,fact in zip(sheet.iter_rows(min_row=2),facts):
  valid&=(row[rawcol].value or '')==fact['VALUE_RAW']
  places=fact['OUTPUT_DECIMAL_PLACES']
  if places is not None:valid&=row[valuecol].number_format==('0'if places==0 else'0.0')
 check('Excel 原值及格式讀回',sheet.max_row==len(facts)+1 and valid,'4415 原值格保留，Excel 0／0.0 格式正確');wb.close()
 datasets=list((out/'data').glob('*.parquet'))
 check('Parquet／JSON 資料量一致',all(pl.read_parquet(p).height==len(json.loads(p.with_suffix('.json').read_text(encoding='utf-8')))for p in datasets),str(len(datasets))+' 組資料集回讀；CSV 編碼為 utf-8-sig')
 with fitz.open(out/'VRN_ProblemsSolved_v0108.pdf')as pdf:text=''.join(p.get_text()for p in pdf)
 check('中文 PDF 可讀且有來源限制','\x00'not in text and 'OPEN_SOURCE_CORRECTION_REQUIRED'in text,'中文成果報告無缺字 NUL；保留來源限制')
 with fitz.open(out/'VRN_AnnualFinancialPages_v0108.pdf')as pdf:
  quarters=[w[4]for page in pdf for w in page.get_text('words')if re.fullmatch(r'(?:[1-4]Q\d{2,4}[AF]?|Q[1-4]\d{2,4}[AF]?|(?:Mar|Jun|Sep)-\d{2}[AF])',w[4])]
  check('年度裁切頁範圍',len(pdf)==12 and not quarters,'12 頁；季度標頭排除，Dec 年度依來源 metadata 認定')
 paths=sorted(out.glob('*.py'))
 for index,path in enumerate(paths):
  source=path.read_text(encoding='utf-8');ast.parse(source);compile(source,str(path),'exec')
  spec=importlib.util.spec_from_file_location('release_module_'+str(index),path);module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module)
 check('AST／compile／import',len(paths)==5,'5 個主模組三輪檢查 PASS；現有引擎原樣共用')
 launcher=(out/'RunAnnual.cmd').read_bytes()
 check('Windows 啟動檔換行',launcher.count(b'\r\n')>=4 and b'\\r\\n'not in launcher,'CRLF 正確；未宣稱 Windows 實機執行')
 return checks


def def_package(out,version):
 import zipfile
 paths=[p for p in out.rglob('*')if p.is_file()and p.suffix!='.zip'and p.name!='SHA256_MANIFEST.json'and '__pycache__'not in p.parts]
 manifest={str(p.relative_to(out)):hashlib.sha256(p.read_bytes()).hexdigest()for p in paths}
 (out/'SHA256_MANIFEST.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2), encoding='utf-8');paths.append(out/'SHA256_MANIFEST.json')
 archive=out/('VRN_AnnualFinancial_'+version+'.zip')
 with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED)as z:
  for path in paths:z.write(path,str(path.relative_to(out)))
 with zipfile.ZipFile(archive)as z:
  def_require(z.testzip()is None,'ZIP CRC failed')
  def_require(all(hashlib.sha256(z.read(k)).hexdigest()==v for k,v in manifest.items()),'ZIP SHA256 failed')
 return {'檢查':'ZIP／SHA256 全檔讀回','結果':'PASS','證據':str(len(manifest))+' 檔；完整成果、預覽與驗證收據均納入'}
