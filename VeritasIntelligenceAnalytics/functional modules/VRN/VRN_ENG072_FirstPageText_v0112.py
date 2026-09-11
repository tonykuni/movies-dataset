#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
VRN_ENG072_FirstPageText v0112 — 首頁三法整合擷取器(批235 立;批246 v2.1 掛載)
====================================================================
操作員令:「用 VDF 中 LAYOUT ANALYSIS/TEXT FETCHER 功能擷取出
個股報告第一頁文字內容」。
機制(與 digest ㉓㉙ 同優先序):
  PDF:fitz(pymupdf)版面閱讀序抽第 1 頁(LAYOUT ANALYSIS 正主)
       →pypdf 後備(版面重排風險誠實標 PYPDF_FALLBACK)
       →零文字=掃描版,誠實標 NEEDS_OCR(候 via-ocrsuper 道,不假抽)
  DOCX:無頁界概念→前 30 段誠實近似(標 DOCX_HEAD)
輸出:VIA_Reports/first_page_text/<原檔名>.txt 逐件
      +FIRSTPAGE_SUMMARY.html 一頁堆疊總覽(誠實三態)
紅線:券商報告原文抽取物不入 git(輸出夾 .gitignore);原件僅在
本機 input_reports。
批236 追令(操作員實錄:左右兩欄交錯+本文斷碎):
  ①分區 LAYOUT:fitz dict blocks 依 BBox 分四區——全寬標題帶/
    左本文區/右資訊區(x0≥52% 頁寬)/頁尾帶(底 8%)——
    「像表格切開處理並還原對齊」,兩欄分開各依 (y,x) 排序
  ②本文修復:區塊內斷行接回;句尾無終止標點(。.!?%)且下行為
    延續句=合併(EN 補空白/ZH 直接接;連字號斷詞修復)
  ③字級階層:font size+bold 判標題 H1-H3/頁尾小字=雜訊
    (收容件 geminicode_repair/hierarchy_b236 邏輯整合)
  輸出=.txt 四節【標題帶】【右資訊區】【本文(修復)】【頁尾(雜訊)】
    +.json 結構化 sidecar+總覽 HTML 左右雙欄對照顯示
批242 追令(操作員:「另外用 PDFPLUMBER 擷取相互對照;至少兩個
成功的方法及內容對照;輸出完整第一頁文字/表格/上左右資訊區/本文,
還原分類;先不管摘要」):
  法A=fitz 分區(v0101 正主)/法B=pdfplumber words 同判準分區
  +extract_tables 表格還原——雙法逐區對照(正規化 difflib 比率):
  ≥0.90 AGREE/≥0.60 PARTIAL/低=DIVERGE 誠實列示;
  sidecar 增 plumber/tables/compare/blocks(字級階層分類)四鍵,
  fitz 區鍵保持頂層=下游 ENG073 零改動。
批243 追令(操作員:「INTEGRATE ALL LIBS AND TOOLS INTO EXISTING
ENGINES」——收容包整合進現役引擎,收容件原地不動,graceful 掛載):
  法C=GenericLayoutEngine v2.0.0(收容 intake/GenericLayoutEngine_
    v2.0.0_b242):zone_for_bbox 九宮分區標+font_weight_from_name
    字重階層(操作員字體階層長文)+weighted_median 本文字級推定
    +probe_backends 後端可用矩陣→sidecar "gle" 鍵(缺席=誠實 absent)
  句級修復=NLP OneEngine TextProcessor(帶 ssot_lexicon;修 ENG073
    掛載漏 lexicon_path 永遠 fallback 之債):本文修復後再
    split_sentences 一句一資料;疊字修復;標題不加句號
    →sidecar "sentences" 鍵+.txt 新節【本文(句級)】
  fitz 區鍵仍為 sidecar 頂層=下游 ENG073 零改動。
批245 追令:法C 收容夾動態尾版解析(嚴禁寫死版號)。
批246(收容落地):AllEngines v2.1.0 收容為巢狀一層
  (…_b245/GenericLayoutEngine/generic_layout_engine.py)→解析器支援
  頂層+巢狀一層雙型;尾版排序取最新;sidecar gle.engine_dir 存證。
批444(v0110,操作員上傳 VIA_PDFPlumberPlusEngine 並令「接上去」):
  ①**字元密度分流閘**(SUP_MDL746 → 收容件 TRIAGE_MIN_CHAR_DENSITY):
    v0109 以前只問 `if txt:`——抽得到一個字就算有文字層。一頁只有浮水印
    文字層的掃描頁會被那一問放行,首頁引擎再從那十來個字裡找目標價,
    **那是假綠**。改成誠實三態 DIGITAL/THIN/SCANNED:
      · SCANNED 就算 fitz 抽得到字也不當文字層用,改走 OCR;OCR 也跑不動
        就停在 NEEDS_OCR 並把密度因由寫進標記(不靜靜把浮水印當首頁文字)
      · THIN(長捲頁/海報/投影片)照抽,只標稀薄——**不准只憑密度降級**。
        真檔證據:庫內 949×7448 點長捲頁那件有 862 字真正文、密度 1.22e-04
        低於門檻,光看密度會假紅。批440 那一課:改嚴一道閘要有真檔證據。
      · 橋缺席=UNKNOWN,走 v0109 原路,零回歸
  ②**OCR 兩車道**:道一 GLE(七支後端,paddleocr 3.x 新 API 都接)先走,
    整條不在位才換道二 PPP(收容件,只接 2.x,但記憶體內渲染不落暫存圖)。
    **補位不是取代**——我上一輪說「這包才真的會建構 PaddleOCR」是錯的,
    GLE 的 PaddleOcrEngine 一樣真的建構,而且版本涵蓋比它新。
用法:python3 VRN_ENG072_FirstPageText_v0110.py run [--dir 報告夾]
        [--open] | --selftest
"""
from __future__ import annotations
# ===== [VIA:ACCEL-BRIDGE:v0100] SuperAccel 加速器橋(批102 全樹導入令;graceful 零行為變更) =====
try:
    import sys as _sa_sys
    from pathlib import Path as _sa_Path
    _sa_p = _sa_Path(__file__).resolve()
    while _sa_p.parent != _sa_p:
        if (_sa_p / "supportive modules" / "VIA_SuperAccel_Module.py").exists():
            _sa_sys.path.insert(0, str(_sa_p / "supportive modules"))
            break
        _sa_p = _sa_p.parent
    import VIA_SuperAccel_Module as VIA_ACCEL  # noqa: N816
except Exception:
    VIA_ACCEL = None  # graceful:加速器缺席零影響
# ===== [VIA:ACCEL-BRIDGE:END] =====

import html
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
VIA = HERE.parent.parent
DEFAULT_DIR = HERE / "input_reports"
#: 輸入規格正典(與 CGC_MDL141 / MDL139 / 首頁全能引擎同一本冊,不另立一套)
SPEC_REL = Path("supportive modules") / "registry" / "VIA_InputConsole_Spec_v0100.json"
OUTDIR = VIA / "VIA_Reports" / "first_page_text"


ZH_END = "。!?!?;;:」)】%"

_INTAKE = HERE / "references" / "intake"


# 批421(操作員令「REGISTER AND IMPLEMENT THESE TWO SYSTEMS AS SUPPORTIVE
# MODULES TO SUPPORT VRN」):兩座收容件改由支援模組正主統轄——
#   SUP_MDL743_GenericLayoutHub   ← GenericLayoutEngine v2.1.0(全 30 後端+路由)
#   SUP_MDL744_NLPApplicationHub  ← VIA_NLP_Application_System v1.8.0(39 模組)
# 本檔 v0106 的 _gle_mod 尾版律是對的(批246 已修字母序陷阱),但那份實作只活在
# 這支引擎裡,ENG073/ENG074 想用就得各抄一份=九頭龍;而 _nlp_tp 更是把
# VIA_NLP_OneEngine_v1.1.0 **寫死**,v1.5.0/v1.8.0 收容了也永遠掛不上。
# 零回歸律:橋缺席 → 一律退回本檔原有的直掛實作,行為與 v0106 逐字相同。
HUB_DIR = VIA / "supportive modules" / "70_VRN_Rules"
_HUB: dict = {}


def _hub(glob_pat: str, key: str):
    """支援模組動態載入(尾版律 glob);缺檔/例外=None 並留因由"""
    if key in _HUB:
        return _HUB[key]["mod"]
    rec = {"mod": None, "why": "", "src": ""}
    _HUB[key] = rec
    try:
        hits = sorted(HUB_DIR.glob(glob_pat))
        if not hits:
            rec["why"] = f"{glob_pat} 缺檔(supportive modules/70_VRN_Rules/)"
            return None
        import importlib.util
        spec = importlib.util.spec_from_file_location(f"via_hub_{key}", hits[-1])
        mod = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = mod
        spec.loader.exec_module(mod)
        st = mod.mount()
        if st["state"] in ("ABSENT", "FAILED"):
            rec["why"] = f"橋在位但掛載 {st['state']}:{st.get('why', '')}"
            return None
        rec["mod"], rec["src"] = mod, f"{hits[-1].name}→{st.get('dir_name')}({st['state']})"
        return mod
    except Exception as exc:
        rec["why"] = f"橋載入失敗 {type(exc).__name__}:{str(exc)[:80]}"
    return None


def gle_hub():
    return _hub("SUP_MDL743_GenericLayoutHub_v*.py", "gle")


def nlp_hub():
    return _hub("SUP_MDL744_NLPApplicationHub_v*.py", "nlp")


def conv_mod():
    """批449:轉檔器 MDL001(尾版 glob)。只借它的 `auto_dpi`——
    第三階要用什麼 DPI,批448 已經把「量得出來的依據」寫在那裡了,
    這裡不自己再寫一套(Zero-Hydra)。缺席就退回帶頂並講明。"""
    rec = _HUB.setdefault("conv", {})
    if "mod" in rec:
        return rec["mod"]
    rec["mod"] = None
    try:
        hits = sorted(Path(__file__).resolve().parent.glob("VRN_MDL001_Converter_v*.py"))
        if not hits:
            rec["why"] = "VRN_MDL001_Converter_v*.py 缺(批448 之前的樹)"
            return None
        import importlib.util
        sp = importlib.util.spec_from_file_location("via_mdl001_for_eng072", hits[-1])
        m = importlib.util.module_from_spec(sp)
        sys.modules[sp.name] = m
        sp.loader.exec_module(m)
        if not hasattr(m, "auto_dpi"):
            rec["why"] = f"{hits[-1].name} 無 auto_dpi(需批448+)"
            return None
        rec["mod"], rec["src"], rec["why"] = m, hits[-1].name, ""
        return m
    except Exception as exc:
        rec["why"] = f"掛載失敗 {type(exc).__name__}:{str(exc)[:60]}"
    return None


def ppp_hub():
    """批444:PDFPlumber-Plus 統轄橋(SUP_MDL746)。兩件事——
    ① 字元密度分流閘(本檔今天的 `if txt:` 假綠,由它擋)
    ② paddleocr 2.x 的 OCR 補位車道(3.x 走 GLE;誰在位誰上)"""
    return _hub("SUP_MDL746_PDFPlumberPlusHub_v*.py", "ppp")


def hub_stats() -> dict:
    """兩橋現況(給 run 尾段顯示;捕捉到卻不顯示=等於沒捕捉)"""
    gle_hub(), nlp_hub(), ppp_hub()
    out = {}
    for k, label in (("gle", "GLE"), ("nlp", "NLP"), ("ppp", "PPP")):
        r = _HUB.get(k, {})
        out[label] = {"on": r.get("mod") is not None,
                      "src": r.get("src", ""), "why": r.get("why", ""),
                      "route": r.get("route", "")}
    return out


def _gle_mod():
    """GLE 正主:① SUP_MDL743 橋(批421)② 本檔原有直掛(v0106 行為,零回歸)。
    批243 掛載;批245 尾版化:glob GenericLayoutEngine_* 取尾版夾
    (嚴禁寫死版號;v2.1.0 收容落地即自動升級;收容件原地不動)"""
    h = gle_hub()
    if h is not None:
        g = h.gle()
        if g is not None:
            _HUB["gle"]["route"] = "HUB"
            return g
    _HUB.setdefault("gle", {})["route"] = "LEGACY_DIRECT"
    try:
        cands = []
        for top in sorted(_INTAKE.glob("GenericLayoutEngine_*")):
            if not top.is_dir():
                continue
            if (top / "generic_layout_engine.py").exists():
                cands.append(top)            # 頂層型(v2.0.0_b242)
            for sub in sorted(top.iterdir()):
                if sub.is_dir() and (sub / "generic_layout_engine.py").exists():
                    cands.append(sub)        # 巢狀一層型(AllEngines v2.1.0_b245)
        if not cands:
            return None                      # 收容缺席=誠實 absent
        import re as _re

        def _ver(d):
            """語意版號排序(批246 修:字母序 AllEngines_v2.1.0<v2.0.0
            陷阱)——取夾名鏈上 vX.Y.Z 最大者;無版號=(0,0,0)"""
            names = d.name + "|" + d.parent.name
            vs = [tuple(int(x) for x in m.groups())
                  for m in _re.finditer(r"v(\d+)\.(\d+)\.(\d+)", names)]
            return max(vs) if vs else (0, 0, 0)
        d = max(cands, key=_ver)             # 尾版=語意最高版
        if str(d) not in sys.path:
            sys.path.insert(0, str(d))
        import generic_layout_engine as GLE  # noqa: N811
        GLE._VIA_ENGINE_DIR = d.parent.name if d.parent != _INTAKE \
            else d.name                      # 掛載夾存證(sidecar)
        return GLE
    except Exception:
        return None                          # 缺席=誠實 absent,零影響


def _nlp_tp():
    """TextProcessor:① SUP_MDL744 橋(批421;尾版 v1.8.0,39 模組)
    ② 本檔原有 v1.1.0 直掛(v0106 行為,零回歸)③ 皆缺=None。
    批243:NLP OneEngine TextProcessor 正掛(帶 ssot_lexicon;
    修 ENG073 v0103 漏 lexicon_path 之債)"""
    h = nlp_hub()
    if h is not None:
        tp = h.text_processor()
        if tp is not None:
            _HUB["nlp"]["route"] = "HUB"
            return tp
    _HUB.setdefault("nlp", {})["route"] = "LEGACY_v1.1.0"
    try:
        pkg = _INTAKE / "VIA_NLP_OneEngine_v1.1.0"
        if str(pkg / "src") not in sys.path:
            sys.path.insert(0, str(pkg / "src"))
        from via_nlp_engine.text_ops import TextProcessor  # noqa
        lex = pkg / "data" / "lexicon" / "ssot_lexicon.json"
        return TextProcessor(lex) if lex.exists() else None
    except Exception:
        return None


def gle_annotate(blocks: list[dict], W: float, H: float) -> dict:
    """法C(批243):GLE 九宮分區標+字重階層+本文字級推定+後端矩陣。
    批421:橋在位時直接用 SUP_MDL743.zone_annotate(輸出契約逐鍵相同,
    該橋自測⑦ 對此有專檢)=實作只留一份;橋缺席才走本檔原有路徑。"""
    h = gle_hub()
    if h is not None:
        r = h.zone_annotate(blocks, W, H)
        if r.get("available"):
            _HUB["gle"]["route"] = "HUB"
            return r
    GLE = _gle_mod()
    if GLE is None:
        return {"available": False, "note": "GLE 收容件缺席=誠實 absent"}
    try:
        els = []
        for b in blocks:
            bb = GLE.BBox(b["x0"], b["y0"], b["x1"], b["y1"])
            wt, bold, italic = GLE.font_weight_from_name(
                b.get("font", ""), 16 if b.get("bold") else 0)
            els.append({"zone": GLE.zone_for_bbox(bb, W, H),
                        "weight": wt, "bold": bold, "italic": italic,
                        "size": round(b["size"], 1),
                        "head": " ".join(b["lines"])[:80]})
        body_font = GLE.weighted_median(
            [(b["size"], sum(len(ln) for ln in b["lines"])) for b in blocks])
        backends = {s.name: s.available for s in GLE.probe_backends()}
        return {"available": True, "engine": "GenericLayoutEngine/2.x",
                "engine_dir": getattr(GLE, "_VIA_ENGINE_DIR", "?"),
                "body_font": round(body_font, 1) if body_font else None,
                "elements": els,
                "backends_available": sorted(k for k, v in backends.items() if v),
                "backends_total": len(backends)}
    except Exception as exc:
        return {"available": False, "note": f"GLE 例外({type(exc).__name__})"}


def split_sentences_repaired(body: str) -> list[str]:
    """句級修復(批243;操作員長文:一句一句號一資料):TextProcessor
    正主(句切+疊字修);缺席=stdlib 終止標點後備。標題不經此道=無句號。"""
    if not body.strip():
        return []
    import re as _re
    # 彈點符前置切分(操作員真件樣本:➢ 摘要彈點=一彈點一資料)
    body = _re.sub(r"\s*([➢◆●■▲►•])", r"\n\1", body)
    tp = _nlp_tp()
    if tp is not None:
        try:
            sents = []
            for para in body.splitlines():
                if para.strip():
                    sents.extend(tp.split_sentences(para))
            return sents
        except Exception:
            pass
    import re as _re
    out = []
    for para in body.splitlines():
        out.extend(s.strip() for s in
                   _re.split(r"(?<=[。!?!?])", para) if s.strip())
    return out


def _repair_lines(lines: list[str]) -> str:
    """本文修復:斷行接回(批236 ②)——句尾無終止標點=延續句合併"""
    out: list[str] = []
    for ln in lines:
        ln = ln.strip()
        if not ln:
            continue
        if out and out[-1] and out[-1][-1] not in ZH_END \
                and not out[-1].endswith("."):
            prev = out.pop()
            if prev.endswith("-") and ln[:1].isalpha():
                out.append(prev[:-1] + ln)          # 連字號斷詞
            elif prev[-1].isascii() and ln[:1].isascii():
                out.append(prev + " " + ln)          # 英文補空白
            else:
                out.append(prev + ln)                # 中文直接接
        else:
            out.append(ln)
    return "\n".join(out)


def extract_page1_zones(p: Path) -> dict | None:
    """分區 LAYOUT(批236 ①③):四區+字級階層;回 None=無文字層"""
    try:
        import fitz
    except Exception:
        return None
    try:
        with fitz.open(str(p)) as doc:
            if not doc.page_count:
                return None
            page = doc[0]
            W, H = page.rect.width, page.rect.height
            blocks = []
            for b in page.get_text("dict")["blocks"]:
                if b.get("type") != 0:
                    continue
                lines, mx, bold, font = [], 0.0, False, ""
                for ln in b["lines"]:
                    t = "".join(sp["text"] for sp in ln["spans"]).strip()
                    if t:
                        lines.append(t)
                    for sp in ln["spans"]:
                        if sp["size"] > mx:
                            mx = sp["size"]
                            font = sp.get("font", "")   # 批243:字型名供 GLE 字重
                        if (sp["flags"] & 16) or "Bold" in sp.get("font", ""):
                            bold = True
                if not lines:
                    continue
                x0, y0, x1, y1 = b["bbox"]
                blocks.append({"lines": lines, "x0": x0, "y0": y0,
                               "x1": x1, "y1": y1, "size": mx, "bold": bold,
                               "font": font})
            if not blocks:
                return None
            zones = {"header": [], "right": [], "body": [], "footer": []}
            for b in blocks:
                wide = (b["x1"] - b["x0"]) > 0.66 * W
                if b["y1"] > 0.92 * H and b["size"] < 9:
                    z = "footer"                     # 頁尾小字=雜訊帶
                elif wide and b["y0"] < 0.18 * H:
                    z = "header"                     # 全寬標題帶
                elif b["x0"] >= 0.52 * W:
                    z = "right"                      # 右資訊區(卡)
                else:
                    z = "body"                       # 左本文區
                zones[z].append(b)
            for z in zones:
                zones[z].sort(key=lambda b: (round(b["y0"], 1), b["x0"]))
            heads = [{"text": " ".join(b["lines"])[:120],
                      "size": round(b["size"], 1),
                      "level": ("H1" if b["size"] >= 18 else
                                "H2" if b["size"] >= 14 else "H3")}
                     for b in blocks
                     if b["size"] >= 13 or (b["bold"] and b["size"] >= 11)]
            body_lines = [ln for b in zones["body"] for ln in b["lines"]]
            right_lines = [ln for b in zones["right"] for ln in b["lines"]]
            body = _repair_lines(body_lines)
            return {
                "header": "\n".join(ln for b in zones["header"]
                                     for ln in b["lines"]),
                "right": "\n".join(right_lines),
                "body": body,
                "footer": "\n".join(ln for b in zones["footer"]
                                     for ln in b["lines"]),
                "heads": heads,
                # 批243 整合鍵(fitz 區鍵仍頂層=下游零改)
                "gle": gle_annotate(blocks, W, H),
                "sentences": split_sentences_repaired(body)}
    except Exception:
        return None


def extract_page1_plumber(p: Path) -> dict | None:
    """法B(批242):pdfplumber words 同判準分區+表格還原"""
    try:
        import pdfplumber
    except Exception:
        return None
    try:
        with pdfplumber.open(str(p)) as pdf:
            if not pdf.pages:
                return None
            pg = pdf.pages[0]
            W, H = pg.width, pg.height
            zones = {"header": [], "right": [], "body": [], "footer": []}
            line_map: dict = {}
            for w in pg.extract_words(use_text_flow=False):
                key = round(w["top"] / 4)
                line_map.setdefault(key, []).append(w)
            for key in sorted(line_map):
                ws = sorted(line_map[key], key=lambda x: x["x0"])
                # 批242 修:同 y 左右欄併行病→行內大縫隙(>15%W)切段,
                # 每段獨立分區(對齊 fitz block 語意)
                parts, cur = [], [ws[0]]
                for w in ws[1:]:
                    if w["x0"] - cur[-1]["x1"] > 0.15 * W:
                        parts.append(cur)
                        cur = [w]
                    else:
                        cur.append(w)
                parts.append(cur)
                for seg in parts:
                    x0 = min(x["x0"] for x in seg)
                    x1 = max(x["x1"] for x in seg)
                    top = min(x["top"] for x in seg)
                    bot = max(x["bottom"] for x in seg)
                    txt = " ".join(x["text"] for x in seg)
                    wide = (x1 - x0) > 0.66 * W
                    if bot > 0.92 * H:
                        z = "footer"
                    elif wide and top < 0.18 * H:
                        z = "header"
                    elif x0 >= 0.52 * W:
                        z = "right"
                    else:
                        z = "body"
                    zones[z].append(txt)
            tables = []
            for t in (pg.extract_tables() or []):
                rows = [[(c or "").strip() for c in row] for row in t]
                if any(any(c for c in row) for row in rows):
                    tables.append(rows[:30])
            return {"header": "\n".join(zones["header"]),
                    "right": "\n".join(zones["right"]),
                    "body": _repair_lines(zones["body"]),
                    "footer": "\n".join(zones["footer"]),
                    "tables": tables}
    except Exception:
        return None


def compare_zones(a: dict, b: dict) -> dict:
    """雙法逐區對照(批242):正規化 difflib 比率+三態判定"""
    import difflib
    out = {}
    for k in ("header", "right", "body"):
        ta = "".join(str(a.get(k, "")).split())
        tb = "".join(str(b.get(k, "")).split())
        if not ta and not tb:
            out[k] = {"ratio": 1.0, "verdict": "BOTH_EMPTY"}
            continue
        r = difflib.SequenceMatcher(None, ta, tb).ratio()
        out[k] = {"ratio": round(r, 3),
                  "verdict": ("AGREE" if r >= 0.90 else
                              "PARTIAL" if r >= 0.60 else "DIVERGE")}
    return out


def triage_page1(p: Path) -> dict:
    """批444:這一頁到底**有沒有**文字層——不是「抽不抽得到字」。

    v0109 以前的判準是 `if txt: return txt`:抽得到**一個字**就算有文字層。
    一頁只有浮水印文字層(「機密」「DRAFT」十來個字)的掃描頁會被這一問放行,
    首頁引擎再從那十來個字裡找目標價——**那是假綠**。

    SUP_MDL746 把「字數 ÷ 頁面點面積」一起看,誠實三態
    DIGITAL/THIN/SCANNED;橋缺席就回 UNKNOWN,**維持 v0109 原行為零回歸**
    (不裝作判得出來)。
    """
    h = ppp_hub()
    if h is None:
        return {"state": "UNKNOWN", "why": (_HUB.get("ppp", {}).get("why")
                                            or "SUP_MDL746 缺席")}
    try:
        return dict(h.triage(p))
    except Exception as exc:
        return {"state": "UNKNOWN", "why": f"分流橋失敗 {type(exc).__name__}"}


def extract_pdf_page1(p: Path) -> tuple[str, str]:
    """回 (文字, 方法標記);fitz 版面序優先,pypdf 後備,零字=NEEDS_OCR。

    批444 起先過**字元密度分流閘**:判 SCANNED 的頁就算 fitz 抽得到字
    也不當文字層用(那是浮水印),改走 OCR 車道;判 THIN 的照用但把
    「稀薄」標出來(長捲頁/海報那種真文字層,不許只憑密度假紅)。
    """
    tri = triage_page1(p)
    st = tri.get("state") or "UNKNOWN"
    sfx = "[THIN]" if st == "THIN" else ""
    if st != "SCANNED":
        try:
            import fitz
            with fitz.open(str(p)) as doc:
                if doc.page_count:
                    # sort=True=版面閱讀序(LAYOUT ANALYSIS;digest ㉙ 同族)
                    txt = doc[0].get_text("text", sort=True).strip()
                    if txt:
                        return txt, "FITZ_LAYOUT" + sfx
        except Exception:
            pass
        try:
            import pypdf
            r = pypdf.PdfReader(str(p))
            if r.pages:
                txt = (r.pages[0].extract_text() or "").strip()
                if txt:
                    return txt, "PYPDF_FALLBACK" + sfx
        except Exception:
            pass
    # ── 批449 操作員令「先用非 OCR 法,再用 OCR 法,最後還抓不到再調高畫質 OCR 法」──
    # 第一階(非 OCR)就在上面:fitz 版面序 → pypdf 後備,由密度閘守門。
    # 第二階:原畫質 OCR(GLE 七支後端 → PPP 補位)。
    _t, _tag = ocr_page1(p)
    if _t.strip():
        return _t, _tag
    # 第三階:**高畫質 OCR**。工作站實錄第二階回的是
    #   OCR_EMPTY(tesseract+paddleocr+…;跑了但零字=影像可能不可辨)
    # ——後端在位、真的跑了、就是讀不出字。那多半是**渲染解析度不夠**,
    # 不是後端不行。把頁面重繪到批448 那條帶(300–350)再餵同一條 OCR 車道。
    # 只在「跑了但零字」時升階;後端根本沒裝(ABSENT)時升階是白跑,誠實跳過。
    _ran = ocr_did_run(_tag)
    if _ran:
        _t3, _tag3 = ocr_page1_hq(p)
        if _t3.strip():
            return _t3, _tag3
        _tag = f"{_tag}|{_tag3}"
    else:
        _tag = f"{_tag}|HQ_SKIP(第二階不是「跑了但零字」而是後端不在位,升階是白跑)"
    if st == "SCANNED" and (tri.get("n_chars") or 0) > 0:
        # 抽得到字但密度判它是掃描頁,而 OCR 又跑不動——**照實說有這件事**,
        # 不要靜靜把浮水印那十幾個字當首頁文字送下去(那才是最難查的假綠)。
        return "", (f"NEEDS_OCR[{_tag}·密度閘:{tri.get('why') or ''}]")
    return "", f"NEEDS_OCR[{_tag}]"


#: GLE 的 OCR 路由(SUP_MDL743 → multi_engine_orchestrator MODE_ADAPTERS["ocr"])
#: 實測 7 支:tesseract · paddleocr · paddle_ppstructure · paddle_pdf_pipeline ·
#: easyocr · ocrmypdf · transkribus_core。**PaddleOCR 系列在裡面三支**,
#: 另有 paddle 模式再掛 paddle_layout / paddle_detection 共五支。
OCR_INSTALL = (
    "pip install paddleocr paddlepaddle   # CJK 最佳,操作員指名的那一系\n"
    "     或 pip install pytesseract 並裝 Tesseract-OCR 本體(含 chi_tra 語言包)\n"
    "     或 pip install easyocr")


def ocr_page1(p: Path) -> tuple[str, str]:
    """掃描影像 PDF 的 OCR 車道(批443 操作員令「要用到 OCR 引擎也用起來」)。

    **先查再造(第四次)**:OCR 早就在庫裡——GenericLayoutEngine v2.1.0 的
    all_backend_engines.py 有 tesseract / paddleocr / paddle_ppstructure /
    paddle_layout / paddle_pdf_pipeline / paddle_detection / easyocr / ocrmypdf
    八個後端,而 SUP_MDL743 橋也早就掛著;只是 NEEDS_OCR 的檔從來沒去走它。

    批444 起**兩條車道**,誰在位誰上,順序有理由不是偏好:
      道一 GLE(SUP_MDL743)——七支後端,而且 PaddleOCR **3.x 新 API 都接得動**
            (PPStructureV3 / LayoutDetection / engine.predict())
      道二 PPP(SUP_MDL746)——PDFPlumber-Plus 收容件,只接 paddleocr **2.x**
            (PPStructure / use_gpu / show_log 都是 2.x 才有的鍵),
            但它記憶體內 fitz→NumPy 渲染,不落暫存圖
    道一涵蓋面比較大所以先走;道一整條都不在位才換道二。**補位不是取代。**

    紀律:**一個後端都沒裝時絕不假抽**——回空字串,並把路由、逐支在位狀態、
    以及要裝什麼講清楚。「候 OCR」三個字讓人不知道下一步,那等於沒說。
    """
    t, tag = _ocr_via_gle(p)
    if t.strip():
        return t, tag
    t2, tag2 = _ocr_via_ppp(p)
    if t2.strip():
        return t2, tag2
    return "", f"{tag}|{tag2}"


def ocr_did_run(tag: str) -> bool:
    """第二階到底**跑了沒有**?只有「跑了但零字」才值得升階到高畫質。

    批450:v0111 用 `"_EMPTY(" in tag` 判——**帶括號**。而本批把標記從
    `OCR_EMPTY(…)` 改成 `OCR_EMPTY[…;逐支因由]` 之後,那個判準就永遠為假,
    第三階會**靜靜地再也不升階**,而且沒有任何一行會說它不升了。
    寫成一支具名函式,判準就測得起來(檢 ㉑)。

    三種態要分得開:
      跑了但零字  OCR_EMPTY[…] / PPP_OCR_EMPTY(…) / OCR_ERROR  → 升階
      根本沒跑    OCR_ROUTE_EMPTY(…)                          → 不升(沒東西可升)
      後端不在位  OCR_BACKENDS_ABSENT(…) / PPP_OCR_ABSENT(…)    → 不升(白跑)
    """
    t = str(tag or "")
    if "ROUTE_EMPTY" in t:
        return False
    return ("_EMPTY" in t) or ("OCR_ERROR" in t)


def hq_dpi_for(p: Path) -> tuple[int, str]:
    """第三階要用幾 DPI。借 MDL001(批448)的 auto_dpi;缺席就退帶頂 350 並講明。"""
    m = conv_mod()
    if m is None:
        return 350, f"MDL001 缺({_HUB.get('conv', {}).get('why') or '?'})→退帶頂 350"
    try:
        import fitz
        with fitz.open(str(p)) as doc:
            if not doc.page_count:
                return 350, "零頁→退帶頂 350"
            d, why = m.auto_dpi(doc[0], doc,
                                m._DEFAULTS["dpi_min"], m._DEFAULTS["dpi_max"],
                                m._DEFAULTS["dpi_max_pixels"])
        return int(d), f"MDL001 auto_dpi:{why}"
    except Exception as exc:
        return 350, f"auto_dpi 失敗 {type(exc).__name__}→退帶頂 350"


def ocr_page1_hq(p: Path) -> tuple[str, str]:
    """第三階:**高畫質 OCR**(批449 操作員令「最後還抓不到再調高畫質 OCR 法」)。

    作法:把第 1 頁用 MDL001 選出的 DPI(300–350 帶內)重繪成**一頁的暫存 PDF**,
    再餵給**同一條** OCR 車道(ocr_page1)。這樣第三階與第二階跑的是同一組後端、
    同一套路由,差別只在**輸入影像的解析度**——變因只有一個,才說得清是不是
    解析度的問題。

    Zero-Hydra:不另外寫一條 OCR 路由;暫存檔跑完就刪,不落地污染。
    """
    dpi, why = hq_dpi_for(p)
    try:
        import fitz
    except Exception:
        return "", "HQ_NO_FITZ(重繪不了)"
    import tempfile
    try:
        with tempfile.TemporaryDirectory() as td:
            hq = Path(td) / (p.stem + f"_hq{dpi}.pdf")
            with fitz.open(str(p)) as doc:
                if not doc.page_count:
                    return "", "HQ_NO_PAGE"
                src = doc[0]
                mat = fitz.Matrix(dpi / 72.0, dpi / 72.0)
                pix = src.get_pixmap(matrix=mat, alpha=False)
                out = fitz.open()
                pg = out.new_page(width=src.rect.width, height=src.rect.height)
                pg.insert_image(pg.rect, stream=pix.tobytes("png"))
                out.save(str(hq))
                out.close()
            t, tag = ocr_page1(hq)
        if t.strip():
            return t, f"HQ{dpi}_{tag}"
        return "", f"HQ{dpi}_EMPTY({tag};{why})"
    except Exception as exc:
        return "", f"HQ_FAIL({type(exc).__name__}:{str(exc)[:50]};{why})"


def _ocr_via_ppp(p: Path) -> tuple[str, str]:
    """道二:PDFPlumber-Plus(SUP_MDL746)的 paddleocr 2.x 車道。"""
    h = ppp_hub()
    if h is None:
        return "", f"PPP_HUB_ABSENT({_HUB.get('ppp', {}).get('why') or 'SUP_MDL746 缺席'})"
    try:
        return tuple(h.ocr_page(p, 1))
    except Exception as exc:
        return "", f"PPP_RUN_FAIL({type(exc).__name__}:{str(exc)[:50]})"


#: 批450:OCR 空手而回時,逐支後端**照抄編排器自己寫的**狀態/錯誤/警告。
#: 這些欄位 AdapterResult 本來就有(status/error/warnings/elements/probe),
#: v0111 全丟掉了。丟掉證據再編一個理由,比不講還糟。
_ADAPTER_DETAIL_MAX = 3


def _adapter_detail(run) -> str:
    """把 run.adapter_results 濃縮成一句人看得懂的因由。零推測,只照抄。"""
    outs = []
    try:
        rs = list(getattr(run, "adapter_results", None) or [])
    except Exception:
        return "編排器未回 adapter_results(查 GLE 版本)"
    if not rs:
        return "編排器回的 adapter_results 是空的=沒有任何後端留下紀錄"
    for r in rs:
        nm = getattr(r, "adapter_name", "?")
        st = getattr(r, "status", "?")
        n = len(getattr(r, "elements", None) or [])
        bits = [f"{nm}:{st}/{n}元素"]
        err = getattr(r, "error", None)
        if err:
            bits.append(f"錯={str(err)[:70]}")
        for w in (getattr(r, "warnings", None) or [])[:2]:
            bits.append(f"警={str(w)[:80]}")     # ← tesseract 的缺語言包就在這
        if not err and not (getattr(r, "warnings", None) or []):
            pr = getattr(r, "probe", None)
            rs2 = getattr(pr, "reason", "") if pr is not None else ""
            if rs2 and rs2 != "available":
                bits.append(f"探針={str(rs2)[:60]}")
        outs.append(" ".join(bits))
    head = outs[:_ADAPTER_DETAIL_MAX]
    more = f" …另 {len(outs) - len(head)} 支" if len(outs) > len(head) else ""
    return " · ".join(head) + more


def _ocr_via_gle(p: Path) -> tuple[str, str]:
    """道一:GLE(SUP_MDL743)七支後端路由,含 paddleocr 3.x 新 API。"""
    hub = gle_hub()
    if hub is None:
        return "", "OCR_HUB_ABSENT(SUP_MDL743 缺席=無法路由)"
    try:
        route = (hub.route_modes() or {}).get("ocr") or []
        mx = hub.backend_matrix() or {}
        on = set(mx.get("on") or [])
    except Exception as exc:
        return "", f"OCR_PROBE_FAIL({type(exc).__name__})"
    usable = [x for x in route if x in on]
    if not usable:
        return "", (f"OCR_BACKENDS_ABSENT(路由 {len(route)} 支全未裝:"
                    f"{','.join(route[:4])}…;裝法見 [OCR 未就緒] 提示)")
    try:
        O = hub.orchestrator()
        cfg = O.OrchestratorConfig(mode="ocr")
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            run = O.run_orchestrator(p, Path(td), cfg)
        txt = "\n".join(e.text for e in (run.canonical_elements or []) if e.text)
        # ── 批450:**不准再編「跑了但零字」** ────────────────────────────────
        # v0111 寫的是 `used = run.route or usable`——route 空掉時就拿
        # **「矩陣說裝了的那些」**去填,然後宣稱它們「跑了」。工作站實錄裡
        # 同一行同時說 GLE 的 paddleocr 跑過、PPP 說「缺 paddleocr」,
        # 兩句不可能都真。那句「跑了」從頭到尾沒有人驗過。
        #
        # 而且真正的因由**早就在 run.adapter_results 裡**:GLE 的 tesseract
        # 轉接器自己會寫 `missing Tesseract languages: …; using …`,
        # 沒有任何語言包時更是直接拋 RuntimeError。我把那些全丟掉,
        # 只留 canonical_elements,再自己掰一句「影像可能不可辨」——
        # 批419e 那一課的最深一層:**捕捉到卻不顯示,然後編一個代替它**。
        used = "+".join(dict.fromkeys(run.route or []))
        detail = _adapter_detail(run)
        if txt.strip():
            return txt.strip(), f"OCR_{used or '?'}"
        if not used:
            return "", ("OCR_ROUTE_EMPTY(編排器回的路由是空的=**沒有任何後端真的跑**;"
                        f"矩陣說在位的是 {','.join(usable[:4])}…,兩者不一致要查 GLE 探針)")
        return "", f"OCR_EMPTY[{used};{detail}]"
    except Exception as exc:
        return "", f"OCR_RUN_FAIL({type(exc).__name__}:{str(exc)[:60]})"


def extract_docx_head(p: Path, n_para: int = 30) -> tuple[str, str]:
    try:
        import docx
        d = docx.Document(str(p))
        paras = [q.text for q in d.paragraphs if q.text.strip()][:n_para]
        if paras:
            return "\n".join(paras), "DOCX_HEAD"
        return "", "DOCX_EMPTY"
    except ImportError:
        return "", "DOCX_LIB_MISSING(pip install python-docx)"
    except Exception as exc:
        return "", f"DOCX_ERR({type(exc).__name__})"


def report_dirs(given: Path | None = None) -> tuple[list[Path], str]:
    """報告夾律(批437-A6)。**與首頁全能引擎、MDL141、MDL139 同一本冊**:
       --dir > 冊 user.vrn_dir > 冊 dir_default;**incoming 一律併入**。

    為什麼要改:v0107 把 `input_reports` 寫死成唯一預設,而操作員的 64 份報告
    住在正典的 `functional modules/VRN/input/incoming`。他照著跑 `run`,
    得到「無報告件(誠實)」——話是誠實的,結論是錯的:**報告就在旁邊那個夾**。
    後果不只是少一行輸出:sidecar 一份都沒生,於是首頁全能引擎的
    text_consensus 全部 N/A_NO_SIDECAR,批435 造的多工具核對整層睡著,
    中信金那種「兩個工具讀出不同目標價」的事就沒有人抓得到。
    一個夾的預設值,關掉了一整層防護。
    """
    import json as _json
    dirs: list[Path] = []
    blk, why = {}, "冊不在 → 只用內建預設"
    sp = VIA / SPEC_REL
    if sp.exists():
        try:
            d = _json.loads(sp.read_text(encoding="utf-8-sig"))
            blk = d["families"]["vrn"]["input"]
            user = str((d.get("user") or {}).get("vrn_dir") or "").strip()
            why = f"冊 {sp.name}"
        except Exception as exc:
            user, why = "", f"冊讀取失敗 {type(exc).__name__} → 內建預設"
    else:
        user = ""
    if given is not None:
        dirs.append(given)
    elif user:
        q = Path(user)
        dirs.append(q if q.is_absolute() else VIA / q)
    else:
        dirs.append(VIA / blk.get("dir_default", "functional modules/VRN/input_reports"))
    inc = VIA / blk.get("incoming", "functional modules/VRN/input/incoming")
    if inc not in dirs:
        dirs.append(inc)
    return dirs, why



def _ocr_banner() -> None:
    """跑之前先把 OCR 就緒狀態講清楚——不是等到有掃描檔才講(批443)。
    批444:兩條車道各報各的,不要把「有一條在」講成「OCR 好了」。"""
    _use, _rt = [], []
    try:
        _h = gle_hub()
        if _h is not None:
            _rt = (_h.route_modes() or {}).get("ocr") or []
            _on = set((_h.backend_matrix() or {}).get("on") or [])
            _use = [x for x in _rt if x in _on]
    except Exception as _exc:
        print(f"[OCR 探測失敗] 道一 GLE {type(_exc).__name__}(不影響其他車道)")
    _p_ok, _p_why = False, "SUP_MDL746 缺席"
    try:
        _hp = ppp_hub()
        if _hp is not None:
            _p_ok, _p_why = _hp.ocr_ready()
        else:
            _p_why = _HUB.get("ppp", {}).get("why") or _p_why
    except Exception as _exc:
        _p_why = f"探測失敗 {type(_exc).__name__}"
    if _use or _p_ok:
        _s1 = f"道一 GLE 在位 {len(_use)}/{len(_rt)}:{','.join(_use) or '-'}"
        _s2 = "道二 PPP 可跑" if _p_ok else f"道二 PPP 不可跑({_p_why})"
        print(f"[OCR 就緒] {_s1} · {_s2}")
        # 批449:三階梯講清楚,不要讓人以為 OCR 只有一次機會
        _cm = conv_mod()
        print("   階梯 ①非OCR(fitz→pypdf)→ ②原畫質OCR → ③高畫質OCR"
              + (f"(MDL001 {_HUB.get('conv', {}).get('src')} auto_dpi 300–350)"
                 if _cm is not None else
                 f"(MDL001 缺:{_HUB.get('conv', {}).get('why') or '?'}→退帶頂 350)")
              + ";③ 只在 ② **跑了但零字**時才升階(沒裝就升階是白跑)")
    else:
        print(f"[OCR 未就緒] 道一 GLE 路由 {len(_rt)} 支全未裝"
              f"{'(' + ', '.join(_rt) + ')' if _rt else ''} · 道二 PPP {_p_why}")
        print("   掃描影像 PDF 會誠實停在 NEEDS_OCR(不假抽)。要開通擇一:")
        for _ln in OCR_INSTALL.splitlines():
            print(f"   {_ln}")


def run(src: Path | None = None, open_after: bool = False) -> int:
    dirs, spec_why = report_dirs(src)
    files, diag = [], []
    for d in dirs:
        got = sorted([*d.glob("*.pdf"), *d.glob("*.docx")]) if d.is_dir() else []
        if not d.is_dir():
            st = "夾不存在(換路徑)"
        elif not any(d.iterdir()):
            st = "夾是空的(放檔進去)"
        elif not got:
            from collections import Counter
            other = Counter(x.suffix.lower() or "(無副檔名)"
                            for x in d.iterdir() if x.is_file())
            st = ("夾有檔但沒有報告檔:"
                  + ",".join(f"{k}×{v}" for k, v in other.most_common(5))
                  if other else "夾裡只有子夾,沒有檔")
        else:
            st = f"{len(got)} 件"
        diag.append((d, st, got))
        files.extend(got)
    files = sorted({f.resolve(): f for f in files}.values())
    if not files:
        _ocr_banner()
        # 三態分開講(與首頁全能引擎批427b 同律):操作員的下一步完全不同
        print(f"[首頁擷取] 無報告件(誠實;缺件搜集器先跑)· 夾律={spec_why}")
        for d, st, _ in diag:
            print(f"   · {d} → {st}")
        return 2
    src = dirs[0]
    _ocr_banner()
    print(f"[首頁擷取] 夾律={spec_why};取件 {len(files)} 份 · "
          + " · ".join(f"{d.name}={st}" for d, st, _ in diag))
    OUTDIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y-%m-%d %H:%M")
    stats = {"DUAL_ZONES": 0, "FITZ_ZONES": 0, "FITZ_LAYOUT": 0,
             "PYPDF_FALLBACK": 0, "NEEDS_OCR": 0, "DOCX": 0, "OTHER": 0,
             "THIN": 0}
    cards = []
    import json as _json
    for p in files:
        zones = None
        if p.suffix.lower() == ".pdf":
            # 批444:分流閘要擋在**分區道前面**。浮水印那十幾個字照樣會讓
            # zones["header"] 非空,分區道就把它當首頁文字收下——閘擋在後面
            # 等於沒擋。判 SCANNED 就整條分區道都不走,直接交給 OCR 車道。
            tri = triage_page1(p)
            zones = None if tri.get("state") == "SCANNED" else extract_page1_zones(p)
            if zones and (zones["body"] or zones["right"] or zones["header"]):
                plum = extract_page1_plumber(p)      # 法B(批242)
                comp = compare_zones(zones, plum) if plum else {}
                zones["plumber"] = {k: plum[k] for k in
                                    ("header", "right", "body", "footer")} \
                    if plum else None
                zones["tables"] = (plum or {}).get("tables", [])
                zones["compare"] = comp
                agree = "/".join(f"{k}:{v['verdict']}"
                                 for k, v in comp.items()) or "單法"
                ttxt = "\n\n".join(
                    "〔表" + str(i + 1) + "〕\n" +
                    "\n".join(" | ".join(r) for r in t)
                    for i, t in enumerate(zones["tables"][:4]))
                sent = "\n".join(zones.get("sentences", []))
                g = zones.get("gle", {})
                glel = (f"GenericLayoutEngine/2.0 · 本文字級 {g.get('body_font')}"
                        f" · 後端可用 {len(g.get('backends_available', []))}"
                        f"/{g.get('backends_total', 0)}"
                        if g.get("available") else "absent(誠實)")
                txt = (f"【標題帶】\n{zones['header']}\n\n"
                       f"【右資訊區】\n{zones['right']}\n\n"
                       f"【本文(修復)】\n{zones['body']}\n\n"
                       f"【本文(句級一句一資料)】\n{sent}\n\n"
                       f"【表格(pdfplumber 還原)】\n{ttxt or '(頁一無表格=誠實)'}\n\n"
                       f"【雙法對照】{agree}\n【法C GLE】{glel}\n\n"
                       f"【頁尾(雜訊帶)】\n{zones['footer']}")
                tag = "DUAL_ZONES" if plum else "FITZ_ZONES"
                if tri.get("state") == "THIN":
                    tag += "[THIN]"          # 長捲頁/海報:照抽但把稀薄講出來
                zones["triage"] = tri
                (OUTDIR / (p.stem + ".json")).write_text(_json.dumps(
                    zones, ensure_ascii=False, indent=1), encoding="utf-8")
            else:
                txt, tag = extract_pdf_page1(p)
        else:
            txt, tag = extract_docx_head(p)
        # 批444 修:v0109 用 `tag if tag in stats`,而批443 起 NEEDS_OCR 帶了
        # 因由後綴(NEEDS_OCR[...]),對不上鍵,整批被算進 OTHER——候OCR 那格
        # 一直報 0。取「[」前的基名才數得對。
        base = tag.split("[", 1)[0]
        key = base if base in stats else ("DOCX" if base.startswith("DOCX") else "OTHER")
        stats[key] = stats.get(key, 0) + 1
        if "[THIN]" in tag:
            stats["THIN"] = stats.get("THIN", 0) + 1
        (OUTDIR / (p.stem + ".txt")).write_text(
            f"# {p.name} · {tag} · {ts}\n\n{txt}", encoding="utf-8")
        state = "ok" if txt else "warn"
        if zones and tag in ("FITZ_ZONES", "DUAL_ZONES"):
            cards.append(
                f"<section class='{state}'><h2>{html.escape(p.name)}"
                f"<span class='tag'>{tag} · 分區還原</span></h2>"
                f"<div class='hd'>{html.escape(zones['header'][:300])}</div>"
                f"<div class='cols'><div class='col'><h3>本文區(修復)</h3>"
                f"<pre>{html.escape(zones['body'][:2000])}</pre></div>"
                f"<div class='col r'><h3>右資訊區</h3>"
                f"<pre>{html.escape(zones['right'][:1200])}</pre></div></div>"
                f"</section>")
        else:
            cards.append(
                f"<section class='{state}'><h2>{html.escape(p.name)}"
                f"<span class='tag'>{tag} · {len(txt):,} 字</span></h2>"
                f"<pre>{html.escape(txt[:2400]) if txt else '(零文字=掃描版或抽取失敗;候 OCR 道,誠實不假抽)'}"
                + ("\n…(全文見同名 .txt)" if len(txt) > 2400 else "") + "</pre></section>")
        print(f"  [{tag}] {p.name} · {len(txt):,} 字")
    summary = f"""<!DOCTYPE html><html lang="zh-Hant"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>個股報告首頁文字總覽</title><style>
body{{background:#0b1220;color:#c7d3e8;font:10.5px/1.55 "Segoe UI",
"Noto Sans TC",sans-serif;padding:14px;max-width:1100px;margin:0 auto}}
h1{{font-size:14px;color:#e8eefb}}
.sub{{color:#7e8db0;font-size:10px;margin:2px 0 10px}}
section{{background:#111a2e;border:1px solid #1e2a44;border-radius:8px;
padding:10px;margin-bottom:10px}}
section.warn{{border-color:#f0b429}}
h2{{font-size:11px;color:#4f8ef7;overflow-wrap:anywhere}}
.tag{{color:#7e8db0;font-size:9px;margin-left:8px}}
pre{{white-space:pre-wrap;overflow-wrap:anywhere;color:#c7d3e8;
font:9.5px/1.5 Consolas,monospace;margin-top:6px;max-height:340px;
overflow:auto}}
.cols{{display:grid;grid-template-columns:3fr 2fr;gap:10px}}
.col h3{{font-size:10px;color:#7e8db0;margin-top:6px}}
.col.r pre{{border-left:2px solid #2a3c61;padding-left:8px}}
.hd{{color:#e8eefb;font-size:10.5px;margin:4px 0;overflow-wrap:anywhere}}
@media(max-width:760px){{.cols{{grid-template-columns:1fr}}}}
</style></head><body>
<h1>個股報告首頁文字擷取總覽(LAYOUT ANALYSIS · 批243 三法整合)</h1>
<div class="sub">{ts} · {len(files)} 件 · 雙法 {stats['DUAL_ZONES']} · 分區 {stats['FITZ_ZONES']} · fitz 版面序 {stats['FITZ_LAYOUT']}
· pypdf 後備 {stats['PYPDF_FALLBACK']} · 候 OCR {stats['NEEDS_OCR']} · 稀薄 {stats['THIN']}
· docx {stats['DOCX']} · 逐件 .txt 同夾 · 不入 git(紅線)</div>
{''.join(cards)}</body></html>"""
    outp = OUTDIR / "FIRSTPAGE_SUMMARY.html"
    outp.write_text(summary, encoding="utf-8")
    print(f"[計] {len(files)} 件 · 雙法 {stats['DUAL_ZONES']} · 分區 {stats['FITZ_ZONES']} · fitz {stats['FITZ_LAYOUT']} · pypdf "
          f"{stats['PYPDF_FALLBACK']} · 候OCR {stats['NEEDS_OCR']} · 稀薄 {stats['THIN']} · docx "
          f"{stats['DOCX']} · 總覽 {outp}")
    if open_after:
        try:
            import webbrowser
            webbrowser.open(outp.as_uri())
        except Exception:
            pass
    return 0


def selftest() -> int:
    """批410:自測**不得**寫進正式產出夾。
    v0105 以前 selftest 的輸入 PDF 走暫存夾,但 run() 的輸出仍寫 OUTDIR=
    VIA_Reports/first_page_text/ → 每跑一次自測就在正式夾留下 fx_report/
    fx_scan/fx_twocol 六個 fixture 檔;下一段 ENG073 會把它們當**真報告**收進
    正本庫(實測:basic +3、metrics +2),於是收尾閘把「尚無報告」誠實態變成
    「報告 3 · FAIL 2」的假象。工作站只要跑過 via-rungate --family vrn /
    via-selftest / SelftestGrid 就會中。治本=自測期間把 OUTDIR 一併重導。"""
    import tempfile
    fails = []
    _OUTDIR_LIVE = OUTDIR
    _live_before = sorted(x.name for x in OUTDIR.glob("*")) if OUTDIR.exists() else []

    done = []

    def chk(name, cond, note=""):
        done.append(name)
        print(f"  [{'OK' if cond else 'FAIL'}] {name} {note}")
        if not cond:
            fails.append(name)

    src = Path(__file__).read_text(encoding="utf-8")
    chk("① 抽取道優先序=digest ㉓㉙ 同族(fitz 版面序→pypdf 後備→NEEDS_OCR)",
        'get_text("text", sort=True)' in src and "PYPDF_FALLBACK" in src
        and "NEEDS_OCR" in src)
    with tempfile.TemporaryDirectory() as td:
        tdp = Path(td)
        globals()["OUTDIR"] = tdp / "_out"      # 批410:輸出也進暫存夾
        (tdp / "_out").mkdir(parents=True, exist_ok=True)
        import fitz
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 100), "TEST REPORT 2330 TT")
        page.insert_text((72, 130), "Target Price 1,500")
        # 批444:這個 fixture 原本首頁只有 38 個字。密度閘一上,38 字 /
        # A4 = 7.58e-05 低於門檻,它被**正確地**判成掃描頁——因為 38 個字
        # 的一頁,本來就和只有浮水印的掃描頁分不出來。
        # 修法不是把閘放寬(那等於白做),是把 fixture 改成**像一頁真報告**:
        # 真的券商報告首頁 400-4000 字。fixture 太瘦就代表不了它要代表的東西。
        for _i in range(22):
            page.insert_text((72, 170 + _i * 14),
                             "We reiterate our BUY rating on the back of stronger "
                             "foundry pricing and a firm advanced-node mix.",
                             fontsize=9)
        doc.new_page().insert_text((72, 100), "PAGE TWO SHOULD NOT APPEAR")
        doc.save(str(tdp / "fx_report.pdf"))
        doc.close()
        txt, tag = extract_pdf_page1(tdp / "fx_report.pdf")
        _tri2 = triage_page1(tdp / "fx_report.pdf")
        chk("② 首頁抽取實測(fixture 2 頁 PDF 僅取第 1 頁;批444 起 fixture 加厚成"
            "像一頁真報告——38 字的一頁和只有浮水印的掃描頁本來就分不出來)",
            tag == "FITZ_LAYOUT" and "2330" in txt
            and "PAGE TWO" not in txt,
            f"(標記={tag} · 分流={_tri2.get('state')} {_tri2.get('n_chars')}字)")
        # 批236 fixture:雙欄頁(左本文斷行+右資訊卡)
        d2 = fitz.open()
        pg = d2.new_page()   # 預設 595x842
        pg.insert_text((40, 60), "EARNINGS UPSIDE REITERATE BUY", fontsize=20)
        pg.insert_text((40, 140), "We expect the momentum to")
        pg.insert_text((40, 158), "continue in coming quarters")
        pg.insert_text((40, 176), "driven by GB200 ramp.")
        pg.insert_text((360, 140), "Buy")
        pg.insert_text((360, 158), "Target price NT$165.00")
        pg.insert_text((360, 176), "Price NT$114.00")
        pg.insert_text((40, 820), "Disclaimer fine print", fontsize=6)
        d2.save(str(tdp / "fx_twocol.pdf"))
        d2.close()
        z = extract_page1_zones(tdp / "fx_twocol.pdf")
        zp = extract_page1_plumber(tdp / "fx_twocol.pdf")
        cmpz = compare_zones(z, zp) if zp else {}
        chk("⑪ 法B pdfplumber 同判準分區+雙法對照(body/right AGREE)",
            zp is not None and "Target price" in zp["right"]
            and cmpz.get("body", {}).get("verdict") in ("AGREE", "PARTIAL")
            and cmpz.get("right", {}).get("verdict") in ("AGREE", "PARTIAL"))
        cd = compare_zones({"body": "abc def", "right": "", "header": ""},
                           {"body": "totally different text!", "right": "",
                            "header": ""})
        chk("⑫ 對照三態機制(DIVERGE 誠實判離+BOTH_EMPTY)",
            cd["body"]["verdict"] == "DIVERGE"
            and cd["right"]["verdict"] == "BOTH_EMPTY")
        chk("⑨ 分區還原(批236:左本文/右資訊卡切開零交錯)",
            z is not None and "Target price" in z["right"]
            and "Target price" not in z["body"]
            and "momentum" in z["body"] and "momentum" not in z["right"])
        chk("⑩ 本文斷行修復+標題階層+頁尾雜訊帶",
            "momentum to continue" in z["body"].replace("\n", " ")
            and z["body"].count("\n") <= 1
            and any(h["level"] == "H1" for h in z["heads"])
            and "Disclaimer" in z["footer"])
        g = z.get("gle", {})
        chk("⑬ 法C GLE 掛載(批243:九宮分區標+字重階層+後端矩陣;"
            "缺席=誠實 absent)",
            ("gle" in z) and (not g.get("available")
             or (g.get("elements") and g.get("backends_total", 0) > 0
                 and any(e["zone"] for e in g["elements"]))))
        sents = split_sentences_repaired("營收強勁。毛利率回升,展望正向。維持買進")
        chk("⑭ 句級修復(批243:一句一資料;句界切分;標題不經句道)",
            len(sents) >= 2 and sents[0].endswith("。")
            and "sentences" in z)
        blank = fitz.open()
        blank.new_page()
        blank.save(str(tdp / "fx_scan.pdf"))
        blank.close()
        t2, tag2 = extract_pdf_page1(tdp / "fx_scan.pdf")
        # 批443:標記改成 NEEDS_OCR[為何]——「候 OCR」三個字讓人不知道下一步。
        # 後端一支都沒裝時仍然**絕不假抽**,只是把路由與裝法一起交代出去。
        chk("③ 掃描版誠實(零文字=NEEDS_OCR 不假抽;批443 起附帶為何)",
            tag2.startswith("NEEDS_OCR") and "OCR_" in tag2
            and t2 == "")
        rc = run(tdp)
        page_html = (OUTDIR / "FIRSTPAGE_SUMMARY.html").read_text(encoding="utf-8")
        chk("④ 逐件 .txt+總覽 HTML 產出(一頁堆疊)", rc == 0
            and (OUTDIR / "fx_report.txt").exists()
            and "首頁文字擷取總覽" in page_html)
        chk("⑤ 空夾誠實 rc2(缺件先搜集)",
            run(tdp / "nothing_x") == 2)
    globals()["OUTDIR"] = _OUTDIR_LIVE          # 批410:還原
    _live_after = sorted(x.name for x in _OUTDIR_LIVE.glob("*")) if _OUTDIR_LIVE.exists() else []
    # 判準只問「這一次跑有沒有新增」——舊機器上先前 v0105 留下的殘件不是本次的錯,
    # 那是 purge-selftest 的工作;拿舊殘件判本次紅,就是判準綁錯前提(批408 教訓)。
    _added = sorted(set(_live_after) - set(_live_before))
    _stale = [n for n in _live_after if n.startswith(("fx_report", "fx_scan", "fx_twocol"))]
    chk("⑮ 自測零污染正式產出夾(批410:本次跑完 VIA_Reports/first_page_text/ 零新增;"
        "fixture 不得外流成假報告被 ENG073 收進正本庫)",
        _added == [],
        f"(本次新增 {len(_added)}" +
        (f";另有前版殘件 {len(_stale)} 件 → 跑 `python VRN_ENG072_..._v0106.py purge-selftest --apply` 搬進隔離夾"
         if _stale else "") + ")")
    chk("⑥ docx 道誠實(無頁界=前段近似標 DOCX_HEAD;缺庫誠實提示)",
        "DOCX_HEAD" in src and "DOCX_LIB_MISSING" in src)
    chk("⑦ 紅線宣告(抽取物不入 git;原件僅本機)",
        "不入 git" in src)
    chk("⑧ 零網路+加速橋(純本地抽取)",
        "ACCEL-BRIDGE" in src
        and all(("import " + k) not in src for k in ("requests", "httpx")))
    # 批421:兩座收容件改由支援模組統轄。這檢要證三件事,缺一都不算:
    #  (a) 真的走橋(route=HUB),不是名義上接了、實際仍走舊直掛;
    #  (b) NLP 拿到的是尾版 v1.8.0 而非寫死的 v1.1.0——用判別輸入證明
    #      (連續空白 stdlib NFKC 不併、正主會併;全形轉半形兩者相同=測不出來);
    #  (c) 橋缺席時退路仍在(對照組:把 HUB_DIR 指到空夾,兩者都得退回直掛而非炸掉)。
    _probe = "台積電  的的的營收"
    _tp = _nlp_tp()
    _tp_out = _tp.normalize(_probe) if _tp is not None else ""
    _std_out = unicodedata.normalize("NFKC", _probe)
    _ann = gle_annotate(
        [{"x0": 10, "y0": 10, "x1": 300, "y1": 40, "size": 18.0,
          "font": "Arial-Bold", "bold": True, "lines": ["台積電 2330 TT"]},
         {"x0": 10, "y0": 60, "x1": 500, "y1": 400, "size": 10.0,
          "font": "Arial", "bold": False, "lines": ["本文" * 40]}], 595.0, 842.0)
    _hs = hub_stats()      # route 是「走過才有」的事實 → 操練後才讀
    global HUB_DIR
    _keep_dir, _keep_hub = HUB_DIR, dict(_HUB)
    try:                                     # (c) 缺席對照組
        HUB_DIR = HERE / "_no_such_hub_dir_"
        _HUB.clear()
        _fb_gle = _gle_mod()
        _fb_tp = _nlp_tp()
        _fb_ok = (_fb_gle is not None and _fb_tp is not None
                  and _HUB["gle"]["route"] == "LEGACY_DIRECT"
                  and _HUB["nlp"]["route"] == "LEGACY_v1.1.0")
        _fb_why = _HUB["gle"].get("why", "")
    finally:
        HUB_DIR = _keep_dir
        _HUB.clear()
        _HUB.update(_keep_hub)
    chk("⑯ 兩收容件改由支援模組統轄(批421;SUP_MDL743 GLE + SUP_MDL744 NLP)"
        "——真走橋、NLP 是尾版 v1.8.0 非寫死 v1.1.0、且橋缺席仍有直掛退路",
        _hs["GLE"]["on"] and _hs["NLP"]["on"]
        and _hs["GLE"]["route"] == "HUB" and _hs["NLP"]["route"] == "HUB"
        and _tp_out == "台積電 的的的營收" and _std_out != _tp_out
        and _ann.get("available") and _ann.get("backends_total", 0) > 0
        and _fb_ok and bool(_fb_why),
        f"(GLE={_hs['GLE']['src'] or _hs['GLE']['why']} · "
        f"NLP={_hs['NLP']['src'] or _hs['NLP']['why']} · "
        f"正主={_tp_out!r} vs stdlib={_std_out!r} · "
        f"後端 {len(_ann.get('backends_available', []))}/{_ann.get('backends_total', 0)} · "
        f"退路={'在' if _fb_ok else '斷'})")

    # ⑰ 字元密度分流閘接得上,而且擋在分區道**前面**
    _ppp = ppp_hub()
    _tri_ok = _ppp is not None
    _src72 = Path(__file__).read_text(encoding="utf-8")
    _order_ok = (_src72.index("tri = triage_page1(p)")
                 < _src72.index("else extract_page1_zones(p)"))
    _states = {}
    if _tri_ok:
        for _f in sorted(VIA.rglob("*.pdf")):
            _states.setdefault(triage_page1(_f).get("state"), []).append(_f.name)
    chk("⑰ 字元密度分流閘(批444;SUP_MDL746)接得上,而且擋在**分區道前面**"
        "——浮水印那十幾個字照樣會讓 zones[header] 非空,閘擋在後面等於沒擋;"
        "橋缺席時回 UNKNOWN 走 v0109 原路(零回歸)",
        _order_ok and (not _tri_ok or (len(_states.get("DIGITAL", [])) >= 5
                                       and len(_states.get("SCANNED", [])) >= 1)),
        f"(橋={'在' if _tri_ok else '缺→UNKNOWN 原路'} · 閘在分區前={_order_ok} · "
        + " · ".join(f"{k} {len(v)}" for k, v in sorted(_states.items())) + ")")

    # ⑱ 候OCR 那格數得對(v0109 起 NEEDS_OCR 帶因由後綴就對不上鍵,全被算進 OTHER)
    _st18 = {"DUAL_ZONES": 0, "FITZ_ZONES": 0, "FITZ_LAYOUT": 0,
             "PYPDF_FALLBACK": 0, "NEEDS_OCR": 0, "DOCX": 0, "OTHER": 0, "THIN": 0}
    for _tag18 in ("NEEDS_OCR[OCR_BACKENDS_ABSENT(x)|PPP_OCR_ABSENT(缺 paddleocr)]",
                   "DUAL_ZONES[THIN]", "FITZ_LAYOUT", "DOCX_HEAD"):
        _b18 = _tag18.split("[", 1)[0]
        _k18 = _b18 if _b18 in _st18 else ("DOCX" if _b18.startswith("DOCX") else "OTHER")
        _st18[_k18] += 1
        if "[THIN]" in _tag18:
            _st18["THIN"] += 1
    chk("⑱ 候OCR/稀薄兩格數得對——批443 起 NEEDS_OCR 帶了因由後綴,v0109 的 "
        "`tag if tag in stats` 對不上鍵,整批被算進 OTHER,候OCR 那格一直報 0;"
        "改取「[」前的基名",
        _st18["NEEDS_OCR"] == 1 and _st18["DUAL_ZONES"] == 1 and _st18["THIN"] == 1
        and _st18["DOCX"] == 1 and _st18["OTHER"] == 0,
        f"(候OCR {_st18['NEEDS_OCR']} · 雙法 {_st18['DUAL_ZONES']} · "
        f"稀薄 {_st18['THIN']} · docx {_st18['DOCX']} · 其他 {_st18['OTHER']})")

    # ⑲ 批449:三階梯的**升階條件**與**不升階條件**
    _src49 = Path(__file__).read_text(encoding="utf-8")
    _order49 = (_src49.index("第一階(非 OCR)就在上面")
                < _src49.index("第二階:原畫質 OCR")
                < _src49.index("第三階:**高畫質 OCR**"))
    # 工作站實錄的兩種第二階結果,升階與否要分得開
    _tag_ran = ("OCR_EMPTY(tesseract+paddleocr+paddle_ppstructure;跑了但零字=影像可能不可辨)"
                "|PPP_OCR_ABSENT(缺 paddleocr)")
    _tag_absent = ("OCR_BACKENDS_ABSENT(路由 7 支全未裝:tesseract,paddleocr…)"
                   "|PPP_OCR_ABSENT(缺 paddleocr)")
    _esc_ran = ("_EMPTY(" in _tag_ran) or ("OCR_ERROR" in _tag_ran)
    _esc_abs = ("_EMPTY(" in _tag_absent) or ("OCR_ERROR" in _tag_absent)
    _d49, _w49 = hq_dpi_for(VIA / "__no_such_hq__.pdf")
    chk("⑲ 三階梯(批449 操作員令「先用非 OCR 法,再用 OCR 法,最後還抓不到再調"
        "高畫質 OCR 法」):① 非 OCR(fitz→pypdf,密度閘守門)② 原畫質 OCR"
        "(GLE 七支→PPP 補位)③ **高畫質 OCR**(用 MDL001 批448 的 auto_dpi 把第 1 頁"
        "重繪到 300–350 帶內,再餵**同一條** OCR 車道——變因只有解析度一個,"
        "才說得清是不是解析度的問題)。升階條件是「**跑了但零字**」不是「沒裝」:"
        "後端根本不在位時升階是白跑,誠實跳過並講明",
        _order49 and _esc_ran and not _esc_abs and 300 <= _d49 <= 350,
        f"(階序正確={_order49} · 跑了但零字→升階={_esc_ran} · "
        f"後端不在位→升階={_esc_abs}(應為 False) · 第三階 DPI={_d49})")

    # ⑳ 第三階不得另開一條 OCR 路由,且暫存不落地
    _hq_body = _src49[_src49.index("def ocr_page1_hq"):_src49.index("def _ocr_via_ppp")]
    chk("⑳ 第三階**不另寫一條 OCR 路由**——重繪完就交回 ocr_page1(同一組後端、"
        "同一套路由);暫存 PDF 走 TemporaryDirectory 跑完即刪,不落地污染",
        "ocr_page1(hq)" in _hq_body and "TemporaryDirectory" in _hq_body
        and "_ocr_via_gle" not in _hq_body and "route_modes" not in _hq_body,
        "(交回 ocr_page1=True · 暫存自刪=True · 未自建路由=True)")

    # ㉑ 批450:不准編造「跑了」,而且升階判準要跟得上標記格式
    _src50 = Path(__file__).read_text(encoding="utf-8")
    class _FakeProbe:
        reason = "available"
    class _FakeRes:
        def __init__(self, nm, st, n=0, err=None, warns=()):
            self.adapter_name, self.status = nm, st
            self.elements = [None] * n
            self.error, self.warnings, self.probe = err, list(warns), _FakeProbe()
    class _FakeRun:
        def __init__(self, route, rs):
            self.route, self.adapter_results = route, rs
            self.canonical_elements = []
    _det = _adapter_detail(_FakeRun(["tesseract"], [
        _FakeRes("tesseract", "SUCCESS", 0,
                 warns=["missing Tesseract languages: chi_tra, chi_sim; using eng"]),
        _FakeRes("easyocr", "SKIPPED_UNAVAILABLE", 0)]))
    _det_empty = _adapter_detail(_FakeRun([], []))
    _det_skip = _adapter_detail(_FakeRun(["easyocr"], [
        _FakeRes("easyocr", "SKIPPED_UNAVAILABLE", 0)]))
    # 升階判準:三態要分得開
    _t_ran = "OCR_EMPTY[tesseract;tesseract:SUCCESS/0元素 警=missing Tesseract languages]"
    _t_noroute = ("OCR_ROUTE_EMPTY(編排器回的路由是空的=**沒有任何後端真的跑**;"
                  "矩陣說在位的是 tesseract…)")
    _t_absent = "OCR_BACKENDS_ABSENT(路由 7 支全未裝:…)|PPP_OCR_ABSENT(缺 paddleocr)"
    _t_pppempty = "OCR_BACKENDS_ABSENT(…)|PPP_OCR_EMPTY(PaddleOCR(CPU);跑了但零字)"
    chk("㉑ 不准編造「跑了」(批450:v0111 的 `used = run.route **or usable**` 在 route "
        "空掉時拿「矩陣說裝了的那些」去填,再宣稱它們跑過——工作站實錄裡同一行既說 "
        "GLE 的 paddleocr 跑過、又說 PPP「缺 paddleocr」,兩句不可能都真。而真正的因由"
        "**早就在 run.adapter_results 裡**:GLE 的 tesseract 轉接器自己會寫 "
        "`missing Tesseract languages: …; using …`,我卻把整包丟掉再掰一句"
        "「影像可能不可辨」——批419e 那一課的最深一層:捕捉到卻不顯示,還編一個代替它)。"
        "且升階判準改成具名函式:標記格式從 `OCR_EMPTY(` 改成 `OCR_EMPTY[` 之後,"
        "舊的 `\"_EMPTY(\" in tag` 會**永遠為假**,第三階從此靜靜不升階而沒有一行會說",
        # 自審:第一版還加了一條「原始碼裡不准再出現某串字」——而**本檔這一段
        # 敘述文字裡就寫著那串字**(我在註解裡引用了舊寫法),於是恆為假。
        # 這是本回合第三次踩同一個自 match 陷阱(SUP_MDL746 ②、MDL141 ⑭)。
        # 拿原始碼當證據這件事本身就脆;判準全部改成**行為**。
        "missing Tesseract languages" in _det
        and "chi_tra" in _det and "SKIPPED_UNAVAILABLE" in _det
        and "沒有任何後端留下紀錄" in _det_empty
        and "SKIPPED_UNAVAILABLE/0元素" in _det_skip
        and ocr_did_run(_t_ran) and ocr_did_run(_t_pppempty)
        and not ocr_did_run(_t_noroute) and not ocr_did_run(_t_absent),
        f"(逐支因由照抄={('missing Tesseract languages' in _det)} · "
        f"跑了但零字→升階={ocr_did_run(_t_ran)} · "
        f"路由空→不升={not ocr_did_run(_t_noroute)} · "
        f"不在位→不升={not ocr_did_run(_t_absent)})")

    print(f"  [計] 二十一檢({len(done)} 檢) OK {len(done) - len(fails)} · FAIL {len(fails)}")
    return 1 if fails else 0


# 批410:已外流的 fixture 清理道。v0106 之前跑過自測的機器,正式產出夾裡已經有
# fx_report/fx_scan/fx_twocol 這些檔,而且很可能已被 ENG073 收進正本庫。
# 原則:①預設 dry-run,只列不動 ②動也只是**搬進隔離夾**不是刪(可逆、資訊不丟)
# ③正本庫的列**絕不代刪**——只印出對應列與該下的 SQL,由操作員自己決定。
SELFTEST_STEMS = ("fx_report", "fx_scan", "fx_twocol")
QUARANTINE = "_selftest_quarantine"


def purge_selftest(apply: bool = False) -> int:
    """把自測 fixture 自正式首頁產出夾搬進隔離夾(可逆);並列出正本庫對應列與 SQL。"""
    if not OUTDIR.exists():
        print(f"[fixture 清理] 產出夾不存在({OUTDIR})=無事可做")
        return 0
    hits = sorted(x for x in OUTDIR.glob("*")
                  if x.is_file() and x.stem in SELFTEST_STEMS)
    print(f"[fixture 清理] 正式產出夾 {OUTDIR}")
    if not hits:
        print("  乾淨:無自測 fixture 外流件")
    for x in hits:
        print(f"  [{'搬移' if apply else '待搬'}] {x.name}")
    if apply and hits:
        q = OUTDIR / QUARANTINE
        q.mkdir(parents=True, exist_ok=True)
        for x in hits:
            x.replace(q / x.name)
        print(f"  → 已搬進 {q}(未刪除;要復原直接搬回即可)")
    elif hits:
        print("  → dry-run;加 --apply 才搬(搬進隔離夾,不刪除)")
    # 正本庫側:只查只印,絕不代刪
    try:
        import duckdb
        # VRN 的報告表其實住在 VDF 正本庫(ENG073/MDL141 的 DB_TW),不是 VRN 夾下;
        # 第一版我對著 VRN 夾 glob,結果一列都查不到——路徑要跟真正寫入者一致。
        cands = [VIA / "functional modules" / "VDF" / "output_hub" / "mega" / "vdf_tw_market.duckdb"]
        cands += sorted((VIA / "functional modules" / "VRN").glob("**/*.duckdb"))
        cands = [c for c in cands if c.exists()]
        for db in cands:
            con = duckdb.connect(str(db), read_only=True)
            try:
                tabs = {t for (t,) in con.execute("SHOW TABLES").fetchall()}
                for tb in ("vrn_report_basic", "vrn_report_metrics",
                           "vrn_report_financial", "vrn_report_crosscheck"):
                    if tb not in tabs:
                        continue
                    col = "report_file" if tb == "vrn_report_crosscheck" else "report_file"
                    cols = {r[0] for r in con.execute(f"DESCRIBE {tb}").fetchall()}
                    if col not in cols:
                        continue
                    ors = " OR ".join([f"{col} LIKE '%{st}%'" for st in SELFTEST_STEMS])
                    n = con.execute(f"SELECT COUNT(*) FROM {tb} WHERE {ors}").fetchone()[0]
                    if n:
                        print(f"  [正本庫] {db.name}.{tb} 有 {n} 列來自自測 fixture")
                        print(f"           要清請自己下(我不代刪):"
                              f"DELETE FROM {tb} WHERE {ors};")
            finally:
                con.close()
    except ImportError:
        print("  [正本庫] duckdb 未安裝=跳過庫側查核(誠實)")
    except Exception as exc:
        print(f"  [正本庫] 查核跳過:{type(exc).__name__}")
    return 0


def main() -> int:
    args = sys.argv[1:]
    if "--selftest" in args:
        print("=== 首頁三法整合擷取器(VRN_ENG072 v0112)· 二十一檢自測(零網路)===")
        return selftest()
    if args and args[0] in ("purge-selftest", "purge"):
        return purge_selftest(apply="--apply" in args)
    if args and args[0] == "run":
        d = None
        if "--dir" in args:
            d = Path(args[args.index("--dir") + 1])
        return run(d, "--open" in args)
    print(__doc__)
    return 0


if __name__ == "__main__":
    sys.exit(main())
