# -*- coding: utf-8 -*-
"""ENG-065 Local attachment metadata/text parser. No OCR."""
from __future__ import annotations
import argparse,csv,hashlib,json,re
from datetime import datetime
from pathlib import Path
HERE=Path(__file__).resolve().parent;ROOT=HERE.parent;OUT=ROOT/"out";ATT=OUT/"attachments";REG=OUT/"attachment_registry.jsonl";RESULT=OUT/"attachment_intelligence.jsonl";LINK=OUT/"attachment_link_registry.jsonl"
def jl(p):
    rows=[]
    if not p.exists():return rows
    for ln in p.read_text(encoding="utf-8",errors="replace").splitlines():
        try:rows.append(json.loads(ln))
        except Exception:pass
    return rows
def attachment_sys_id(message_graph_id,attachment_id):
    key=f"{message_graph_id}|{attachment_id}"
    for r in jl(LINK):
        if r.get("key")==key:return r.get("sys_attachment_id")
    from workops_autocode import next_code
    aid=next_code("ATT")
    with LINK.open("a",encoding="utf-8") as f:f.write(json.dumps({"key":key,"sys_attachment_id":aid,"ts":datetime.now().astimezone().isoformat(timespec="seconds")},ensure_ascii=False)+"\n")
    return aid
def sha(p):
    h=hashlib.sha256()
    with p.open("rb") as f:
        for b in iter(lambda:f.read(1048576),b""):h.update(b)
    return h.hexdigest()
def text_of(p):
    ext=p.suffix.lower()
    if ext in (".txt",".csv",".json",".html",".htm"):return p.read_text(encoding="utf-8",errors="replace")
    if ext==".xlsx":
        import openpyxl
        wb=openpyxl.load_workbook(p,read_only=True,data_only=True);parts=[]
        for ws in wb.worksheets:
            parts.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                parts.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(parts)
    if ext==".docx":
        from docx import Document
        return "\n".join(x.text for x in Document(p).paragraphs)
    if ext==".pdf":
        import fitz
        doc=fitz.open(p);return "\n".join(page.get_text("text") for page in doc)
    return ""
def build():
    meta=jl(REG);rows=[]
    for m in meta:
        p=Path(m.get("local_path") or "")
        rec={**m,"sys_attachment_id":attachment_sys_id(m.get("message_graph_id"),m.get("attachment_id")),"parsed":False,"text_excerpt":"","parse_error":""}
        if p.exists() and p.is_file():
            rec["sha256"]=sha(p)
            try:
                txt=text_of(p);rec["parsed"]=bool(txt);rec["text_excerpt"]=txt[:12000]
                rec["product_codes"]=sorted(set(re.findall(r"\b[A-Z]{2,5}[-_ ]?\d{3,10}\b",txt,re.I)))[:50]
            except Exception as e:rec["parse_error"]=str(e)
        rows.append(rec)
    with RESULT.open("w",encoding="utf-8") as f:
        for r in rows:f.write(json.dumps(r,ensure_ascii=False)+"\n")
    return {"engine_id":"ENG-065","gate":"PASS","attachments":len(rows),"parsed":sum(x["parsed"] for x in rows),"warnings":sum(bool(x["parse_error"]) for x in rows)}
def main():
    argparse.ArgumentParser().parse_args();print(json.dumps(build(),ensure_ascii=False,indent=2));return 0
if __name__=="__main__":raise SystemExit(main())
